# -*- coding: utf-8 -*-
"""
captura_origen.py  ·  P1 · Captura íntegra de origen (pasiva)
================================================================

Regla obligatoria: NORMALIZAR NUNCA DEBE DESTRUIR INFORMACIÓN DE ORIGEN.

Este módulo genera un archivo independiente, ORIGEN.xlsx, con cuatro hojas:

    DATOS_ORIGINALES     una fila por CELDA NO VACÍA de TODAS las hojas del extracto
    METADATOS_EXTRACTO   una fila por extracto (tabla fija)
    METADATOS_EXTRA      cabecera / pie / filas especiales / hojas extra, como etiqueta→valor
    MAPA_ORIGEN          una fila por movimiento normalizado: CLAVE TRANSACCIÓN ↔ ID_ORIGEN

Garantías de diseño
-------------------
* Es PASIVO: no importa el motor, no llama a ninguna función normalizar_*, no modifica
  ningún DataFrame ni ningún archivo de entrada. Lee el archivo bancario por su cuenta
  (xlrd / openpyxl / calamine) con las coordenadas reales de Excel.
* No depende de pandas para LEER las celdas (pandas convierte "NA", "N/A", "null"… en
  vacío y ajusta tipos). Aquí cada celda se guarda como texto exacto + tipo original.
* IDs determinísticos: ID_EXTRACTO = SHA-256 de los bytes del archivo;
  ID_ORIGEN = "{ID_EXTRACTO[:16]}|{HOJA}|R{FILA_EXCEL}". No dependen del lote ni de la hora.
* Un fallo de captura NUNCA debe romper la normalización: `generar_origen_seguro`
  atrapa cualquier excepción y devuelve el estado; el motor sigue produciendo
  NORMALIZADO.xlsx y LISTS.csv exactamente igual que antes.

Enlace movimiento ↔ fila original
---------------------------------
`normalizar_*` conserva el índice de la tabla leída (pandas, header=fila_encabezado), de
modo que  FILA_EXCEL = fila_encabezado_0based + 2 + índice.  El módulo recibe las tablas
normalizadas de cada archivo (con su índice original, antes del `pd.concat`) y aplica esa
regla; después verifica que cada fila apuntada exista, no esté vacía y no sea la fila de
encabezado. Cualquier inconsistencia levanta ValueError (y el ORIGEN no se escribe).

Roles de fila (ROL_FILA)
------------------------
Hoja de movimientos: CABECERA (antes del encabezado) · ENCABEZADO_TABLA · MOVIMIENTO (la fila
fue normalizada) · SALDO_INICIAL / SALDO_CIERRE / TOTAL_PIE (fila no normalizada con un texto que
empieza con «saldo inicial», «saldo al cierre», «total»…) · FILA_ESPECIAL (otra fila no normalizada antes del último
movimiento) · PIE (fila no normalizada después del último movimiento).
Cualquier otra hoja: HOJA_EXTRA.
"""
import datetime as _dt
import hashlib
import math
import os
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

VERSION_CAPTURA = "P1-1.0"
NOMBRE_ORIGEN = "ORIGEN.xlsx"

COLS_DATOS = [
    "ID_EXTRACTO", "ID_ORIGEN", "ID_CELDA", "ARCHIVO ORIGEN", "HOJA", "FILA_EXCEL",
    "COLUMNA_EXCEL", "ROL_FILA", "CAMPO_ORIGINAL", "CAMPO_CANONICO", "TIPO_CELDA",
    "VALOR_ORIGINAL",
]
COLS_META = [
    "ID_EXTRACTO", "ARCHIVO ORIGEN", "TAMAÑO_BYTES", "LOTE DE CARGA", "FECHA DE CARGA",
    "VERSION_CAPTURA", "FORMATO", "BANCO", "CUENTA BANCARIA", "MONEDA",
    "HOJAS_EN_ARCHIVO", "HOJA_LEIDA", "FILA_ENCABEZADO", "PUNTAJE_ENCABEZADO",
    "COLUMNAS_ORIGINALES", "COLUMNAS_SIN_DATOS",
    "CELDAS_NO_VACIAS", "FILAS_TABLA", "FILAS_MOVIMIENTO", "FILAS_ESPECIALES", "FILAS_PIE",
    "MOVIMIENTOS_CREDITO", "MOVIMIENTOS_DEBITO",
    "SALDO_INICIAL_MOTOR", "SALDO_FINAL_MOTOR", "ESTADO_VALIDACION", "DIFERENCIA",
    "OBSERVACIONES",
]
COLS_EXTRA = [
    "ID_EXTRACTO", "HOJA", "ZONA", "FILA_EXCEL", "CELDA_ETIQUETA", "ETIQUETA",
    "CELDA_VALOR", "VALOR_ORIGINAL", "CAMPO_CANONICO",
]
COLS_MAPA = [
    "CLAVE TRANSACCIÓN", "ID_ORIGEN", "ID_EXTRACTO", "ARCHIVO ORIGEN", "HOJA",
    "FILA_EXCEL", "FORMATO", "LOTE DE CARGA",
]

_ILEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# =====================================================================
# 1. LECTURA NATIVA DE CELDAS (coordenadas reales de Excel)
# =====================================================================

def _numero_a_texto(v):
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return repr(v) if isinstance(v, float) else str(v)


def clasificar_valor(v):
    """Valor de Python -> (TIPO_CELDA, texto exacto) o None si la celda está vacía."""
    if v is None:
        return None
    if isinstance(v, bool):
        return ("BOOLEANO", "TRUE" if v else "FALSE")
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None
        return ("NUMERO", _numero_a_texto(v))
    if isinstance(v, _dt.datetime):
        return ("FECHA", v.isoformat())
    if isinstance(v, _dt.date):
        return ("FECHA", v.isoformat())
    if isinstance(v, _dt.time):
        return ("HORA", v.isoformat())
    if isinstance(v, _dt.timedelta):
        s = int(v.total_seconds())
        return ("HORA", f"{s // 3600:02d}:{(s % 3600) // 60:02d}:{s % 60:02d}")
    s = str(v)
    if s == "":
        return None
    return ("TEXTO", s)


def _leer_xls_xlrd(ruta):
    import xlrd
    libro = xlrd.open_workbook(ruta, on_demand=False)
    hojas = {}
    orden = []
    for sh in libro.sheets():
        celdas = {}
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                ct = sh.cell_type(r, c)
                if ct in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                v = sh.cell_value(r, c)
                if ct == xlrd.XL_CELL_DATE:
                    if 0 <= v < 1:
                        v = xlrd.xldate.xldate_as_datetime(v, libro.datemode).time()
                    else:
                        v = xlrd.xldate.xldate_as_datetime(v, libro.datemode)
                elif ct == xlrd.XL_CELL_ERROR:
                    celdas[(r + 1, c + 1)] = ("ERROR", xlrd.error_text_from_code.get(v, str(v)))
                    continue
                elif ct == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                k = clasificar_valor(v)
                if k:
                    celdas[(r + 1, c + 1)] = k
        orden.append(sh.name)
        hojas[sh.name] = celdas
    return orden, hojas


def _leer_calamine(ruta):
    from python_calamine import CalamineWorkbook
    libro = CalamineWorkbook.from_path(ruta)
    hojas, orden = {}, []
    for nombre in libro.sheet_names:
        datos = libro.get_sheet_by_name(nombre).to_python(skip_empty_area=False)
        celdas = {}
        for r, fila in enumerate(datos, start=1):
            for c, v in enumerate(fila, start=1):
                k = clasificar_valor(v)
                if k:
                    celdas[(r, c)] = k
        orden.append(nombre)
        hojas[nombre] = celdas
    return orden, hojas


def _leer_xlsx_openpyxl(ruta):
    libro = load_workbook(ruta, data_only=True)
    hojas, orden = {}, []
    for ws in libro.worksheets:
        celdas = {}
        for fila in ws.iter_rows():
            for cel in fila:
                if cel.value is None:
                    continue
                if cel.data_type == "e":
                    celdas[(cel.row, cel.column)] = ("ERROR", str(cel.value))
                    continue
                k = clasificar_valor(cel.value)
                if k:
                    celdas[(cel.row, cel.column)] = k
        orden.append(ws.title)
        hojas[ws.title] = celdas
    return orden, hojas


def leer_celdas(ruta):
    """
    Devuelve (orden_de_hojas, {hoja: {(fila, col): (TIPO_CELDA, texto)}}) con filas y
    columnas 1-based, solo celdas NO vacías. Elige el lector por el contenido del archivo
    (no por la extensión) y reintenta con calamine si xlrd/openpyxl fallan.
    """
    with open(ruta, "rb") as f:
        cabecera = f.read(8)
    intentos = []
    if cabecera.startswith(b"PK"):
        intentos = [_leer_xlsx_openpyxl, _leer_calamine]
    else:
        intentos = [_leer_xls_xlrd, _leer_calamine]
    ultimo = None
    for lector in intentos:
        try:
            return lector(ruta)
        except Exception as e:  # noqa: BLE001 - se reintenta con el siguiente lector
            ultimo = e
    raise ValueError(f"No se pudo leer '{ruta}' con ningún lector: {ultimo!r}")


# =====================================================================
# 2. UTILIDADES
# =====================================================================

def sha256_archivo(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloque)
    return h.hexdigest()


def id_origen(id_extracto, hoja, fila):
    return f"{id_extracto[:16]}|{hoja}|R{fila}"


def _norm(txt):
    t = unicodedata.normalize("NFD", str(txt))
    t = "".join(ch for ch in t if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", t).strip().upper()


_ROLES_ESPECIALES = [
    ("SALDO_INICIAL", re.compile(r"^SALDO (INICIAL|ANTERIOR|AL INICIO)")),
    ("SALDO_CIERRE", re.compile(r"^SALDO (FINAL|CIERRE|AL CIERRE|ACTUAL)")),
    ("TOTAL_PIE", re.compile(r"^TOTAL")),
]


def _rol_especial(celdas_fila):
    """Rol descriptivo de una fila NO normalizada según sus textos (sin lógica por banco)."""
    for col in sorted(celdas_fila):
        tipo, txt = celdas_fila[col]
        if tipo == "TEXTO" and txt.strip():
            t = _norm(txt)
            for rol, rx in _ROLES_ESPECIALES:
                if rx.match(t):
                    return rol
    return None


# =====================================================================
# 3. CAPTURA DE UN EXTRACTO
# =====================================================================

def capturar_extracto(
    ruta, nombre_origen, formato, contrato, indices_movimiento, tabla=None,
    lote="", fecha_carga=None, validacion=None,
):
    """
    Captura un extracto completo. No transforma nada.

    contrato : dict con
        "hojas_validas"            {formato: hoja de movimientos}
        "encabezados_esperados"    {formato: [encabezados]}
        "encontrar_fila_encabezado" callable(df_raw, formato) -> (fila0, puntaje, total)
        "version_motor"            texto (opcional)
    indices_movimiento : índices (pandas) de las filas que el motor normalizó
    tabla              : DataFrame normalizado del archivo (opcional; solo se leen
                         BANCO / CUENTA BANCARIA / MONEDA / CLAVE TRANSACCIÓN / TIPO MOVIMIENTO)
    validacion         : dict/Series de validar_archivo para este archivo (opcional)
    Devuelve dict con listas de filas: datos, meta (1 fila), extra, mapa.
    """
    import pandas as pd  # solo para reutilizar encontrar_fila_encabezado del motor

    id_ext = sha256_archivo(ruta)
    orden_hojas, hojas = leer_celdas(ruta)
    hoja_mov = contrato["hojas_validas"][formato]
    if hoja_mov not in hojas:
        raise ValueError(f"{nombre_origen}: no existe la hoja de movimientos '{hoja_mov}'")

    # -- fila de encabezado, con las mismas reglas que el motor pero en coordenadas nativas
    celdas_mov = hojas[hoja_mov]
    max_r = max((r for r, _ in celdas_mov), default=0)
    max_c = max((c for _, c in celdas_mov), default=0)
    grid = [[None] * max_c for _ in range(max_r)]
    for (r, c), (_, txt) in celdas_mov.items():
        grid[r - 1][c - 1] = txt
    raw = pd.DataFrame(grid, dtype=object)
    fila0, puntaje, total_esp = contrato["encontrar_fila_encabezado"](raw, formato)
    fila_hdr = int(fila0) + 1  # 1-based
    encabezado = {c: txt for (r, c), (_, txt) in celdas_mov.items() if r == fila_hdr}
    if not encabezado:
        raise ValueError(f"{nombre_origen}: la fila de encabezado {fila_hdr} está vacía")

    # -- filas de movimiento (índice pandas -> fila de Excel)
    indices = [int(i) for i in indices_movimiento]
    if len(set(indices)) != len(indices):
        raise ValueError(f"{nombre_origen}: índices de movimiento repetidos")
    filas_mov = [fila_hdr + 1 + i for i in indices]
    conj_mov = set(filas_mov)
    for r in filas_mov:
        if r <= fila_hdr or not any(rr == r for (rr, _) in celdas_mov):
            raise ValueError(f"{nombre_origen}: el movimiento apunta a la fila {r}, que no existe o está vacía")
    ultima_mov = max(filas_mov) if filas_mov else None

    filas_por_hoja = {}
    for hoja, celdas in hojas.items():
        d = {}
        for (r, c), v in celdas.items():
            d.setdefault(r, {})[c] = v
        filas_por_hoja[hoja] = d

    # -- roles por fila de la hoja de movimientos
    roles = {}
    for r, fila in filas_por_hoja[hoja_mov].items():
        if r < fila_hdr:
            roles[r] = "CABECERA"
        elif r == fila_hdr:
            roles[r] = "ENCABEZADO_TABLA"
        elif r in conj_mov:
            roles[r] = "MOVIMIENTO"
        else:
            rol = _rol_especial(fila)
            if rol is None:
                rol = "PIE" if (ultima_mov is None or r > ultima_mov) else "FILA_ESPECIAL"
            roles[r] = rol

    datos, extra = [], []
    for hoja in orden_hojas:
        for r in sorted(filas_por_hoja[hoja]):
            fila = filas_por_hoja[hoja][r]
            es_mov_hoja = hoja == hoja_mov
            rol = roles[r] if es_mov_hoja else "HOJA_EXTRA"
            idr = id_origen(id_ext, hoja, r)
            for c in sorted(fila):
                tipo, txt = fila[c]
                campo = ""
                if es_mov_hoja and r >= fila_hdr:
                    campo = encabezado.get(c, "")
                letra = get_column_letter(c)
                datos.append([
                    id_ext, idr, f"{idr}|{letra}", nombre_origen, hoja, r, letra, rol,
                    campo, "", tipo, txt,
                ])
            if rol not in ("MOVIMIENTO", "ENCABEZADO_TABLA"):
                extra.extend(_pares_etiqueta_valor(id_ext, hoja, r, fila, rol))

    # -- columnas del formato (incluye las que no traen datos)
    cols_orig = [f"{get_column_letter(c)}={encabezado[c]}" for c in sorted(encabezado)]
    con_datos = set()
    for r in conj_mov:
        con_datos.update(filas_por_hoja[hoja_mov][r].keys())
    sin_datos = [f"{get_column_letter(c)}={encabezado[c]}" for c in sorted(encabezado) if c not in con_datos]

    n_esp = sum(1 for r, rol in roles.items() if rol in ("FILA_ESPECIAL", "SALDO_INICIAL", "SALDO_CIERRE", "TOTAL_PIE"))
    n_pie = sum(1 for rol in roles.values() if rol == "PIE")

    banco = cuenta = moneda = ""
    cred = deb = ""
    obs = []
    if tabla is not None and len(tabla) > 0:
        banco = str(tabla["BANCO"].iloc[0]); cuenta = str(tabla["CUENTA BANCARIA"].iloc[0]); moneda = str(tabla["MONEDA"].iloc[0])
        cred = int((tabla["TIPO MOVIMIENTO"] == "CRÉDITO").sum())
        deb = int((tabla["TIPO MOVIMIENTO"] == "DÉBITO").sum())
    elif tabla is not None:
        cred = deb = 0
        obs.append("sin movimientos: BANCO/CUENTA/MONEDA no disponibles hasta existir el registro de cuentas (P3)")
    v = validacion or {}

    def _v(k):
        x = v.get(k, "")
        return "" if x is None or (isinstance(x, float) and math.isnan(x)) else x

    meta = [[
        id_ext, nombre_origen, os.path.getsize(ruta), lote, fecha_carga,
        VERSION_CAPTURA, formato, banco, cuenta, moneda,
        " | ".join(orden_hojas), hoja_mov, fila_hdr, f"{puntaje}/{total_esp}",
        " | ".join(cols_orig), " | ".join(sin_datos),
        sum(len(c) for c in hojas.values()),
        sum(1 for r in filas_por_hoja[hoja_mov] if r > fila_hdr),
        len(filas_mov), n_esp, n_pie, cred, deb,
        _v("SALDO INICIAL"), _v("SALDO FINAL"), _v("ESTADO"), _v("DIFERENCIA"),
        "; ".join(obs),
    ]]

    mapa = []
    if tabla is not None and len(tabla) > 0:
        claves = tabla["CLAVE TRANSACCIÓN"].tolist()
        for clave, r in zip(claves, filas_mov):
            mapa.append([clave, id_origen(id_ext, hoja_mov, r), id_ext, nombre_origen,
                         hoja_mov, r, formato, lote])
    return {"datos": datos, "meta": meta, "extra": extra, "mapa": mapa}


def _pares_etiqueta_valor(id_ext, hoja, r, fila, rol):
    """
    Interpretación derivada (no evidencia) de una fila fuera de la tabla de movimientos.
    Regla general, sin lógica por banco: un texto que termina en ':' es ETIQUETA y su valor
    es la siguiente celda no vacía a la derecha; toda otra celda queda como valor sin
    etiqueta. Cada celda no vacía de la fila aparece exactamente una vez.
    """
    cols = sorted(fila)
    zona = rol
    out = []
    i = 0
    while i < len(cols):
        c = cols[i]
        tipo, txt = fila[c]
        if tipo == "TEXTO" and txt.rstrip().endswith(":") and i + 1 < len(cols):
            c2 = cols[i + 1]
            out.append([id_ext, hoja, zona, r, f"{get_column_letter(c)}{r}", txt,
                        f"{get_column_letter(c2)}{r}", fila[c2][1], ""])
            i += 2
        else:
            out.append([id_ext, hoja, zona, r, "", "", f"{get_column_letter(c)}{r}", txt, ""])
            i += 1
    return out


# =====================================================================
# 4. ESCRITURA DE ORIGEN.xlsx
# =====================================================================

def _celda_texto(ws, valor):
    """Celda de texto puro (nunca fórmula, nunca número) y sin caracteres ilegales de XML."""
    s = _ILEGAL.sub(lambda m: "_x%04X_" % ord(m.group(0)), str(valor))
    if len(s) > 32767:
        raise ValueError("valor de celda mayor que el máximo de Excel (32.767 caracteres)")
    c = WriteOnlyCell(ws, value=s)
    c.data_type = "s"
    return c


def escribir_origen_xlsx(ruta_salida, datos, meta, extra, mapa):
    wb = Workbook(write_only=True)
    hojas = [
        ("DATOS_ORIGINALES", COLS_DATOS, datos, {"FILA_EXCEL"}),
        ("METADATOS_EXTRACTO", COLS_META, meta, {"TAMAÑO_BYTES", "FILA_ENCABEZADO", "CELDAS_NO_VACIAS",
                                                "FILAS_TABLA", "FILAS_MOVIMIENTO", "FILAS_ESPECIALES",
                                                "FILAS_PIE", "MOVIMIENTOS_CREDITO", "MOVIMIENTOS_DEBITO"}),
        ("METADATOS_EXTRA", COLS_EXTRA, extra, {"FILA_EXCEL"}),
        ("MAPA_ORIGEN", COLS_MAPA, mapa, {"FILA_EXCEL"}),
    ]
    for nombre, cols, filas, numericas in hojas:
        ws = wb.create_sheet(nombre)
        ws.append([_celda_texto(ws, c) for c in cols])
        for fila in filas:
            out = []
            for col, v in zip(cols, fila):
                if v is None or v == "":
                    out.append(None)
                elif col in numericas and isinstance(v, int):
                    out.append(v)
                elif col == "FECHA DE CARGA":
                    out.append(str(v))
                else:
                    out.append(_celda_texto(ws, v))
            ws.append(out)
    carpeta = os.path.dirname(os.path.abspath(ruta_salida))
    if carpeta:
        os.makedirs(carpeta, exist_ok=True)
    wb.save(ruta_salida)


# =====================================================================
# 5. PUNTO DE ENTRADA PARA EL MOTOR
# =====================================================================

def generar_origen(mapa_archivos, df_deteccion_final, tablas, df_validacion, lote,
                   fecha_carga, ruta_origen, contrato):
    """
    mapa_archivos       {nombre_original: ruta}
    df_deteccion_final  DataFrame con columnas ARCHIVO, FORMATO (mismo orden que `tablas`)
    tablas              lista de DataFrames normalizados por archivo, CON su índice original
    df_validacion       DataFrame de validar_archivo (una fila por archivo) o None
    """
    datos, meta, extra, mapa = [], [], [], []
    val = {}
    if df_validacion is not None and len(df_validacion):
        for _, f in df_validacion.iterrows():
            val[f["ARCHIVO"]] = f.to_dict()
    filas_det = df_deteccion_final.to_dict("records")
    if len(filas_det) != len(tablas):
        raise ValueError("captura_origen: desalineación entre detecciones y tablas normalizadas")
    for det, tabla in zip(filas_det, tablas):
        nombre, formato = det["ARCHIVO"], det["FORMATO"]
        r = capturar_extracto(
            mapa_archivos[nombre], nombre, formato, contrato, list(tabla.index), tabla=tabla,
            lote=lote, fecha_carga=fecha_carga, validacion=val.get(nombre),
        )
        datos += r["datos"]; meta += r["meta"]; extra += r["extra"]; mapa += r["mapa"]
    # verificaciones de integridad del enlace (1:1 dentro del lote)
    claves = [m[0] for m in mapa]
    ids = [m[1] for m in mapa]
    if len(set(claves)) != len(claves) or len(set(ids)) != len(ids):
        raise ValueError("captura_origen: el enlace CLAVE TRANSACCIÓN ↔ ID_ORIGEN no es 1:1")
    escribir_origen_xlsx(ruta_origen, datos, meta, extra, mapa)
    return {
        "ruta_origen": ruta_origen, "celdas": len(datos), "extractos": len(meta),
        "movimientos_mapeados": len(mapa),
    }


def generar_origen_seguro(*args, **kwargs):
    """Igual que generar_origen, pero jamás propaga una excepción: devuelve el estado."""
    try:
        info = generar_origen(*args, **kwargs)
        info["estado"] = "OK"
        return info
    except Exception as e:  # noqa: BLE001 - la captura no debe romper la normalización
        return {"estado": "ERROR", "error": f"{type(e).__name__}: {e}", "ruta_origen": None}
