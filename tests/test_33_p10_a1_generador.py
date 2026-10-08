"""P10-A.1 — generador histórico mensual (EXTRACTO + AUDITORIA) sobre los 12 extractos REALES y el motor REAL de P0.

Las entradas son las salidas de `ejecutar_motor` (fixture de sesión `corrida_lote`: LISTS.csv, ORIGEN.xlsx) y un
snapshot SIMULADO de Depositos_Activos. Las expectativas se escriben aquí, a partir de LISTS.csv / ORIGEN.xlsx /
snapshot crudos, sin llamar a las funciones de cruce de p10: un error de p10 no puede "pasar" comparándose consigo mismo.
"""
import ast
import csv
import datetime as dt
import hashlib
import json
import math
import shutil
import sys
import unicodedata
import zipfile
from collections import defaultdict
from pathlib import Path
from types import SimpleNamespace

import pytest

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

import historico as H                                                  # noqa: E402
from openpyxl import load_workbook                                      # noqa: E402
from openpyxl.utils import range_boundaries                             # noqa: E402

from helpers import COLUMNAS_LISTS_CONTRATO                             # noqa: E402
from p10 import contrato as C, empaque, generador as G, motor_p0, snapshot_simulado as SS, validacion  # noqa: E402
from p10.comparar import diferencias                                    # noqa: E402
from p10.snapshot import SnapshotError, cargar_snapshot                 # noqa: E402

pytestmark = pytest.mark.p10

BNB_MN = "EXTRACTO_HISTORICO_BNB_3000100152_BOB_2026-08.xlsx"
OPERATIVAS_EXTRACTO = ["ESTADO", "CONFIRMADO POR", "FECHA DE CONFIRMACIÓN", "OBSERVACIONES"]
OPERATIVAS_26 = ["ESTADO", "ESTUDIANTE", "SOLICITADO POR", "SEDE SOLICITANTE", "CONFIRMADO POR",
                 "FECHA CONFIRMACIÓN", "OBSERVACIÓN"]
ADICIONALES = ["CODIGO_ESTUDIANTE", "ULTIMA_REVERSION_ID"]          # columnas REALES de Depositos_Activos sin equivalente en P0
OPERATIVAS_AUDITORIA = OPERATIVAS_26 + ADICIONALES
# columnas del banco por cuenta, según el diseño aprobado (DISENO_TRES_CAPAS.md §6.4)
COLUMNAS_BANCO_ESPERADAS = {"BNB_MN": 11, "BNB_ME": 11, "BNB_AHORRO": 11, "BNB_CLINICA": 11, "BCP_MN": 9, "BCP_ME": 9,
                            "BISA_MN": 12, "ECO_CTA_CTE": 7, "ECO_AHORRO": 7, "BMSC": 21, "UNION_MN": 6}


# ------------------------------------------------------------------ fixtures
def _sha(ruta):
    return hashlib.sha256(Path(ruta).read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def p0(corrida_lote):
    _, salida = corrida_lote
    s = SimpleNamespace(lists=salida / "LISTS.csv", origen=salida / "ORIGEN.xlsx", norm=salida / "NORMALIZADO.xlsx")
    s.huellas = {k: _sha(getattr(s, k)) for k in ("lists", "origen", "norm")}
    return s


@pytest.fixture(scope="module")
def filas(p0):
    with open(p0.lists, encoding="utf-8-sig", newline="") as f:      # lectura propia: no usa p10 ni historico
        return list(csv.DictReader(f))


@pytest.fixture(scope="module")
def snap0(filas):
    return SS.snapshot_simulado(filas)


@pytest.fixture(scope="module")
def t0(p0, snap0, tmp_path_factory):
    salida = tmp_path_factory.mktemp("p10_t0")
    info = G.generar_libros(p0.origen, p0.lists, snap0, salida)
    return info, Path(salida)


@pytest.fixture(scope="module")
def libros(t0):
    """{nombre_archivo: Workbook} abiertos una sola vez."""
    info, salida = t0
    return {a["nombre_archivo"]: load_workbook(salida / a["nombre_archivo"]) for a in info["archivos"]}


def _filas_de(wb, hoja, tabla):
    c0, r0, c1, r1 = range_boundaries(wb[hoja].tables[tabla].ref)
    ws = wb[hoja]
    titulos = [ws.cell(r0, c).value for c in range(c0, c1 + 1)]
    return titulos, list(ws.iter_rows(min_row=r0 + 1, max_row=r1, min_col=c0, max_col=c1)), r0


def _por_clave(wb, hoja):
    tabla = C.TABLA_EXTRACTO if hoja == "EXTRACTO" else C.TABLA_AUDITORIA
    titulos, filas_, _ = _filas_de(wb, hoja, tabla)
    k = titulos.index("CLAVE TRANSACCIÓN")
    return titulos, {f[k].value: dict(zip(titulos, f)) for f in filas_}


def _grupos_p0(filas):
    g = defaultdict(list)
    for r in filas:
        g[(r["BANCO"], r["CUENTA BANCARIA"], r["MONEDA"], r["FECHA MOVIMIENTO"][:7])].append(r)
    return g


def _nombre(banco, cuenta, moneda, periodo):
    """Nombre esperado del archivo, escrito aquí a partir del ejemplo aprobado (BNB_3000100152_BOB_2026-08)."""
    def seguro(t):
        t = "".join(ch for ch in unicodedata.normalize("NFD", t) if unicodedata.category(ch) != "Mn").upper()
        return "".join(ch if (ch.isascii() and ch.isalnum()) or ch == "-" else "_" for ch in t)
    return f"EXTRACTO_HISTORICO_{seguro(banco)}_{seguro(cuenta)}_{seguro(moneda)}_{periodo}.xlsx"


# ====================================================================== 1. REUTILIZACIÓN DE P0 (no se duplica)
def test_p10_no_define_logica_de_clave_ni_de_normalizacion():
    prohibidas = {"crear_clave", "valor_clave_numero", "finalizar_dataframe", "normalizar_extracto", "normalizar_fecha",
                  "normalizar_hora", "normalizar_texto", "codigo_texto", "numero", "extraer_nombre_bnb",
                  "ecuacion_saldo", "validar_extracto", "normalizar", "validar"}
    importadas_prohibidas = {"motor_generico", "deteccion_registro", "captura_origen"}
    for py in sorted((RAIZ / "p10").glob("*.py")):
        arbol = ast.parse(py.read_text(encoding="utf-8"))
        definidas = {n.name for n in ast.walk(arbol) if isinstance(n, (ast.FunctionDef, ast.ClassDef))}
        assert not (definidas & prohibidas), (py.name, definidas & prohibidas)
        for n in ast.walk(arbol):
            mods = [a.name for a in n.names] if isinstance(n, ast.Import) else [n.module or ""] if isinstance(n, ast.ImportFrom) else []
            assert not (set(m.split(".")[0] for m in mods) & importadas_prohibidas), (py.name, mods)


def test_p10_usa_el_motor_real_y_sus_26_columnas_por_identidad():
    import motor_control_depositos_cbba as motor
    assert motor_p0.motor_real() is motor
    assert C.COLUMNAS_LISTS is motor.COLUMNAS_LISTS                    # el MISMO objeto, no una copia
    assert list(motor.COLUMNAS_LISTS) == COLUMNAS_LISTS_CONTRATO       # y coincide con el contrato congelado de las pruebas


def test_modulos_de_p0_p3b_p9_sin_cambios():
    """P10 no toca ningún módulo previo: mismas huellas que exige P9, con la excepción de P0 ya autorizada allí."""
    from test_16_asignacion_p9 import EXCEPCIONES_AUTORIZADAS          # única fuente de las excepciones autorizadas
    huellas = json.loads((RAIZ / "p9/evidencias/huellas_base_p8_5.json").read_text(encoding="utf-8"))["archivos"]
    for f in ("motor_control_depositos_cbba.py", "motor_generico.py", "deteccion_registro.py", "captura_origen.py",
              "historico.py", "registro_bancos.json", "adaptador_m365.py"):
        assert _sha(RAIZ / f) in (huellas[f], EXCEPCIONES_AUTORIZADAS.get(f)), f


def test_p10_solo_lee_las_salidas_de_p0(t0, p0):
    assert {k: _sha(getattr(p0, k)) for k in ("lists", "origen", "norm")} == p0.huellas


def test_claves_de_auditoria_son_las_de_p0_y_se_recalculan_con_crear_clave_de_p0(libros, filas):
    import motor_control_depositos_cbba as motor
    claves_p0 = {r["CLAVE TRANSACCIÓN"] for r in filas}
    vistas = []
    for wb in libros.values():
        titulos, por_clave = _por_clave(wb, "AUDITORIA")
        for clave, fila in por_clave.items():
            fila = {k: c.value for k, c in fila.items()}
            # la fórmula de la clave sigue siendo UNA sola: la de P0, ejecutada sobre los valores del libro
            assert motor.crear_clave(fila) == clave
            vistas.append(clave)
    assert len(vistas) == len(set(vistas)) == len(claves_p0) == 4064
    assert set(vistas) == claves_p0


# ====================================================================== 2. UNIDAD DE SALIDA Y ESTRUCTURA
def test_un_archivo_por_banco_cuenta_moneda_mes(t0, filas):
    info, salida = t0
    esperados = {_nombre(b, c, m, p) for (b, c, m, p) in _grupos_p0(filas)}
    assert {a["nombre_archivo"] for a in info["archivos"]} == esperados
    assert len(esperados) == 11 and BNB_MN in esperados
    assert info["estado"] == "OK" and info["omitidos"] == []
    assert sorted(p.name for p in salida.iterdir()) == sorted(esperados)   # ni temporales ni archivos de más
    assert info["extractos_sin_movimientos"] == ["bisa_me_2.xls"]          # BISA_ME no tiene movimientos: no genera archivo


def test_cada_libro_tiene_exactamente_dos_hojas_visibles(libros):
    for nombre, wb in libros.items():
        assert wb.sheetnames == ["EXTRACTO", "AUDITORIA"], nombre
        assert [w.sheet_state for w in wb.worksheets] == ["visible", "visible"], nombre


def test_auditoria_tiene_las_26_columnas_de_p0_en_orden_y_las_2_reales_adicionales_al_final(libros):
    for nombre, wb in libros.items():
        ws = wb["AUDITORIA"]
        cab = [c.value for c in ws[1]]
        assert cab[:26] == COLUMNAS_LISTS_CONTRATO, nombre                 # las 26 de P0, exactas y en el mismo orden
        assert cab[26:] == ADICIONALES and len(cab) == 28, nombre          # + 2 columnas reales de la lista, al final
        assert ws.max_column == 28 and ws.tables["tblAUDITORIA"].ref == f"A1:AB{ws.max_row}", nombre
        assert not any(t.startswith(("P10", "SNAPSHOT", "HASH")) for t in cab)   # ninguna columna técnica de P10


def test_cantidad_de_movimientos_igual_a_p0_por_cuenta_y_mes(libros, filas):
    por_nombre = {_nombre(*k): len(v) for k, v in _grupos_p0(filas).items()}
    for nombre, wb in libros.items():
        assert wb["AUDITORIA"].max_row - 1 == por_nombre[nombre], nombre
        _, f_ext, _ = _filas_de(wb, "EXTRACTO", "tblEXTRACTO")
        assert len(f_ext) == por_nombre[nombre], nombre
    assert sum(por_nombre.values()) == 4064


def test_ninguna_clave_duplicada_en_ninguna_hoja_ni_entre_libros(libros):
    todas = []
    for nombre, wb in libros.items():
        for hoja, tabla in (("EXTRACTO", C.TABLA_EXTRACTO), ("AUDITORIA", C.TABLA_AUDITORIA)):
            titulos, filas_, _ = _filas_de(wb, hoja, tabla)
            claves = [f[titulos.index("CLAVE TRANSACCIÓN")].value for f in filas_]
            assert len(claves) == len(set(claves)) and all(claves), (nombre, hoja)
            if hoja == "AUDITORIA":
                todas += claves
    assert len(todas) == len(set(todas)) == 4064


def test_tabla_filtros_y_clave_oculta(libros):
    for nombre, wb in libros.items():
        we = wb["EXTRACTO"]
        t = we.tables["tblEXTRACTO"]
        c0, r0, c1, r1 = range_boundaries(t.ref)
        assert t.autoFilter is not None, nombre
        assert we.cell(r0, c1).value == "CLAVE TRANSACCIÓN" and we.column_dimensions[we.cell(r0, c1).column_letter].hidden
        assert [we.cell(r0, c).value for c in range(c1 - 4, c1)] == OPERATIVAS_EXTRACTO, nombre
        assert wb["AUDITORIA"].tables["tblAUDITORIA"].autoFilter is not None, nombre


def test_ningun_libro_tiene_paneles_inmovilizados_ni_en_el_xml_crudo(t0):
    """Decisión de Gabriel: desplazamiento libre en EXTRACTO y AUDITORIA. Se comprueba el XML, no solo openpyxl."""
    info, salida = t0
    assert len(info["archivos"]) == 11
    for a in info["archivos"]:
        with zipfile.ZipFile(salida / a["nombre_archivo"]) as z:
            hojas = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
            assert len(hojas) == 2, a["nombre_archivo"]
            for n in hojas:
                xml = z.read(n).decode("utf-8")
                assert "<pane" not in xml and 'pane="' not in xml and "state=\"frozen" not in xml, (a["nombre_archivo"], n)
        wb = load_workbook(salida / a["nombre_archivo"])
        for ws in wb.worksheets:
            assert ws.freeze_panes is None and ws.sheet_view.pane is None, (a["nombre_archivo"], ws.title)


def test_columnas_propias_de_cada_banco_completas_y_operativas_al_final(libros, filas):
    registro = json.loads((RAIZ / "registro_bancos.json").read_text(encoding="utf-8"))
    cuentas = {(c["banco"], c["cuenta"], c["moneda"]): c for c in registro["CUENTAS"]}
    for nombre, wb in libros.items():
        titulos, _, _ = _filas_de(wb, "EXTRACTO", "tblEXTRACTO")
        banco = titulos[:-5]
        grupo = next(k for k in _grupos_p0(filas) if _nombre(*k) == nombre)
        cuenta = cuentas[grupo[:3]]
        assert len(banco) == COLUMNAS_BANCO_ESPERADAS[cuenta["id"]], (cuenta["id"], banco)
        specs = registro["FORMATOS"][cuenta["formato"]]["historico"]["columnas"]
        esperadas = [s.get("titulo") or s["origen"] for s in specs if not s.get("opcional")]
        assert all(t in banco for t in esperadas), (nombre, [t for t in esperadas if t not in banco])
        assert titulos[-5:] == OPERATIVAS_EXTRACTO + ["CLAVE TRANSACCIÓN"]


def test_bnb_mn_columnas_literales_del_diseno_aprobado(libros):
    titulos, _, _ = _filas_de(libros[BNB_MN], "EXTRACTO", "tblEXTRACTO")
    assert titulos == ["Fecha", "Hora", "Oficina", "Descripción", "Referencia", "Código de transacción", "ITF", "Débitos",
                       "Créditos", "Saldo", "Adicionales", "ESTADO", "CONFIRMADO POR", "FECHA DE CONFIRMACIÓN",
                       "OBSERVACIONES", "CLAVE TRANSACCIÓN"]


# ====================================================================== 3. VALORES = LO QUE ENTREGA P0
def test_importes_debitos_creditos_y_saldos_iguales_a_p0_en_ambas_hojas(libros, filas):
    p0 = {r["CLAVE TRANSACCIÓN"]: r for r in filas}
    num = lambda s: float(s) if s != "" else None
    total = 0
    for nombre, wb in libros.items():
        _, aud = _por_clave(wb, "AUDITORIA")
        t_ext, ext = _por_clave(wb, "EXTRACTO")
        for clave, f in aud.items():
            r = p0[clave]
            for col in ("IMPORTE", "DÉBITO", "CRÉDITO", "SALDO"):
                assert f[col].value == num(r[col]), (clave, col)
            total += 1
        # columnas bancarias numéricas de EXTRACTO: las que salen de NORMALIZADO deben ser las mismas cifras
        for clave, f in ext.items():
            r = p0[clave]
            for titulo, col in (("Débitos", "DÉBITO"), ("Débito", "DÉBITO"), ("Créditos", "CRÉDITO"),
                                ("Crédito", "CRÉDITO"), ("Saldo", "SALDO")):
                if titulo in f:
                    assert f[titulo].value == num(r[col]), (nombre, clave, titulo)
            for titulo in ("Importe", "Monto"):
                if titulo in f:
                    imp = num(r["IMPORTE"])
                    assert f[titulo].value == (-imp if r["TIPO MOVIMIENTO"] == "DÉBITO" else imp), (nombre, clave)
        # sumas del archivo = sumas de P0 para esa cuenta-mes
        grupo = [r for r in filas if _nombre(r["BANCO"], r["CUENTA BANCARIA"], r["MONEDA"], r["FECHA MOVIMIENTO"][:7]) == nombre]
        for col in ("CRÉDITO", "DÉBITO"):
            assert math.isclose(math.fsum(f[col].value or 0 for f in aud.values()),
                                math.fsum(num(r[col]) or 0 for r in grupo), abs_tol=1e-6), (nombre, col)
    assert total == 4064


def test_las_19_columnas_no_operativas_de_auditoria_son_exactamente_las_de_p0(libros, filas):
    p0 = {r["CLAVE TRANSACCIÓN"]: r for r in filas}
    for nombre, wb in libros.items():
        _, aud = _por_clave(wb, "AUDITORIA")
        for clave, f in aud.items():
            r = p0[clave]
            for col in C.COLUMNAS_FIJAS:
                v = f[col].value
                if col == "FECHA MOVIMIENTO":
                    assert v == dt.datetime.strptime(r[col], "%Y-%m-%d"), (clave, col)
                elif col == "FECHA DE CARGA":
                    assert v == dt.datetime.fromisoformat(r[col]).replace(microsecond=0), (clave, col)
                elif col in ("IMPORTE", "DÉBITO", "CRÉDITO", "SALDO"):
                    assert v == (float(r[col]) if r[col] != "" else None), (clave, col)
                else:
                    assert (v if v is not None else "") == r[col], (clave, col, v, r[col])   # texto EXACTO, sin recortes


def test_fechas_y_horas_correctas(libros, filas):
    p0 = {r["CLAVE TRANSACCIÓN"]: r for r in filas}
    registro = json.loads((RAIZ / "registro_bancos.json").read_text(encoding="utf-8"))
    formato_hora = {c["id"]: registro["FORMATOS"][c["formato"]]["historico"].get("formato_hora", "hh:mm:ss")
                    for c in registro["CUENTAS"]}
    id_de = {_nombre(c["banco"], c["cuenta"], c["moneda"], "2026-08"): c["id"] for c in registro["CUENTAS"]}
    id_de[_nombre("BCP", "301-5005425-2-71", "USD", "2026-07")] = "BCP_ME"
    for nombre, wb in libros.items():
        titulos, ext = _por_clave(wb, "EXTRACTO")
        for clave, f in ext.items():
            r = p0[clave]
            fecha = f.get("Fecha") or f.get("Fecha Movimiento")
            assert fecha.value == dt.datetime.strptime(r["FECHA MOVIMIENTO"], "%Y-%m-%d") and fecha.number_format == "dd/mm/yyyy"
            if "Hora" in f:
                esperada = dt.time.fromisoformat(r["HORA MOVIMIENTO"]) if r["HORA MOVIMIENTO"] else None
                assert f["Hora"].value == esperada and f["Hora"].number_format == formato_hora[id_de[nombre]], (nombre, clave)


def test_importes_y_saldos_son_numericos_y_los_codigos_son_texto(libros):
    registro = json.loads((RAIZ / "registro_bancos.json").read_text(encoding="utf-8"))
    tipos = {}
    for fmt in registro["FORMATOS"].values():
        for s in fmt.get("historico", {}).get("columnas", []):
            tipos[s.get("titulo") or s["origen"]] = s["tipo"]
    codigos = importes = 0
    for nombre, wb in libros.items():
        titulos, ext = _por_clave(wb, "EXTRACTO")
        for clave, f in ext.items():
            for t, c in f.items():
                tipo = tipos.get(t)
                if tipo == "codigo" and c.value is not None:
                    assert isinstance(c.value, str) and c.number_format == "@", (nombre, t, c.value)
                    codigos += 1
                if tipo == "importe" and c.value is not None and t in ("Saldo", "Débitos", "Créditos", "Débito", "Crédito"):
                    assert isinstance(c.value, (int, float)) and not isinstance(c.value, bool), (nombre, t, c.value)
                    importes += 1
    assert codigos > 5000 and importes > 5000


def test_ceros_iniciales_y_textos_completos_bnb_mn(libros, filas, p0):
    titulos, ext = _por_clave(libros[BNB_MN], "EXTRACTO")
    # caso del diseño aprobado: un "Debito Pago Cheque Efectivo" conserva Referencia 0008125 y su Oficina
    d = [f for f in ext.values() if f["Referencia"].value == "0008125"]
    assert d and d[0]["Descripción"].value == "Debito Pago Cheque Efectivo" and d[0]["Oficina"].value
    assert any(f["Referencia"].value == "0000000004" for f in ext.values())            # ceros a la izquierda
    # texto completo: el Adicionales más largo de P0 está íntegro en AUDITORIA y en EXTRACTO (sin recortes)
    bnb = [r for r in filas if r["CUENTA BANCARIA"] == "3000100152"]
    mas_largo = max(bnb, key=lambda r: len(r["INFORMACIÓN ADICIONAL"]))
    assert len(mas_largo["INFORMACIÓN ADICIONAL"].strip()) > 120
    _, aud = _por_clave(libros[BNB_MN], "AUDITORIA")
    assert aud[mas_largo["CLAVE TRANSACCIÓN"]]["INFORMACIÓN ADICIONAL"].value == mas_largo["INFORMACIÓN ADICIONAL"]
    assert ext[mas_largo["CLAVE TRANSACCIÓN"]]["Adicionales"].value == mas_largo["INFORMACIÓN ADICIONAL"].strip()


def test_columnas_originales_del_banco_igual_a_origen_xlsx(libros, p0):
    """Texto de ORIGEN.xlsx (DATOS_ORIGINALES, sin recorte salvo espacios de relleno) = celda del EXTRACTO."""
    datos, mapa = H.leer_origen(p0.origen)
    clave_a_origen = {r["CLAVE TRANSACCIÓN"]: r["ID_ORIGEN"] for r in mapa}
    mov = defaultdict(dict)
    enc = {}
    for r in datos:
        if r["ROL_FILA"] == "MOVIMIENTO":
            mov[r["ID_ORIGEN"]][r["COLUMNA_EXCEL"]] = (r["VALOR_ORIGINAL"] or "")
        elif r["ROL_FILA"] == "ENCABEZADO_TABLA":
            enc.setdefault(r["ID_EXTRACTO"], {})[r["COLUMNA_EXCEL"]] = r["VALOR_ORIGINAL"]
    wb = libros[BNB_MN]
    _, ext = _por_clave(wb, "EXTRACTO")
    por_titulo = {"Oficina": "Oficina", "Descripción": "Descripción", "Referencia": "Referencia",
                  "Código de transacción": "Código de transacción", "Adicionales": "Adicionales"}
    id_ext = next(iter({m["ID_EXTRACTO"] for m in mapa if m["CLAVE TRANSACCIÓN"] in ext}))
    col_de = {" ".join(str(t).split()).upper(): c for c, t in enc[id_ext].items()}
    comprobadas = 0
    for clave, f in ext.items():
        celdas = mov[clave_a_origen[clave]]
        for titulo, original in por_titulo.items():
            texto = celdas.get(col_de[original.upper()], "").strip()
            assert (f[titulo].value or "") == texto, (clave, titulo)
            comprobadas += 1
    assert comprobadas == 5 * 1184


# ====================================================================== 3b. AUDITORIA CUBRE TODA LA LISTA REAL
# Columnas PROPIAS de Depositos_Activos vistas en el selector de columnas del tenant (captura aportada por Gabriel, 2026-10-08).
# «Última reversión aplicada» es el nombre para mostrar de ULTIMA_REVERSION_ID. No incluye las columnas de sistema de
# SharePoint (ID, Creado, Creado por, Modificado, Modificado por, Título, Tipo, Datos adjuntos, etiquetas de retención...).
COLUMNAS_PROPIAS_TENANT_2026_10_08 = {
    "CLAVE_TRANSACCION", "BANCO", "FECHA_MOVIMIENTO", "IMPORTE", "ESTADO_ASIGNACION", "CODIGO_ASIGNACION",
    "CUENTA_BANCARIA", "HORA_MOVIMIENTO", "SALDO", "DESCRIPCION", "ESTUDIANTE", "SOLICITADO_POR", "CODIGO_ESTUDIANTE",
    "SEDE_ASIGNACION", "FECHA_HORA_ASIGNACION", "OBSERVACION", "ULTIMA_REVERSION_ID", "MONEDA", "DEBITO", "CREDITO",
    "TIPO_MOVIMIENTO", "DEPOSITANTE_ORIGINANTE", "INFORMACION_ADICIONAL", "MOTOR_ESTADO", "MOTOR_ESTUDIANTE",
    "MOTOR_SOLICITADO_POR", "MOTOR_SEDE_SOLICITANTE", "MOTOR_CONFIRMADO_POR", "MOTOR_FECHA_CONFIRMACION",
    "MOTOR_OBSERVACION", "TEXTO_BUSQUEDA", "ARCHIVO_ORIGEN", "LOTE_CARGA", "FECHA_CARGA", "USUARIO_ASIGNACION"}


def _esquema_real_depositos_activos():
    p8 = json.loads((RAIZ / "p8/esquema_listas_p8.json").read_text(encoding="utf-8"))["Depositos_Activos"]["columnas"]
    p9 = json.loads((RAIZ / "p9/reversion/esquema_reversiones.json").read_text(encoding="utf-8"))["Depositos_Activos"]["columnas"]
    return [c["nombre_tecnico"] for c in p8], [c["nombre_tecnico"] for c in p9]


def test_el_esquema_del_repo_coincide_con_las_35_columnas_propias_del_tenant():
    p8, p9 = _esquema_real_depositos_activos()
    assert len(p8) == 34 and p9 == ["ULTIMA_REVERSION_ID"]
    assert set(p8) | set(p9) == COLUMNAS_PROPIAS_TENANT_2026_10_08 and len(COLUMNAS_PROPIAS_TENANT_2026_10_08) == 35


def test_auditoria_cubre_cada_columna_real_de_depositos_activos():
    """
    Si la lista evoluciona (columna nueva en el esquema del repo) y AUDITORIA no la conserva, esta prueba FALLA:
    en P10-B, al eliminar un registro, AUDITORIA debe llevar toda la información persistida relevante.
    Cada columna real es: (a) una de las 19 de P0 sin cambio de nombre, (b) la columna operativa viva que reemplaza a una
    de las 7 de P0 (cuyo homónimo MOTOR_* es una constante reservada que ningún flujo escribe), o (c) una adicional.
    """
    import adaptador_m365 as A
    p8, p9 = _esquema_real_depositos_activos()
    interno = dict(zip(A.COLUMNAS_LISTS_P6, A.COLUMNAS_TECNICAS))               # nombre P0 -> nombre interno en la lista
    fijas = {interno[c] for c in C.COLUMNAS_FIJAS}
    vivas = set(C.OPERATIVOS.values())
    adicionales = set(C.ADICIONALES_AUDITORIA.values())
    reservadas = {interno[c] for c in C.OPERATIVOS}                              # MOTOR_*: constantes de P0
    assert all(r.startswith("MOTOR_") for r in reservadas) and len(reservadas) == 7
    reales = set(p8) | set(p9)
    assert reales == fijas | vivas | adicionales | reservadas                    # nada de la lista queda sin clasificar
    assert not (fijas & vivas) and not (vivas & adicionales) and not (reservadas & (fijas | vivas | adicionales))
    # y AUDITORIA realmente trae todo lo conservable: 19 fijas + 7 vivas (en las columnas P0 homónimas) + 2 adicionales
    assert len(fijas) + len(vivas) + len(adicionales) == 28 == len(C.COLUMNAS_AUDITORIA)
    # P9 solo escribe estas columnas: ninguna MOTOR_* cambia nunca (por eso no se duplican en AUDITORIA)
    from p9.contrato import CAMPOS_ESCRITOS
    from p9.reversion.contrato import CAMPOS_ESCRITOS as ESCRITOS_REVERSION
    escritos = set(CAMPOS_ESCRITOS) | set(ESCRITOS_REVERSION)
    assert escritos <= vivas | adicionales and not (escritos & reservadas)


# ====================================================================== 4. CRUCE CON EL SNAPSHOT (por CLAVE)
def _esperado_snapshot(snapshot_dict):
    """Expectativa literal de las 7 columnas operativas, escrita aquí (no importada de p10)."""
    out = {}
    for f in snapshot_dict["filas"]:
        asignado = f["ESTADO_ASIGNACION"] == "ASIGNADO"
        out[f["CLAVE_TRANSACCION"]] = {
            "ESTADO": "CONFIRMADO" if asignado else "DISPONIBLE",
            "ESTUDIANTE": f["ESTUDIANTE"], "SOLICITADO POR": f["SOLICITADO_POR"],
            "SEDE SOLICITANTE": f["SEDE_ASIGNACION"], "CONFIRMADO POR": f["USUARIO_ASIGNACION"],
            "FECHA CONFIRMACIÓN": dt.datetime.fromisoformat(f["FECHA_HORA_ASIGNACION"]) if f["FECHA_HORA_ASIGNACION"] else None,
            "OBSERVACIÓN": f["OBSERVACION"],
            "CODIGO_ESTUDIANTE": f["CODIGO_ESTUDIANTE"], "ULTIMA_REVERSION_ID": f["ULTIMA_REVERSION_ID"]}
    return out


def test_estados_cruzados_por_clave_en_las_dos_hojas(libros, snap0):
    esperado = _esperado_snapshot(snap0.a_dict())
    n_conf = n_disp = 0
    for nombre, wb in libros.items():
        _, aud = _por_clave(wb, "AUDITORIA")
        _, ext = _por_clave(wb, "EXTRACTO")
        for clave, f in aud.items():
            e = esperado[clave]
            for col in OPERATIVAS_AUDITORIA:
                assert f[col].value == e[col], (nombre, clave, col, f[col].value, e[col])
            x = ext[clave]
            assert (x["ESTADO"].value, x["CONFIRMADO POR"].value, x["FECHA DE CONFIRMACIÓN"].value,
                    x["OBSERVACIONES"].value) == (e["ESTADO"], e["CONFIRMADO POR"], e["FECHA CONFIRMACIÓN"],
                                                  e["OBSERVACIÓN"]), (nombre, clave)
            n_conf += e["ESTADO"] == "CONFIRMADO"
            n_disp += e["ESTADO"] == "DISPONIBLE"
    assert n_conf > 200 and n_disp > 3000 and n_conf + n_disp == 4064


def test_solo_creditos_confirmados_y_fecha_de_confirmacion_posterior_al_movimiento(libros):
    for nombre, wb in libros.items():
        _, aud = _por_clave(wb, "AUDITORIA")
        for clave, f in aud.items():
            if f["ESTADO"].value == "CONFIRMADO":
                assert f["TIPO MOVIMIENTO"].value == "CRÉDITO"
                assert f["FECHA CONFIRMACIÓN"].value >= f["FECHA MOVIMIENTO"].value
                assert f["CONFIRMADO POR"].value and f["ESTUDIANTE"].value and f["SOLICITADO POR"].value


def test_snapshot_simulado_es_coherente_con_el_modelo_de_p9(snap0, filas):
    tipo = {r["CLAVE TRANSACCIÓN"]: r for r in filas}
    asignados = [(k, v) for k, v in snap0.filas.items() if v["ESTADO_ASIGNACION"] == "ASIGNADO"]
    assert len(asignados) > 200 and not snap0.advertencias
    for clave, v in asignados:
        assert tipo[clave]["TIPO MOVIMIENTO"] == "CRÉDITO"                       # la app solo confirma créditos
        assert dt.datetime.fromisoformat(tipo[clave]["FECHA MOVIMIENTO"]) <= v["FECHA_HORA_ASIGNACION"] <= snap0.fecha_corte
        assert "SIMULADO" in v["ESTUDIANTE"] and v["USUARIO_ASIGNACION"].endswith("@example.invalid")
    assert all(v["ESTADO_ASIGNACION"] in ("DISPONIBLE", "ASIGNADO") for v in snap0.filas.values())
    assert len(snap0) == len(filas) == 4064


def test_observaciones_se_alimentan_de_observacion_y_se_conservan_tal_cual(libros, snap0):
    esperado = _esperado_snapshot(snap0.a_dict())
    _, ext = _por_clave(libros[BNB_MN], "EXTRACTO")
    _, aud = _por_clave(libros[BNB_MN], "AUDITORIA")
    textos = {e["OBSERVACIÓN"] for k, e in esperado.items() if k in ext and e["OBSERVACIÓN"]}
    assert any(t.startswith("=") for t in textos) and any("\n" in t for t in textos) and any("ñ" in t or "ó" in t for t in textos)
    assert any(t.startswith("SIMULADO 0004521") for t in textos)
    for clave, x in ext.items():
        e = esperado[clave]["OBSERVACIÓN"]
        assert x["OBSERVACIONES"].value == e and aud[clave]["OBSERVACIÓN"].value == e
        if e:
            assert x["OBSERVACIONES"].data_type == "s" and aud[clave]["OBSERVACIÓN"].data_type == "s"   # "=..." NO es fórmula


def test_columnas_reales_adicionales_van_solo_en_auditoria_con_su_valor_del_snapshot(libros, snap0):
    filas_snap = snap0.a_dict()["filas"]
    codigos = {f["CODIGO_ESTUDIANTE"] for f in filas_snap if f["CODIGO_ESTUDIANTE"]}
    marcadores = {f["ULTIMA_REVERSION_ID"] for f in filas_snap if f["ULTIMA_REVERSION_ID"]}
    assert codigos and any(c.startswith("0") for c in codigos) and len(marcadores) > 50
    wb = libros[BNB_MN]
    _, aud = _por_clave(wb, "AUDITORIA")
    vistos_codigo = {f["CODIGO_ESTUDIANTE"].value for f in aud.values() if f["CODIGO_ESTUDIANTE"].value}
    vistos_marcador = {f["ULTIMA_REVERSION_ID"].value for f in aud.values() if f["ULTIMA_REVERSION_ID"].value}
    assert vistos_codigo and vistos_marcador
    for f in aud.values():
        for col in ADICIONALES:
            c = f[col]
            assert c.value is None or (isinstance(c.value, str) and c.data_type == "s" and c.number_format == "@")
    assert any(v.startswith("0") for v in vistos_codigo)                           # ceros a la izquierda conservados
    # un depósito revertido que volvió a DISPONIBLE conserva el marcador (como en P9) y tiene los 7 campos limpios
    revertidos = [f for f in aud.values() if f["ESTADO"].value == "DISPONIBLE" and f["ULTIMA_REVERSION_ID"].value]
    assert revertidos and all(f["ESTUDIANTE"].value is None and f["CODIGO_ESTUDIANTE"].value is None for f in revertidos)
    # EXTRACTO NO cambia: no recibe ninguna de las dos columnas
    titulos, ext = _por_clave(wb, "EXTRACTO")
    assert not (set(ADICIONALES) & set(titulos))
    visto_en_extracto = {str(c.value) for f in ext.values() for c in f.values() if c.value is not None}
    assert not (codigos & visto_en_extracto) and not (marcadores & visto_en_extracto)


def test_manifiesto_resume_el_cruce(t0, snap0):
    info, _ = t0
    s = info["snapshot"]
    assert s["origen"] == "SIMULADO" and s["sha256"] == snap0.sha256 and s["snapshot_filas"] == 4064
    assert s["con_fila_en_snapshot"] == 4064 and s["sin_fila_en_snapshot"] == 0 and s["snapshot_fuera_del_extracto"] == 0
    assert all(a["validacion"] == "OK" for a in info["archivos"])
    assert sum(a["movimientos"] for a in info["archivos"]) == 4064


# ====================================================================== 5. DETERMINISMO Y REGENERACIÓN
def test_dos_ejecuciones_producen_el_mismo_archivo(p0, snap0, tmp_path):
    a, b = tmp_path / "a", tmp_path / "b"
    ia = G.generar_libros(p0.origen, p0.lists, snap0, a, solo_cuentas={"BNB_MN", "BCP_ME"})
    ib = G.generar_libros(p0.origen, p0.lists, snap0, b, solo_cuentas={"BNB_MN", "BCP_ME"})
    assert [f["nombre_archivo"] for f in ia["archivos"]] == [f["nombre_archivo"] for f in ib["archivos"]]
    for fa in ia["archivos"]:
        assert diferencias(a / fa["nombre_archivo"], b / fa["nombre_archivo"]) == []        # mismo contenido lógico
        assert _sha(a / fa["nombre_archivo"]) == _sha(b / fa["nombre_archivo"])             # y los mismos bytes
    G.escribir_manifiesto(ia, a / "m.json")
    G.escribir_manifiesto(ib, b / "m.json")
    assert (a / "m.json").read_bytes() == (b / "m.json").read_bytes()


@pytest.fixture(scope="module")
def escenarios(p0, filas, snap0, t0, tmp_path_factory):
    """T0 -> T1 (4 DISPONIBLE pasan a CONFIRMADO) -> T2 (2 reversiones: una de T0 y una de T1) sobre BNB MN."""
    cuenta = "3000100152"
    nuevas = SS.elegir_a_confirmar(snap0, filas, cuenta, 4)
    t1 = SS.confirmar(snap0, filas, nuevas)
    previa = SS.elegir_a_revertir(snap0, cuenta, filas, 1)
    revertidas = previa + [nuevas[0]]
    t2 = SS.revertir(t1, revertidas)
    sal = {}
    for nombre, snap in (("t1", t1), ("t2", t2)):
        d = tmp_path_factory.mktemp("p10_" + nombre)
        info = G.generar_libros(p0.origen, p0.lists, snap, d, solo_cuentas={"BNB_MN"})
        assert info["estado"] == "OK" and [a["nombre_archivo"] for a in info["archivos"]] == [BNB_MN]
        sal[nombre] = d / BNB_MN
    return SimpleNamespace(nuevas=nuevas, previa=previa, revertidas=revertidas, t1=t1, t2=t2,
                           ruta_t0=t0[1] / BNB_MN, ruta_t1=sal["t1"], ruta_t2=sal["t2"])


def test_disponible_a_confirmado_cambia_solo_los_campos_operativos_de_esas_claves(escenarios):
    e = escenarios
    difs = diferencias(e.ruta_t0, e.ruta_t1)
    assert difs and all(d["tipo"] == "CELDA" for d in difs)                  # misma estructura, claves, orden y zona superior
    assert {d["clave"] for d in difs} == set(e.nuevas)
    assert {d["columna"] for d in difs if d["hoja"] == "EXTRACTO"} <= set(OPERATIVAS_EXTRACTO)
    assert {d["columna"] for d in difs if d["hoja"] == "AUDITORIA"} <= set(OPERATIVAS_AUDITORIA)
    for hoja, col in (("EXTRACTO", "ESTADO"), ("AUDITORIA", "ESTADO")):
        cambios = [d for d in difs if d["hoja"] == hoja and d["columna"] == col]
        assert len(cambios) == 4 and all((d["antes"], d["despues"]) == ("DISPONIBLE", "CONFIRMADO") for d in cambios)
    # y no hay filas nuevas ni repetidas: mismo número de movimientos y claves únicas
    _, a0 = _por_clave(load_workbook(e.ruta_t0), "AUDITORIA")
    _, a1 = _por_clave(load_workbook(e.ruta_t1), "AUDITORIA")
    assert list(a0) == list(a1) and len(a1) == 1184


def test_reversion_representada_con_el_payload_de_p9_vuelve_a_disponible_y_limpia_los_7_campos(escenarios):
    e = escenarios
    d12 = diferencias(e.ruta_t1, e.ruta_t2)
    assert {d["clave"] for d in d12 if d["tipo"] == "CELDA"} == set(e.revertidas) and all(d["tipo"] == "CELDA" for d in d12)
    _, aud = _por_clave(load_workbook(e.ruta_t2), "AUDITORIA")
    for clave in e.revertidas:
        f = aud[clave]
        assert f["ESTADO"].value == "DISPONIBLE"
        assert all(f[c].value is None for c in OPERATIVAS_26 + ["CODIGO_ESTUDIANTE"] if c != "ESTADO")
        assert f["ULTIMA_REVERSION_ID"].value                       # el marcador de la reversión SÍ queda (P9)
    # Lo recién confirmado y revertido queda IGUAL a T0 en EXTRACTO y, en AUDITORIA, solo conserva el marcador de la reversión
    # (P9 nunca borra ULTIMA_REVERSION_ID).
    d02 = diferencias(e.ruta_t0, e.ruta_t2)
    claves_extracto = {d["clave"] for d in d02 if d["hoja"] == "EXTRACTO"}
    assert claves_extracto == (set(e.nuevas) - {e.nuevas[0]}) | set(e.previa)
    assert {d["columna"] for d in d02 if d["hoja"] == "AUDITORIA" and d["clave"] == e.nuevas[0]} == {"ULTIMA_REVERSION_ID"}


def test_confirmar_y_revertir_lo_mismo_reconstruye_el_extracto_de_t0_y_deja_solo_el_marcador_en_auditoria(
        p0, filas, snap0, t0, tmp_path):
    nuevas = SS.elegir_a_confirmar(snap0, filas, "3000100152", 3)
    ida_vuelta = SS.revertir(SS.confirmar(snap0, filas, nuevas), nuevas)
    G.generar_libros(p0.origen, p0.lists, ida_vuelta, tmp_path, solo_cuentas={"BNB_MN"})
    difs = diferencias(t0[1] / BNB_MN, tmp_path / BNB_MN)
    assert difs and all(d["tipo"] == "CELDA" for d in difs)
    assert not [d for d in difs if d["hoja"] == "EXTRACTO"]                          # EXTRACTO: idéntico a T0
    assert {(d["hoja"], d["columna"], d["clave"]) for d in difs} == {("AUDITORIA", "ULTIMA_REVERSION_ID", k) for k in nuevas}


def test_la_reversion_deja_el_marcador_en_auditoria_y_nada_nuevo_en_extracto(escenarios):
    wb = load_workbook(escenarios.ruta_t2)
    _, aud = _por_clave(wb, "AUDITORIA")
    nuevos = {escenarios.t2.get(k)["ULTIMA_REVERSION_ID"] for k in escenarios.revertidas}
    assert len(nuevos) == 2 and all(nuevos)
    assert {aud[k]["ULTIMA_REVERSION_ID"].value for k in escenarios.revertidas} == nuevos
    visto_en_extracto = {str(c.value) for fila in _por_clave(wb, "EXTRACTO")[1].values() for c in fila.values()
                         if c.value is not None}
    assert not (nuevos & visto_en_extracto)


def test_el_estado_visible_no_cambia_la_zona_superior_ni_las_columnas(escenarios):
    wa, wb_ = load_workbook(escenarios.ruta_t0), load_workbook(escenarios.ruta_t2)
    for hoja in ("EXTRACTO", "AUDITORIA"):
        assert wa[hoja].dimensions == wb_[hoja].dimensions and wa[hoja].freeze_panes == wb_[hoja].freeze_panes


# ====================================================================== 6. SNAPSHOT: validaciones y casos límite
def _snap(*filas_extra, **kw):
    base = {"version": C.VERSION_SNAPSHOT, "origen": "SIMULADO", "filas": list(filas_extra)}
    base.update(kw)
    return base


def _fila(clave="K1", estado="DISPONIBLE", **kw):
    f = {"CLAVE_TRANSACCION": clave, "ESTADO_ASIGNACION": estado}
    f.update(kw)
    return f


def test_snapshot_rechaza_clave_duplicada_estado_desconocido_y_fecha_con_zona():
    with pytest.raises(SnapshotError, match="repetida"):
        cargar_snapshot(_snap(_fila("A"), _fila("A")))
    with pytest.raises(SnapshotError, match="estado vigente"):
        cargar_snapshot(_snap(_fila("A", "ANULADO")))
    with pytest.raises(SnapshotError, match="falta CLAVE_TRANSACCION"):
        cargar_snapshot(_snap({"ESTADO_ASIGNACION": "DISPONIBLE"}))
    with pytest.raises(SnapshotError, match="zona"):
        cargar_snapshot(_snap(_fila("A", "ASIGNADO", FECHA_HORA_ASIGNACION="2026-08-12T10:15:00Z")))
    with pytest.raises(SnapshotError, match="no es ISO"):
        cargar_snapshot(_snap(_fila("A", "ASIGNADO", FECHA_HORA_ASIGNACION="12/08/2026")))
    with pytest.raises(SnapshotError, match="versión"):
        cargar_snapshot(_snap(version="otra"))


def test_snapshot_incompleto_se_informa_pero_no_bloquea():
    s = cargar_snapshot(_snap(_fila("A", "ASIGNADO"), _fila("B", "DISPONIBLE", ESTUDIANTE="X")))
    assert len(s.advertencias) == 2 and "ASIGNADO sin" in s.advertencias[0] and "residuales" in s.advertencias[1]


def test_snapshot_no_recorta_ni_cambia_la_clave():
    s = cargar_snapshot(_snap(_fila(" BNB|x ")))
    assert list(s.filas) == [" BNB|x "] and s.get("BNB|x") is None


def test_movimientos_sin_fila_en_el_snapshot_conservan_lo_de_p0_y_se_cuentan(p0, filas, snap0, tmp_path):
    mitad = snap0.a_dict()
    mitad["filas"] = mitad["filas"][::2]
    mitad["filas"].append(_fila("CLAVE-QUE-NO-ESTA-EN-EL-EXTRACTO"))
    info = G.generar_libros(p0.origen, p0.lists, cargar_snapshot(mitad), tmp_path, solo_cuentas={"BNB_CLINICA"})
    s = info["snapshot"]
    assert info["estado"] == "OK" and s["sin_fila_en_snapshot"] > 0 and s["snapshot_fuera_del_extracto"] == 1
    assert s["con_fila_en_snapshot"] + s["sin_fila_en_snapshot"] == 4064


def test_sin_snapshot_el_libro_muestra_lo_que_entrega_p0(p0, tmp_path):
    info = G.generar_libros(p0.origen, p0.lists, None, tmp_path, solo_cuentas={"BNB_CLINICA"})
    wb = load_workbook(tmp_path / info["archivos"][0]["nombre_archivo"])
    _, aud = _por_clave(wb, "AUDITORIA")
    assert {f["ESTADO"].value for f in aud.values()} == {"DISPONIBLE"}
    assert {f["SEDE SOLICITANTE"].value for f in aud.values()} == {"COCHABAMBA"}      # el valor constante de P0


def test_leer_normalizado_xlsx_da_el_mismo_libro_que_lists_csv(p0, snap0, tmp_path):
    a, b = tmp_path / "csv", tmp_path / "xlsx"
    G.generar_libros(p0.origen, p0.lists, snap0, a, solo_cuentas={"BNB_CLINICA"})
    ib = G.generar_libros(p0.origen, p0.norm, snap0, b, solo_cuentas={"BNB_CLINICA"})
    assert ib["estado"] == "OK"
    nombre = ib["archivos"][0]["nombre_archivo"]
    assert diferencias(a / nombre, b / nombre) == []


# ====================================================================== 7. LA VALIDACIÓN DETECTA ERRORES REALES
@pytest.fixture(scope="module")
def entorno_validacion(p0, snap0):
    datos, mapa = H.leer_origen(p0.origen)
    filas_p0 = H.leer_normalizado(p0.lists)
    registro = H.cargar_registro()
    libros_, _ = G.construir_libros(datos, mapa, filas_p0, registro, snap0)
    libro = next(l for l in libros_ if l.id_cuenta == "BNB_CLINICA")
    return SimpleNamespace(libro=libro, filas=filas_p0, registro=registro, indice=validacion.indice_origen(datos, mapa))


def _validar(env, snap, ruta):
    return validacion.validar_libro(ruta, env.libro, env.filas, snap, env.indice, env.registro)


@pytest.fixture()
def copia(t0, entorno_validacion, tmp_path):
    origen = t0[1] / entorno_validacion.libro.nombre_archivo
    destino = tmp_path / origen.name
    shutil.copy(origen, destino)
    return destino


def test_validador_acepta_el_libro_correcto(entorno_validacion, snap0, copia):
    assert _validar(entorno_validacion, snap0, copia) == []


def _manipular(ruta, fn):
    wb = load_workbook(ruta)
    fn(wb)
    wb.save(ruta)


@pytest.mark.parametrize("nombre,fn,fragmento", [
    ("importe", lambda wb: wb["AUDITORIA"].cell(2, 8).__setattr__("value", 1.5), "IMPORTE"),
    ("saldo_extracto", lambda wb: wb["EXTRACTO"].cell(11, 10).__setattr__("value", 1.0), "Saldo"),
    ("estado_extracto", lambda wb: wb["EXTRACTO"].cell(11, 12).__setattr__("value", "ANULADO"), "ESTADO"),
    ("estado_auditoria", lambda wb: wb["AUDITORIA"].cell(2, 16).__setattr__("value", "ANULADO"), "ESTADO"),
    ("hoja_extra", lambda wb: wb.create_sheet("TECNICA"), "hojas"),
    ("hoja_oculta", lambda wb: wb["AUDITORIA"].__setattr__("sheet_state", "hidden"), "visible"),
    ("columna_renombrada", lambda wb: wb["AUDITORIA"].cell(1, 3).__setattr__("value", "BANCOS"), "26"),
    ("clave_duplicada", lambda wb: wb["AUDITORIA"].cell(3, 1).__setattr__("value", wb["AUDITORIA"].cell(2, 1).value), "duplicada"),
    ("codigo_como_numero", lambda wb: wb["AUDITORIA"].cell(2, 2).__setattr__("value", 123), "CÓDIGO"),
    ("codigo_estudiante_ajeno", lambda wb: wb["AUDITORIA"].cell(2, 27).__setattr__("value", "9999999"), "CODIGO_ESTUDIANTE"),
    ("marcador_ajeno", lambda wb: wb["AUDITORIA"].cell(2, 28).__setattr__("value", "uid-que-no-existe"), "ULTIMA_REVERSION_ID"),
    ("panel_en_extracto", lambda wb: setattr(wb["EXTRACTO"], "freeze_panes", "A11"), "paneles"),
    ("panel_en_auditoria", lambda wb: setattr(wb["AUDITORIA"], "freeze_panes", "A2"), "paneles"),
    ("fecha_como_texto", lambda wb: wb["AUDITORIA"].cell(2, 6).__setattr__("value", "2026-08-01"), "FECHA MOVIMIENTO"),
])
def test_validador_detecta_manipulaciones(entorno_validacion, snap0, copia, nombre, fn, fragmento):
    _manipular(copia, fn)
    errores = _validar(entorno_validacion, snap0, copia)
    assert errores and any(fragmento in e for e in errores), (nombre, errores)


def test_validador_detecta_fila_eliminada(entorno_validacion, snap0, copia):
    _manipular(copia, lambda wb: wb["AUDITORIA"].delete_rows(5))
    assert _validar(entorno_validacion, snap0, copia)


def test_validador_compara_contra_el_snapshot_que_se_le_entrega(entorno_validacion, snap0, copia):
    """El mismo archivo deja de ser válido si el snapshot de referencia dice otra cosa para alguna CLAVE."""
    libro = entorno_validacion.libro
    clave = next(k for k in libro.claves if snap0.get(k)["ESTADO_ASIGNACION"] == "DISPONIBLE")
    otro = SS.confirmar(snap0, entorno_validacion.filas, [clave])
    errores = _validar(entorno_validacion, otro, copia)
    assert errores and any(clave in e and "ESTADO" in e for e in errores)


def test_un_libro_que_no_valida_no_se_acepta(p0, snap0, tmp_path, monkeypatch):
    monkeypatch.setattr(validacion, "validar_libro", lambda *a, **k: ["falla simulada"])
    info = G.generar_libros(p0.origen, p0.lists, snap0, tmp_path, solo_cuentas={"BNB_CLINICA"})
    assert info["estado"] == "CON_ERRORES" and info["archivos"] == []
    assert info["omitidos"][0]["errores"] == ["falla simulada"]
    assert list(tmp_path.iterdir()) == []                      # ni archivo ni temporales


def test_clave_repetida_en_p0_detiene_el_cruce(p0, filas, snap0):
    datos, mapa = H.leer_origen(p0.origen)
    with pytest.raises(ValueError, match="repetida"):
        G.construir_libros(datos, mapa, filas + [dict(filas[0])], H.cargar_registro(), snap0)


# ====================================================================== 8. FIDELIDAD CON LA CAPA 4 APROBADA Y EMPAQUE
def test_hoja_extracto_identica_a_la_que_escribe_historico_mas_las_4_operativas(p0, snap0, tmp_path):
    """La hoja EXTRACTO no se reinventa: es la de historico.py (mismo escritor) con OBSERVACIONES y un ancho."""
    datos, mapa = H.leer_origen(p0.origen)
    libros_, _ = G.construir_libros(datos, mapa, H.leer_normalizado(p0.lists), H.cargar_registro(), snap0)
    libro = next(l for l in libros_ if l.id_cuenta == "BNB_MN")
    solo = tmp_path / "solo_historico.xlsx"
    H.escribir_extracto_historico(libro.resultado, str(solo))
    final = tmp_path / "final.xlsx"
    G.escribir_libro(libro, str(final), snap0)
    a, b = load_workbook(solo)["EXTRACTO"], load_workbook(final)["EXTRACTO"]
    assert a.freeze_panes is not None and b.freeze_panes is None       # único cambio de comportamiento: sin paneles
    assert a.dimensions == b.dimensions and a.print_title_rows == b.print_title_rows
    assert sorted(map(str, a.merged_cells.ranges)) == sorted(map(str, b.merged_cells.ranges))
    assert a.tables["tblEXTRACTO"].ref == b.tables["tblEXTRACTO"].ref
    for fa, fb in zip(a.iter_rows(), b.iter_rows()):
        for ca, cb in zip(fa, fb):
            assert (ca.value, ca.number_format, ca.data_type, ca.font.b, ca.font.color and ca.font.color.rgb,
                    ca.fill.fgColor.rgb, ca.alignment.horizontal, ca.border.bottom.style) == (
                cb.value, cb.number_format, cb.data_type, cb.font.b, cb.font.color and cb.font.color.rgb,
                cb.fill.fgColor.rgb, cb.alignment.horizontal, cb.border.bottom.style)
    anchos = lambda ws: {k: v.width for k, v in ws.column_dimensions.items() if v.width}
    ocultas = lambda ws: {k for k, v in ws.column_dimensions.items() if v.hidden}
    assert ocultas(a) == ocultas(b)
    difs = {k for k in anchos(a) if anchos(a)[k] != anchos(b)[k]}
    assert len(difs) == 1 and anchos(b)[difs.pop()] == C.ANCHO_CONFIRMADO_POR      # el único ajuste visual de P10


def test_propiedades_deterministas_y_sin_marcas_de_tiempo_del_proceso(t0, snap0):
    wb = load_workbook(t0[1] / BNB_MN)
    assert wb.properties.created == wb.properties.modified == snap0.fecha_corte
    assert "SIMULADO" in wb.properties.description and snap0.sha256[:12] in wb.properties.description
    with zipfile.ZipFile(t0[1] / BNB_MN) as z:
        assert {i.date_time for i in z.infolist()} == {(1980, 1, 1, 0, 0, 0)}


def test_zip_determinista_y_con_los_archivos_exactos(t0, tmp_path):
    info, salida = t0
    rutas = [a["ruta"] for a in info["archivos"]]
    sha1 = empaque.crear_zip(rutas, tmp_path / "a.zip")
    sha2 = empaque.crear_zip(list(reversed(rutas)), tmp_path / "b.zip")
    assert sha1 == sha2 == _sha(tmp_path / "a.zip")
    with zipfile.ZipFile(tmp_path / "a.zip") as z:
        assert z.namelist() == sorted(a["nombre_archivo"] for a in info["archivos"])
        assert z.testzip() is None
        assert hashlib.sha256(z.read(BNB_MN)).hexdigest() == _sha(salida / BNB_MN)
