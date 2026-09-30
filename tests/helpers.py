"""Utilidades compartidas de la suite. No contiene logica de negocio del motor."""
import importlib.util
import os
import re
from pathlib import Path

import pandas as pd

RAIZ = Path(__file__).resolve().parent
EXTRACTOS = RAIZ / "fixtures" / "extractos"
GOLDEN = RAIZ / "golden"
MOTOR_PATH = Path(os.environ.get("MOTOR_PATH", RAIZ.parent / "motor_control_depositos_cbba.py"))

# formato -> archivo real de fixture
FIXTURES = {
    "BCP_ME": "bcp_me_1.xls",
    "BCP_MN": "bcp_mn_3.xls",
    "BISA_ME": "bisa_me_2.xls",
    "BISA_MN": "bisa_mn_2.xls",
    "BNB_AHORRO": "bnb_ahorro_2.xls",
    "BNB_ME": "bnb_me_1_1.xls",
    "BNB_MN": "bnb_mn_3.xls",
    "BNB_CLINICA": "clinica_1.xls",
    "ECO_CTA_CTE": "economico_1.xlsx",
    "ECO_AHORRO": "economico_ahorro.xlsx",
    "BMSC": "mercantil_1.xls",
    "UNION_MN": "union_mn_2.xls",
}
# Formatos cuyo fixture es un extracto valido y normalizable hoy.
FORMATOS_OK = list(FIXTURES)  # los 12 formatos con extracto real con datos o vacio valido
# UNION_ME: formato UNION_FECHAS_V1 valido (estructura confirmada por codigo legado + captura real). Comportamiento con movimientos: pendiente de fixture real.
NEGATIVOS = RAIZ / "fixtures" / "negativos"
UNION_ULTIMOS12 = NEGATIVOS / "union_ultimos12_no_valido.xls"
UNION_ME_VACIO_REAL = EXTRACTOS / "union_me_vacio.xls"   # pendiente de subir por el usuario

# Contrato esperado (banco, cuenta, moneda). Copiado del negocio, NO del motor.
CONTRATO = {
    "BCP_ME": ("BCP", "301-5005425-2-71", "USD"),
    "BCP_MN": ("BCP", "301-5005684-3-97", "BOB"),
    "BISA_ME": ("BISA", "0696872023", "USD"),
    "BISA_MN": ("BISA", "0696870039", "BOB"),
    "BNB_AHORRO": ("BNB", "3501936692", "BOB"),
    "BNB_ME": ("BNB", "3400041236", "USD"),
    "BNB_MN": ("BNB", "3000100152", "BOB"),
    "BNB_CLINICA": ("BNB", "3000100705", "BOB"),
    "ECO_CTA_CTE": ("BANCO ECONÓMICO", "3041210569", "BOB"),
    "ECO_AHORRO": ("BANCO ECONÓMICO", "3051446946", "BOB"),
    "BMSC": ("BMSC", "1000872489", "BOB"),
    "UNION_ME": ("BANCO UNIÓN", "20000003224544", "USD"),
    "UNION_MN": ("BANCO UNIÓN", "10000003224552", "BOB"),
}

COLUMNAS_LISTS_CONTRATO = [
    "CLAVE TRANSACCIÓN", "CÓDIGO DE ASIGNACIÓN", "BANCO", "CUENTA BANCARIA", "MONEDA",
    "FECHA MOVIMIENTO", "HORA MOVIMIENTO", "IMPORTE", "DÉBITO", "CRÉDITO",
    "TIPO MOVIMIENTO", "SALDO", "DESCRIPCIÓN", "DEPOSITANTE / ORIGINANTE",
    "INFORMACIÓN ADICIONAL", "ESTADO", "ESTUDIANTE", "SOLICITADO POR",
    "SEDE SOLICITANTE", "CONFIRMADO POR", "FECHA CONFIRMACIÓN", "OBSERVACIÓN",
    "TEXTO DE BÚSQUEDA", "ARCHIVO ORIGEN", "LOTE DE CARGA", "FECHA DE CARGA",
]
COLS_VOLATILES = ["LOTE DE CARGA", "FECHA DE CARGA"]


def cargar_motor():
    spec = importlib.util.spec_from_file_location("motor_bajo_prueba", MOTOR_PATH)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def a_texto(df):
    """Serializacion determinista (sin columnas volatiles) usada para las referencias doradas."""
    d = df.drop(columns=[c for c in COLS_VOLATILES if c in df.columns])
    return d.to_csv(index=False, lineterminator="\n")


def leer_csv_texto(ruta_o_texto):
    from io import StringIO
    if isinstance(ruta_o_texto, Path):
        ruta_o_texto = ruta_o_texto.read_text(encoding="utf-8-sig")
    d = pd.read_csv(StringIO(ruta_o_texto), dtype=str, keep_default_na=False)
    return d.drop(columns=[c for c in COLS_VOLATILES if c in d.columns])


# ---------- Oraculo independiente (no reutiliza la logica de normalizacion) ----------
DATE_RE = re.compile(r"^\s*\d{1,2}/(\d{1,2}|[A-Za-z]{3})/\d{2,4}\s*$")


def contar_movimientos_independiente(motor, ruta, formato):
    hoja = motor.HOJAS_VALIDAS[formato]
    raw = motor.leer_excel_robusto(ruta, sheet_name=hoja, header=None)
    fila_hdr = None
    for i in range(len(raw)):
        celdas = [str(x).strip().upper() for x in raw.iloc[i].tolist() if pd.notna(x) and str(x).strip()]
        if len(celdas) >= 3 and celdas[0].startswith("FECHA"):
            fila_hdr = i
            break
    assert fila_hdr is not None, "el oraculo no encontro fila de encabezado"
    n = 0
    for i in range(fila_hdr + 1, len(raw)):
        fila = raw.iloc[i]
        texto = " ".join(str(x) for x in fila.tolist() if pd.notna(x)).upper()
        if any(t in texto for t in ("SALDO INICIAL", "SALDO AL CIERRE", "TOTAL")):
            continue
        for x in fila.tolist()[:6]:
            if isinstance(x, (pd.Timestamp,)) or (isinstance(x, str) and DATE_RE.match(x)):
                n += 1
                break
    return n


def _cadena_exacta(s, d, tol):
    asc = all(abs(s[i] - s[i - 1] - d[i]) <= tol for i in range(1, len(s)))
    desc = all(abs(s[i - 1] - s[i] - d[i - 1]) <= tol for i in range(1, len(s)))
    return "asc" if asc else "desc" if desc else None


def _cadena_tolerante(claves, s, d, tol):
    """Tolera movimientos con la MISMA fecha+hora cuyo orden en el archivo no sigue el orden del saldo."""
    grupos = []
    for k, sv, dv in zip(claves, s, d):
        if grupos and grupos[-1][0] == k:
            grupos[-1][1].append(sv); grupos[-1][2] += dv
        else:
            grupos.append([k, [sv], dv])
    factibles = set(grupos[0][1])
    for _, saldos, delta in grupos[1:]:
        nuevos = {x for x in saldos if any(abs(p + delta - x) <= tol for p in factibles)}
        if not nuevos:
            return False
        factibles = nuevos
    return True


def saldo_encadenado_ok(df, tol=0.011):
    """Propiedad independiente: cada SALDO = SALDO anterior +/- movimiento.
    Devuelve (ok, modo): 'asc'/'desc' (exacta fila a fila) o 'tolerante-empates' (misma fecha+hora)."""
    s = df["SALDO"].astype(float).tolist()
    d = (df["CRÉDITO"].fillna(0).astype(float) - df["DÉBITO"].fillna(0).astype(float)).tolist()
    if len(s) < 2:
        return True, "n/a"
    modo = _cadena_exacta(s, d, tol)
    if modo:
        return True, modo
    claves = list(zip(df["FECHA MOVIMIENTO"].astype(str), df["HORA MOVIMIENTO"].astype(str)))
    if _cadena_tolerante(claves, s, d, tol):
        return True, "tolerante-empates"
    if _cadena_tolerante(claves[::-1], s[::-1], d[::-1], tol):
        return True, "tolerante-empates"
    return False, "ninguno"


# ---------- Constructores de extractos sinteticos (para defectos) ----------
def crear_xlsx_bnb(ruta, cuenta, filas, hoja="Hoja 1"):
    """Extracto BNB minimo: fila 1 = cuenta, fila 2 = encabezado (igual que los reales)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws.append(["Cuenta:", cuenta])
    ws.append(["Fecha", "Hora", "Oficina", "Descripción", "Referencia", "Código de transacción",
               "ITF", "Débitos", "Créditos", "Saldo", "Adicionales"])
    for f in filas:
        ws.append([f.get("fecha"), f.get("hora", "10:00:00"), "OF", f.get("desc", "DEPOSITO"),
                   "REF", f.get("cod", "77"), 0, f.get("deb"), f.get("cred"), f.get("saldo"),
                   f.get("adic", "Nombre Originante: PRUEBA;")])
    wb.save(ruta)
    return ruta


def crear_xlsx_bmsc_otra_cuenta(ruta, cuenta="9999999999"):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Excel"
    ws.append(["Banco Mercantil Santa Cruz"])
    ws.append(["Cuenta:", cuenta])
    ws.append(["Fecha", "Hora", "Cod. Bca.", "Débito", "Crédito", "Saldo"])
    ws.append(["01/08/2026", "10:00:00", "123", None, 100, 1100])
    wb.save(ruta)
    return ruta


def crear_xlsx_union_me_vacio(ruta, con_verificasion=True):
    """PROXY SINTETICO (no es el archivo real) de un UNION_ME por fechas sin movimientos.
    Copia la disposicion del UNION_MN real (etiqueta 'Cuenta:' fila 8, encabezado fila 16, columnas
    combinadas) y agrega 'Nro de verificasion' como declaro el usuario. La posicion de esa columna es
    una SUPOSICION: REQUIERE ARCHIVO REAL."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "ExtractoMovimientosFechas"
    ws["N2"] = "Extracto de Movimientos"
    ws["B6"] = "UNIVERSIDAD PRIVADA DEL VALLE S.A."
    ws["B8"] = "Cuenta:"; ws["E8"] = "20000003224544"
    ws["B10"] = "Producto:"; ws["E10"] = "UNICUENTA ESPECIAL PERSONA JURIDICA M/E"
    ws["E12"] = "Desde: 01/08/2026 Hasta: 21/08/2026"
    for col, txt in {"B": "Fecha Movimiento", "D": "AG", "H": "Descripción", "U": "Nro Documento",
                     "Z": "Monto", "AD": "Saldo"}.items():
        ws[f"{col}16"] = txt
    if con_verificasion:
        ws["AG16"] = "Nro de verificasion"
    ws["B18"] = "Total Créditos:"; ws["I18"] = "0.00"
    ws["B20"] = "Total Débitos:"; ws["I20"] = "0.00"
    wb.save(ruta)
    return ruta


def crear_xlsx_union_me_con_movimiento(ruta):
    """PROXY SINTETICO (no es un extracto real): el vacio de UNION_ME + UNA fila de movimiento inventada.
    Sirve SOLO para probar la mecanica de captura (que 'Nro de verificasion' viaja a ORIGEN aunque el motor
    legado no la lea). No valida ningun comportamiento real de UNION_ME con movimientos."""
    from openpyxl import load_workbook
    crear_xlsx_union_me_vacio(ruta)
    wb = load_workbook(ruta)
    ws = wb.active
    for col, v in {"B": "10/08/2026", "D": "201", "H": "DEPOSITO SINTETICO", "U": "12345", "Z": "100.00",
                   "AD": "100.00", "AG": "V-0001"}.items():
        ws[f"{col}17"] = v
    wb.save(ruta)
    return ruta


# ---------- P3b: huella determinista de un EXTRACTO_HISTORICO (doradas sin copiar los datos) ----------
def leer_historico(ruta):
    """Lee un EXTRACTO_HISTORICO: devuelve dict con hojas, fila de encabezado de la tabla, columnas,
    columnas ocultas, filas (valores) y la zona superior (etiqueta -> valor)."""
    from openpyxl import load_workbook
    wb = load_workbook(ruta)
    ws = wb.worksheets[0]
    tabla = next(iter(ws.tables.values()))
    ini, fin = tabla.ref.split(":")
    hr = int(re.sub(r"[A-Z]", "", ini))
    ultima = int(re.sub(r"[A-Z]", "", fin))
    ncol = ws.max_column
    columnas = [ws.cell(hr, j).value for j in range(1, ncol + 1)]
    ocultas = [ws.column_dimensions[ws.cell(hr, j).column_letter].hidden for j in range(1, ncol + 1)]
    filas = [[ws.cell(i, j).value for j in range(1, ncol + 1)] for i in range(hr + 1, ultima + 1)]
    celdas = [[ws.cell(i, j) for j in range(1, ncol + 1)] for i in range(hr + 1, ultima + 1)]
    zona = {}
    for i in range(2, hr):
        vals = [(j, ws.cell(i, j).value) for j in range(1, ncol + 1) if ws.cell(i, j).value is not None]
        vals = [(j, v) for j, v in vals if not str(v).startswith(("DATOS DE LA CUENTA", "SALDOS Y TOTALES"))]
        for (j1, et), (j2, va) in zip(vals[0::2], vals[1::2]):
            zona[str(et)] = va
    return {"wb": wb, "ws": ws, "tabla": tabla, "hr": hr, "columnas": columnas, "ocultas": ocultas,
            "visibles": [c for c, o in zip(columnas, ocultas) if not o], "filas": filas, "celdas": celdas,
            "zona": zona, "titulo": ws.cell(1, 1).value}


def huella_historico(ruta):
    """Resumen estable de un EXTRACTO_HISTORICO: columnas, cantidad de filas, zona superior y SHA-256
    de todos los valores de la tabla (incluida la CLAVE oculta) y de sus formatos de numero."""
    import hashlib
    h = leer_historico(ruta)
    m = hashlib.sha256()
    for fila in h["celdas"]:
        m.update(repr([(c.value, c.number_format) for c in fila]).encode("utf-8"))
    return {"hojas": h["wb"].sheetnames, "titulo": h["titulo"], "columnas": h["columnas"],
            "ocultas": [c for c, o in zip(h["columnas"], h["ocultas"]) if o], "filas": len(h["filas"]),
            "zona": {k: (round(v, 2) if isinstance(v, float) else v) for k, v in h["zona"].items()},
            "sha256_tabla": m.hexdigest()}
