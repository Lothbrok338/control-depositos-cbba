"""P5: la NORMALIZACIÓN productiva sale del motor genérico gobernado por registro_bancos.json.

    archivo → detección por registro (P4) → normalización genérica (P5) → salida productiva

Reglas que estas pruebas protegen:
  * `ejecutar_motor` no usa ningún normalizar_* / validar_archivo / encontrar_fila_encabezado legado: con
    todos ellos inutilizados, los 12 extractos reales producen exactamente la salida aprobada;
  * para los 12 extractos reales la ruta productiva es idéntica a la referencia legada (26 columnas, índice,
    tipos, valores y validación de saldos, sin tolerancia);
  * una cuenta nueva de un formato existente se normaliza SOLO con su entrada en CUENTAS: sin normalizar_*
    específico, sin plantilla legada y sin estar en HOJAS_VALIDAS / ENCABEZADOS_ESPERADOS;
  * UNION_ME se normaliza con UNION_FECHAS_V1 (contrato confirmado); lo que depende de movimientos reales
    queda SKIP hasta tener el fixture;
  * la normalización lee la misma hoja y la misma fila de encabezado que comprobó la detección.
Los normalizar_* legados siguen en el motor como REFERENCIA (paso 16, en sombra) hasta P6.
"""
import ast
import copy
import json
import shutil

import pandas as pd
import pytest
from openpyxl import load_workbook

from helpers import (COLUMNAS_LISTS_CONTRATO, CONTRATO, EXTRACTOS, FIXTURES, FORMATOS_OK, GOLDEN, RAIZ, crear_xlsx_bnb,
                     crear_xlsx_union_me_con_movimiento, leer_csv_texto)

pytestmark = pytest.mark.normalizacion

REPO = RAIZ.parent
TS = pd.Timestamp("2026-01-01 00:00:00")
FIJO = pd.Timestamp("2026-08-21 09:00:00")
PENDIENTE_MOV = "REQUIERE MUESTRA REAL CON MOVIMIENTOS (UNION_ME)"

LEGADO = ("normalizar_archivo", "normalizar_bnb", "normalizar_bcp", "normalizar_union", "normalizar_economico",
          "normalizar_bisa", "normalizar_bmsc", "validar_archivo", "leer_tabla_movimientos",
          "encontrar_fila_encabezado", "aplicar_identidad_registro")

VARIABLES_IDENTIDAD = ["BANCO", "CUENTA BANCARIA", "MONEDA", "CLAVE TRANSACCIÓN", "TEXTO DE BÚSQUEDA",
                       "ARCHIVO ORIGEN"]


# ------------------------------- utilidades -------------------------------
@pytest.fixture(scope="session")
def datos_registro():
    return json.loads((REPO / "registro_bancos.json").read_text(encoding="utf-8"))


def _registro(datos, tmp_path, mutar=None, *cuentas):
    d = copy.deepcopy(datos)
    d["CUENTAS"].extend(cuentas)
    if mutar:
        mutar(d)
    ruta = tmp_path / "registro_p5.json"
    ruta.write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")
    return ruta


def _sin_legado(mp, motor):
    """Inutiliza TODA la normalización / validación legada: si la producción la llamara, fallaría."""
    def prohibido(nombre):
        def f(*a, **k):
            raise AssertionError(f"la producción usó el legado: {nombre}")
        return f
    for nombre in LEGADO:
        mp.setattr(motor, nombre, prohibido(nombre))


def _correr(motor, entrada, salida, registro=None, sin_legado=False, sombra=False):
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(pd.Timestamp, "now", classmethod(lambda cls, tz=None: FIJO))
        mp.setenv("CBBA_MOTOR_SOMBRA", "1" if sombra else "0")
        if registro is None:
            mp.delenv("CBBA_REGISTRO_BANCOS", raising=False)
        else:
            mp.setenv("CBBA_REGISTRO_BANCOS", str(registro))
        if sin_legado:
            _sin_legado(mp, motor)
        return motor.ejecutar_motor(str(entrada), str(salida / "NORMALIZADO.xlsx"))


def _carpeta(tmp_path, *archivos):
    ent = tmp_path / "in"
    ent.mkdir(parents=True, exist_ok=True)
    for a in archivos:
        shutil.copy(a, ent / a.name)
    return ent


def _lote_12(tmp_path):
    return _carpeta(tmp_path, *[EXTRACTOS / FIXTURES[fm] for fm in FORMATOS_OK])


def _hoja_golden(nombre):
    return pd.read_csv(GOLDEN / f"LOTE_12_NORMALIZADO__{nombre}.csv", dtype=str, keep_default_na=False)


def _hoja_salida(salida, nombre):
    return pd.read_excel(salida / "NORMALIZADO.xlsx", sheet_name=nombre, dtype=str, keep_default_na=False)


def _salida_aprobada(salida, res, ignorar_formato=False):
    """LISTS.csv y las 4 hojas de NORMALIZADO.xlsx = doradas del motor original (sin LOTE/FECHA DE CARGA)."""
    pd.testing.assert_frame_equal(leer_csv_texto(salida / "LISTS.csv"), leer_csv_texto(GOLDEN / "LOTE_12_LISTS.csv"))
    volatiles = ["LOTE DE CARGA", "FECHA DE CARGA"]
    x = _hoja_salida(salida, "LISTS").drop(columns=volatiles)
    e = _hoja_golden("LISTS").drop(columns=volatiles, errors="ignore")
    assert list(x.columns) == COLUMNAS_LISTS_CONTRATO[:-2] and len(x) == len(e) == 4064
    assert x["CLAVE TRANSACCIÓN"].tolist() == e["CLAVE TRANSACCIÓN"].tolist()
    for hoja in ("VALIDACION", "RESUMEN", "DIAGNOSTICO"):
        x, e = _hoja_salida(salida, hoja), _hoja_golden(hoja)
        if ignorar_formato and "FORMATO" in x.columns:
            x, e = x.drop(columns="FORMATO"), e.drop(columns="FORMATO")
        assert list(x.columns) == list(e.columns), hoja
        assert len(x) == len(e), hoja
        # mismos números (las doradas guardan el texto de pandas; se comparan como número si lo son)
        for c in x.columns:
            a, b = pd.to_numeric(x[c], errors="coerce"), pd.to_numeric(e[c], errors="coerce")
            if a.notna().all() and b.notna().all():
                assert (a - b).abs().max() < 1e-9, (hoja, c)
            elif c not in ("PRIMERA_FECHA", "ULTIMA_FECHA"):
                assert x[c].tolist() == e[c].tolist(), (hoja, c)
    assert set(res["df_validacion"]["ESTADO"]) <= {"OK", "SIN MOVIMIENTOS"}


# =========================== 1. LA PRODUCCIÓN NO USA EL LEGADO ===========================
def test_p5_ejecutar_motor_no_nombra_la_normalizacion_legada():
    """El texto de ejecutar_motor ya no llama normalizar_* / validar_archivo / encabezados legados; normaliza y
    valida con normalizar_extracto / validar_extracto (motor genérico + registro)."""
    src = (REPO / "motor_control_depositos_cbba.py").read_text(encoding="utf-8")
    f = next(n for n in ast.parse(src).body if isinstance(n, ast.FunctionDef) and n.name == "ejecutar_motor")
    nombres = {n.id for n in ast.walk(f) if isinstance(n, ast.Name)} | \
              {n.attr for n in ast.walk(f) if isinstance(n, ast.Attribute)}
    assert nombres & (set(LEGADO) | {"HOJAS_VALIDAS", "ENCABEZADOS_ESPERADOS", "formato_legado"}) == set()
    assert {"normalizar_extracto", "validar_extracto", "normalizador_registro", "contrato_origen_registro"} <= nombres


def test_p5_lote_real_completo_sin_ningun_normalizador_legado(motor, tmp_path):
    """Los 12 extractos reales con TODO el legado de normalización inutilizado: la salida es la aprobada."""
    res = _correr(motor, _lote_12(tmp_path), tmp_path / "out", sin_legado=True)
    _salida_aprobada(tmp_path / "out", res)
    assert res["origen_estado"]["estado"] == "OK" and res["origen_estado"]["movimientos_mapeados"] == 4064
    assert res["sombra_estado"]["estado"] == "DESACTIVADO"


# =========================== 2. LOS 12 EXTRACTOS REALES: PRODUCCIÓN = REFERENCIA LEGADA ===========================
@pytest.mark.parametrize("formato", FORMATOS_OK)
def test_p5_formato_real_identico_a_la_referencia_legada(normalizado, normalizado_legado, formato):
    """26 columnas, orden, índice de fila, tipos y valores idénticos; validación de saldos idéntica (sin
    tolerancia)."""
    df, val = normalizado[formato]
    df_l, val_l = normalizado_legado[formato]
    assert list(df.columns) == COLUMNAS_LISTS_CONTRATO
    if len(df_l):
        pd.testing.assert_frame_equal(df, df_l, check_exact=True)
    else:
        assert len(df) == 0
    assert val.keys() == val_l.keys()
    for k in val:
        assert (pd.isna(val[k]) and pd.isna(val_l[k])) or val[k] == val_l[k], (k, val[k], val_l[k])


@pytest.mark.parametrize("formato", FORMATOS_OK)
def test_p5_formato_real_clasificacion_banco_cuenta_moneda(motor, ruta_fixture, formato):
    """Identidad del registro también en un extracto válido sin movimientos (BISA_ME)."""
    det = motor.detectar_extracto(ruta_fixture(formato))
    df, ctx = motor.normalizar_extracto(ruta_fixture(formato), det, "L", TS, nombre_origen=FIXTURES[formato])
    ident = ctx["identidad"]
    assert (ident["BANCO"], ident["CUENTA BANCARIA"], ident["MONEDA"]) == CONTRATO[formato]
    assert ident["CUENTA_ID"] == formato
    if len(df):
        assert {tuple(x) for x in df[["BANCO", "CUENTA BANCARIA", "MONEDA"]].values} == {CONTRATO[formato]}
    assert (ctx["hoja"], ctx["fila_encabezado"]) == (det.hoja, det.fila_encabezado)


def test_p5_lote_real_referencia_legada_sin_diferencias_ni_observaciones(corrida_lote):
    res, _ = corrida_lote
    s = res["sombra_estado"]
    assert (s["estado"], s["modo"]) == ("SIN_DIFERENCIAS", "REFERENCIA_LEGADO")
    assert (s["archivos_comparados"], s["archivos_coinciden"], s["diferencias_total"], s["observaciones_total"]) == \
           (12, 12, 0, 0)


def test_p5_la_normalizacion_lee_la_misma_hoja_y_fila_que_la_deteccion(motor, ruta_fixture):
    det = motor.detectar_extracto(ruta_fixture("BNB_ME"))
    det.fila_encabezado += 1
    with pytest.raises(ValueError, match="la detección comprobó"):
        motor.normalizar_extracto(ruta_fixture("BNB_ME"), det, "L", TS, nombre_origen="x")


def test_p5_registro_de_la_version_de_normalizacion(datos_registro):
    assert datos_registro["version_normalizacion"] == "P5-1"
    assert "NORMALIZACION" in datos_registro["_ayuda"]["modo"] and "P5" in datos_registro["_ayuda"]["modo"]


# =========================== 3. CUENTA NUEVA SOLO POR CONFIGURACIÓN ===========================
def test_p5_las_12_cuentas_reales_como_cuentas_nuevas_solo_por_configuracion(motor, datos_registro, tmp_path):
    """Cada cuenta del registro cambia de id (p. ej. BNB_ME → BNB_ME_CFG): el legado no conoce ninguna (no
    están en HOJAS_VALIDAS / ENCABEZADOS_ESPERADOS ni tienen normalizar_* propio). Con el legado inutilizado,
    los 12 extractos reales dan exactamente las mismas filas; solo cambia el id en FORMATO."""
    def renombrar(d):
        for c in d["CUENTAS"]:
            c["id"] += "_CFG"
    reg = _registro(datos_registro, tmp_path, renombrar)
    assert not {c["id"] for c in json.loads(reg.read_text(encoding="utf-8"))["CUENTAS"]} & set(motor.HOJAS_VALIDAS)
    res = _correr(motor, _lote_12(tmp_path), tmp_path / "out", registro=reg, sin_legado=True)
    _salida_aprobada(tmp_path / "out", res, ignorar_formato=True)
    assert sorted(res["df_deteccion_final"]["FORMATO"]) == sorted(f"{fm}_CFG" for fm in FORMATOS_OK)
    det = res["deteccion_estado"]["archivos"]
    assert all(d["CUENTA_NUEVA"] is True for d in det.values())
    assert res["origen_estado"]["estado"] == "OK" and res["origen_estado"]["movimientos_mapeados"] == 4064


NUEVA_BNB = {"id": "BNB_NUEVA_USD", "formato": "BNB_EXTRACTO_V1", "banco": "BNB", "cuenta": "3999000111",
             "moneda": "USD", "activa": True}
FILAS_BNB = [{"fecha": "01/08/2026", "hora": "09:00:00", "cred": 100.0, "saldo": 1100.0, "cod": "11"},
             {"fecha": "02/08/2026", "hora": "10:00:00", "deb": 30.0, "saldo": 1070.0, "cod": "12",
              "adic": "Nombre: ACME; Doc.ID:1"}]


def _eco_con_otra_cuenta(origen, destino, cuenta_nueva):
    """Copia REAL de un extracto del Económico (.xlsx) con otro número en la celda rotulada 'Cuenta:'."""
    wb = load_workbook(origen)
    ws = wb["Extracto"]
    celda = ws["G4"]
    assert ws["E4"].value == "Cuenta:" and celda.value.startswith("C")
    viejo = celda.value.split(":")[1].split("(")[0].strip()
    celda.value = celda.value.replace(viejo, cuenta_nueva)
    wb.save(destino)
    return destino


@pytest.mark.parametrize("fm", ["ECO_AHORRO", "ECO_CTA_CTE"])
def test_p5_cuenta_nueva_con_extracto_real_del_economico(motor, datos_registro, normalizado, tmp_path, fm):
    """Extracto REAL con otro número de cuenta (solo cambia la celda de cabecera) + una entrada en CUENTAS:
    mismas filas, importes, saldos y validación que la cuenta real; solo cambia la identidad."""
    nueva = _eco_con_otra_cuenta(EXTRACTOS / FIXTURES[fm], tmp_path / "eco_nueva.xlsx", "3059990001")
    reg = _registro(datos_registro, tmp_path, None, {"id": "ECO_NUEVA", "formato": "ECO_EXTRACTO_V1",
                                                     "banco": "BANCO ECONÓMICO", "cuenta": "3059990001",
                                                     "moneda": "BOB", "activa": True})
    with pytest.MonkeyPatch.context() as mp:
        _sin_legado(mp, motor)
        det = motor.detector_registro(str(reg)).detectar(str(nueva))
        assert det.estado == "OK" and det.cuenta_id == "ECO_NUEVA" and det.cuenta_nueva is True
        norm = motor.normalizador_registro(str(reg))
        df, ctx = motor.normalizar_extracto(str(nueva), det, "LOTE_TEST", TS, nombre_origen=FIXTURES[fm],
                                            normalizador=norm)
        val = motor.validar_extracto(str(nueva), det, df, ctx, normalizador=norm)
    real, val_real = normalizado[fm]
    assert set(df["CUENTA BANCARIA"]) == {"3059990001"} and set(df["BANCO"]) == {"BANCO ECONÓMICO"}
    assert all(c.startswith("BANCO ECONÓMICO|3059990001|") for c in df["CLAVE TRANSACCIÓN"])
    pd.testing.assert_frame_equal(df.drop(columns=VARIABLES_IDENTIDAD), real.drop(columns=VARIABLES_IDENTIDAD))
    assert val == val_real and val["ESTADO"] == "OK"


def test_p5_cuenta_nueva_bnb_de_punta_a_punta_sin_legado(motor, datos_registro, tmp_path):
    """NORMALIZADO, LISTS y ORIGEN con el legado inutilizado; identidad y CLAVE de la cuenta nueva."""
    ruta = crear_xlsx_bnb(tmp_path / "nueva.xlsx", "3999000111", FILAS_BNB)
    reg = _registro(datos_registro, tmp_path, None, NUEVA_BNB)
    res = _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out", registro=reg, sin_legado=True)
    df = res["df_final"]
    assert df["IMPORTE"].tolist() == [100.0, 30.0] and df["TIPO MOVIMIENTO"].tolist() == ["CRÉDITO", "DÉBITO"]
    assert set(df["CUENTA BANCARIA"]) == {"3999000111"} and set(df["MONEDA"]) == {"USD"}
    assert df["DEPOSITANTE / ORIGINANTE"].tolist() == ["PRUEBA", "ACME"]
    assert df["CLAVE TRANSACCIÓN"].tolist() == df.apply(motor.crear_clave, axis=1).tolist()
    assert res["df_validacion"]["ESTADO"].tolist() == ["OK"]
    lists = pd.read_csv(tmp_path / "out" / "LISTS.csv", encoding="utf-8-sig", dtype=str)
    assert list(lists.columns) == COLUMNAS_LISTS_CONTRATO and set(lists["CUENTA BANCARIA"]) == {"3999000111"}
    assert res["origen_estado"]["estado"] == "OK" and res["origen_estado"]["movimientos_mapeados"] == 2


def test_p5_cuenta_nueva_bnb_en_hoja_alternativa_sin_plantilla_legada(motor, datos_registro, tmp_path):
    """La hoja alternativa del formato ('Hoja') se lee por el registro, no por la plantilla BNB_AHORRO."""
    ruta = crear_xlsx_bnb(tmp_path / "nueva.xlsx", "3999000111", FILAS_BNB, hoja="Hoja")
    reg = _registro(datos_registro, tmp_path, None, dict(NUEVA_BNB, moneda="BOB"))
    res = _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out", registro=reg, sin_legado=True)
    assert set(res["df_final"]["MONEDA"]) == {"BOB"} and len(res["df_final"]) == 2
    assert res["deteccion_estado"]["archivos"]["nueva.xlsx"]["HOJA"] == "Hoja"


ENCABEZADO_BMSC = ["Fecha", "Hora", "Cod. Bca.", "Nro.Cheque", "Nro/Nom.Plantilla", "Cod.Dep.Num", "Doc.Depositante",
                   "Nombre/Denominación Depositante", "Tipo transact", "Descripción", "Oficina", "Banco", "Tipo dep",
                   "Nom.Destinatario", "Glosa", "Originador", "Originador ACH", "Ciudad Origen", "Débito", "Crédito",
                   "Saldo"]


def _bmsc_sintetico(ruta, cuenta):
    """BMSC con la disposición del extracto real (cabecera, 'Nro de Cuenta:', encabezado de 21 columnas)."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Excel"
    for fila in (["Extracto de cuenta"], [], ["Banco Mercantil Santa Cruz"], [], ["Titular"], ["UNIVALLE"],
                 ["Nro de Cuenta:", cuenta, "Fecha de emisión:", "21/08/2026"], [], ["Movimientos"],
                 ["01/08/2026 - 21/08/2026"], ENCABEZADO_BMSC):
        ws.append(fila)
    fila = dict.fromkeys(ENCABEZADO_BMSC, None)
    ws.append(list({**fila, "Fecha": "12/08/2026", "Hora": "17:31:00", "Cod. Bca.": "TT1", "Doc.Depositante": "41",
                    "Nombre/Denominación Depositante": "ROSARIO", "Descripción": "DEPOSITO DE EFECTIVO",
                    "Crédito": " 3,718.00", "Saldo": " 198,142.79"}.values()))
    ws.append(list({**fila, "Fecha": "13/08/2026", "Hora": "09:00:00", "Cod. Bca.": "TT2",
                    "Descripción": "PAGO A PROVEEDOR", "Débito": " 142.79", "Saldo": " 198,000.00"}.values()))
    wb.save(ruta)
    return ruta


def test_p5_cuenta_nueva_bmsc_sin_legado(motor, datos_registro, tmp_path):
    ruta = _bmsc_sintetico(tmp_path / "bmsc_nueva.xlsx", "4000123456")
    reg = _registro(datos_registro, tmp_path, None, {"id": "BMSC_2", "formato": "BMSC_EXCEL_V1", "banco": "BMSC",
                                                     "cuenta": "4000123456", "moneda": "BOB", "activa": True})
    res = _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out", registro=reg, sin_legado=True)
    df = res["df_final"]
    assert set(df["CUENTA BANCARIA"]) == {"4000123456"} and set(df["BANCO"]) == {"BMSC"}
    assert df[["IMPORTE", "TIPO MOVIMIENTO", "CÓDIGO DE ASIGNACIÓN"]].values.tolist() == [
        [3718.0, "CRÉDITO", "TT1"], [142.79, "DÉBITO", "TT2"]]
    assert res["df_validacion"]["ESTADO"].tolist() == ["OK"]
    # la referencia legada (plantilla BMSC + identidad del registro) llega a lo mismo
    ref = _correr(motor, tmp_path / "in", tmp_path / "out_ref", registro=reg, sombra=True)
    assert ref["sombra_estado"]["estado"] == "SIN_DIFERENCIAS"


def test_p5_cuenta_nueva_union_junto_al_lote_real_sin_legado(motor, datos_registro, tmp_path):
    """Otro formato (UNION_FECHAS_V1) junto a extractos reales, con el legado inutilizado: las cuentas reales
    salen igual que sus doradas y la nueva con su identidad."""
    ruta = crear_xlsx_union_me_con_movimiento(tmp_path / "union_nueva.xlsx")
    wb = load_workbook(ruta)
    wb.active["E8"] = "30000009990001"
    wb.save(ruta)
    reg = _registro(datos_registro, tmp_path, None, {"id": "UNION_NUEVA", "formato": "UNION_FECHAS_V1",
                                                     "banco": "BANCO UNIÓN", "cuenta": "30000009990001",
                                                     "moneda": "USD", "activa": True})
    ent = _carpeta(tmp_path, ruta, EXTRACTOS / FIXTURES["BNB_ME"], EXTRACTOS / FIXTURES["UNION_MN"])
    res = _correr(motor, ent, tmp_path / "out", registro=reg, sin_legado=True)
    df = res["df_final"]
    nueva = df[df["ARCHIVO ORIGEN"] == "union_nueva.xlsx"]
    assert len(nueva) == 1 and set(nueva["CUENTA BANCARIA"]) == {"30000009990001"} and set(nueva["MONEDA"]) == {"USD"}
    for fm in ("BNB_ME", "UNION_MN"):
        real = df[df["ARCHIVO ORIGEN"] == FIXTURES[fm]]
        esperado = pd.read_csv(GOLDEN / f"{fm}.csv", dtype=str, keep_default_na=False)
        assert real["CLAVE TRANSACCIÓN"].tolist() == esperado["CLAVE TRANSACCIÓN"].tolist(), fm


def test_p5_cuenta_nueva_la_referencia_legada_coincide(motor, datos_registro, tmp_path):
    """Con el legado disponible, la referencia en sombra (plantilla + identidad del registro) coincide con la
    producción genérica para la cuenta nueva."""
    ruta = crear_xlsx_bnb(tmp_path / "nueva.xlsx", "3999000111", FILAS_BNB)
    reg = _registro(datos_registro, tmp_path, None, NUEVA_BNB)
    res = _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out", registro=reg, sombra=True)
    assert res["sombra_estado"]["estado"] == "SIN_DIFERENCIAS" and res["sombra_estado"]["archivos_coinciden"] == 1
    informe = json.loads((tmp_path / "out" / "SOMBRA_REPORTE.json").read_text(encoding="utf-8"))
    assert informe["archivos"][0]["formato_legado"] == "BNB_MN"          # plantilla, solo en la referencia


def test_p5_cuenta_sin_registrar_sigue_rechazada(motor, tmp_path):
    ruta = crear_xlsx_bnb(tmp_path / "nueva.xlsx", "3999000111", FILAS_BNB)
    with pytest.raises(ValueError, match="CUENTA_NO_REGISTRADA"):
        _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out")
    assert not (tmp_path / "out").exists()


# =========================== 4. UNION_ME (contrato UNION_FECHAS_V1) ===========================
def test_p5_union_me_proxy_sintetico_se_normaliza_por_registro(motor, tmp_path):
    """PROXY SINTÉTICO (no es un extracto real): UNION_ME usa la misma configuración UNION_FECHAS_V1 que
    UNION_MN con su propia identidad; 'Nro de verificasion' no entra a las 26 columnas."""
    ruta = crear_xlsx_union_me_con_movimiento(tmp_path / "u.xlsx")
    with pytest.MonkeyPatch.context() as mp:
        _sin_legado(mp, motor)
        det = motor.detectar_extracto(str(ruta))
        df, ctx = motor.normalizar_extracto(str(ruta), det, "L", TS, nombre_origen="u.xlsx")
        val = motor.validar_extracto(str(ruta), det, df, ctx)
    assert det.cuenta_id == "UNION_ME" and list(df.columns) == COLUMNAS_LISTS_CONTRATO
    assert (df["BANCO"].iloc[0], df["CUENTA BANCARIA"].iloc[0], df["MONEDA"].iloc[0]) == CONTRATO["UNION_ME"]
    assert df[["IMPORTE", "TIPO MOVIMIENTO", "CÓDIGO DE ASIGNACIÓN"]].values.tolist() == [[100.0, "CRÉDITO", "12345"]]
    assert df["INFORMACIÓN ADICIONAL"].iloc[0].startswith("AG: 201")   # el proxy guarda AG numérico (201 → 201.0)
    assert not df.astype(str).apply(lambda c: c.str.contains("V-0001")).any().any()
    assert val["ESTADO"] == "OK"
    # la referencia legada llega a lo mismo
    legado = motor.normalizar_archivo(str(ruta), "UNION_ME", "L", TS, nombre_origen="u.xlsx")
    pd.testing.assert_frame_equal(df, legado)


@pytest.mark.skip(reason=PENDIENTE_MOV + ": normalización genérica de movimientos reales (26 columnas, importes "
                                         "con signo, AG, saldos) contra la referencia legada y una dorada")
def test_p5_union_me_real_con_movimientos_se_normaliza_por_registro():
    pass


@pytest.mark.skip(reason=PENDIENTE_MOV + ": validación de saldos productiva (RECONSTRUIDO / ULTIMA_FILA) y "
                                         "débitos reales (Monto negativo) de UNION_ME")
def test_p5_union_me_real_validacion_y_debitos():
    pass
