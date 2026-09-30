"""P3b · Capa 4: EXTRACTO_HISTORICO (historico.py + bloque `historico` de registro_bancos.json).

Un Excel amigable por banco / cuenta / mes, construido SOLO con DATOS_ORIGINALES + MAPA_ORIGEN (ORIGEN.xlsx)
+ NORMALIZADO (NORMALIZADO.xlsx o LISTS.csv) + registro. Las expectativas de columnas se copian del diseño
aprobado (DISENO_TRES_CAPAS.md §6.4), no del registro; los valores se contrastan con un lector independiente
(pandas) de ORIGEN.xlsx y NORMALIZADO.xlsx.
"""
import builtins
import datetime as dt
import hashlib
import importlib.util
import io
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from helpers import (COLUMNAS_LISTS_CONTRATO, EXTRACTOS, FIXTURES, GOLDEN, MOTOR_PATH, RAIZ, crear_xlsx_bnb,
                     crear_xlsx_union_me_con_movimiento, crear_xlsx_union_me_vacio, huella_historico,
                     leer_csv_texto, leer_historico)

pytestmark = pytest.mark.historico
REPO = RAIZ.parent
PENDIENTE_MOV = "REQUIERE MUESTRA REAL CON MOVIMIENTOS (UNION_ME)"
OPERATIVAS = ["ESTADO", "CONFIRMADO POR", "FECHA DE CONFIRMACIÓN"]

# Unidad de salida esperada (diseño §6.2): 12 extractos válidos -> 11 archivos; BISA_ME sin movimientos no genera.
ARCHIVOS = {
    "BNB_MN": "EXTRACTO_HISTORICO_BNB_3000100152_2026-08.xlsx",
    "BNB_ME": "EXTRACTO_HISTORICO_BNB_3400041236_2026-08.xlsx",
    "BNB_AHORRO": "EXTRACTO_HISTORICO_BNB_3501936692_2026-08.xlsx",
    "BNB_CLINICA": "EXTRACTO_HISTORICO_BNB_3000100705_2026-08.xlsx",
    "BCP_MN": "EXTRACTO_HISTORICO_BCP_301-5005684-3-97_2026-08.xlsx",
    "BCP_ME": "EXTRACTO_HISTORICO_BCP_301-5005425-2-71_2026-07.xlsx",
    "BISA_MN": "EXTRACTO_HISTORICO_BISA_0696870039_2026-08.xlsx",
    "ECO_CTA_CTE": "EXTRACTO_HISTORICO_BANCO_ECONOMICO_3041210569_2026-08.xlsx",
    "ECO_AHORRO": "EXTRACTO_HISTORICO_BANCO_ECONOMICO_3051446946_2026-08.xlsx",
    "BMSC": "EXTRACTO_HISTORICO_BMSC_1000872489_2026-08.xlsx",
    "UNION_MN": "EXTRACTO_HISTORICO_BANCO_UNION_10000003224552_2026-08.xlsx",
}
BNB = ["Fecha", "Hora", "Oficina", "Descripción", "Referencia", "Código de transacción", "ITF", "Débitos",
       "Créditos", "Saldo", "Adicionales"]
BCP = ["Fecha", "Hora", "Glosa", "Tipo", "Suc. Age.", "Usuario", "Importe", "Saldo", "Nro. Operación"]
BISA = ["Fecha", "Hora", "Nro. Cheque", "Descripción", "Importe", "Saldo", "Info. Complementaria", "Sucursal",
        "Canal", "Nro. Ref.", "Codigo", "Nro. Lote"]
ECO = ["Fecha", "Hora", "Nro Trn./Cheque", "Transacción", "Nota", "Monto", "Saldo"]
BMSC = ["Fecha", "Hora", "Cod. Bca.", "Nro.Cheque", "Nro/Nom.Plantilla", "Cod.Dep.Num", "Doc.Depositante",
        "Nombre/Denominación Depositante", "Tipo transact", "Descripción", "Oficina", "Banco", "Tipo dep",
        "Nom.Destinatario", "Glosa", "Originador", "Originador ACH", "Ciudad Origen", "Débito", "Crédito", "Saldo"]
UNION_MN = ["Fecha Movimiento", "AG", "Descripción", "Nro Documento", "Monto", "Saldo"]
UNION_ME = UNION_MN + ["Nro de verificasion"]
COLUMNAS = {"BNB_MN": BNB, "BNB_ME": BNB, "BNB_AHORRO": BNB, "BNB_CLINICA": BNB, "BCP_MN": BCP, "BCP_ME": BCP,
            "BISA_MN": BISA, "ECO_CTA_CTE": ECO, "ECO_AHORRO": ECO, "BMSC": BMSC, "UNION_MN": UNION_MN}
# columnas que vienen de NORMALIZADO (valores validados) y no del texto del banco
DE_NORMALIZADO = {"Fecha", "Hora", "Fecha Movimiento", "Débitos", "Créditos", "Débito", "Crédito", "Saldo",
                  "Importe", "Monto"}
IMPORTE_CON_SIGNO = {"Importe", "Monto"}


# ------------------------------------------------------------------ utilidades y fixtures
def cargar(nombre, archivo):
    spec = importlib.util.spec_from_file_location(nombre, REPO / archivo)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def hist():
    return cargar("historico_bajo_prueba", "historico.py")


@pytest.fixture(scope="module")
def cap():
    return cargar("captura_para_historico", "captura_origen.py")


def sha(ruta):
    return hashlib.sha256(Path(ruta).read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def lote(corrida_lote, hist, tmp_path_factory):
    """Históricos de la corrida de 12 fixtures, generados en una carpeta PROPIA (no en la de producción)."""
    _, salida = corrida_lote
    antes = {n: sha(salida / n) for n in ("NORMALIZADO.xlsx", "LISTS.csv", "ORIGEN.xlsx")}
    carpeta = tmp_path_factory.mktemp("historico_lote")
    info = hist.generar_extractos_historicos(str(salida / "ORIGEN.xlsx"), str(salida / "NORMALIZADO.xlsx"),
                                             str(carpeta))
    despues = {n: sha(salida / n) for n in ("NORMALIZADO.xlsx", "LISTS.csv", "ORIGEN.xlsx")}
    return {"info": info, "carpeta": carpeta, "salida": salida, "antes": antes, "despues": despues}


@pytest.fixture(scope="module")
def hojas(lote):
    return {fm: leer_historico(lote["carpeta"] / nombre) for fm, nombre in ARCHIVOS.items()}


@pytest.fixture(scope="module")
def indep(lote):
    """Lector independiente (pandas, texto exacto) de ORIGEN.xlsx y NORMALIZADO.xlsx."""
    o = pd.read_excel(lote["salida"] / "ORIGEN.xlsx", sheet_name=["DATOS_ORIGINALES", "MAPA_ORIGEN"], dtype=str,
                      keep_default_na=False)
    lists = pd.read_excel(lote["salida"] / "NORMALIZADO.xlsx", sheet_name="LISTS")
    val = pd.read_excel(lote["salida"] / "NORMALIZADO.xlsx", sheet_name="VALIDACION")
    d = o["DATOS_ORIGINALES"]
    mov = d[d["ROL_FILA"] == "MOVIMIENTO"]
    celdas = {}
    for i, campo, v in zip(mov["ID_ORIGEN"], mov["CAMPO_ORIGINAL"], mov["VALOR_ORIGINAL"]):
        celdas.setdefault(i, {})[re.sub(r"\s+", " ", campo).strip()] = v
    mapa = o["MAPA_ORIGEN"].set_index("CLAVE TRANSACCIÓN")
    return {"datos": d, "celdas": celdas, "mapa": mapa, "lists": lists.set_index("CLAVE TRANSACCIÓN"),
            "validacion": val.set_index("ARCHIVO")}


def claves_de(h):
    i = h["columnas"].index("CLAVE TRANSACCIÓN")
    return [f[i] for f in h["filas"]]


def col(h, nombre):
    i = h["columnas"].index(nombre)
    return [f[i] for f in h["filas"]]


def capas_sinteticas(motor, cap, ruta, formato, nombre):
    """ORIGEN (en memoria) + NORMALIZADO de un extracto sintético, con el mismo motor y la misma captura."""
    df = motor.normalizar_archivo(str(ruta), formato, "L", pd.Timestamp("2026-01-01"), nombre_origen=nombre)
    contrato = {"hojas_validas": motor.HOJAS_VALIDAS, "encabezados_esperados": motor.ENCABEZADOS_ESPERADOS,
                "encontrar_fila_encabezado": motor.encontrar_fila_encabezado}
    r = cap.capturar_extracto(str(ruta), nombre, formato, contrato, list(df.index), tabla=df)
    datos = [dict(zip(cap.COLS_DATOS, f)) for f in r["datos"]]
    mapa = [dict(zip(cap.COLS_MAPA, f)) for f in r["mapa"]]
    return datos, mapa, df


# =========================================================== 1. archivo correcto por cuenta / mes
def test_once_archivos_uno_por_cuenta_y_mes_con_nombre_esperado(lote):
    info = lote["info"]
    assert info["estado"] == "OK" and info["omitidos"] == [], info["omitidos"]
    assert sorted(a["nombre_archivo"] for a in info["archivos"]) == sorted(ARCHIVOS.values())
    assert sorted(p.name for p in lote["carpeta"].iterdir()) == sorted(ARCHIVOS.values())
    assert {a["cuenta"]: a["nombre_archivo"] for a in info["archivos"]} == ARCHIVOS
    assert info["extractos_sin_movimientos"] == [FIXTURES["BISA_ME"]]   # BISA_ME no genera archivo
    assert info["advertencias"] == []


def test_movimientos_por_archivo_igual_a_normalizado(lote, indep):
    lists = indep["lists"]
    por_cuenta = lists.groupby("CUENTA BANCARIA").size().to_dict()
    for a in lote["info"]["archivos"]:
        cuenta = a["nombre_archivo"].split("_")[-2]
        assert a["movimientos"] == por_cuenta[cuenta]
    assert sum(a["movimientos"] for a in lote["info"]["archivos"]) == len(lists) == 4064


def test_extracto_que_cruza_dos_meses_genera_dos_archivos(motor, cap, hist, tmp_path):
    ruta = crear_xlsx_bnb(tmp_path / "bnb.xlsx", "3000100152", [
        {"fecha": "31/07/2026", "hora": "23:59:00", "cred": 100, "saldo": 1100, "cod": "A1"},
        {"fecha": "01/08/2026", "hora": "08:00:00", "deb": 50, "saldo": 1050, "cod": "A2"},
        {"fecha": "02/08/2026", "hora": "09:00:00", "cred": 10, "saldo": 1060, "cod": "A3"}])
    datos, mapa, df = capas_sinteticas(motor, cap, ruta, "BNB_MN", "bnb.xlsx")
    res = hist.construir_extractos_historicos(datos, mapa, df, hist.cargar_registro())
    assert [(r["nombre_archivo"], r["totales"]["movimientos"]) for r in res] == [
        ("EXTRACTO_HISTORICO_BNB_3000100152_2026-07.xlsx", 1), ("EXTRACTO_HISTORICO_BNB_3000100152_2026-08.xlsx", 2)]
    assert all(not r["errores"] for r in res)


# =========================================================== 2. una sola hoja visible EXTRACTO
@pytest.mark.parametrize("fm", list(ARCHIVOS))
def test_una_sola_hoja_visible_llamada_extracto(fm, hojas):
    wb = hojas[fm]["wb"]
    assert wb.sheetnames == ["EXTRACTO"]
    assert wb["EXTRACTO"].sheet_state == "visible"
    assert len(wb.defined_names) == 0
    assert list(hojas[fm]["ws"].tables) == ["tblEXTRACTO"]


# =========================================================== 3. columnas bancarias correctas por formato
@pytest.mark.parametrize("fm", list(ARCHIVOS))
def test_columnas_visibles_son_las_del_banco_en_su_orden_mas_tres_operativas(fm, hojas):
    h = hojas[fm]
    assert h["visibles"] == COLUMNAS[fm] + OPERATIVAS
    assert h["columnas"] == COLUMNAS[fm] + OPERATIVAS + ["CLAVE TRANSACCIÓN"]
    assert h["ocultas"] == [False] * (len(COLUMNAS[fm]) + 3) + [True]


@pytest.mark.parametrize("fm", list(ARCHIVOS))
def test_orden_de_columnas_igual_al_orden_original_del_banco(fm, indep, lote):
    """El orden visible es el orden de COLUMNAS_ORIGINALES que la captura registró del archivo real."""
    meta = pd.read_excel(lote["salida"] / "ORIGEN.xlsx", sheet_name="METADATOS_EXTRACTO", dtype=str,
                         keep_default_na=False).set_index("ARCHIVO ORIGEN")
    orig = [re.sub(r"\s+", " ", c.split("=", 1)[1]).strip()
            for c in meta.loc[FIXTURES[fm], "COLUMNAS_ORIGINALES"].split(" | ")]
    assert orig == COLUMNAS[fm]


def test_registro_declara_el_bloque_historico_de_todos_los_formatos_aceptados():
    reg = json.loads((REPO / "registro_bancos.json").read_text(encoding="utf-8"))
    assert reg["version_registro"] == "P3-1" and reg["version_historico"] == "P3b-1"
    for fid, f in reg["FORMATOS"].items():
        if f.get("aceptado", True):
            assert f.get("historico", {}).get("columnas"), fid
        else:
            assert "historico" not in f
    esperado = {"BNB_EXTRACTO_V1": BNB, "BCP_EXTRACTO_V1": BCP, "BISA_EXTRACTO_V1": BISA, "ECO_EXTRACTO_V1": ECO,
                "BMSC_EXCEL_V1": BMSC, "UNION_FECHAS_V1": UNION_ME}
    for fid, cols in esperado.items():
        assert [c["origen"] for c in reg["FORMATOS"][fid]["historico"]["columnas"]] == cols


def test_registro_sigue_siendo_valido_para_el_motor_generico():
    mg = cargar("mg_para_historico", "motor_generico.py")
    assert mg.Registro.cargar(REPO / "registro_bancos.json").validar() == []


# =========================================================== 4. conservación de referencias, cheques, códigos, glosas
@pytest.mark.parametrize("fm", list(ARCHIVOS))
def test_sin_perdida_cada_columna_bancaria_conserva_el_valor_del_banco(fm, hojas, indep):
    """Todo valor bancario de cada movimiento (DATOS_ORIGINALES) aparece en su columna, solo sin los espacios de
    relleno; las columnas de importe/fecha/hora vienen de NORMALIZADO (se prueban aparte)."""
    h = hojas[fm]
    for fila, clave in zip(h["filas"], claves_de(h)):
        celdas = indep["celdas"][indep["mapa"].loc[clave, "ID_ORIGEN"]]
        assert set(celdas) <= set(COLUMNAS[fm]), f"columna del banco sin mostrar: {set(celdas) - set(COLUMNAS[fm])}"
        for nombre in COLUMNAS[fm]:
            if nombre in DE_NORMALIZADO:
                continue
            v = fila[h["columnas"].index(nombre)]
            original = celdas.get(nombre, "").strip()
            if nombre == "ITF":
                assert (v is None and original == "") or v == pytest.approx(float(original.replace(",", "")))
            else:
                assert v == (original or None), f"{clave} · {nombre}: {v!r} != {original!r}"


def test_codigos_cheques_y_referencias_son_texto_y_conservan_ceros(hojas):
    ref = hojas["BNB_ME"]
    i = ref["columnas"].index("Referencia")
    celdas = [f[i] for f in ref["celdas"] if f[i].value == "0000335"]
    assert len(celdas) == 1 and celdas[0].data_type == "s" and celdas[0].number_format == "@"
    bisa = hojas["BISA_MN"]
    fila = dict(zip(bisa["columnas"], bisa["celdas"][0]))
    assert fila["Codigo"].value == "034" and fila["Codigo"].data_type == "s"
    assert fila["Nro. Cheque"].value == "0" and fila["Nro. Cheque"].data_type == "s"
    assert fila["Nro. Ref."].value == "53519247217" and fila["Nro. Ref."].data_type == "s"
    bcp = hojas["BCP_MN"]
    tipos = set(col(bcp, "Tipo"))
    assert {"4401", "2401", "3001"} <= tipos and all(isinstance(t, str) for t in tipos)
    assert set(col(bcp, "Suc. Age.")) >= {"201204", "301314"}          # en LISTS quedó degradado a '301314.0'
    assert all(isinstance(x, str) for x in col(hojas["ECO_CTA_CTE"], "Nro Trn./Cheque"))
    assert all(isinstance(x, str) for x in col(hojas["UNION_MN"], "Nro Documento"))


def test_glosas_nombres_y_datos_de_contraparte_se_conservan(hojas):
    assert any(g.startswith("CHEQUE 00015496") for g in col(hojas["BCP_MN"], "Glosa"))
    bmsc = hojas["BMSC"]
    assert col(bmsc, "Doc.Depositante") == ["4153606TJ", "12581144", "13320679CB"]
    assert "CASTRILLO  DE ALTAMIRANO ROSARIO" in col(bmsc, "Nombre/Denominación Depositante")  # espacio interno intacto
    assert col(bmsc, "Originador ACH") == [None, "MONTERO CORTEZ LUCIANA NICOLE", None]
    assert col(bmsc, "Ciudad Origen") == ["TA", "LP", "CB"]
    assert col(bmsc, "Nom.Destinatario") == [None, None, None]        # columna del banco vacía: existe y queda vacía
    assert col(hojas["UNION_MN"], "AG")[:3] == ["SUC", "SUC", "ORU"]
    info = col(hojas["BISA_MN"], "Info. Complementaria")[0]
    assert info.startswith("Nombre:FRIAS NATALIA Doc.ID:CCB-3816500")


def test_debitos_conservan_todas_las_columnas_del_banco(hojas):
    bnb = hojas["BNB_MN"]
    deb = [dict(zip(bnb["columnas"], f)) for f in bnb["filas"] if f[bnb["columnas"].index("Débitos")] is not None]
    assert len(deb) == 620
    assert all(d["Créditos"] is None and d["Oficina"] and d["Código de transacción"] for d in deb)
    cheque = [d for d in deb if d["Referencia"] == "0008125"]
    assert len(cheque) == 1 and cheque[0]["Descripción"] == "Debito Pago Cheque Efectivo" and cheque[0]["Oficina"]
    me = hojas["BNB_ME"]
    d = next(dict(zip(me["columnas"], f)) for f in me["filas"] if f[me["columnas"].index("Referencia")] == "0000335")
    assert d["Descripción"] == "Debito Pago Cheque Efectivo" and d["Oficina"] == "COCHABAMBA-AGENCIA AMERICA"
    assert d["Débitos"] == 1601.60 and d["Créditos"] is None and d["ITF"] == 0.0
    assert d["Adicionales"] == "Nombre:MELVY ROXANA CRUZ MOBO; Doc.ID:4511776"
    bcp = hojas["BCP_MN"]
    deb = [dict(zip(bcp["columnas"], f)) for f in bcp["filas"] if f[bcp["columnas"].index("Importe")] < 0]
    assert len(deb) == 445
    assert all(x["Tipo"] and x["Suc. Age."] and x["Usuario"] and x["Nro. Operación"] and x["Glosa"] for x in deb)
    eco = hojas["ECO_CTA_CTE"]
    deb = [dict(zip(eco["columnas"], f)) for f in eco["filas"] if f[eco["columnas"].index("Monto")] < 0]
    assert len(deb) == 3 and {x["Transacción"] for x in deb} == {"CARGO OPERACIONES CONTABILIDAD",
                                                                 "DEBIT.AUTOMATICO TARJETA DE CREDITO M/N"}
    assert any(x["Nota"] == "COMISIONES UNIVALLE JULIO 2026" and x["Monto"] == -6.0 for x in deb)


# =========================================================== 5. números y fechas desde NORMALIZADO
@pytest.mark.parametrize("fm", list(ARCHIVOS))
def test_importes_saldos_fecha_y_hora_son_los_de_normalizado(fm, hojas, indep):
    h = hojas[fm]
    lists = indep["lists"]
    for fila, clave in zip(h["filas"], claves_de(h)):
        n = lists.loc[clave]
        d = dict(zip(h["columnas"], fila))
        fecha = d.get("Fecha", d.get("Fecha Movimiento"))
        assert isinstance(fecha, dt.datetime) and fecha.date() == n["FECHA MOVIMIENTO"].date()
        if "Hora" in d:
            assert d["Hora"].strftime("%H:%M:%S") == n["HORA MOVIMIENTO"]
        assert d["Saldo"] == pytest.approx(n["SALDO"], abs=1e-9)
        for c in IMPORTE_CON_SIGNO & set(d):
            signo = -1 if n["TIPO MOVIMIENTO"] == "DÉBITO" else 1
            assert d[c] == pytest.approx(signo * n["IMPORTE"], abs=1e-9)
        for c, k in (("Débitos", "DÉBITO"), ("Débito", "DÉBITO"), ("Créditos", "CRÉDITO"), ("Crédito", "CRÉDITO")):
            if c in d:
                assert (d[c] is None and pd.isna(n[k])) or d[c] == pytest.approx(n[k], abs=1e-9)


@pytest.mark.parametrize("fm", list(ARCHIVOS))
def test_formato_visible_de_fecha_dd_mm_yyyy_y_de_importes(fm, hojas):
    h = hojas[fm]
    nombre_fecha = "Fecha Movimiento" if fm == "UNION_MN" else "Fecha"
    i = h["columnas"].index(nombre_fecha)
    assert {f[i].number_format for f in h["celdas"]} == {"dd/mm/yyyy"}
    i = h["columnas"].index("Saldo")
    assert {f[i].number_format for f in h["celdas"]} == {"#,##0.00"}
    i = h["columnas"].index("FECHA DE CONFIRMACIÓN")
    assert {f[i].number_format for f in h["celdas"]} == {"dd/mm/yyyy hh:mm"}


def test_hora_bisa_respeta_el_formato_del_banco_hh_mm(hojas):
    i = hojas["BISA_MN"]["columnas"].index("Hora")
    assert hojas["BISA_MN"]["celdas"][0][i].number_format == "hh:mm"
    assert "Hora" not in hojas["UNION_MN"]["columnas"]                  # Unión no entrega hora


# =========================================================== 6. columnas operativas
@pytest.mark.parametrize("fm", list(ARCHIVOS))
def test_estado_confirmado_por_y_fecha_de_confirmacion_al_final(fm, hojas):
    h = hojas[fm]
    assert h["visibles"][-3:] == OPERATIVAS
    assert set(col(h, "ESTADO")) == {"DISPONIBLE"}
    assert set(col(h, "CONFIRMADO POR")) == {None} and set(col(h, "FECHA DE CONFIRMACIÓN")) == {None}
    ws, hr = h["ws"], h["hr"]
    colores = [ws.cell(hr, j + 1).fill.fgColor.rgb[-6:] for j in range(len(h["visibles"]))]
    assert set(colores[:-3]) == {"7B1E2B"} and set(colores[-3:]) == {"3F4447"}   # bloque banco vs. operativo


def test_tabla_estados_opcional_reemplaza_los_valores_operativos(lote, hist, tmp_path):
    datos, mapa = hist.leer_origen(lote["salida"] / "ORIGEN.xlsx")
    norm = hist.leer_normalizado(lote["salida"] / "NORMALIZADO.xlsx")
    res = {r["id_cuenta"]: r for r in hist.construir_extractos_historicos(datos, mapa, norm, hist.cargar_registro())}
    clave = res["BMSC"]["claves"][1]
    estados = {clave: {"ESTADO": "CONFIRMADO", "CONFIRMADO POR": "Ingresos CBBA",
                       "FECHA DE CONFIRMACIÓN": "2026-08-22T10:30:00"}}
    r = next(x for x in hist.construir_extractos_historicos(datos, mapa, norm, hist.cargar_registro(), estados)
             if x["id_cuenta"] == "BMSC")
    hist.escribir_extracto_historico(r, str(tmp_path / r["nombre_archivo"]))
    h = leer_historico(tmp_path / r["nombre_archivo"])
    assert col(h, "ESTADO") == ["DISPONIBLE", "CONFIRMADO", "DISPONIBLE"]
    assert col(h, "CONFIRMADO POR") == [None, "Ingresos CBBA", None]
    assert col(h, "FECHA DE CONFIRMACIÓN")[1] == dt.datetime(2026, 8, 22, 10, 30)


# =========================================================== 7. orden de movimientos
@pytest.mark.parametrize("fm", list(ARCHIVOS))
def test_orden_cronologico_y_empates_en_el_orden_original_del_archivo(fm, hojas, indep):
    h = hojas[fm]
    mapa = indep["mapa"]
    fechas = col(h, "Fecha Movimiento" if fm == "UNION_MN" else "Fecha")
    horas = col(h, "Hora") if "Hora" in h["columnas"] else [None] * len(fechas)
    filas_excel = [int(mapa.loc[c, "FILA_EXCEL"]) for c in claves_de(h)]
    claves = [(f, hh or dt.time(0), fe) for f, hh, fe in zip(fechas, horas, filas_excel)]
    assert claves == sorted(claves)


def test_empates_del_mismo_segundo_en_bnb_mn_siguen_el_orden_del_archivo(hojas, indep):
    h = hojas["BNB_MN"]
    mapa = indep["mapa"]
    vistos = {}
    for f, hh, c in zip(col(h, "Fecha"), col(h, "Hora"), claves_de(h)):
        vistos.setdefault((f, hh), []).append(int(mapa.loc[c, "FILA_EXCEL"]))
    empates = [v for v in vistos.values() if len(v) > 1]
    assert empates and all(v == sorted(v) for v in empates)


# =========================================================== 8. zona superior y saldos declarados
def test_zona_superior_muestra_lo_que_trae_cada_banco(hojas):
    z = {fm: h["zona"] for fm, h in hojas.items()}
    for fm, h in hojas.items():
        assert z[fm]["Banco"] and z[fm]["Cuenta"] == ARCHIVOS[fm].split("_")[-2] and z[fm]["Moneda"] in ("BOB", "USD")
        assert z[fm]["Periodo"] in ("Agosto 2026", "Julio 2026")
    u = z["UNION_MN"]
    assert u["Congelado"] == 40578.52 and u["Disponible"] == 339068.09 and u["Total (saldo final del extracto)"] == 379646.61
    assert u["Producto"] == "UNICUENTA NORMAL PERSONA JURIDICA M/N" and u["Período del extracto"] == "01/08/2026 al 21/08/2026"
    e = z["ECO_CTA_CTE"]
    assert e["Retenciones judiciales"] == 0.0 and e["Fondos reservados"] == 0.0 and e["Depósitos por confirmar"] == 0.0
    assert e["Producto"] == "BASICA-PERSONA JURIDICA" and e["Estado de la cuenta"] == "ACTIVA"
    b = z["BISA_MN"]
    assert b["Producto"] == "Bisa Gestion" and b["Saldo inicial del extracto"] == 62687.52 and b["Total créditos"] == 1139.0
    assert z["BMSC"]["Tipo de producto"] == "CUENTA CORRIENTE" and z["BMSC"]["Saldo a la fecha de emisión"] == 201662.79
    assert z["BCP_ME"]["Saldo inicial del extracto (07/07/2026)"] == 22323.54
    assert z["BCP_ME"]["Saldo al cierre del extracto (31/07/2026)"] == 27888.54
    # BNB no declara saldos: no se muestra ninguno (nada se inventa)
    assert not any("aldo" in k or "otal" in k for k in z["BNB_MN"])
    # lo que el diseño excluye: usuario de consulta de BISA y filtros del reporte de Económico
    todo = " ".join(str(v) for h in hojas.values() for v in list(h["zona"]) + list(h["zona"].values()))
    assert "lvelasquezs26" not in todo and "Monto mínimo" not in todo and "Busqueda de texto" not in todo


def test_saldos_declarados_coinciden_con_la_validacion_del_motor(hojas, indep):
    v = indep["validacion"]
    casos = [("BISA_MN", "Saldo inicial del extracto", "SALDO INICIAL"),
             ("ECO_CTA_CTE", "Saldo inicial del extracto", "SALDO INICIAL"),
             ("ECO_CTA_CTE", "Saldo final del extracto", "SALDO FINAL"),
             ("ECO_AHORRO", "Saldo final del extracto", "SALDO FINAL"),
             ("BCP_MN", "Saldo inicial del extracto (01/08/2026)", "SALDO INICIAL"),
             ("BCP_MN", "Saldo al cierre del extracto (20/08/2026)", "SALDO FINAL"),
             ("BMSC", "Saldo a la fecha de emisión", "SALDO FINAL"),
             ("UNION_MN", "Total (saldo final del extracto)", "SALDO FINAL")]
    for fm, etiqueta, campo in casos:
        assert hojas[fm]["zona"][etiqueta] == pytest.approx(v.loc[FIXTURES[fm], campo], abs=0.005), (fm, etiqueta)


def test_un_saldo_declarado_que_no_cuadra_bloquea_el_archivo(lote, hist, tmp_path):
    datos, mapa = hist.leer_origen(lote["salida"] / "ORIGEN.xlsx")
    norm = hist.leer_normalizado(lote["salida"] / "NORMALIZADO.xlsx")
    alterado = [dict(d) for d in datos]
    for d in alterado:
        if d["ARCHIVO ORIGEN"] == FIXTURES["ECO_CTA_CTE"] and str(d["VALOR_ORIGINAL"]).strip() == "647,240.28":
            d["VALOR_ORIGINAL"] = "647,240.99"
    res = {r["id_cuenta"]: r for r in hist.construir_extractos_historicos(alterado, mapa, norm, hist.cargar_registro())}
    assert any("Saldo final del extracto" in e and "no cuadra" in e for e in res["ECO_CTA_CTE"]["errores"])
    assert res["ECO_CTA_CTE"]["filas"] == [] and not res["BNB_MN"]["errores"]
    with pytest.raises(ValueError, match="no escrito"):
        hist.escribir_extracto_historico(res["ECO_CTA_CTE"], str(tmp_path / "x.xlsx"))
    assert not (tmp_path / "x.xlsx").exists()


def test_movimiento_sin_fila_de_origen_bloquea_solo_su_archivo(lote, hist):
    datos, mapa = hist.leer_origen(lote["salida"] / "ORIGEN.xlsx")
    norm = hist.leer_normalizado(lote["salida"] / "NORMALIZADO.xlsx")
    quitada = next(m for m in mapa if m["FORMATO"] == "BMSC")
    res = {r["id_cuenta"] or r["nombre_archivo"]: r for r in hist.construir_extractos_historicos(
        datos, [m for m in mapa if m is not quitada], norm,
        hist.cargar_registro())}
    bmsc = next(r for r in res.values() if r["nombre_archivo"] == ARCHIVOS["BMSC"])
    assert any("sin fila de origen" in e for e in bmsc["errores"]) and bmsc["filas"] == []
    assert sum(1 for r in res.values() if r["errores"]) == 1


# =========================================================== 9. reconstrucción sin releer el banco
def volcado(ruta):
    h = leer_historico(ruta)
    ws = h["ws"]
    return {"hojas": h["wb"].sheetnames, "merge": sorted(str(m) for m in ws.merged_cells.ranges),
            "tabla": h["tabla"].ref, "congelado": ws.freeze_panes,
            "anchos": {k: (v.width, v.hidden) for k, v in ws.column_dimensions.items()},
            "celdas": {(c.row, c.column): (c.value, c.number_format, c.data_type)
                       for f in ws.iter_rows() for c in f if c.value is not None}}


@pytest.fixture
def sin_banco(monkeypatch):
    """Cualquier intento de abrir un extracto bancario (.xls, fixtures) falla en la prueba."""
    abrir = builtins.open

    def vigilado(archivo, *a, **k):
        p = str(archivo)
        if p.lower().endswith(".xls") or str(EXTRACTOS) in p or any(n in p for n in FIXTURES.values()):
            raise AssertionError(f"intentó abrir el archivo bancario {p}")
        return abrir(archivo, *a, **k)

    monkeypatch.setattr(builtins, "open", vigilado)
    monkeypatch.setattr(io, "open", vigilado)
    import xlrd
    monkeypatch.setattr(xlrd, "open_workbook", lambda *a, **k: (_ for _ in ()).throw(AssertionError("xlrd")))
    import python_calamine
    monkeypatch.setattr(python_calamine.CalamineWorkbook, "from_path",
                        classmethod(lambda cls, *a, **k: (_ for _ in ()).throw(AssertionError("calamine"))))


@pytest.mark.parametrize("fuente", ["NORMALIZADO.xlsx", "LISTS.csv"])
def test_reconstruccion_sin_el_banco_es_identica_celda_a_celda(fuente, lote, hist, tmp_path, sin_banco):
    """Solo ORIGEN.xlsx + NORMALIZADO (xlsx o csv) + registro, en otra carpeta y sin poder abrir ningún extracto."""
    aislada = tmp_path / "solo_capas"
    aislada.mkdir()
    for n in ("ORIGEN.xlsx", fuente):
        shutil.copy(lote["salida"] / n, aislada / n)
    shutil.copy(REPO / "registro_bancos.json", aislada / "registro_bancos.json")
    info = hist.generar_extractos_historicos(str(aislada / "ORIGEN.xlsx"), str(aislada / fuente), str(tmp_path / "out"),
                                             str(aislada / "registro_bancos.json"))
    assert info["estado"] == "OK" and len(info["archivos"]) == len(ARCHIVOS)
    for nombre in ARCHIVOS.values():
        assert volcado(tmp_path / "out" / nombre) == volcado(lote["carpeta"] / nombre), nombre


def test_historico_no_depende_del_motor_ni_de_la_captura_ni_de_lectores_de_xls():
    src = (REPO / "historico.py").read_text(encoding="utf-8")
    codigo = "\n".join(l for l in src.splitlines() if not l.lstrip().startswith("#"))
    codigo = re.sub(r'"""[\s\S]*?"""', "", codigo)
    for prohibido in ("import xlrd", "python_calamine", "captura_origen", "motor_control_depositos_cbba",
                      "motor_generico", "import pandas"):
        assert prohibido not in codigo, prohibido


def test_generacion_determinista_y_huella_dorada(lote, hist, tmp_path):
    info = hist.generar_extractos_historicos(str(lote["salida"] / "ORIGEN.xlsx"),
                                             str(lote["salida"] / "NORMALIZADO.xlsx"), str(tmp_path))
    for a in info["archivos"]:
        assert volcado(a["ruta"]) == volcado(lote["carpeta"] / a["nombre_archivo"])
    dorada = json.loads((GOLDEN / "historico" / "MANIFEST_HISTORICO.json").read_text(encoding="utf-8"))
    assert sorted(dorada) == sorted(ARCHIVOS.values())
    for nombre, esperado in dorada.items():
        assert json.loads(json.dumps(huella_historico(lote["carpeta"] / nombre), ensure_ascii=False)) == esperado, nombre


# =========================================================== 10. sin IDs / hashes técnicos visibles
TECNICO = re.compile(r"\b(?=[0-9a-f]*[a-f])(?=[0-9a-f]*\d)[0-9a-f]{16,}\b|\|R\d+|ID_ORIGEN|ID_EXTRACTO|ID_CELDA|FILA_EXCEL|COLUMNA_EXCEL|ROL_FILA|"
                     r"CAMPO_CANONICO|LOTE DE CARGA|FECHA DE CARGA|VERSION|ARCHIVO ORIGEN|CLAVE TRANSACCI|"
                     r"^[A-Z ÁÉÍÓÚÑ]+\|[\d-]+\|\d{8}\|")


@pytest.mark.parametrize("fm", list(ARCHIVOS))
def test_ningun_id_hash_ni_campo_tecnico_es_visible(fm, hojas):
    h = hojas[fm]
    ws = h["ws"]
    ocultas = {ws.cell(h["hr"], j + 1).column_letter for j, o in enumerate(h["ocultas"]) if o}
    assert len(ocultas) == 1                                    # solo la CLAVE, oculta
    for fila in ws.iter_rows():
        for c in fila:
            if c.value is None or c.column_letter in ocultas:
                continue
            assert not TECNICO.search(str(c.value)), f"{c.coordinate}: {c.value!r}"
    assert FIXTURES[fm] not in " ".join(str(c.value) for f in ws.iter_rows() for c in f if c.value is not None
                                        and c.column_letter not in ocultas)


def test_clave_oculta_permite_sincronizar_con_lists(hojas, indep):
    for fm, h in hojas.items():
        claves = claves_de(h)
        assert len(set(claves)) == len(claves)
        assert set(claves) <= set(indep["lists"].index)


# =========================================================== 11. columna nueva / encabezado del banco
def test_columna_nueva_del_banco_aparece_antes_de_las_operativas_con_advertencia(motor, cap, hist, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja 1"
    ws.append(["Número De cuenta", "3000100152"])
    ws.append(["Fecha", "Hora", "Oficina", "Descripción", "Referencia", "Código de transacción", "ITF", "Débitos",
               "Créditos", "Saldo", "Adicionales", "Canal Nuevo"])
    ws.append(["03/08/2026", "10:00:00", "OF", "Deposito", "0000777", "X1", "0.00", None, "100.00", "1,100.00",
               "Nombre:PRUEBA", "APP-007"])
    ruta = tmp_path / "bnb_nueva.xlsx"
    wb.save(ruta)
    datos, mapa, df = capas_sinteticas(motor, cap, ruta, "BNB_MN", "bnb_nueva.xlsx")
    assert list(df.columns) == COLUMNAS_LISTS_CONTRATO                 # el operativo no cambia
    r = hist.construir_extractos_historicos(datos, mapa, df, hist.cargar_registro())[0]
    titulos = [c["titulo"] for c in r["columnas"] if not c.get("oculta")]
    assert titulos == BNB + ["Canal Nuevo"] + OPERATIVAS
    assert r["filas"][0][titulos.index("Canal Nuevo")] == "APP-007"
    assert r["filas"][0][titulos.index("Referencia")] == "0000777"
    assert any("Canal Nuevo" in a and "columna nueva" in a for a in r["advertencias"]) and not r["errores"]


# =========================================================== 12. UNION_ME (estructura confirmada; movimientos pendientes)
def test_union_me_proxy_sintetico_con_movimiento_muestra_nro_de_verificasion(motor, cap, hist, tmp_path):
    """PROXY SINTÉTICO (no es un extracto real): prueba solo la mecánica estructural de la capa 4."""
    ruta = crear_xlsx_union_me_con_movimiento(tmp_path / "u.xlsx")
    datos, mapa, df = capas_sinteticas(motor, cap, ruta, "UNION_ME", "u.xlsx")
    assert "V-0001" not in " ".join(str(x) for x in df.iloc[0])      # el operativo (26 columnas) no la lleva
    res = hist.construir_extractos_historicos(datos, mapa, df, hist.cargar_registro())
    # el proxy agrega un movimiento de 100.00 pero conserva el pie "Total Créditos: 0.00": se bloquea
    assert len(res) == 1 and any("Total créditos" in e and "no cuadra" in e for e in res[0]["errores"])
    from openpyxl import load_workbook
    wb = load_workbook(ruta)
    wb.active["I18"] = "100.00"                                     # pie coherente con el movimiento sintético
    wb.save(ruta)
    datos, mapa, df = capas_sinteticas(motor, cap, ruta, "UNION_ME", "u.xlsx")
    res = hist.construir_extractos_historicos(datos, mapa, df, hist.cargar_registro())
    assert len(res) == 1 and not res[0]["errores"], res[0]["errores"]
    r = res[0]
    assert r["nombre_archivo"] == "EXTRACTO_HISTORICO_BANCO_UNION_20000003224544_2026-08.xlsx"
    hist.escribir_extracto_historico(r, str(tmp_path / r["nombre_archivo"]))
    h = leer_historico(tmp_path / r["nombre_archivo"])
    assert h["visibles"] == UNION_ME + OPERATIVAS
    i = h["columnas"].index("Nro de verificasion")
    assert h["filas"][0][i] == "V-0001" and h["celdas"][0][i].data_type == "s"
    assert h["zona"]["Moneda"] == "USD" and h["zona"]["Producto"] == "UNICUENTA ESPECIAL PERSONA JURIDICA M/E"


def test_union_me_vacio_no_genera_archivo_ni_bloquea(motor, cap, hist, tmp_path):
    """PROXY SINTÉTICO vacío: sin movimientos no hay archivo del mes y nada falla (no se espera muestra real)."""
    ruta = crear_xlsx_union_me_vacio(tmp_path / "u.xlsx")
    contrato = {"hojas_validas": motor.HOJAS_VALIDAS, "encabezados_esperados": motor.ENCABEZADOS_ESPERADOS,
                "encontrar_fila_encabezado": motor.encontrar_fila_encabezado}
    r = cap.capturar_extracto(str(ruta), "u.xlsx", "UNION_ME", contrato, [],
                              tabla=pd.DataFrame(columns=COLUMNAS_LISTS_CONTRATO))
    datos = [dict(zip(cap.COLS_DATOS, f)) for f in r["datos"]]
    assert hist.construir_extractos_historicos(datos, [], [], hist.cargar_registro()) == []
    assert hist.extractos_sin_movimientos(datos, []) == ["u.xlsx"]


@pytest.mark.skip(reason=PENDIENTE_MOV + ": EXTRACTO_HISTORICO real de UNION_ME (columnas, Nro de verificasion real, "
                         "orden, saldos y totales de pie)")
def test_union_me_real_historico_con_movimientos():
    ...


@pytest.mark.skip(reason=PENDIENTE_MOV + ": débitos reales de UNION_ME en el histórico (Monto negativo, AG, "
                         "Nro Documento, Nro de verificasion)")
def test_union_me_real_historico_debitos():
    ...


# =========================================================== 13. NORMALIZADO.xlsx y LISTS.csv sin cambios
def test_generar_historicos_no_modifica_normalizado_lists_ni_origen(lote):
    assert lote["antes"] == lote["despues"]


def test_normalizado_y_lists_siguen_identicos_a_sus_doradas(lote):
    salida = lote["salida"]
    pd.testing.assert_frame_equal(leer_csv_texto(salida / "LISTS.csv"), leer_csv_texto(GOLDEN / "LOTE_12_LISTS.csv"))
    for hoja in ("LISTS", "VALIDACION", "RESUMEN", "DIAGNOSTICO"):
        d = pd.read_excel(salida / "NORMALIZADO.xlsx", sheet_name=hoja, dtype=str, keep_default_na=False)
        d = d.drop(columns=[c for c in ("LOTE DE CARGA", "FECHA DE CARGA") if c in d.columns])
        assert d.to_csv(index=False, lineterminator="\n") == \
            (GOLDEN / f"LOTE_12_NORMALIZADO__{hoja}.csv").read_text(encoding="utf-8")


def test_el_motor_no_cambio_y_no_escribe_historicos(corrida_lote):
    """P3b no toca el motor: la carpeta productiva conserva exactamente sus archivos y el retorno sus claves."""
    res, salida = corrida_lote
    src = MOTOR_PATH.read_text(encoding="utf-8")
    assert "historico.py" not in src and "historico_estado" not in src
    assert not any(p.name.startswith("EXTRACTO_HISTORICO") for p in salida.iterdir())
    assert "historico_estado" not in res


# =========================================================== 14. CLI
def test_cli_genera_desde_origen_y_lists_csv(lote, tmp_path):
    p = subprocess.run([sys.executable, "-W", "ignore", str(REPO / "historico.py"), str(lote["salida"] / "ORIGEN.xlsx"),
                        str(lote["salida"] / "LISTS.csv"), str(tmp_path)], capture_output=True, text=True,
                       encoding="utf-8")
    assert p.returncode == 0, p.stdout + p.stderr
    assert "Estado: OK · 11 archivo(s)" in p.stdout and f"SIN MOVIMIENTOS {FIXTURES['BISA_ME']}" in p.stdout
    assert sorted(x.name for x in tmp_path.iterdir()) == sorted(ARCHIVOS.values())
