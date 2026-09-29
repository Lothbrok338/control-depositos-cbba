"""Corrida completa (ejecutar_motor) con los 12 extractos validos y sus salidas."""
import pandas as pd
import pytest
from openpyxl import load_workbook
from helpers import COLUMNAS_LISTS_CONTRATO, GOLDEN, leer_csv_texto

pytestmark = pytest.mark.regresion


def test_lote_csv_igual_a_dorada(corrida_lote):
    _, salida = corrida_lote
    actual = leer_csv_texto(salida / "LISTS.csv")
    esperado = leer_csv_texto(GOLDEN / "LOTE_12_LISTS.csv")
    pd.testing.assert_frame_equal(actual, esperado)


def test_lote_csv_utf8_bom_y_columnas(corrida_lote):
    _, salida = corrida_lote
    raw = (salida / "LISTS.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    assert list(pd.read_csv(salida / "LISTS.csv", encoding="utf-8-sig", nrows=1).columns) == COLUMNAS_LISTS_CONTRATO


def test_lote_xlsx_hojas_tabla_y_formatos(corrida_lote):
    res, salida = corrida_lote
    wb = load_workbook(salida / "NORMALIZADO.xlsx")
    assert wb.sheetnames == ["LISTS", "VALIDACION", "RESUMEN", "DIAGNOSTICO"]
    ws = wb["LISTS"]
    assert [c.value for c in ws[1]] == COLUMNAS_LISTS_CONTRATO
    assert "tblLISTS" in ws.tables
    assert ws.tables["tblLISTS"].ref == f"A1:Z{ws.max_row}"
    enc = {c.value: c.column for c in ws[1]}
    assert ws.cell(2, enc["FECHA MOVIMIENTO"]).number_format == "mm/dd/yy"
    assert ws.cell(2, enc["IMPORTE"]).number_format == "0.00"
    assert ws.cell(2, enc["CUENTA BANCARIA"]).number_format == "@"
    assert isinstance(ws.cell(2, enc["CUENTA BANCARIA"]).value, str)


def test_lote_totales_y_orden(corrida_lote):
    res, _ = corrida_lote
    df = res["df_final"]
    assert len(df) == 4064  # 8+2331+0+1+95+7+1184+14+388+1+3+32
    assert df["FECHA MOVIMIENTO"].is_monotonic_increasing
    assert df["CUENTA BANCARIA"].nunique() == 11  # BISA_ME sin movimientos
    assert set(res["df_validacion"]["ESTADO"]) == {"OK"}
    assert len(res["df_validacion"]) == 12


def test_xlsx_y_csv_tienen_mismas_filas_y_claves(corrida_lote):
    _, salida = corrida_lote
    x = pd.read_excel(salida / "NORMALIZADO.xlsx", sheet_name="LISTS", dtype=str)
    c = pd.read_csv(salida / "LISTS.csv", encoding="utf-8-sig", dtype=str)
    assert len(x) == len(c)
    assert x["CLAVE TRANSACCIÓN"].tolist() == c["CLAVE TRANSACCIÓN"].tolist()


def test_lote_bloquea_formato_repetido(motor, tmp_path):
    import shutil
    from helpers import EXTRACTOS, FIXTURES
    ent = tmp_path / "in"; ent.mkdir()
    shutil.copy(EXTRACTOS / FIXTURES["BNB_ME"], ent / "a.xls")
    shutil.copy(EXTRACTOS / FIXTURES["BNB_ME"], ent / "b.xls")
    with pytest.raises(ValueError, match="misma"):
        motor.ejecutar_motor(str(ent), str(tmp_path / "o" / "NORMALIZADO.xlsx"))


def test_lote_ignora_archivos_normalizado_y_no_excel(motor):
    import os
    from helpers import EXTRACTOS
    d = motor.descubrir_archivos(str(EXTRACTOS))
    assert len(d) == 12 and all(n.lower().endswith((".xls", ".xlsx")) for n in d)


@pytest.mark.parametrize("nombre", [
    "NORMALIZADO.xlsx", "normalizado.XLSX", "NORMALIZADO_2026.xlsx",
    "ORIGEN.xlsx", "origen.xlsx", "ORIGEN_2026-08-21.xlsx", "ORIGEN (1).xlsx",
    "EXTRACTO_HISTORICO_BNB_3000100152_2026-08.xlsx", "extracto_historico_x.XLSX",
])
def test_d20_archivos_del_sistema_no_son_extractos(motor, tmp_path, nombre):
    """D-20 (corregido): las salidas del propio sistema nunca se descubren como extractos de entrada;
    los extractos reales siguen descubriendose igual."""
    import shutil
    from helpers import EXTRACTOS, FIXTURES
    for fm in ("BNB_ME", "ECO_CTA_CTE"):
        shutil.copy(EXTRACTOS / FIXTURES[fm], tmp_path / FIXTURES[fm])
    (tmp_path / nombre).write_bytes(b"PK")
    assert list(motor.descubrir_archivos(str(tmp_path))) == sorted(FIXTURES[f] for f in ("BNB_ME", "ECO_CTA_CTE"))


def test_d20_extractos_reales_con_nombres_parecidos_siguen_incluidos(motor, tmp_path):
    """No se sobre-excluye: solo prefijos exactos ORIGEN / EXTRACTO_HISTORICO_ y solo .xlsx."""
    import shutil
    from helpers import EXTRACTOS, FIXTURES
    for n in ("extracto_bnb.xls", "MI_ORIGEN.xlsx", "Extracto historico bcp.xlsx", "ORIGEN_banco.xls"):
        shutil.copy(EXTRACTOS / FIXTURES["BNB_ME"], tmp_path / n)
    assert set(motor.descubrir_archivos(str(tmp_path))) == {"extracto_bnb.xls", "MI_ORIGEN.xlsx", "Extracto historico bcp.xlsx", "ORIGEN_banco.xls"}


def test_d20_corrida_completa_con_salidas_en_la_carpeta_de_entrada(motor, tmp_path):
    """Escenario real de D-20: salida == entrada y se ejecuta dos veces; la segunda no se bloquea."""
    import shutil
    from helpers import EXTRACTOS, FIXTURES
    for fm in ("BNB_ME", "ECO_CTA_CTE"):
        shutil.copy(EXTRACTOS / FIXTURES[fm], tmp_path / FIXTURES[fm])
    r1 = motor.ejecutar_motor(str(tmp_path), str(tmp_path / "NORMALIZADO.xlsx"))
    assert (tmp_path / "ORIGEN.xlsx").exists() and r1["origen_estado"]["estado"] == "OK"
    r2 = motor.ejecutar_motor(str(tmp_path), str(tmp_path / "NORMALIZADO.xlsx"))
    assert len(r2["df_final"]) == len(r1["df_final"])
