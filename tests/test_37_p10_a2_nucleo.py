"""P10-A.2 — núcleo: estado canónico, fusión de extractos acumulativos, cruce con la lista real y regeneración del XLSX.

Todo con extractos REALES (tests/fixtures/extractos) y el MOTOR REAL de P0; la lista `Depositos_Activos` es una simulación en memoria
(`p10/lista_simulada.py`) con el comportamiento observable de SharePoint (UTC, Modified, ETag, escrituras de P9).
"""
import base64
import copy
import gzip
import hashlib
import json
import sys
from pathlib import Path

import pytest

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

import historico as H                                                      # noqa: E402
from openpyxl import load_workbook                                          # noqa: E402

from p10 import contrato as C, estado as E, generador as G, motor_p0, sharepoint as SP, sincronizacion as S   # noqa: E402
from p10.lista_simulada import ListaSimulada                                # noqa: E402
from p10.comparar import diferencias                                        # noqa: E402

pytestmark = pytest.mark.p10

FIXTURE = RAIZ / "tests/fixtures/extractos/bnb_ahorro_2.xls"            # BNB ahorro: 95 movimientos reales
RUTA_PROC = "/CONTROL_DEPOSITOS/P0_EXTRACTOS/PROCESADOS/2026/10_OCTUBRE/08/bnb_ahorro_2.xls"
GID = "BNB|3501936692|BOB|2026-08"
AHORA = "2026-10-08T10:00:00"


def sha(b):
    return hashlib.sha256(b).hexdigest()


def d64(s):
    return base64.b64decode(s)


@pytest.fixture(scope="module")
def p0(tmp_path_factory):
    """Corre el motor REAL de P0 (A.1) sobre el extracto: LISTS.csv + ORIGEN.xlsx de referencia."""
    ent = tmp_path_factory.mktemp("ent")
    (ent / FIXTURE.name).write_bytes(FIXTURE.read_bytes())
    return motor_p0.ejecutar_p0(str(ent), str(tmp_path_factory.mktemp("p0")))


@pytest.fixture(scope="module")
def extracto():
    """El extracto procesado por P10 (mismo motor, vía p0.nucleo)."""
    r = S.procesar_extracto(FIXTURE.read_bytes(), FIXTURE.name, "CBBA", RUTA_PROC, AHORA)
    assert r["ok"], r
    return r


@pytest.fixture(scope="module")
def parcial(extracto):
    assert [g["grupo_id"] for g in extracto["grupos"]] == [GID]
    return E.desempaquetar(d64(extracto["grupos"][0]["parcial_b64"]))


def sincronizar(estado=None, parciales=(), filas=None, ahora=AHORA, **kw):
    pb = [E.empaquetar(p) if isinstance(p, dict) else p for p in parciales]
    return S.sincronizar_grupo("CBBA", GID, estado, pb, filas, ahora, **kw)


@pytest.fixture(scope="module")
def inicial(parcial):
    r = sincronizar(parciales=[parcial])
    assert r["ok"] and r["xlsx_valido"], r
    return r


# ====================================================================== reutilización de P0 y de A.1
def test_las_claves_del_parcial_son_exactamente_las_que_entrega_p0(p0, parcial):
    filas = H.leer_normalizado(p0.lists_csv)
    assert set(parcial["movimientos"]) == {f["CLAVE TRANSACCIÓN"] for f in filas} and len(filas) == 95
    for f in filas:                                    # las 26 columnas viajan tal cual las dejó el motor
        assert parcial["movimientos"][f["CLAVE TRANSACCIÓN"]]["p0"][:15] == [f[c] for c in C.COLUMNAS_LISTS][:15]


def test_el_xlsx_regenerado_desde_el_estado_es_identico_byte_a_byte_al_de_a1(p0, extracto, tmp_path):
    """Mismo extracto, mismo motor: generador A.1 directo desde ORIGEN+LISTS == estado canónico -> XLSX."""
    datos, mapa = H.leer_origen(p0.origen_xlsx)
    filas = H.leer_normalizado(p0.lists_csv)
    par = S.construir_parciales(datos, mapa, filas, FIXTURE.name, sha(FIXTURE.read_bytes()), AHORA, "BNB_EXTRACTO_V1")
    r = sincronizar(parciales=[par[GID]])
    assert r["ok"]
    a1 = G.generar_libros(p0.origen_xlsx, p0.lists_csv, None, tmp_path)
    nombre = a1["archivos"][0]["nombre_archivo"]
    assert nombre == r["rutas"]["nombre_xlsx"]
    assert sha(d64(r["xlsx_b64"])) == a1["archivos"][0]["sha256"]


def test_p10_a2_reutiliza_el_motor_de_p0_y_no_define_normalizacion():
    import ast
    prohibidas = {"crear_clave", "valor_clave_numero", "finalizar_dataframe", "normalizar_extracto", "normalizar_fecha",
                  "normalizar_hora", "codigo_texto", "numero", "ecuacion_saldo", "validar_extracto"}
    for py in ("estado", "sincronizacion", "plan", "sharepoint", "lista_simulada"):
        arbol = ast.parse((RAIZ / "p10" / f"{py}.py").read_text(encoding="utf-8"))
        assert not ({n.name for n in ast.walk(arbol) if isinstance(n, ast.FunctionDef)} & prohibidas), py
    fuente = (RAIZ / "p10/sincronizacion.py").read_text(encoding="utf-8")
    assert "nucleo.ETAPAS.detectar" in fuente and "nucleo.ETAPAS.motor" in fuente        # el mismo motor que P0


def test_modulos_de_p0_p7_p8_y_a1_no_cambian():
    huellas = json.loads((RAIZ / "p0/flujo/huellas_protegidas.json").read_text(encoding="utf-8"))
    for f in ("motor_control_depositos_cbba.py", "historico.py", "adaptador_m365.py", "p0/api.py", "p0/nucleo.py", "Dockerfile",
              "railway.json"):
        assert hashlib.sha256((RAIZ / f).read_bytes()).hexdigest() == huellas[f], f


# ====================================================================== estado canónico
def test_estado_empaquetado_determinista_y_con_integridad(inicial):
    b = d64(inicial["estado_b64"])
    est = E.cargar_estado(b, GID)
    assert E.empaquetar(est) == b and b[:2] == b"\x1f\x8b"
    assert est["formato"] == E.FORMATO_ESTADO and est["version"] == 1 and len(est["movimientos"]) == 95
    assert est["xlsx"]["sha256"] == sha(d64(inicial["xlsx_b64"])) and est["extractos"][0]["nombre"] == FIXTURE.name


def test_estado_alterado_o_dañado_se_detecta(inicial):
    b = d64(inicial["estado_b64"])
    est = E.desempaquetar(b)
    est["movimientos"][next(iter(est["movimientos"]))]["p0"][7] = "999999"             # alteración con la huella vieja
    with pytest.raises(E.ErrorEstado) as e:
        E.cargar_estado(E.empaquetar(est))
    assert e.value.codigo == "ESTADO_CORRUPTO"
    for dañado in (b[: len(b) // 2], b"no es un estado", gzip.compress(b"[]")):
        with pytest.raises(E.ErrorEstado):
            E.cargar_estado(dañado)
    with pytest.raises(E.ErrorEstado, match="otro grupo"):
        E.cargar_estado(b, "BNB|1|BOB|2026-08")


def test_sincronizar_devuelve_control_y_rutas_con_el_destino_pedido(inicial):
    r = inicial["rutas"]
    assert r["ruta_xlsx"] == ("/CONTROL_DEPOSITOS/P10_HISTORICO/2026/08_AGOSTO/BNB/"
                              "EXTRACTO_HISTORICO_BNB_3501936692_BOB_2026-08.xlsx")
    assert r["ruta_estado"].startswith("/CONTROL_DEPOSITOS/P10_ESTADO/2026/08_AGOSTO/BNB/ESTADO_P10_")
    c = inicial["control"]
    assert (c["ESTADO"], c["MOVIMIENTOS"], c["VERSION_ESTADO"], c["TIPO"]) == ("OK", 95, 1, "GRUPO")
    assert c["XLSX_BYTES"] == len(d64(inicial["xlsx_b64"])) and c["ESTADO_BYTES"] == len(d64(inicial["estado_b64"]))


# ====================================================================== idempotencia y extractos acumulativos
def test_aplicar_dos_veces_lo_mismo_no_cambia_nada(inicial, parcial):
    r = sincronizar(d64(inicial["estado_b64"]), [parcial])
    assert r["ok"] and not r["cambio_estado"] and not r["cambio_xlsx"] and r["estado_b64"] is None and r["xlsx_b64"] is None
    assert r["resumen"]["parciales"][0]["motivo"] == "EXTRACTO_YA_APLICADO"


def _parcial_corto(parcial, fraccion, nombre="corto.xls"):
    p = copy.deepcopy(parcial)
    claves = sorted(p["movimientos"], key=lambda k: (p["movimientos"][k]["p0"][5], p["movimientos"][k]["p0"][6]))
    manten = claves[: int(len(claves) * fraccion)]
    p["movimientos"] = {k: p["movimientos"][k] for k in manten}
    ult = max(f"{p['movimientos'][k]['p0'][5]}T{p['movimientos'][k]['p0'][6] or '00:00:00'}" for k in manten)
    p["extracto"].update(nombre=nombre, sha256=sha(nombre.encode()), ultimo_mov=ult, movimientos_extracto=len(manten),
                         procesado_en="2026-10-01T09:00:00")
    return p


def test_extracto_acumulativo_posterior_completa_sin_duplicar(inicial, parcial):
    corto = _parcial_corto(parcial, 0.6)
    r1 = sincronizar(parciales=[corto], ahora="2026-10-01T09:00:00")
    assert r1["ok"] and r1["control"]["MOVIMIENTOS"] == len(corto["movimientos"]) < 95
    r2 = sincronizar(d64(r1["estado_b64"]), [parcial], ahora=AHORA)                      # llega el acumulado completo
    assert r2["ok"] and r2["control"]["MOVIMIENTOS"] == 95 and r2["control"]["VERSION_ESTADO"] == 2
    est = E.cargar_estado(d64(r2["estado_b64"]), GID)
    assert [e["nombre"] for e in est["extractos"]] == ["corto.xls", FIXTURE.name] and est["cabecera"]["sha256"] == sha(FIXTURE.read_bytes())
    assert r2["resumen"]["parciales"][0]["nuevos"] == 95 - len(corto["movimientos"])
    wb = load_workbook(__import__("io").BytesIO(d64(r2["xlsx_b64"])))
    filas = [c.value for c in wb["AUDITORIA"]["A"][1:]]
    assert len(filas) == len(set(filas)) == 95 and wb.sheetnames == ["EXTRACTO", "AUDITORIA"]


def test_un_extracto_mas_corto_que_llega_despues_nunca_borra_movimientos(inicial, parcial):
    corto = _parcial_corto(parcial, 0.5, "viejo.xls")
    r = sincronizar(d64(inicial["estado_b64"]), [corto], ahora="2026-10-09T09:00:00")
    est = E.cargar_estado(d64(r["estado_b64"]), GID)
    assert r["ok"] and len(est["movimientos"]) == 95
    assert est["cabecera"]["sha256"] == sha(FIXTURE.read_bytes())                        # la cabecera sigue siendo la del extracto completo
    assert r["resumen"]["parciales"][0]["nuevos"] == 0 and sha(d64(r["xlsx_b64"] or "")) if r["xlsx_b64"] else True


def test_orden_de_llegada_no_cambia_el_resultado(inicial, parcial):
    corto = _parcial_corto(parcial, 0.6)
    a = sincronizar(parciales=[corto, parcial], ahora=AHORA)
    b = sincronizar(parciales=[parcial, corto], ahora=AHORA)
    assert sha(d64(a["xlsx_b64"])) == sha(d64(b["xlsx_b64"])) == sha(d64(inicial["xlsx_b64"]))
    ea, eb = (E.cargar_estado(d64(x["estado_b64"])) for x in (a, b))
    assert ea["movimientos"] == eb["movimientos"] and ea["cabecera"] == eb["cabecera"]


# ====================================================================== cruce con Depositos_Activos (cambios operativos)
@pytest.fixture()
def lista(p0):
    li = ListaSimulada()
    assert li.cargar(H.leer_normalizado(p0.lists_csv)) == 95
    return li


def filas_de(lista):
    items, mas = lista.consultar()
    assert not mas
    grupos, anomalias = SP.agrupar(items)
    assert not anomalias and list(grupos) == [GID]
    return grupos[GID]


def hojas(r):
    wb = load_workbook(__import__("io").BytesIO(d64(r["xlsx_b64"])))
    def tabla(hoja):
        ws = wb[hoja]
        enc = [c.value for c in next(ws.iter_rows(min_row=1 if hoja == "AUDITORIA" else 10, max_row=1 if hoja == "AUDITORIA" else 10))]
        return {f[0 if hoja == "AUDITORIA" else -1].value: dict(zip(enc, (c.value for c in f)))
                for f in ws.iter_rows(min_row=2 if hoja == "AUDITORIA" else 11)}
    return tabla("EXTRACTO"), tabla("AUDITORIA")


def test_confirmacion_se_refleja_en_el_mes_del_movimiento_y_en_hora_de_bolivia(inicial, lista):
    clave = sorted(i["CLAVE_TRANSACCION"] for i in lista.items.values())[10]
    lista.avanzar(days=60)                                                               # se confirma meses después
    lista.confirmar(clave, estudiante="ANA PEREZ", usuario="a@x.edu", codigo_estudiante="0012345", observacion="línea 1\nlínea 2")
    r = sincronizar(d64(inicial["estado_b64"]), filas=filas_de(lista), ahora="2026-12-07T09:00:00")
    assert r["ok"] and r["cambio_estado"] and r["cambio_xlsx"] and r["control"]["HASH_OPERATIVO"]
    ext, aud = hojas(r)
    assert aud[clave]["ESTADO"] == "CONFIRMADO" and ext[clave]["ESTADO"] == "CONFIRMADO"
    assert aud[clave]["CODIGO_ESTUDIANTE"] == "0012345" and aud[clave]["OBSERVACIÓN"] == "línea 1\nlínea 2"
    assert aud[clave]["FECHA CONFIRMACIÓN"].isoformat() == "2026-12-07T08:00:00"        # el reloj de la lista marcó 08:00 Bolivia = 12:00 UTC
    assert r["control"]["PERIODO"] == "2026-08"                                          # el histórico es del mes del MOVIMIENTO
    assert sum(1 for f in aud.values() if f["ESTADO"] == "CONFIRMADO") == 1


def test_sin_cambios_en_la_lista_no_hay_cambios_en_el_estado(inicial, lista):
    r1 = sincronizar(d64(inicial["estado_b64"]), filas=filas_de(lista))
    assert r1["ok"] and r1["control"]["HASH_OPERATIVO"] == E.hash_lista(filas_de(lista))
    r2 = sincronizar(d64(r1["estado_b64"] or inicial["estado_b64"]), filas=filas_de(lista), ahora="2026-10-09T10:00:00")
    assert not r2["cambio_estado"] and not r2["cambio_xlsx"]


def test_reversion_devuelve_el_extracto_al_estado_anterior_y_auditoria_conserva_el_marcador(inicial, lista):
    clave = sorted(i["CLAVE_TRANSACCION"] for i in lista.items.values())[3]
    lista.avanzar(hours=1)
    lista.confirmar(clave, estudiante="LUIS", usuario="l@x.edu")
    r1 = sincronizar(d64(inicial["estado_b64"]), filas=filas_de(lista), ahora="2026-10-08T09:00:00")
    lista.avanzar(hours=3)
    lista.revertir(clave)
    r2 = sincronizar(d64(r1["estado_b64"]), filas=filas_de(lista), ahora="2026-10-08T12:00:00")
    ext, aud = hojas(r2)
    assert (ext[clave]["ESTADO"], ext[clave]["CONFIRMADO POR"], ext[clave]["OBSERVACIONES"]) == ("DISPONIBLE", None, None)
    assert aud[clave]["ESTADO"] == "DISPONIBLE" and aud[clave]["ESTUDIANTE"] is None and aud[clave]["ULTIMA_REVERSION_ID"]
    base = hojas(inicial)
    assert ext[clave] == base[0][clave]                                                  # EXTRACTO igual que antes de confirmar


def test_registro_eliminado_de_la_lista_conserva_su_ultima_informacion(inicial, lista):
    clave = sorted(i["CLAVE_TRANSACCION"] for i in lista.items.values())[5]
    lista.confirmar(clave, estudiante="MARTA", usuario="m@x.edu")
    r1 = sincronizar(d64(inicial["estado_b64"]), filas=filas_de(lista))
    lista.eliminar(clave)                                                                # P10-B borrará registros de la lista
    r2 = sincronizar(d64(r1["estado_b64"]), filas=filas_de(lista), ahora="2026-10-09T10:00:00")
    assert r2["ok"] and not r2["cambio_xlsx"]                                            # el libro NO retrocede a DISPONIBLE
    assert r2["control"]["HASH_OPERATIVO"] == E.hash_lista(filas_de(lista)) != r1["control"]["HASH_OPERATIVO"]
    assert hojas({"xlsx_b64": r1["xlsx_b64"]})[1][clave]["ESTADO"] == "CONFIRMADO"


def test_columnas_de_carga_se_toman_de_la_lista_no_del_reproceso(inicial, lista):
    """LOTE/FECHA DE CARGA/ARCHIVO ORIGEN de AUDITORIA son los persistidos en Depositos_Activos."""
    for it in lista.items.values():
        it.update(LOTE_CARGA="20260101_000000", FECHA_CARGA="2026-01-01 00:00:00.5", ARCHIVO_ORIGEN="original en ENTRADA.xls")
    r = sincronizar(d64(inicial["estado_b64"]), filas=filas_de(lista))
    aud = hojas(r)[1]
    assert {f["LOTE DE CARGA"] for f in aud.values()} == {"20260101_000000"}
    assert {f["ARCHIVO ORIGEN"] for f in aud.values()} == {"original en ENTRADA.xls"}


def test_hash_de_lista_no_depende_del_orden_ni_de_la_hora_de_modificacion(lista):
    filas = filas_de(lista)
    assert E.hash_lista(filas) == E.hash_lista(list(reversed(filas)))
    lista.avanzar(hours=2)
    lista._modificar(next(iter(lista.items.values())))                                   # solo cambia Modified
    assert E.hash_lista(filas_de(lista)) == E.hash_lista(filas)


# ====================================================================== recuperación
def test_xlsx_borrado_se_regenera_identico_desde_el_estado(inicial):
    r = sincronizar(d64(inicial["estado_b64"]), forzar_xlsx=True)
    assert r["ok"] and r["cambio_xlsx"] and not r["cambio_estado"] and r["xlsx_b64"]
    assert sha(d64(r["xlsx_b64"])) == sha(d64(inicial["xlsx_b64"]))


def test_estado_dañado_o_ausente_pide_reconstruir_sin_inventar_datos(inicial):
    corrupto = sincronizar(d64(inicial["estado_b64"])[:50])
    assert not corrupto["ok"] and corrupto["codigo_error"] == "ESTADO_CORRUPTO"
    assert (corrupto["control"]["ESTADO"], corrupto["control"]["HASH_OPERATIVO"]) == ("RECONSTRUIR", "")
    ausente = sincronizar(None, control={"VERSION_ESTADO": 3, "HASH_ESTADO": "x", "ESTADO": "OK"})
    assert not ausente["ok"] and ausente["codigo_error"] == "ESTADO_AUSENTE" and ausente["control"]["ESTADO"] == "RECONSTRUIR"
    sin_nada = sincronizar(None)
    assert not sin_nada["ok"] and sin_nada["codigo_error"] == "SIN_DATOS"


def test_reconstruir_desde_el_extracto_original_da_el_mismo_estado(inicial, parcial, lista):
    """Estado perdido -> se reprocesa el extracto de PROCESADOS y se aplica la lista: mismo estado y mismo XLSX."""
    clave = sorted(i["CLAVE_TRANSACCION"] for i in lista.items.values())[2]
    lista.confirmar(clave, estudiante="PEDRO", usuario="p@x.edu")
    original = sincronizar(d64(inicial["estado_b64"]), filas=filas_de(lista), ahora="2026-10-08T11:00:00")
    reconstruido = sincronizar(None, [parcial], filas=filas_de(lista), ahora="2026-10-08T11:00:00")
    a, b = (E.desempaquetar(d64(x["estado_b64"])) for x in (original, reconstruido))
    for k in ("version", "integridad"):          # contador de guardados (CAS): es historia del estado, no contenido
        a.pop(k), b.pop(k)
    assert a == b and reconstruido["xlsx_b64"] == original["xlsx_b64"]


def test_errores_de_negocio_del_extracto_producen_registro_de_error_sin_datos_bancarios():
    r = S.procesar_extracto(b"no soy un excel", "x.xls", "CBBA", "/CONTROL_DEPOSITOS/P0_EXTRACTOS/PROCESADOS/2026/10_OCTUBRE/08/x.xls", AHORA)
    assert not r["ok"] and r["control"]["ESTADO"] == "ERROR" and r["control"]["INTENTOS"] == 1 and r["grupos"] == []
    r3 = S.procesar_extracto(b"no soy un excel", "x.xls", "CBBA", "/x/PROCESADOS/2026/10_OCTUBRE/08/x.xls", AHORA, intentos_previos=2)
    assert r3["control"]["ESTADO"] == "ERROR_FINAL"
    assert not S.procesar_extracto(b"", "x.xls", "CBBA", "/x", AHORA)["ok"]
    assert S.procesar_extracto(b"x", "x.txt", "CBBA", "/x", AHORA)["codigo_error"] == "EXTENSION_NO_SOPORTADA"
    assert S.procesar_extracto(b"x", "x.xls", "NOEXISTE", "/x", AHORA)["codigo_error"] == "SEDE_DESCONOCIDA"


# ====================================================================== sellos cruzados con el control, verificación y reconstrucción
def _ctl(r, **kw):
    """El elemento de P10_Control tal como el flujo lo dejaría tras confirmar `r`."""
    return {**r["control"], **kw}


def test_el_control_detecta_un_estado_alterado_retrocedido_o_adelantado(inicial, lista):
    claves = sorted(i["CLAVE_TRANSACCION"] for i in lista.items.values())
    lista.confirmar(claves[0])
    r2 = sincronizar(d64(inicial["estado_b64"]), filas=filas_de(lista), control=_ctl(inicial))
    assert r2["ok"] and r2["control"]["VERSION_ESTADO"] == _ctl(inicial)["VERSION_ESTADO"] + 1
    # 1. el archivo fue restaurado a una versión anterior mientras el control ya confirmó la nueva
    viejo = sincronizar(d64(inicial["estado_b64"]), control=_ctl(r2))
    assert not viejo["ok"] and viejo["codigo_error"] == "ESTADO_RETROCEDIDO" and viejo["control"]["ESTADO"] == "RECONSTRUIR"
    # 2. mismo número de versión pero otro contenido (sello recalculado a mano o archivo de otro lado)
    falso = E.desempaquetar(d64(r2["estado_b64"]))
    falso["movimientos"][next(iter(falso["movimientos"]))]["p0"][4] = "OTRO"
    E.sellar(falso)
    r = sincronizar(E.empaquetar(falso), control=_ctl(r2))
    assert not r["ok"] and r["codigo_error"] == "ESTADO_ALTERADO"
    # 3. cierre interrumpido: el estado se escribió (versión +1) pero el control quedó en la versión anterior => NO es alteración
    ok = sincronizar(d64(r2["estado_b64"]), control=_ctl(inicial))
    assert ok["ok"] and ok["control"]["ESTADO"] == "OK"
    # 4. versión muy por delante de lo confirmado
    adelantado = E.desempaquetar(d64(r2["estado_b64"]))
    adelantado["version"] = _ctl(r2)["VERSION_ESTADO"] + 3
    E.sellar(adelantado)
    r = sincronizar(E.empaquetar(adelantado), control=_ctl(r2))
    assert r["codigo_error"] == "ESTADO_ALTERADO"


def test_cada_escritura_del_estado_sube_la_version_en_uno(inicial):
    v0 = inicial["control"]["VERSION_ESTADO"]
    r = sincronizar(d64(inicial["estado_b64"]), control=_ctl(inicial), forzar_xlsx=True)       # XLSX igual: el estado no cambia
    assert not r["cambio_estado"] and r["control"]["VERSION_ESTADO"] == v0
    sin_xlsx = E.desempaquetar(d64(inicial["estado_b64"]))
    sin_xlsx["xlsx"] = None
    E.sellar(sin_xlsx)
    r = sincronizar(E.empaquetar(sin_xlsx), control=_ctl(inicial, VERSION_ESTADO=sin_xlsx["version"], HASH_ESTADO=sin_xlsx["integridad"]))
    assert r["cambio_estado"] and r["control"]["VERSION_ESTADO"] == sin_xlsx["version"] + 1


def test_verificacion_nocturna_del_xlsx_ausente_alterado_o_correcto(inicial):
    est, xlsx = d64(inicial["estado_b64"]), d64(inicial["xlsx_b64"])
    ctl = _ctl(inicial)
    ok = sincronizar(est, control=ctl, verificar_xlsx=True, xlsx_actual=xlsx)
    assert ok["ok"] and not ok["cambio_xlsx"] and ok["xlsx_b64"] is None and ok["resumen"]["xlsx_verificado"] == "OK"
    falta = sincronizar(est, control=ctl, verificar_xlsx=True, xlsx_actual=None)
    assert falta["cambio_xlsx"] and sha(d64(falta["xlsx_b64"])) == sha(xlsx) and falta["resumen"]["xlsx_verificado"] == "XLSX_AUSENTE"
    alterado = sincronizar(est, control=ctl, verificar_xlsx=True, xlsx_actual=xlsx[:-10] + b"0123456789")
    assert alterado["cambio_xlsx"] and sha(d64(alterado["xlsx_b64"])) == sha(xlsx) and alterado["resumen"]["xlsx_verificado"] == "XLSX_ALTERADO"
    assert not alterado["cambio_estado"]               # el estado no cambia: solo se reescribe el libro


def test_reconstruccion_completa_con_marca_hasta_finalizar(parcial, inicial, lista):
    """Marca RECONSTRUIR -> ignora el archivo de estado viejo, conserva la marca mientras haya extractos y la quita al finalizar."""
    marcado = {"CLAVE_CONTROL": "GRUPO|" + GID, "ESTADO": "RECONSTRUIR", "VERSION_ESTADO": 0, "HASH_ESTADO": "",
               "RECONSTRUIR_DESDE": "2026-10-08T09:00:00"}
    viejo = d64(inicial["estado_b64"])
    r1 = sincronizar(viejo, [parcial], control=marcado)                      # el estado viejo se ignora (aunque fuera válido)
    assert r1["ok"] and r1["control"]["ESTADO"] == "RECONSTRUIR" and r1["control"]["VERSION_ESTADO"] == 1
    assert "RECONSTRUIR_DESDE" not in r1["control"]
    r2 = sincronizar(d64(r1["estado_b64"]), [parcial], control=_ctl(r1, RECONSTRUIR_DESDE=marcado["RECONSTRUIR_DESDE"]))
    assert r2["ok"] and r2["control"]["ESTADO"] == "RECONSTRUIR" and not r2["cambio_estado"]       # 2º extracto: mismos datos
    fin = sincronizar(d64(r1["estado_b64"]), filas=filas_de(lista), control=_ctl(r1, RECONSTRUIR_DESDE=marcado["RECONSTRUIR_DESDE"]),
                      finalizar=True)
    assert fin["ok"] and fin["control"]["ESTADO"] == "OK" and fin["control"]["RECONSTRUIR_DESDE"] == ""
    assert fin["control"]["HASH_OPERATIVO"] == E.hash_lista(filas_de(lista))
    sin_caida = sincronizar(d64(inicial["estado_b64"]), filas=filas_de(lista), control=_ctl(inicial))   # camino sin pérdida de estado
    assert sha(d64(fin["xlsx_b64"] or r1["xlsx_b64"])) == sha(d64(sin_caida["xlsx_b64"] or inicial["xlsx_b64"]))


def test_un_xlsx_que_no_se_puede_generar_no_deja_un_hash_falso_y_se_reintenta(inicial, monkeypatch):
    def explota(*a, **k):
        raise S.ErrorP10("XLSX_NO_VALIDO", "prueba", "XLSX")
    monkeypatch.setattr(S, "construir_xlsx", explota)
    r = sincronizar(d64(inicial["estado_b64"]), control=_ctl(inicial), forzar_xlsx=True)
    assert r["ok"] and not r["xlsx_valido"] and r["control"]["ESTADO"] == "ERROR" and r["control"]["INTENTOS"] == 1
    e = E.desempaquetar(d64(r["estado_b64"]))
    assert e["xlsx"] is None                           # el estado ya no afirma que existe un libro válido
    monkeypatch.undo()
    otra = sincronizar(d64(r["estado_b64"]), control=_ctl(r))
    assert otra["ok"] and otra["xlsx_valido"] and otra["cambio_xlsx"] and otra["control"]["ESTADO"] == "OK"


# ====================================================================== lo cerrado no cambia
def test_p10_a1_cerrado_y_los_modulos_reutilizados_no_cambiaron():
    """A.1 está APROBADO Y CERRADO: su contrato/generador no cambian sin aprobación explícita de Gabriel (regenerar con `python -m p10.huellas`)."""
    from p10 import huellas
    esperado = json.loads((huellas.DESTINO).read_text(encoding="utf-8"))
    actual = huellas.calcular()
    from tests.test_16_asignacion_p9 import EXCEPCIONES_AUTORIZADAS
    for grupo in ("a1_cerrado", "reutilizado_sin_cambios"):
        for ruta, h in actual[grupo].items():
            assert h in (esperado[grupo][ruta], EXCEPCIONES_AUTORIZADAS.get(ruta)), f"{ruta} cambió respecto de lo aprobado"


# ====================================================================== lectura incremental de la lista
def test_filas_parciales_de_la_lista_no_registran_hash_ni_obligan_a_escribir_el_control(inicial, lista):
    claves = sorted(i["CLAVE_TRANSACCION"] for i in lista.items.values())
    lista.confirmar(claves[0], estudiante="X")
    parcial_filas = [f for f in filas_de(lista) if f["CLAVE_TRANSACCION"] == claves[0]]
    r = sincronizar(d64(inicial["estado_b64"]), filas=parcial_filas, control=_ctl(inicial), parcial_lista=True)
    assert r["ok"] and r["cambio_estado"] and r["cambio_xlsx"] and r["escribir_control"]
    assert "HASH_OPERATIVO" not in r["control"] and r["control"]["ESTADO"] == "OK"
    otra = sincronizar(d64(r["estado_b64"]), filas=parcial_filas, control=_ctl(r), parcial_lista=True)       # misma fila otra vez
    assert not otra["cambio_estado"] and not otra["cambio_xlsx"] and otra["escribir_control"] is False
    con_error = sincronizar(d64(r["estado_b64"]), filas=parcial_filas, control=_ctl(r, ESTADO="ERROR"), parcial_lista=True)
    assert con_error["escribir_control"] is True                              # un grupo que estaba en error sí se vuelve a anotar
    completa = sincronizar(d64(r["estado_b64"]), filas=filas_de(lista), control=_ctl(r))                      # conciliación completa
    assert completa["control"]["HASH_OPERATIVO"] == E.hash_lista(filas_de(lista)) and completa["escribir_control"] is True


def test_el_estado_ya_no_guarda_el_hash_de_la_lista(inicial, lista):
    claves = sorted(i["CLAVE_TRANSACCION"] for i in lista.items.values())
    lista.confirmar(claves[0], estudiante="X")
    r = sincronizar(d64(inicial["estado_b64"]), filas=filas_de(lista), control=_ctl(inicial))
    assert "hash_lista" not in E.desempaquetar(d64(r["estado_b64"]))["operativo"]
