"""P4: la DETECCIÓN productiva (banco, cuenta, moneda, formato) sale de registro_bancos.json.

Reglas que estas pruebas protegen:
  * los 12 extractos reales se detectan exactamente en su formato actual (mismo id, misma hoja y fila de
    encabezado que leía su normalizar_* legado, retirado en P6) y las salidas válidas no cambian;
  * una cuenta nueva de un formato conocido entra con UNA entrada en CUENTAS, sin código;
  * la cuenta se lee solo en la celda rotulada de la cabecera: una cuenta dentro de una glosa no cambia nada;
  * la cuenta se valida en todos los formatos, BMSC incluido;
  * encabezado insuficiente, cabecera ambigua o reporte Unión «Últimos 12» = error claro, nunca una elección
    silenciosa;
  * (P4) la normalización seguía en los normalizar_* legados. DESDE P5 la normaliza el motor genérico con el
    registro (test_09_normalizacion_p5.py). P6 retiró `formato_legado`, `cuenta_nueva`, la plantilla
    `legado.plantilla_por_hoja` y `aplicar_identidad_registro`; las reglas de detección no cambian.
"""
import copy
import importlib.util
import json
import shutil
import sys

import pandas as pd
import pytest
from openpyxl import Workbook

from helpers import (CONTRATO, ENCABEZADOS, EXTRACTOS, FIXTURES, FORMATOS_OK, GOLDEN, HOJAS, RAIZ, UNION_ULTIMOS12,
                     crear_xlsx_bmsc_otra_cuenta, crear_xlsx_bnb, crear_xlsx_union_me_con_movimiento,
                     leer_csv_texto)

pytestmark = pytest.mark.deteccion

REPO = RAIZ.parent
TS = pd.Timestamp("2026-01-01 00:00:00")
FIJO = pd.Timestamp("2026-08-21 09:00:00")


# ------------------------------- fixtures y utilidades -------------------------------
@pytest.fixture(scope="session")
def dr():
    spec = importlib.util.spec_from_file_location("deteccion_registro_bajo_prueba", REPO / "deteccion_registro.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="session")
def datos_registro():
    return json.loads((REPO / "registro_bancos.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="session")
def detector(motor):
    return motor.detector_registro(str(REPO / "registro_bancos.json"))


def _registro_con(datos, tmp_path, *cuentas, mutar=None):
    """Copia del registro + cuentas nuevas (solo configuración), escrita en un JSON aparte."""
    d = copy.deepcopy(datos)
    d["CUENTAS"].extend(cuentas)
    if mutar:
        mutar(d)
    ruta = tmp_path / "registro_p4.json"
    ruta.write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")
    return ruta


def _correr(motor, entrada, salida, registro=None):
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(pd.Timestamp, "now", classmethod(lambda cls, tz=None: FIJO))
        if registro is None:
            mp.delenv("CBBA_REGISTRO_BANCOS", raising=False)
        else:
            mp.setenv("CBBA_REGISTRO_BANCOS", str(registro))
        return motor.ejecutar_motor(str(entrada), str(salida / "NORMALIZADO.xlsx"))


def _carpeta(tmp_path, *archivos):
    ent = tmp_path / "in"
    ent.mkdir(parents=True, exist_ok=True)
    for a in archivos:
        shutil.copy(a, ent / a.name)
    return ent


FILAS_BNB = [{"fecha": "01/08/2026", "hora": "09:00:00", "cred": 100.0, "saldo": 1100.0, "cod": "11"},
             {"fecha": "02/08/2026", "hora": "10:00:00", "deb": 30.0, "saldo": 1070.0, "cod": "12",
              "adic": "Nombre: ACME; Doc.ID:1"}]


# =========================== 1. LOS 12 FORMATOS REALES ===========================
@pytest.mark.parametrize("formato", FORMATOS_OK)
def test_p4_formato_real_se_detecta_exactamente_como_hoy(motor, detector, ruta_fixture, formato):
    r = detector.detectar(ruta_fixture(formato))
    assert r.estado == "OK", r.motivo
    assert r.cuenta_id == formato and r.formato_motor == formato
    assert (r.banco, r.cuenta, r.moneda) == CONTRATO[formato]
    # misma hoja y misma fila de encabezado que leía el normalizador legado (contrato congelado en helpers)
    assert r.hoja == HOJAS[formato]
    raw = motor.leer_todas_hojas(ruta_fixture(formato))[r.hoja]
    fila, puntaje, total = motor.normalizador_registro()._fila_encabezado(raw, ENCABEZADOS[formato])
    assert (r.fila_encabezado, puntaje) == (fila, total)
    # la cuenta salió de la celda rotulada de la cabecera, antes del encabezado
    assert r.fila_encabezado > 0 and CONTRATO[formato][1].replace("-", "") in r.cuenta_leida.replace("-", "")
    assert r.observaciones == []


@pytest.mark.parametrize("formato", FORMATOS_OK)
def test_p4_detectar_formato_conserva_su_contrato(motor, ruta_fixture, formato):
    assert motor.detectar_formato(ruta_fixture(formato)) == formato


def test_p4_lote_12_salidas_identicas_y_deteccion_por_registro(corrida_lote):
    """NORMALIZADO.xlsx y LISTS.csv del lote real = doradas del motor original; detección 12/12 por registro."""
    res, salida = corrida_lote
    pd.testing.assert_frame_equal(leer_csv_texto(salida / "LISTS.csv"), leer_csv_texto(GOLDEN / "LOTE_12_LISTS.csv"))
    d = res["deteccion_estado"]
    assert d["version_deteccion"] == "P4-1" and len(d["sha256_registro"]) == 64
    assert {a["CUENTA_ID"] for a in d["archivos"].values()} == set(FORMATOS_OK)
    assert all(a["ESTADO"] == "OK" and a["HOJA"] == HOJAS[a["CUENTA_ID"]] for a in d["archivos"].values())
    assert all("FORMATO_LEGADO" not in a and "CUENTA_NUEVA" not in a for a in d["archivos"].values())   # P6
    assert list(res["df_deteccion_final"].columns) == ["ARCHIVO", "FORMATO", "ESTADO"]
    assert set(res["df_deteccion_final"]["ESTADO"]) == {"OK"}


# =========================== 2. REGISTRO ===========================
def test_p4_registro_valido_para_detectar(dr, datos_registro):
    assert dr.validar_registro_deteccion(datos_registro) == []
    assert datos_registro["version_deteccion"] == "P4-1"
    for fid, f in datos_registro["FORMATOS"].items():
        assert "legado" not in f, fid                                   # P6: sin plantillas legadas
        if f.get("aceptado", True) is False:
            continue
        assert f["deteccion"]["zona_cuenta"] == "CABECERA" and f["deteccion"]["etiquetas_cuenta"], fid


MUTACIONES = {
    "sin_etiquetas_de_cuenta": lambda d: d["FORMATOS"]["BMSC_EXCEL_V1"]["deteccion"].update(etiquetas_cuenta=[]),
    "mismo_numero_en_dos_formatos": lambda d: d["CUENTAS"].append(
        {"id": "X", "formato": "BMSC_EXCEL_V1", "banco": "BMSC", "cuenta": "3000100152", "moneda": "BOB"}),
    "id_repetido": lambda d: d["CUENTAS"].append(dict(d["CUENTAS"][0], cuenta="5555555555")),
    "formato_inexistente": lambda d: d["CUENTAS"][0].update(formato="NO_EXISTE"),
    "cuenta_sin_moneda": lambda d: d["CUENTAS"][0].update(moneda=""),
    "cuenta_en_formato_rechazado": lambda d: d["CUENTAS"].append(
        {"id": "X", "formato": "UNION_ULTIMOS12_V1", "banco": "U", "cuenta": "20000009999999", "moneda": "USD"}),
    "cuenta_sin_digitos": lambda d: d["CUENTAS"].append(
        {"id": "X", "formato": "BNB_EXTRACTO_V1", "banco": "BNB", "cuenta": "ABC", "moneda": "BOB"}),
    "prohibida_sin_mensaje": lambda d: d["FORMATOS"]["UNION_FECHAS_V1"]["deteccion"]["cabecera_prohibida"][0].pop(
        "mensaje"),
}


@pytest.mark.parametrize("nombre", list(MUTACIONES))
def test_p4_registro_invalido_se_rechaza_al_cargar(motor, dr, datos_registro, tmp_path, nombre):
    d = copy.deepcopy(datos_registro)
    MUTACIONES[nombre](d)
    ruta = tmp_path / "r.json"
    ruta.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ValueError, match="registro de bancos inválido"):
        motor.detector_registro(str(ruta))


# =========================== 3. CUENTA NUEVA SOLO POR CONFIGURACIÓN ===========================
NUEVA_BNB = {"id": "BNB_NUEVA_USD", "formato": "BNB_EXTRACTO_V1", "banco": "BNB", "cuenta": "3999000111",
             "moneda": "USD", "activa": True}


def test_p4_cuenta_nueva_sin_registrar_se_rechaza_con_motivo(motor, tmp_path):
    ruta = crear_xlsx_bnb(tmp_path / "nueva.xlsx", "3999000111", FILAS_BNB)
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "CUENTA_NO_REGISTRADA" and "3999000111" in r.motivo and "CUENTAS" in r.motivo
    with pytest.raises(ValueError, match="CUENTA_NO_REGISTRADA") as e:
        _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out")
    assert "nueva.xlsx" in str(e.value)
    assert not (tmp_path / "out").exists()


def test_p4_cuenta_nueva_se_incorpora_solo_con_una_entrada_en_cuentas(motor, datos_registro, tmp_path):
    """Mismo motor, mismo código: solo el registro gana una entrada. La cuenta se procesa de punta a punta
    (NORMALIZADO, LISTS, ORIGEN) con BANCO / CUENTA / MONEDA del registro."""
    ruta = crear_xlsx_bnb(tmp_path / "nueva.xlsx", "3999000111", FILAS_BNB)
    reg = _registro_con(datos_registro, tmp_path, NUEVA_BNB)
    res = _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out", registro=reg)

    df = res["df_final"]
    assert len(df) == 2
    assert set(df["BANCO"]) == {"BNB"} and set(df["CUENTA BANCARIA"]) == {"3999000111"}
    assert set(df["MONEDA"]) == {"USD"}                      # del registro (la plantilla BNB_MN diría BOB)
    assert df["CLAVE TRANSACCIÓN"].tolist() == df.apply(motor.crear_clave, axis=1).tolist()
    assert all(c.startswith("BNB|3999000111|2026080") for c in df["CLAVE TRANSACCIÓN"])
    assert res["df_deteccion_final"]["FORMATO"].tolist() == ["BNB_NUEVA_USD"]
    assert res["df_validacion"]["ESTADO"].tolist() == ["OK"]
    det = res["deteccion_estado"]["archivos"]["nueva.xlsx"]
    assert (det["CUENTA_ID"], det["FORMATO_REGISTRO"], det["HOJA"]) == ("BNB_NUEVA_USD", "BNB_EXTRACTO_V1", "Hoja 1")

    lists = pd.read_csv(tmp_path / "out" / "LISTS.csv", encoding="utf-8-sig", dtype=str)
    assert set(lists["CUENTA BANCARIA"]) == {"3999000111"} and set(lists["MONEDA"]) == {"USD"}
    assert res["origen_estado"]["estado"] == "OK"            # la captura (P1) también la procesa


def test_p4_cuenta_nueva_tambien_llega_al_historico_p3b(motor, datos_registro, tmp_path):
    """P3b sigue funcionando: la cuenta nueva genera su EXTRACTO_HISTORICO con el mismo registro."""
    ruta = crear_xlsx_bnb(tmp_path / "nueva.xlsx", "3999000111", FILAS_BNB)
    reg = _registro_con(datos_registro, tmp_path, NUEVA_BNB)
    _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out", registro=reg)
    spec = importlib.util.spec_from_file_location("historico_p4", REPO / "historico.py")
    h = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(h)
    r = h.generar_extractos_historicos(str(tmp_path / "out" / "ORIGEN.xlsx"), str(tmp_path / "out" / "LISTS.csv"),
                                       str(tmp_path / "hist"), str(reg))
    assert r["estado"] == "OK" and r["advertencias"] == [] and r["omitidos"] == []
    assert [(a["nombre_archivo"], a["cuenta"], a["movimientos"]) for a in r["archivos"]] == [
        ("EXTRACTO_HISTORICO_BNB_3999000111_2026-08.xlsx", "BNB_NUEVA_USD", 2)]


def test_p4_cuenta_nueva_en_hoja_alternativa_se_detecta_en_esa_hoja(motor, datos_registro, tmp_path):
    ruta = crear_xlsx_bnb(tmp_path / "nueva.xlsx", "3999000111", FILAS_BNB, hoja="Hoja")
    reg = _registro_con(datos_registro, tmp_path, dict(NUEVA_BNB, moneda="BOB"))
    res = _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out", registro=reg)
    assert res["deteccion_estado"]["archivos"]["nueva.xlsx"]["HOJA"] == "Hoja"
    assert set(res["df_final"]["CUENTA BANCARIA"]) == {"3999000111"} and set(res["df_final"]["MONEDA"]) == {"BOB"}


def test_p4_cuenta_nueva_de_union_junto_al_lote_real(motor, datos_registro, tmp_path):
    """Otro formato (UNION_FECHAS_V1) y junto a extractos reales: las cuentas reales no cambian."""
    ruta = crear_xlsx_union_me_con_movimiento(tmp_path / "union_nueva.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(ruta)
    wb.active["E8"] = "30000009990001"
    wb.save(ruta)
    reg = _registro_con(datos_registro, tmp_path, {"id": "UNION_NUEVA", "formato": "UNION_FECHAS_V1",
                                                   "banco": "BANCO UNIÓN", "cuenta": "30000009990001",
                                                   "moneda": "USD", "activa": True})
    ent = _carpeta(tmp_path, ruta, EXTRACTOS / FIXTURES["BNB_ME"], EXTRACTOS / FIXTURES["UNION_MN"])
    res = _correr(motor, ent, tmp_path / "out", registro=reg)
    df = res["df_final"]
    nueva = df[df["ARCHIVO ORIGEN"] == "union_nueva.xlsx"]
    assert len(nueva) == 1 and set(nueva["CUENTA BANCARIA"]) == {"30000009990001"} and set(nueva["MONEDA"]) == {"USD"}
    assert res["deteccion_estado"]["archivos"]["union_nueva.xlsx"]["FORMATO_REGISTRO"] == "UNION_FECHAS_V1"
    for fm in ("BNB_ME", "UNION_MN"):
        real = df[df["ARCHIVO ORIGEN"] == FIXTURES[fm]]
        esperado = pd.read_csv(GOLDEN / f"{fm}.csv", dtype=str, keep_default_na=False)
        assert real["CLAVE TRANSACCIÓN"].tolist() == esperado["CLAVE TRANSACCIÓN"].tolist(), fm


def test_p4_cuenta_nueva_no_modifica_el_registro_productivo(datos_registro):
    """Las pruebas de cuenta nueva trabajan sobre copias: el registro del repositorio sigue con 13 cuentas."""
    assert len(datos_registro["CUENTAS"]) == 13
    assert json.loads((REPO / "registro_bancos.json").read_text(encoding="utf-8"))["CUENTAS"] == datos_registro["CUENTAS"]


# =========================== 4. LA CUENTA SOLO SE LEE EN LA CABECERA ===========================
def test_p4_cuenta_en_glosa_no_altera_la_deteccion(motor, tmp_path):
    """BNB_CLINICA con la cuenta de BNB_MN (y de BMSC) en descripción y adicionales: sigue siendo BNB_CLINICA."""
    ruta = crear_xlsx_bnb(tmp_path / "glosa.xlsx", "3000100705", [
        {"fecha": "01/08/2026", "cred": 100.0, "saldo": 1100.0, "desc": "TRASPASO CTA 3000100152",
         "adic": "Transferencia desde cuenta 3000100152 / 1000872489"},
        {"fecha": "02/08/2026", "cred": 50.0, "saldo": 1150.0, "adic": "Cuenta: 3400041236"}])
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "OK" and r.cuenta_id == "BNB_CLINICA" and r.cuenta_leida == "3000100705"
    res = _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out")
    assert set(res["df_final"]["CUENTA BANCARIA"]) == {"3000100705"}


def test_p4_cuenta_no_registrada_con_glosa_que_menciona_una_registrada_se_rechaza(motor, tmp_path):
    """El caso peligroso de D-10: una cuenta ajena cuyo movimiento menciona una cuenta registrada NO se
    procesa como la registrada."""
    ruta = crear_xlsx_bnb(tmp_path / "ajena.xlsx", "3777000999", [
        {"fecha": "01/08/2026", "cred": 100.0, "saldo": 1100.0, "adic": "Abono desde cuenta 3000100152"}])
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "CUENTA_NO_REGISTRADA" and "3777000999" in r.motivo


def test_p4_cuenta_fuera_de_la_cabecera_no_cuenta(motor, tmp_path):
    """Sin cuenta en la cabecera (solo en una fila de datos) = SIN_CUENTA, aunque sea una cuenta registrada."""
    wb = Workbook(); ws = wb.active; ws.title = "Hoja 1"
    ws.append(["Fecha", "Hora", "Oficina", "Descripción", "Referencia", "Código de transacción",
               "ITF", "Débitos", "Créditos", "Saldo", "Adicionales"])
    ws.append(["01/08/2026", "10:00:00", "OF", "DEP", "R", "77", 0, None, 100, 1100, "x"])
    ws.append(["Cuenta:", "3000100152"])
    ruta = tmp_path / "sin_cabecera.xlsx"; wb.save(ruta)
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "SIN_CUENTA" and "no trae el número de cuenta" in r.motivo
    assert motor.detectar_formato(str(ruta)) == "NO_RECONOCIDO"


def test_p4_los_extractos_reales_tienen_otras_cuentas_en_sus_movimientos_y_no_confunden(motor, detector, ruta_fixture):
    """BNB_AHORRO real trae decenas de 'Cuenta Origen: ...' en adicionales: la detección no las mira."""
    raw = motor.leer_todas_hojas(ruta_fixture("BNB_AHORRO"))["Hoja"]
    assert raw.astype(str).apply(lambda c: c.str.contains("Cuenta Origen", na=False)).to_numpy().sum() > 10
    assert detector.detectar(ruta_fixture("BNB_AHORRO")).cuenta_id == "BNB_AHORRO"


# =========================== 5. BMSC VALIDA LA CUENTA ===========================
def test_p4_bmsc_con_cuenta_incorrecta_se_rechaza(motor, tmp_path):
    ruta = crear_xlsx_bmsc_otra_cuenta(tmp_path / "bmsc.xlsx", "9999999999")
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "CUENTA_NO_REGISTRADA" and r.formato_id == "BMSC_EXCEL_V1" and "9999999999" in r.motivo
    with pytest.raises(ValueError, match="hay archivos no reconocidos") as e:
        _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out")
    assert "bmsc.xlsx: CUENTA_NO_REGISTRADA" in str(e.value)
    assert not (tmp_path / "out").exists()


def test_p4_bmsc_con_cuenta_de_otro_banco_se_rechaza_y_lo_dice(motor, tmp_path):
    ruta = crear_xlsx_bmsc_otra_cuenta(tmp_path / "bmsc.xlsx", "3000100152")
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "CUENTA_NO_REGISTRADA" and "BNB_MN de otro formato" in r.motivo


def test_p4_bmsc_con_su_cuenta_registrada_se_acepta(motor, tmp_path):
    ruta = crear_xlsx_bmsc_otra_cuenta(tmp_path / "bmsc.xlsx", "1000872489")
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "OK" and r.cuenta_id == "BMSC"


def test_p4_bmsc_sin_cuenta_en_la_cabecera_se_rechaza(motor, tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "Excel"
    ws.append(["Banco Mercantil Santa Cruz"])
    ws.append(["Fecha", "Hora", "Cod. Bca.", "Débito", "Crédito", "Saldo"])
    ws.append(["01/08/2026", "10:00:00", "123", None, 100, 1100])
    ruta = tmp_path / "bmsc.xlsx"; wb.save(ruta)
    assert motor.detectar_extracto(str(ruta)).estado == "SIN_CUENTA"


# =========================== 6. CABECERA AMBIGUA ===========================
def test_p4_dos_cuentas_en_la_celda_de_cuenta_es_ambiguo(motor, tmp_path):
    ruta = crear_xlsx_bnb(tmp_path / "a.xlsx", "3000100152 y 3000100705", FILAS_BNB)
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "AMBIGUO" and set(r.candidatos) == {"BNB_MN", "BNB_CLINICA"}
    assert "más de un número de cuenta" in r.motivo


def test_p4_dos_rotulos_de_cuenta_distintos_es_ambiguo(motor, tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "Hoja 1"
    ws.append(["Número De cuenta", "3000100152"])
    ws.append(["Cuenta:", "3000100705"])
    ws.append(["Fecha", "Hora", "Oficina", "Descripción", "Referencia", "Código de transacción",
               "ITF", "Débitos", "Créditos", "Saldo", "Adicionales"])
    ws.append(["01/08/2026", "10:00:00", "OF", "DEP", "R", "77", 0, None, 100, 1100, "x"])
    ruta = tmp_path / "dos.xlsx"; wb.save(ruta)
    assert motor.detectar_extracto(str(ruta)).estado == "AMBIGUO"


def test_p4_una_cuenta_registrada_y_otra_desconocida_en_la_cabecera_es_ambiguo(motor, tmp_path):
    ruta = crear_xlsx_bnb(tmp_path / "a.xlsx", "3000100152 / 3777000999", FILAS_BNB)
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "AMBIGUO" and r.candidatos == ["BNB_MN"]


def test_p4_archivo_con_estructura_de_dos_formatos_es_ambiguo(motor, tmp_path):
    """Un libro con una hoja BNB y una hoja BMSC válidas: no se elige por orden de evaluación."""
    wb = Workbook(); ws = wb.active; ws.title = "Hoja 1"
    ws.append(["Cuenta:", "3000100152"])
    ws.append(["Fecha", "Hora", "Oficina", "Descripción", "Referencia", "Código de transacción",
               "ITF", "Débitos", "Créditos", "Saldo", "Adicionales"])
    ws2 = wb.create_sheet("Excel")
    ws2.append(["Banco Mercantil Santa Cruz"])
    ws2.append(["Nro de Cuenta:", "1000872489"])
    ws2.append(["Fecha", "Hora", "Cod. Bca.", "Débito", "Crédito", "Saldo"])
    ruta = tmp_path / "doble.xlsx"; wb.save(ruta)
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "AMBIGUO" and set(r.candidatos) == {"BNB_EXTRACTO_V1", "BMSC_EXCEL_V1"}


def test_p4_cabecera_ambigua_detiene_el_lote_sin_escribir(motor, tmp_path):
    ruta = crear_xlsx_bnb(tmp_path / "a.xlsx", "3000100152 y 3000100705", FILAS_BNB)
    ent = _carpeta(tmp_path, ruta, EXTRACTOS / FIXTURES["BNB_ME"])
    with pytest.raises(ValueError, match="a.xlsx: AMBIGUO"):
        _correr(motor, ent, tmp_path / "out")
    assert not (tmp_path / "out").exists()


# =========================== 7. ENCABEZADO INSUFICIENTE ===========================
def test_p4_encabezado_incompleto_se_rechaza_y_dice_que_falta(motor, tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "Hoja 1"
    ws.append(["Cuenta:", "3000100152"])
    ws.append(["Fecha", "Hora", "Codigo de transaccion", "Adicionales"])       # faltan 4 de 7
    ws.append(["01/08/2026", "10:00:00", "77", "x"])
    ruta = tmp_path / "inc.xlsx"; wb.save(ruta)
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "ENCABEZADO_INCOMPLETO"
    assert "(3/7)" in r.motivo and all(x in r.motivo for x in ("DESCRIPCION", "DEBITOS", "CREDITOS", "SALDO"))
    with pytest.raises(ValueError, match="encabezado incompleto"):
        _correr(motor, _carpeta(tmp_path, ruta), tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_p4_encabezado_con_una_sola_columna_faltante_tambien_se_rechaza(motor, tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "Hoja 1"
    ws.append(["Cuenta:", "3000100152"])
    ws.append(["Fecha", "Hora", "Oficina", "Descripción", "Referencia", "Código de transacción",
               "ITF", "Débitos", "Créditos", "Adicionales"])                   # falta Saldo
    ruta = tmp_path / "inc.xlsx"; wb.save(ruta)
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "ENCABEZADO_INCOMPLETO" and r.motivo.endswith("faltan: SALDO")


def test_p4_hoja_sin_encabezado_reconocible_no_se_asigna_a_ningun_formato(motor, tmp_path):
    """Escenario D-09 en la ruta productiva: hoja aceptada sin ningún encabezado = error, no 'fila cualquiera'."""
    wb = Workbook(); ws = wb.active; ws.title = "Hoja 1"
    ws.append(["Reporte", "x"]); ws.append(["nada", "util"])
    ruta = tmp_path / "vacio.xlsx"; wb.save(ruta)
    r = motor.detectar_extracto(str(ruta))
    assert r.estado == "SIN_FORMATO" and "ninguna firma" in r.motivo
    assert motor.detectar_formato(str(ruta)) == "NO_RECONOCIDO"


# =========================== 8. UNIÓN «ÚLTIMOS 12 MOVIMIENTOS» ===========================
def test_p4_union_ultimos12_real_se_rechaza_en_la_deteccion(motor):
    r = motor.detectar_extracto(str(UNION_ULTIMOS12))
    assert r.estado == "RECHAZADO" and "Ultimos 12 movimientos" in r.motivo and r.cuenta_id is None
    assert motor.detectar_formato(str(UNION_ULTIMOS12)) == "NO_RECONOCIDO"


def test_p4_union_ultimos12_detiene_el_lote_con_su_mensaje(motor, tmp_path):
    ent = _carpeta(tmp_path, UNION_ULTIMOS12, EXTRACTOS / FIXTURES["UNION_MN"])
    with pytest.raises(ValueError, match="RECHAZADO") as e:
        _correr(motor, ent, tmp_path / "out")
    assert "Descargar el extracto por rango de fechas" in str(e.value)
    assert not (tmp_path / "out").exists()


def _ultimos12_en_hoja(motor, tmp_path, agregar_saldo):
    """El contenido REAL de 'Últimos 12' copiado a un libro cuya hoja se llama como la de UNION_FECHAS_V1."""
    raw = motor.leer_todas_hojas(str(UNION_ULTIMOS12))["ExtractoMovimientosUltimos"]
    wb = Workbook(); ws = wb.active; ws.title = "ExtractoMovimientosFechas"
    for fila in raw.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in fila])
    if agregar_saldo:
        ws.cell(16, 36).value = "Saldo"               # fila del encabezado real (16), columna libre
    ruta = tmp_path / "ultimos12_renombrado.xlsx"; wb.save(ruta)
    return ruta


def test_p4_union_ultimos12_con_la_hoja_renombrada_tampoco_pasa_como_union_me(motor, tmp_path):
    """No cumple UNION_FECHAS_V1: le falta la columna Saldo."""
    r = motor.detectar_extracto(str(_ultimos12_en_hoja(motor, tmp_path, agregar_saldo=False)))
    assert r.estado == "ENCABEZADO_INCOMPLETO" and "faltan: SALDO" in r.motivo


def test_p4_union_ultimos12_aunque_traiga_saldo_se_rechaza_por_su_cabecera(motor, tmp_path):
    r = motor.detectar_extracto(str(_ultimos12_en_hoja(motor, tmp_path, agregar_saldo=True)))
    assert r.estado == "RECHAZADO" and "no cumple el contrato UNION_FECHAS_V1" in r.motivo


def test_p4_union_me_proxy_sintetico_sigue_detectandose(motor, tmp_path):
    """UNION_ME conserva sus pruebas estructurales/sintéticas hasta tener movimientos reales."""
    r = motor.detectar_extracto(str(crear_xlsx_union_me_con_movimiento(tmp_path / "u.xlsx")))
    assert r.estado == "OK" and r.cuenta_id == "UNION_ME" and r.hoja == HOJAS["UNION_ME"]
    assert (r.banco, r.cuenta, r.moneda) == CONTRATO["UNION_ME"]


# =========================== 9. P4 NO TOCA LA NORMALIZACIÓN ===========================
def test_p4_la_deteccion_no_contiene_logica_de_normalizacion(dr):
    """deteccion_registro.py no nombra primitivas de importes, fechas, saldos ni la CLAVE."""
    import ast
    arbol = ast.parse((REPO / "deteccion_registro.py").read_text(encoding="utf-8"))
    nombres = {n.id for n in ast.walk(arbol) if isinstance(n, ast.Name)} | \
              {n.attr for n in ast.walk(arbol) if isinstance(n, ast.Attribute)}
    prohibidos = {"numero", "normalizar_fecha", "normalizar_hora", "codigo_texto", "crear_clave",
                  "finalizar_dataframe", "validar_archivo", "ecuacion_saldo", "normalizar_archivo",
                  "texto_de_archivo", "COLUMNAS_LISTS"}
    assert nombres & prohibidos == set()
