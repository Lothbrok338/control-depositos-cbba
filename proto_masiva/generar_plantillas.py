"""Genera las plantillas oficiales del prototipo P9 MASIVA (solo datos FICTICIOS).

    python proto_masiva/generar_plantillas.py

Salida (determinista: mismo SHA-256 en cada ejecución):
    proto_masiva/xlsx/Plantilla_Confirmacion_Masiva_P9.xlsx         tabla con 3 filas de ejemplo
    proto_masiva/xlsx/Plantilla_Confirmacion_Masiva_P9_VACIA.xlsx   misma tabla, sin registros

Contrato: una sola tabla `tblConfirmacionMasiva` con 9 columnas, en este orden. TIPO_CAMBIO NO existe
(el tipo de cambio se pide solo al imprimir el comprobante, como hoy).
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

SALIDA = Path(__file__).resolve().parent / "xlsx"
TABLA = "tblConfirmacionMasiva"
HOJA = "CARGA"
ENCABEZADOS = ("BANCO", "CUENTA_BANCARIA", "CODIGO_ASIGNACION", "IMPORTE", "MONEDA",
               "ESTUDIANTE", "SOLICITADO_POR", "SEDE", "OBSERVACION")
# Ejemplos 100 % inventados (cuentas con el formato de las reales, pero ficticias).
EJEMPLOS = (
    ("BCP", "999-9999990-9-99", "900001", 1000.00, "BOB", "ESTUDIANTE FICTICIO 1", "SOLICITANTE FICTICIO", "COCHABAMBA", ""),
    ("BISA", "0999999991", "900002", 250.50, "BOB", "ESTUDIANTE FICTICIO 2", "SOLICITANTE FICTICIO", "COCHABAMBA", "ejemplo"),
    ("BNB", "9999999992", "3P00000003", 500.00, "USD", "ESTUDIANTE FICTICIO 3", "SOLICITANTE FICTICIO", "COCHABAMBA", ""),
)
FILAS_FORMATO = 500  # formatos y validación preconfigurados hasta esta fila (la tabla se amplía sola)
GRANATE = "7B1632"
LEEME = (
    "PLANTILLA DE IMPORTACIÓN MASIVA · P9 (PROTOTIPO)",
    "",
    "1. Trabaja SOLO en la hoja CARGA, dentro de la tabla tblConfirmacionMasiva. No cambies el nombre de la tabla ni los encabezados.",
    "2. Borra las 3 filas de ejemplo y escribe tus filas. No dejes filas en blanco en medio.",
    "3. CUENTA_BANCARIA y CODIGO_ASIGNACION son TEXTO (conservan guiones y ceros a la izquierda). No los conviertas a número.",
    "4. MONEDA: BOB o USD. IMPORTE: número con 2 decimales.",
    "5. ESTUDIANTE, SOLICITADO_POR y SEDE son obligatorios. OBSERVACION es opcional.",
    "6. Guarda como .xlsx y ciérralo antes de adjuntarlo en la app.",
    "",
    "Esta plantilla no contiene tipo de cambio: se solicita solo al imprimir el comprobante de un depósito USD.",
)
FECHA_FIJA = datetime(2026, 10, 5, 0, 0, 0)
FECHA_ZIP = (2026, 10, 5, 0, 0, 0)


def _anchos(ws):
    for col, ancho in zip("ABCDEFGHI", (14, 22, 22, 14, 11, 28, 26, 16, 34)):
        ws.column_dimensions[col].width = ancho


def construir(con_ejemplos: bool) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    ws.append(list(ENCABEZADOS))
    filas = EJEMPLOS if con_ejemplos else ()
    for fila in filas:
        ws.append(list(fila))
    n = max(len(filas), 1)  # una tabla de Excel siempre conserva al menos una fila de datos (en blanco si está vacía)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=GRANATE)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for fila in range(2, FILAS_FORMATO + 2):
        ws.cell(fila, 2).number_format = "@"          # CUENTA_BANCARIA = Texto
        ws.cell(fila, 3).number_format = "@"          # CODIGO_ASIGNACION = Texto
        ws.cell(fila, 4).number_format = "#,##0.00"   # IMPORTE
    tabla = Table(displayName=TABLA, ref=f"A1:I{1 + n}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tabla)
    moneda = DataValidation(type="list", formula1='"BOB,USD"', allow_blank=True, showErrorMessage=True,
                            errorTitle="MONEDA", error="Use BOB o USD.")
    moneda.add(f"E2:E{FILAS_FORMATO + 1}")
    ws.add_data_validation(moneda)
    ws.freeze_panes = "A2"
    _anchos(ws)
    leeme = wb.create_sheet("LEEME")
    for i, texto in enumerate(LEEME, 1):
        leeme.cell(i, 1, texto)
    leeme["A1"].font = Font(bold=True, size=13, color=GRANATE)
    leeme.column_dimensions["A"].width = 120
    wb.properties.creator = "P9 PROTOTIPO"
    wb.properties.created = wb.properties.modified = FECHA_FIJA
    return _zip_determinista(wb)


def _zip_determinista(wb) -> bytes:
    crudo = io.BytesIO()
    wb.save(crudo)
    salida = io.BytesIO()
    with zipfile.ZipFile(crudo) as origen, zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as destino:
        for info in sorted(origen.infolist(), key=lambda i: i.filename):
            nuevo = zipfile.ZipInfo(info.filename, date_time=FECHA_ZIP)
            nuevo.compress_type, nuevo.external_attr = zipfile.ZIP_DEFLATED, 0o644 << 16
            datos = origen.read(info.filename)
            if info.filename == "docProps/core.xml":  # openpyxl reescribe «modified» con la hora actual al guardar
                datos = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-10-05T00:00:00Z\g<2>", datos)
            destino.writestr(nuevo, datos)
    return salida.getvalue()


def generar() -> dict[str, Path]:
    SALIDA.mkdir(exist_ok=True)
    rutas = {"ejemplo": SALIDA / "Plantilla_Confirmacion_Masiva_P9.xlsx",
             "vacia": SALIDA / "Plantilla_Confirmacion_Masiva_P9_VACIA.xlsx"}
    rutas["ejemplo"].write_bytes(construir(True))
    rutas["vacia"].write_bytes(construir(False))
    return rutas


if __name__ == "__main__":
    import hashlib
    for nombre, ruta in generar().items():
        print(hashlib.sha256(ruta.read_bytes()).hexdigest(), ruta.name)
