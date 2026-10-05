"""Genera los XLSX ficticios de la Prueba A (solo infraestructura; sin datos reales).

Ejecutar: python proto_masiva/generar_xlsx.py
Plantilla V1 definitiva: 9 columnas (sin TIPO_CAMBIO).
Todos los valores son inventados, incluidas las cuentas (con el mismo formato que las reales: guiones, cero inicial).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

SALIDA = Path(__file__).resolve().parent / "xlsx"
ENCABEZADOS = ["BANCO", "CUENTA_BANCARIA", "CODIGO_ASIGNACION", "IMPORTE", "MONEDA",
               "ESTUDIANTE", "SOLICITADO_POR", "SEDE", "OBSERVACION"]
FILAS = [
    ["BCP", "999-9999990-9-99", "900001", 1000.00, "BOB", "ESTUDIANTE FICTICIO 1", "SOLICITANTE FICTICIO", "COCHABAMBA", ""],
    ["BISA", "0999999991", "900002", 250.50, "BOB", "ESTUDIANTE FICTICIO 2", "SOLICITANTE FICTICIO", "COCHABAMBA", "prueba"],
    ["BNB", "9999999992", "3P00000003", 500.00, "USD", "ESTUDIANTE FICTICIO 3", "SOLICITANTE FICTICIO", "COCHABAMBA", ""],
]


def libro(nombre, encabezados=ENCABEZADOS, filas=FILAS, tabla="tblConfirmacionMasiva", ref_filas=None):
    """ref_filas: nº de filas de datos que abarca la tabla (None = len(filas)); 0 = tabla solo con encabezado."""
    wb = Workbook()
    ws = wb.active
    ws.title = "CARGA"
    ws.append(encabezados)
    for f in filas:
        ws.append(f)
    # CUENTA_BANCARIA y CODIGO_ASIGNACION como Texto (ceros a la izquierda)
    for fila in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2), min_col=2, max_col=3):
        for c in fila:
            c.number_format = "@"
    n = len(filas) if ref_filas is None else ref_filas
    if tabla:
        t = Table(displayName=tabla, ref=f"A1:{chr(64 + len(encabezados))}{1 + n}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    SALIDA.mkdir(exist_ok=True)
    wb.save(SALIDA / nombre)


if __name__ == "__main__":
    libro("01_OK_3filas.xlsx")
    libro("02a_TABLA_VACIA_1fila_en_blanco.xlsx", filas=[[""] * 9], ref_filas=1)   # lo que Excel crea al insertar una tabla vacía
    libro("02b_TABLA_SOLO_ENCABEZADO.xlsx", filas=[], ref_filas=0)                 # puede abrirse con aviso de reparación
    libro("03_SIN_TABLA.xlsx", tabla=None)                                          # datos en el rango pero sin tabla
    libro("04_TABLA_NOMBRE_DISTINTO.xlsx", tabla="Tabla1")                         # tabla existe, otro nombre
    h = list(ENCABEZADOS)
    h[1] = "CUENTA"                                                                  # encabezado cambiado
    libro("05_ENCABEZADO_CAMBIADO.xlsx", encabezados=h)
    print(*sorted(p.name for p in SALIDA.glob("*.xlsx")), sep="\n")
