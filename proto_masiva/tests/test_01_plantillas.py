"""Plantillas XLSX del prototipo: contrato de 9 columnas, tabla única, formatos, datos 100 % ficticios, determinismo."""
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from proto_masiva import generar_plantillas as G

RAIZ = Path(__file__).resolve().parents[1]
REPO = RAIZ.parent
XLSX = RAIZ / "xlsx"
PLANTILLAS = ("Plantilla_Confirmacion_Masiva_P9.xlsx", "Plantilla_Confirmacion_Masiva_P9_VACIA.xlsx")
NUEVE = ["BANCO", "CUENTA_BANCARIA", "CODIGO_ASIGNACION", "IMPORTE", "MONEDA", "ESTUDIANTE", "SOLICITADO_POR", "SEDE", "OBSERVACION"]
FIXTURES = ("01_OK_3filas.xlsx", "02a_TABLA_VACIA_1fila_en_blanco.xlsx", "02b_TABLA_SOLO_ENCABEZADO.xlsx",
            "03_SIN_TABLA.xlsx", "04_TABLA_NOMBRE_DISTINTO.xlsx", "05_ENCABEZADO_CAMBIADO.xlsx")


def cargar(nombre):
    wb = load_workbook(XLSX / nombre)
    return wb, wb["CARGA"]


@pytest.mark.parametrize("nombre", PLANTILLAS)
def test_una_sola_tabla_con_el_nombre_exacto_y_9_columnas_en_orden(nombre):
    wb, ws = cargar(nombre)
    assert [h.value for h in ws[1] if h.value is not None] == NUEVE == list(G.ENCABEZADOS)
    assert list(ws.tables) == ["tblConfirmacionMasiva"]
    assert all(not hoja.tables for hoja in wb.worksheets if hoja.title != "CARGA")
    assert ws.tables["tblConfirmacionMasiva"].ref.startswith("A1:I")


@pytest.mark.parametrize("nombre", PLANTILLAS)
def test_TIPO_CAMBIO_no_existe_en_ningun_sitio_del_archivo(nombre):
    with zipfile.ZipFile(XLSX / nombre) as z:
        contenido = b"".join(z.read(n) for n in z.namelist())
    assert b"TIPO_CAMBIO" not in contenido


def test_ejemplo_tiene_3_filas_y_la_vacia_ninguna():
    _, ejemplo = cargar(PLANTILLAS[0])
    assert ejemplo.tables["tblConfirmacionMasiva"].ref == "A1:I4"
    assert [ejemplo.cell(r, 3).value for r in (2, 3, 4)] == ["900001", "900002", "3P00000003"]
    _, vacia = cargar(PLANTILLAS[1])
    # Excel no admite una tabla sin fila de datos: «vacía» = una fila en blanco
    assert vacia.tables["tblConfirmacionMasiva"].ref == "A1:I2"
    assert all(c.value is None for c in vacia[2])


@pytest.mark.parametrize("nombre", PLANTILLAS)
def test_formatos_texto_numero_y_validacion_de_moneda(nombre):
    _, ws = cargar(nombre)
    for fila in (2, 3, 4, 100, 500):
        assert ws.cell(fila, 2).number_format == "@" and ws.cell(fila, 3).number_format == "@"
        assert ws.cell(fila, 4).number_format == "#,##0.00"
    validaciones = ws.data_validations.dataValidation
    assert len(validaciones) == 1 and validaciones[0].type == "list" and validaciones[0].formula1 == '"BOB,USD"'
    assert str(validaciones[0].sqref).startswith("E2")


def test_el_ejemplo_conserva_cuentas_como_texto_con_cero_inicial():
    _, ws = cargar(PLANTILLAS[0])
    assert ws["B3"].value == "0999999991" and isinstance(ws["B3"].value, str)
    assert ws["D2"].value == 1000 and ws["D3"].value == 250.5


@pytest.mark.parametrize("nombre", PLANTILLAS)
def test_hoja_leeme_no_agrega_tablas_y_dice_que_no_hay_tipo_de_cambio(nombre):
    wb, _ = cargar(nombre)
    assert wb.sheetnames == ["CARGA", "LEEME"]
    texto = " ".join(str(c.value) for fila in wb["LEEME"].iter_rows() for c in fila if c.value)
    assert "tipo de cambio" in texto and "comprobante" in texto


# --------------------------------------------------------------------------- datos ficticios
def cuentas_reales():
    texto = (REPO / "p9/reversion/powerapps/Main_Screen.yaml").read_text(encoding="utf-8")
    return set(re.findall(r'Cuenta: "([^"]+)"', texto))


def test_hay_13_cuentas_reales_de_referencia_y_ninguna_aparece_en_los_xlsx():
    reales = cuentas_reales()
    assert len(reales) == 13
    for ruta in sorted(XLSX.glob("*.xlsx")):
        wb = load_workbook(ruta)
        valores = {str(c.value) for hoja in wb.worksheets for fila in hoja.iter_rows() for c in fila if c.value is not None}
        assert not (valores & reales), ruta.name


def test_ninguna_cuenta_real_en_el_resto_de_archivos_del_prototipo():
    reales = cuentas_reales()
    # Main_Screen_CON_BOTON… es una copia del Main_Screen real (que ya contiene las cuentas en su selector): se excluye.
    for ruta in RAIZ.rglob("*"):
        if ruta.suffix in (".md", ".py", ".json", ".yaml", ".txt") and ruta.name != "Main_Screen_CON_BOTON_IMPORTACION_MASIVA.yaml" \
                and "__pycache__" not in ruta.parts:
            texto = ruta.read_text(encoding="utf-8")
            assert not [c for c in reales if c in texto], ruta.name


@pytest.mark.parametrize("nombre", PLANTILLAS)
def test_los_ejemplos_son_inventados(nombre):
    _, ws = cargar(nombre)
    for fila in ws.iter_rows(min_row=2, max_row=4, values_only=True):
        if fila[0]:
            assert fila[1].startswith(("999", "0999", "9999")) and "FICTICIO" in fila[5] and "FICTICIO" in fila[6]


# --------------------------------------------------------------------------- determinismo
def test_las_plantillas_versionadas_son_exactamente_la_salida_del_generador():
    assert (XLSX / PLANTILLAS[0]).read_bytes() == G.construir(True)
    assert (XLSX / PLANTILLAS[1]).read_bytes() == G.construir(False)
    assert G.construir(True) == G.construir(True)


# --------------------------------------------------------------------------- fixtures de escenarios (A-E)
@pytest.mark.parametrize("nombre", FIXTURES)
def test_fixtures_tienen_9_columnas_sin_TIPO_CAMBIO(nombre):
    _, ws = cargar(nombre)
    encabezados = [c.value for c in ws[1] if c.value is not None]
    assert len(encabezados) == 9 and "TIPO_CAMBIO" not in encabezados
    if nombre.startswith("05_"):
        assert encabezados[1] == "CUENTA" and encabezados != NUEVE  # única variante intencionalmente distinta
    else:
        assert encabezados == NUEVE
