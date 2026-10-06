"""Archivos y protocolo de MEDICIÓN en el tenant: los archivos son válidos y ficticios; el documento no contiene medidas inventadas."""
import hashlib
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from proto_masiva import catalogo_p9 as K
from proto_masiva import contrato_plantilla as P
from proto_masiva import generar_medicion as M
from proto_masiva import generar_plantillas_produccion as G

RAIZ = Path(__file__).resolve().parents[1]
CARPETA = RAIZ / "xlsx" / "medicion"
CAT = K.extraer()
TAMANOS = (10, 50, 100, 250, 500, 1000, 2000)


def test_hay_exactamente_los_siete_archivos_de_la_escala():
    assert M.TAMANOS == TAMANOS
    assert sorted(p.name for p in CARPETA.glob("*.xlsx")) == [f"Filas_{n:04d}.xlsx" for n in TAMANOS]


@pytest.mark.parametrize("n", TAMANOS)
def test_cada_archivo_tiene_n_filas_validas_y_ficticias_con_las_9_columnas(n):
    ws = load_workbook(CARPETA / f"Filas_{n:04d}.xlsx")[P.HOJA_CARGA]
    assert tuple(c.value for c in ws[G.FILA_ENC] if c.value is not None) == P.ENCABEZADOS and list(ws.tables) == [P.NOMBRE_TABLA]
    filas = [[c.value for c in ws[r][:9]] for r in range(G.PRIMERA, G.PRIMERA + n)]
    assert ws.tables[P.NOMBRE_TABLA].ref == f"A{G.FILA_ENC}:I{G.FILA_ENC + n}" and all(f[0] for f in filas)
    assert not any(c.value for c in ws[G.PRIMERA + n][:9])  # y ninguna fila más
    validas = {(c["banco"], c["cuenta"], c["moneda"]) for c in CAT["cuentas"]}
    assert all((f[0], f[1], f[4]) in validas for f in filas)  # banco/cuenta/moneda del catálogo: cada fila pasa sus propias listas
    codigos = [f[2] for f in filas]
    assert codigos == [f"MEDICION-{i:05d}" for i in range(1, n + 1)] and len(set(codigos)) == n
    assert all(isinstance(f[3], (int, float)) and f[3] > 0 for f in filas)
    assert all("MEDICION" in f[5] and f[6] == "SOLICITANTE MEDICION" for f in filas)


def test_los_archivos_versionados_son_la_salida_del_generador_y_pesan_poco():
    for n in TAMANOS:
        ruta = CARPETA / f"Filas_{n:04d}.xlsx"
        assert ruta.read_bytes() == G.construir(False, CAT, filas=list(M.filas(n, CAT)))
        assert ruta.stat().st_size < 200_000  # son pequeños: lo que se mide es el flujo, no la subida
    assert hashlib.sha256((CARPETA / "Filas_0010.xlsx").read_bytes()).hexdigest() != hashlib.sha256((CARPETA / "Filas_0050.xlsx").read_bytes()).hexdigest()


def test_el_generador_rechaza_un_numero_de_filas_inconsistente():
    datos = (CARPETA / "Filas_0010.xlsx").read_bytes()
    with pytest.raises(AssertionError, match="se esperaban 11 filas"):
        G.verificar(datos, CAT, 11)
    G.verificar(datos, CAT, 10)


# --------------------------------------------------------------------------- el protocolo no trae medidas inventadas
DOC = (RAIZ / "MEDICION_TENANT.md").read_text(encoding="utf-8")


def test_el_documento_de_medicion_esta_sin_medir_y_no_inventa_limites():
    assert "NO MEDIDO" in DOC
    filas = [l for l in DOC.splitlines() if re.match(r"\| (10|50|100|250|500|1000|2000) \| NO MEDIDO", l)]
    assert len(filas) == 7  # una fila por tamaño y todas sin medir
    assert "No se ha supuesto ningún límite" in DOC and "fijar un máximo de filas por archivo" in DOC
    assert not re.search(r"\b\d+(?:[.,]\d+)?\s*(segundos|seg\b|minutos|min\b)", DOC.lower())  # ningún tiempo de espera inventado


def test_los_hashes_y_tamanos_del_documento_son_los_de_los_archivos():
    for n in TAMANOS:
        ruta = CARPETA / f"Filas_{n:04d}.xlsx"
        corto = hashlib.sha256(ruta.read_bytes()).hexdigest()[:12]
        assert re.search(rf"`{ruta.name}` \| {n} \| \d+ KB \| `{corto}…`", DOC), ruta.name


def test_el_protocolo_cubre_los_casos_pedidos():
    for caso in ("tabla vacía", "Sin tabla", "Encabezado cambiado", "Otra tabla", "Archivo abierto", "FLUJO_SIN_RESPUESTA", "TMP_",
                 "PARAM_MAX_FILAS", "infraestructura de lotes"):
        assert caso.lower() in DOC.lower() or caso in DOC, caso
