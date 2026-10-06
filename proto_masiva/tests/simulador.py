"""Tenant SIMULADO para el flujo directo P9_MASIVA_PROTO_PREVALIDAR.

NO es Power Automate, SharePoint ni Excel Online: es un intérprete WDL local (el de P8, solo lectura) más conectores falsos.
Prueba la LÓGICA del flujo y el CONTRATO de sus llamadas; no certifica nombres de operaciones, tiempos ni códigos HTTP reales
del conector Excel (supuestos marcados [SUPUESTO]). El reloj simulado avanza en cada lectura, solo para ejercitar el cálculo de tiempos.
"""
from __future__ import annotations

import base64
import io
import re
import uuid
from datetime import datetime, timedelta

from openpyxl import load_workbook

from p8.ensayo_wdl import EnsayoWDL, FalloConector
from proto_masiva.flows import construir as F

INICIO = datetime(2026, 10, 5, 12, 0, 0)
ISO = "%Y-%m-%dT%H:%M:%S.%fZ"


class TenantSimulado:
    """Una biblioteca de documentos (copia temporal) y Excel Online. Nada más: el flujo no tiene listas."""

    def __init__(self):
        self.nombre = "TENANT_SIMULADO"
        self.biblioteca = {}      # id -> (carpeta, nombre, bytes)
        self.creados, self.borrados = [], []
        self.llamadas = []        # (operationId, acción)
        self.fallos = {}          # {"crear" | "excel" | "borrar": (estado, http)}
        self.bloqueados = set()   # nombres de copia que Excel ve bloqueados (423)

    def operacion(self, operacion, parametros, nombre_accion, accion):
        self.llamadas.append((operacion, nombre_accion))
        clave = {"Crear_archivo": "crear", "Leer_tabla_Excel": "excel", "Borrar_copia_temporal": "borrar"}.get(nombre_accion)
        if clave in self.fallos:
            raise FalloConector(*self.fallos[clave])
        if operacion == "CreateFile":
            assert parametros["folderPath"] == F.CARPETA_TEMP and re.fullmatch(r"TMP_[0-9a-f-]{36}\.xlsx", parametros["name"])
            assert parametros["body"]["$content-type"]
            ident = f"file-{uuid.uuid4().hex[:8]}"
            self.biblioteca[ident] = (parametros["folderPath"], parametros["name"], base64.b64decode(parametros["body"]["$content"]))
            self.creados.append(ident)
            return {"statusCode": 201, "body": {"Id": ident, "Name": parametros["name"]}}
        if operacion == "GetItems":
            return self._excel(parametros, accion)
        if operacion == "DeleteFile":
            if parametros["id"] not in self.biblioteca:
                raise FalloConector("Failed", 404)
            del self.biblioteca[parametros["id"]]
            self.borrados.append(parametros["id"])
            return {"statusCode": 200, "body": {}}
        raise AssertionError(f"operación no prevista en el prototipo: {operacion}")

    def _excel(self, p, accion):
        """Imita «Enumerar filas presentes en una tabla».
        [SUPUESTO] tabla inexistente -> 404; archivo bloqueado -> 423; celdas vacías -> ""; claves = encabezados;
        con paginación, devuelve como máximo el umbral configurado."""
        carpeta, nombre, datos = self.biblioteca[p["file"]]
        assert p["table"] == F.TABLA
        if nombre in self.bloqueados:
            raise FalloConector("Failed", 423)
        try:
            libro = load_workbook(io.BytesIO(datos))
        except Exception as e:
            raise FalloConector("Failed", 400) from e
        umbral = accion["runtimeConfiguration"]["paginationPolicy"]["minimumItemCount"]
        for hoja in libro.worksheets:
            for tabla in hoja.tables.values():
                if tabla.displayName == p["table"]:
                    celdas = list(hoja[tabla.ref])
                    cabeceras = [str(c.value) for c in celdas[0]]
                    filas = [{h: ("" if c.value is None else c.value) for h, c in zip(cabeceras, fila)} for fila in celdas[1:]]
                    return {"statusCode": 200, "body": {"value": [{**f, "@odata.etag": "", "ItemInternalId": str(uuid.uuid4())}
                                                                    for f in filas[:umbral]]}}
        raise FalloConector("Failed", 404)


class EnsayoDirecto(EnsayoWDL):
    def __init__(self, definicion, tenant, nombre, contenido, paso_ms=7):
        super().__init__(definicion, tenant, scope_estricto=False)
        archivo = {} if nombre is None else {"name": nombre, "contentBytes": base64.b64encode(contenido or b"").decode()}
        self.trigger = {"body": {"file": archivo} if archivo else {}}
        self.reloj, self.paso = INICIO, timedelta(milliseconds=paso_ms)
        self.respuesta = None

    def _ahora(self):
        self.reloj += self.paso
        return self.reloj.strftime(ISO)

    def _nodo(self, n):
        if n[0] == "funcion" and n[1] in ("toLower", "trim", "first", "createArray", "take", "utcNow", "ticks", "div",
                                          "greaterOrEquals", "base64ToBinary", "guid"):
            args = [self._nodo(a) for a in n[2]]
            f = n[1]
            if f == "toLower":
                return args[0].lower()
            if f == "trim":
                return args[0].strip()
            if f == "first":
                return args[0][0] if args[0] else None
            if f == "createArray":
                return list(args)
            if f == "take":
                return args[0][:args[1]]
            if f == "utcNow":
                return self._ahora()
            if f == "ticks":
                return int((datetime.strptime(args[0], ISO) - datetime(1, 1, 1)).total_seconds() * 10_000_000)
            if f == "div":
                return args[0] // args[1]
            if f == "greaterOrEquals":
                return args[0] >= args[1]
            if f == "base64ToBinary":
                return {"$content-type": "application/octet-stream", "$content": args[0]}
            if f == "guid":
                return str(uuid.uuid4())
        return super()._nodo(n)

    def _accion(self, nombre, accion):
        if accion["type"] == "Response":
            entradas = self.evaluar(accion["inputs"])
            assert entradas["statusCode"] == 200 and accion["kind"] == "PowerApp"
            assert set(entradas["body"]) == set(F.SALIDAS) and all(isinstance(v, str) for v in entradas["body"].values())
            self.respuesta = entradas["body"]
            self.salidas[nombre] = {"statusCode": 200}
            return "Succeeded"
        if accion["type"] == "OpenApiConnection":
            entradas = self.evaluar(accion["inputs"])
            self.salidas[nombre] = self.sp.operacion(entradas["host"]["operationId"], entradas["parameters"], nombre, accion)
            return "Succeeded"
        return super()._accion(nombre, accion)

    def ejecutar(self):
        self.estado_final = self._bloque(self.definicion["actions"])
        return self


def configurar(definicion, ubicacion="me", biblioteca="b!BIBLIOTECA-FICTICIA"):
    """Sustituye los marcadores <CONFIGURAR_...> como haría el usuario en el diseñador."""
    import json
    texto = json.dumps(definicion).replace(F.PLACEHOLDER_UBICACION, ubicacion).replace(F.PLACEHOLDER_BIBLIOTECA, biblioteca)
    return json.loads(texto)


def procesar(nombre, contenido, tenant=None, definicion=None, paso_ms=7):
    """Ejecuta el flujo como lo haría Power Apps: con un archivo {name, contentBytes} (o sin archivo si nombre es None)."""
    tenant = tenant or TenantSimulado()
    definicion = definicion or configurar(F.construir_definicion())
    ensayo = EnsayoDirecto(definicion, tenant, nombre, contenido, paso_ms).ejecutar()
    return tenant, ensayo
