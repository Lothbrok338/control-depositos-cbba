"""Plantillas definitivas de trabajo: Plantilla (vacía) y Ejemplo (3 filas ficticias).

Las comprobaciones leen los XLSX con openpyxl de forma INDEPENDIENTE del generador (que además se verifica a sí mismo antes de
guardar). Los bancos/cuentas esperados se contrastan con el texto del Main_Screen del checkpoint 3b407e2 leído desde git.
Las pruebas con LibreOffice evalúan las fórmulas de las listas; NO son Excel y se omiten si LibreOffice Calc no está instalado.
"""
import csv
import hashlib
import io
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from proto_masiva import catalogo_p9 as K
from proto_masiva import contrato_plantilla as P
from proto_masiva import generar_plantillas_produccion as G

RAIZ = Path(__file__).resolve().parents[1]
REPO = RAIZ.parent
XLSX = RAIZ / "xlsx"
RUTA = {"plantilla": XLSX / "Plantilla_Confirmacion_Masiva_P9.xlsx", "ejemplo": XLSX / "Ejemplo_Confirmacion_Masiva_P9.xlsx"}

# Oráculo independiente del extractor: bancos y las 13 cuentas del Main_Screen validado, en el orden de la fuente.
BANCOS = ["BNB", "BCP", "BISA", "BANCO UNIÓN", "BANCO ECONÓMICO", "BMSC"]
CUENTAS = [("BNB", "3000100152"), ("BNB", "3400041236"), ("BNB", "3501936692"), ("BNB", "3000100705"),
           ("BCP", "301-5005684-3-97"), ("BCP", "301-5005425-2-71"), ("BISA", "0696870039"), ("BISA", "0696872023"),
           ("BANCO UNIÓN", "10000003224552"), ("BANCO UNIÓN", "20000003224544"),
           ("BANCO ECONÓMICO", "3041210569"), ("BANCO ECONÓMICO", "3051446946"), ("BMSC", "1000872489")]
NUEVE = ["BANCO", "CUENTA_BANCARIA", "CODIGO_ASIGNACION", "IMPORTE", "MONEDA", "ESTUDIANTE", "SOLICITADO_POR", "SEDE", "OBSERVACION"]
CATALOGO = K.extraer()


def cargar(clave):
    wb = load_workbook(RUTA[clave])
    return wb, wb[P.HOJA_CARGA]


def filas_con_datos(ws):
    fin = int(re.sub(r"\D", "", ws.tables[P.NOMBRE_TABLA].ref.split(":")[1]))
    return [[c.value for c in ws[r][:9]] for r in range(G.PRIMERA, fin + 1) if any(c.value not in (None, "") for c in ws[r][:9])]


# --------------------------------------------------------------------------- archivos
def test_nombres_de_archivo_exactos_y_el_antiguo_VACIA_ya_no_existe():
    assert {p.name for p in RUTA.values()} == {"Plantilla_Confirmacion_Masiva_P9.xlsx", "Ejemplo_Confirmacion_Masiva_P9.xlsx"}
    assert all(p.is_file() and p.suffix == ".xlsx" for p in RUTA.values())
    assert not (XLSX / "Plantilla_Confirmacion_Masiva_P9_VACIA.xlsx").exists()


@pytest.mark.parametrize("clave", RUTA)
def test_abre_con_openpyxl_y_es_xlsx_estandar_sin_macros_ni_vinculos(clave):
    z = zipfile.ZipFile(RUTA[clave])
    assert z.testzip() is None
    nombres = z.namelist()
    assert not [n for n in nombres if "vba" in n.lower() or n.lower().endswith(".bin") or n.startswith("xl/externalLinks")]
    assert "macroEnabled" not in z.read("[Content_Types].xml").decode()
    assert b"<connections" not in b"".join(z.read(n) for n in nombres)
    wb = load_workbook(RUTA[clave])  # sin keep_vba: un .xlsx
    assert wb.sheetnames == ["CARGA", "_CATALOGOS"] and wb["_CATALOGOS"].sheet_state == "hidden" and wb.active.title == "CARGA"


# --------------------------------------------------------------------------- tabla y columnas
@pytest.mark.parametrize("clave", RUTA)
def test_tabla_unica_con_nombre_exacto_9_columnas_en_orden_y_filtros(clave):
    _, ws = cargar(clave)
    assert list(ws.tables) == ["tblConfirmacionMasiva"] == [P.NOMBRE_TABLA]
    assert [c.value for c in ws[G.FILA_ENC] if c.value is not None] == NUEVE == list(P.ENCABEZADOS)
    tabla = ws.tables["tblConfirmacionMasiva"]
    assert tabla.ref.startswith(f"A{G.FILA_ENC}:I") and tabla.autoFilter is not None
    assert [c.name for c in tabla.tableColumns] == NUEVE


@pytest.mark.parametrize("clave", RUTA)
def test_TIPO_CAMBIO_no_existe_en_ninguna_parte(clave):
    z = zipfile.ZipFile(RUTA[clave])
    assert b"TIPO_CAMBIO" not in b"".join(z.read(n) for n in z.namelist())


@pytest.mark.parametrize("clave", RUTA)
def test_titulo_instruccion_congelado_y_obligatorias_visibles(clave):
    _, ws = cargar(clave)
    assert ws["A1"].value == "P9 — CONFIRMACIÓN MASIVA DE DEPÓSITOS"
    assert "Complete una fila por depósito. No modifique los encabezados." in ws["A2"].value
    assert ws.freeze_panes == f"A{G.PRIMERA}"
    for celda in ws[G.FILA_ENC][:9]:
        esperado = "00" + (G.GRANATE if celda.value != "OBSERVACION" else G.GRIS)
        assert celda.fill.fgColor.rgb in (esperado, "FF" + (G.GRANATE if celda.value != "OBSERVACION" else G.GRIS)), celda.value
    assert P.OPCIONALES == ("OBSERVACION",) and len(P.OBLIGATORIAS) == 8
    assert all(ws.column_dimensions[c].width >= 11 for c in "ABCDEFGHI")


# --------------------------------------------------------------------------- formatos
@pytest.mark.parametrize("clave", RUTA)
def test_formatos_texto_numero_y_validaciones_hasta_la_ultima_fila_preconfigurada(clave):
    _, ws = cargar(clave)
    for r in (G.PRIMERA, G.PRIMERA + 2, 300, G.ULTIMA):
        assert ws[f"B{r}"].number_format == "@" and ws[f"C{r}"].number_format == "@"
        assert ws[f"D{r}"].number_format == "#,##0.00"
    val = {str(v.sqref): v for v in ws.data_validations.dataValidation}
    assert set(val) == {f"{c}{G.PRIMERA}:{c}{G.ULTIMA}" for c in "ABDEFGH"}  # CODIGO (C) y OBSERVACION (I): solo formato / libre
    banco, cuenta = val[f"A{G.PRIMERA}:A{G.ULTIMA}"], val[f"B{G.PRIMERA}:B{G.ULTIMA}"]
    assert (banco.type, banco.formula1) == ("list", "'_CATALOGOS'!$A$2:$A$7")
    assert cuenta.type == "list" and cuenta.formula1.startswith("IF(ISNUMBER(MATCH($A6,") and "OFFSET(" in cuenta.formula1
    assert "INDIRECT" not in cuenta.formula1 and "VBA" not in cuenta.formula1
    moneda, importe = val[f"E{G.PRIMERA}:E{G.ULTIMA}"], val[f"D{G.PRIMERA}:D{G.ULTIMA}"]
    assert (moneda.type, moneda.formula1) == ("list", '"BOB,USD"')
    assert (importe.type, importe.operator, importe.formula1) == ("decimal", "greaterThan", "0")
    for col in "FGH":
        v = val[f"{col}{G.PRIMERA}:{col}{G.ULTIMA}"]
        assert (v.type, v.operator, v.formula1, v.formula2) == ("textLength", "between", "1", "255")
    assert all(v.showErrorMessage and v.errorStyle == "stop" for v in val.values())
    assert "I" not in {k[0] for k in val}  # OBSERVACION: texto libre y opcional


@pytest.mark.parametrize("clave", RUTA)
def test_SEDE_es_texto_libre_porque_el_repo_no_tiene_catalogo_de_sedes(clave):
    _, ws = cargar(clave)
    v = next(v for v in ws.data_validations.dataValidation if str(v.sqref).startswith("H"))
    assert v.type == "textLength"  # no es una lista
    texto = (REPO / "DISENO_LISTA_DEPOSITOS_ACTIVOS.md").read_text(encoding="utf-8")
    assert "Lista de sedes no definida" in texto


# --------------------------------------------------------------------------- catálogos
@pytest.mark.parametrize("clave", RUTA)
def test_catalogo_oculto_con_6_bancos_y_13_cuentas_exactas_sin_duplicados(clave):
    wb, _ = cargar(clave)
    cat = wb[P.HOJA_CATALOGOS]
    bancos = [c.value for c in cat["A"][1:] if c.value]
    cuentas = [(cat[f"E{r}"].value, cat[f"F{r}"].value) for r in range(2, cat.max_row + 1) if cat[f"F{r}"].value]
    assert bancos == BANCOS and "(Todos)" not in bancos and len(bancos) == 6
    assert cuentas == CUENTAS and len(cuentas) == 13 and len({c for _, c in cuentas}) == 13
    assert [c.value for c in cat["C"][1:] if c.value] == ["BOB", "USD"]
    assert all(cat[f"F{r}"].number_format == "@" for r in range(2, 15))


def test_el_oraculo_coincide_con_el_texto_del_Main_Screen_del_checkpoint_en_git():
    fuente = subprocess.run(["git", "show", f"{K.CHECKPOINT_COMMIT}:p9/reversion/powerapps/Main_Screen.yaml"], cwd=REPO,
                            capture_output=True, text=True, check=True).stdout
    for banco in BANCOS:
        assert f'"{banco}"' in fuente
    for banco, cuenta in CUENTAS:
        assert f'Cuenta: "{cuenta}"' in fuente
    assert len(re.findall(r'Cuenta: "[^"]+"', fuente)) == 13  # exactamente 13 cuentas no vacías en el selector
    assert 'Table({Label: "", Cuenta: ""})' in fuente


def test_catalogo_extraido_y_json_generado_coinciden_y_traen_la_procedencia():
    assert CATALOGO["bancos"] == BANCOS
    assert [(c["banco"], c["cuenta"]) for c in CATALOGO["cuentas"]] == CUENTAS
    assert CATALOGO["fuente"]["checkpoint_commit"] == "3b407e202ca53e8827151d9e03a41c300bb8d84a"
    import json
    assert json.loads((RAIZ / "catalogo_bancos_p9.json").read_text(encoding="utf-8")) == CATALOGO
    monedas = {c["cuenta"]: c["moneda"] for c in CATALOGO["cuentas"]}
    assert monedas["3400041236"] == "USD" and monedas["301-5005684-3-97"] == "BOB" and sum(m == "USD" for m in monedas.values()) == 4


def test_el_extractor_falla_si_la_fuente_cambia(tmp_path):
    for rel in (K.FUENTE_MAIN_SCREEN, K.FUENTE_MONEDA):
        destino = tmp_path / rel
        destino.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(REPO / rel, destino)
    K.extraer(tmp_path)  # copia exacta: funciona
    ms = tmp_path / K.FUENTE_MAIN_SCREEN
    ms.write_text(ms.read_text(encoding="utf-8").replace("3000100152", "3000100153"), encoding="utf-8")
    with pytest.raises(K.CatalogoInvalido):
        K.extraer(tmp_path)


def test_ni_el_generador_ni_el_flujo_repiten_los_numeros_de_cuenta():
    numeros = [c for _, c in CUENTAS]
    for ruta in list(RAIZ.glob("*.py")) + list((RAIZ / "flows").glob("*.py")) + list((RAIZ / "powerapps").glob("*.py")):
        texto = ruta.read_text(encoding="utf-8")
        assert not [n for n in numeros if n in texto], ruta.name


# --------------------------------------------------------------------------- Plantilla vacía vs Ejemplo
def test_plantilla_de_produccion_no_tiene_ninguna_fila_de_datos():
    _, ws = cargar("plantilla")
    assert ws.tables[P.NOMBRE_TABLA].ref == f"A{G.FILA_ENC}:I{G.FILA_ENC + 1}"  # una tabla de Excel conserva 1 fila: en blanco
    assert filas_con_datos(ws) == []
    celdas = [c.value for fila in ws.iter_rows(min_row=G.PRIMERA, max_row=G.ULTIMA, max_col=9) for c in fila]
    assert all(v is None for v in celdas)


def test_ejemplo_tiene_3_filas_ficticias_validas_y_solo_el_ejemplo_las_tiene():
    _, ws = cargar("ejemplo")
    filas = filas_con_datos(ws)
    assert len(filas) == 3 and ws.tables[P.NOMBRE_TABLA].ref == f"A{G.FILA_ENC}:I{G.FILA_ENC + 3}"
    por_cuenta = {c["cuenta"]: c for c in CATALOGO["cuentas"]}
    for i, f in enumerate(filas, 1):
        banco, cuenta, codigo, importe, moneda, estudiante, solicitante, sede, obs = f
        assert por_cuenta[cuenta]["banco"] == banco and por_cuenta[cuenta]["moneda"] == moneda  # la fila pasa sus propias listas
        assert codigo == f"EJEMPLO-{i:04d}" and isinstance(cuenta, str) and isinstance(codigo, str)
        assert "EJEMPLO" in estudiante and "EJEMPLO" in solicitante and "EJEMPLO" in obs
        assert isinstance(importe, (int, float)) and importe > 0 and moneda in ("BOB", "USD")
    assert "FICTICIOS" in ws["A2"].value


def test_los_ejemplos_no_traen_nombres_reales_ni_codigos_que_parezcan_transacciones():
    _, ws = cargar("ejemplo")
    for f in filas_con_datos(ws):
        assert f[2].startswith("EJEMPLO-") and f[5].startswith("ESTUDIANTE EJEMPLO") and f[6] == "SOLICITANTE EJEMPLO"


# --------------------------------------------------------------------------- generador: reproducible y auto-verificado
def test_los_archivos_versionados_son_exactamente_la_salida_del_generador_y_es_determinista():
    assert RUTA["plantilla"].read_bytes() == G.construir(False, CATALOGO) == G.construir(False, CATALOGO)
    assert RUTA["ejemplo"].read_bytes() == G.construir(True, CATALOGO) == G.construir(True, CATALOGO)
    assert hashlib.sha256(RUTA["plantilla"].read_bytes()).hexdigest() != hashlib.sha256(RUTA["ejemplo"].read_bytes()).hexdigest()


def _mutar(clave, cambio):
    wb = load_workbook(RUTA[clave])
    cambio(wb, wb[P.HOJA_CARGA], wb[P.HOJA_CATALOGOS])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _sin_validacion(col):
    def f(wb, ws, cats):
        ws.data_validations.dataValidation = [v for v in ws.data_validations.dataValidation if not str(v.sqref).startswith(col)]
    return f


MUTANTES = {
    "encabezado renombrado": lambda wb, ws, c: ws.cell(G.FILA_ENC, 2, "CUENTA"),
    "columna TIPO_CAMBIO": lambda wb, ws, c: ws.cell(G.FILA_ENC, 10, "TIPO_CAMBIO"),
    "cuenta como numero": lambda wb, ws, c: setattr(ws[f"B{G.PRIMERA}"], "number_format", "General"),
    "codigo como numero": lambda wb, ws, c: setattr(ws[f"C{G.PRIMERA}"], "number_format", "0"),
    "sin lista de banco": _sin_validacion("A"),
    "sin lista de moneda": _sin_validacion("E"),
    "sin lista de cuenta": _sin_validacion("B"),
    "tabla renombrada (displayName)": lambda wb, ws, c: setattr(ws.tables[P.NOMBRE_TABLA], "displayName", "Tabla1"),
    "tabla renombrada (name)": lambda wb, ws, c: setattr(ws.tables[P.NOMBRE_TABLA], "name", "Tabla1"),
    "columna de tabla renombrada": lambda wb, ws, c: setattr(ws.tables[P.NOMBRE_TABLA].tableColumns[1], "name", "CUENTA"),
    "cuenta duplicada": lambda wb, ws, c: c.__setitem__("F3", c["F2"].value),
    "banco (Todos)": lambda wb, ws, c: c.__setitem__("A8", "(Todos)"),
    "cuenta de menos": lambda wb, ws, c: c.__setitem__("F14", None),
    "catalogo visible": lambda wb, ws, c: setattr(c, "sheet_state", "visible"),
}


@pytest.mark.parametrize("nombre", MUTANTES)
def test_la_verificacion_del_generador_rechaza_plantillas_defectuosas(nombre):
    with pytest.raises(AssertionError, match="Plantilla inválida"):
        G.verificar(_mutar("plantilla", MUTANTES[nombre]), CATALOGO, False)


def test_la_verificacion_rechaza_filas_en_la_plantilla_de_produccion_y_un_ejemplo_incompleto():
    with pytest.raises(AssertionError, match="trae filas con datos"):
        G.verificar(RUTA["ejemplo"].read_bytes(), CATALOGO, False)
    with pytest.raises(AssertionError, match="3 filas"):
        G.verificar(RUTA["plantilla"].read_bytes(), CATALOGO, True)


# --------------------------------------------------------------------------- LibreOffice (evalúa fórmulas; NO es Excel)
def _calc_disponible(tmp_path):
    if not shutil.which("soffice"):
        return False
    p = tmp_path / "t.csv"
    p.write_text("a\n1\n")
    r = subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "o"), str(p)],
                       capture_output=True, text=True, env={"HOME": str(tmp_path), "PATH": "/usr/bin:/bin"}, timeout=120)
    return (tmp_path / "o" / "t.xlsx").exists()


def _a_csv(tmp_path, libro: Path):
    r = subprocess.run(["soffice", "--headless", "--convert-to", "csv", "--outdir", str(tmp_path / "csv"), str(libro)],
                       capture_output=True, text=True, env={"HOME": str(tmp_path), "PATH": "/usr/bin:/bin"}, timeout=120)
    salida = tmp_path / "csv" / (libro.stem + ".csv")
    assert salida.exists(), r.stdout + r.stderr
    return list(csv.reader(salida.open(encoding="utf-8")))


@pytest.fixture
def calc(tmp_path):
    if not _calc_disponible(tmp_path):
        pytest.skip("LibreOffice Calc no está instalado")
    return tmp_path


@pytest.mark.parametrize("clave", RUTA)
def test_libreoffice_abre_ambos_archivos(calc, clave):
    filas = _a_csv(calc, RUTA[clave])
    assert filas[0][0] == "P9 — CONFIRMACIÓN MASIVA DE DEPÓSITOS" and [c for c in filas[G.FILA_ENC - 1] if c] == NUEVE


def test_libreoffice_evalua_la_lista_dependiente_de_cuentas_por_banco(calc):
    wb = Workbook()
    ws = wb.active
    ws.title = "PRUEBA"
    G.hoja_catalogos(wb, CATALOGO)
    casos = BANCOS + ["", "NO EXISTE"]
    for i, banco in enumerate(casos, 1):
        f = G.formula_cuentas(CATALOGO, f"$A{i}")
        ws[f"A{i}"] = banco
        ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"] = f"=COUNTA({f})", f"=INDEX({f},1)", f"=INDEX({f},COUNTA({f}))"
    libro = calc / "formula.xlsx"
    wb.save(libro)
    filas = _a_csv(calc, libro)
    esperado = {b: [c for bb, c in CUENTAS if bb == b] for b in BANCOS}
    for banco, fila in zip(BANCOS, filas):
        assert int(fila[1]) == len(esperado[banco]) and fila[2] == esperado[banco][0] and fila[3] == esperado[banco][-1], banco
    assert [int(f[1]) for f in filas[6:8]] == [13, 13]  # banco vacío o desconocido: se ofrecen las 13


def test_libreoffice_evalua_las_reglas_de_ayuda_visual(calc):
    _, ws = cargar("plantilla")
    reglas = {str(r.sqref): rule.formula[0] for r in ws.conditional_formatting for rule in r.rules}
    assert len(reglas) == 3
    obligatoria, cuenta_ajena, moneda_ajena = (reglas[k] for k in (f"A{G.PRIMERA}:H{G.ULTIMA}", f"B{G.PRIMERA}:B{G.ULTIMA}",
                                                                   f"E{G.PRIMERA}:E{G.ULTIMA}"))
    # cada fórmula (escrita para la celda superior izquierda de su rango, fila PRIMERA) se evalúa en K1/L1/M1 de un libro por caso
    casos = [  # BANCO, CUENTA, CODIGO, IMPORTE, MONEDA, esperado (obligatoria A6, cuenta ajena, moneda ajena)
        ("BNB", "3000100152", "X", 1, "BOB", (False, False, False)),
        ("BNB", "301-5005684-3-97", "X", 1, "BOB", (False, True, False)),
        ("BNB", "3400041236", "X", 1, "BOB", (False, False, True)),
        ("", "3000100152", "X", 1, "BOB", (True, False, False)),
        ("", "", "", "", "", (False, False, False)),
    ]
    for i, (b, c, k, imp, m, esp) in enumerate(casos):
        w = Workbook()
        h = w.active
        h.title = "PRUEBA"
        G.hoja_catalogos(w, CATALOGO)
        for col, v in zip("ABCDE", (b, c, k, imp, m)):
            h[f"{col}{G.PRIMERA}"] = v if v != "" else None
        h["K1"], h["L1"], h["M1"] = f"={obligatoria}", f"={cuenta_ajena}", f"={moneda_ajena}"
        libro = calc / f"regla{i}.xlsx"
        w.save(libro)
        fila = _a_csv(calc, libro)[0]
        assert [x.upper() for x in fila[10:13]] == [str(v).upper() for v in esp], (b, c, m, fila[10:13])


# --------------------------------------------------------------------------- documentación
def test_documentacion_de_la_descarga_desde_power_apps():
    texto = (RAIZ / "PLANTILLA_PRODUCCION.md").read_text(encoding="utf-8")
    for fragmento in ("Plantilla_Confirmacion_Masiva_P9.xlsx", "download.aspx?UniqueId=84b7ef88-43aa-43d2-8d1c-0f0b682dafbd", "btnDescargarPlantillaP9", "P9_Confirmacion_Masiva",
                      "PASOS PARA GABRIEL", "tblConfirmacionMasiva"):
        assert fragmento in texto, fragmento
    assert texto.count("\n1. ") + texto.count("\n2. ") >= 2
