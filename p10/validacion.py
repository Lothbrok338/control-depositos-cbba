# -*- coding: utf-8 -*-
"""
p10/validacion.py · P10-A.1 · validación del .xlsx YA ESCRITO, antes de aceptarlo.

Reabre el archivo con openpyxl y lo compara, celda por celda, contra las FUENTES (no contra las estructuras
internas que lo produjeron):
  * LISTS.csv de P0 (importes, débitos, créditos, saldos, fechas, horas, claves, textos de las 26 columnas),
  * ORIGEN.xlsx de P0 (texto original de cada columna propia del banco),
  * el snapshot (los 7 campos operativos, cruzados por CLAVE).
Devuelve la lista de errores (vacía = el archivo se acepta). Nunca corrige nada.
"""
import datetime as dt

import historico as H
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from . import contrato as C
from . import generador as G


def indice_origen(datos_originales, mapa_origen):
    """CLAVE -> ID_ORIGEN y celdas de movimiento por ID_ORIGEN, leídos de ORIGEN.xlsx (texto exacto)."""
    clave_a_origen = {H._texto(r.get(C.COLUMNA_CLAVE)).strip(): H._texto(r.get("ID_ORIGEN"))
                      for r in H._registros(mapa_origen)}
    movimientos = {}
    for r in H._registros(datos_originales):
        if H._texto(r.get("ROL_FILA")) == "MOVIMIENTO":
            col = column_index_from_string(H._texto(r.get("COLUMNA_EXCEL")))
            movimientos.setdefault(H._texto(r.get("ID_ORIGEN")), {})[col] = H._texto(r.get("VALOR_ORIGINAL"))
    return {"clave_a_origen": clave_a_origen, "movimientos": movimientos}


def _del_grupo(libro, filas_p0):
    out = {}
    for r in filas_p0:
        if (H._texto(r["BANCO"]).strip(), H._texto(r["CUENTA BANCARIA"]).strip(),
                H._texto(r["MONEDA"]).strip()) != (libro.banco, libro.cuenta, libro.moneda):
            continue
        if H._fecha(r["FECHA MOVIMIENTO"]).strftime("%Y-%m") != libro.periodo:
            continue
        out[r[C.COLUMNA_CLAVE]] = r
    return out


def _vacio(v):
    return v is None


def _cmp_numero(esperado, celda):
    e = H._numero(esperado)
    v = celda.value
    if e is None:
        return None if _vacio(v) else f"debía estar vacío y trae {v!r}"
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return f"no es numérico ({v!r}); esperado {e!r}"
    return None if float(v) == e else f"{v!r} distinto de {e!r}"


def _cmp_texto(esperado, celda, exigir_formato_texto=False):
    e = "" if esperado is None else str(esperado)
    v = celda.value
    if e == "":
        return None if _vacio(v) else f"debía estar vacío y trae {v!r}"
    if not isinstance(v, str) or celda.data_type != "s":
        return f"debía ser texto {e!r} y es {v!r} ({celda.data_type})"
    if v != e:
        return f"{v!r} distinto de {e!r}"
    if exigir_formato_texto and celda.number_format != "@":
        return f"formato {celda.number_format!r}; los códigos deben ir con formato texto '@'"
    return None


def _cmp_fecha(esperado, celda):
    e = H._fecha(esperado)
    v = celda.value
    if e is None:
        return None if _vacio(v) else f"debía estar vacío y trae {v!r}"
    ok = isinstance(v, dt.datetime) and v == dt.datetime(e.year, e.month, e.day)
    return None if ok else f"{v!r} distinto de {e}"


def _cmp_fecha_hora(esperado, celda):
    e = H._fecha_hora(esperado)          # mismo criterio que el libro: segundos (la lista guarda la marca del motor)
    v = celda.value
    if e is None:
        return None if _vacio(v) else f"debía estar vacío y trae {v!r}"
    return None if (isinstance(v, dt.datetime) and v == e) else f"{v!r} distinto de {e}"


def _cmp_hora(esperado, celda):
    e = H._hora(esperado)
    v = celda.value
    if e is None:
        return None if _vacio(v) else f"debía estar vacía y trae {v!r}"
    return None if (isinstance(v, dt.time) and v == e) else f"{v!r} distinto de {e}"


def _cmp_auditoria(col, esperado, celda):
    tipo = C.TIPO_AUDITORIA[col]
    if tipo == "importe":
        return _cmp_numero(esperado, celda)
    if tipo == "fecha":
        return _cmp_fecha(esperado, celda)
    if tipo in ("fecha_hora", "fecha_hora_seg"):
        return _cmp_fecha_hora(esperado, celda)
    return _cmp_texto(esperado, celda, exigir_formato_texto=(tipo == "codigo"))


def validar_libro(ruta, libro, filas_p0, snapshot, indice, registro, max_errores=30):
    """Lista de errores del archivo `ruta` respecto de sus fuentes; vacía = aceptable."""
    errores = []

    def err(msg):
        if len(errores) < max_errores:
            errores.append(msg)

    wb = load_workbook(ruta)
    if wb.sheetnames != list(C.HOJAS):
        return [f"hojas {wb.sheetnames}; deben ser exactamente {list(C.HOJAS)}"]
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            err(f"la hoja {ws.title} no está visible")

    esperadas = _del_grupo(libro, filas_p0)
    n = len(libro.claves)
    if len(set(libro.claves)) != n:
        err("CLAVE TRANSACCIÓN duplicada entre los movimientos del libro")
    if set(libro.claves) != set(esperadas):
        err(f"los movimientos del libro ({n}) no coinciden con los de P0 para esta cuenta y mes ({len(esperadas)})")

    # ------------------------------------------------------------ AUDITORIA
    wa = wb[C.HOJA_AUDITORIA]
    if [c.value for c in wa[1]] != list(C.COLUMNAS_AUDITORIA) or wa.max_column != len(C.COLUMNAS_AUDITORIA):
        err("AUDITORIA: las columnas no son las 26 de COLUMNAS_LISTS en su orden + las 2 reales de Depositos_Activos")
    filas_a = list(wa.iter_rows(min_row=2, max_col=len(C.COLUMNAS_AUDITORIA)))
    if len(filas_a) != n:
        err(f"AUDITORIA tiene {len(filas_a)} filas y el libro {n} movimientos")
    claves_a = [f[0].value for f in filas_a]
    if len(set(claves_a)) != len(claves_a):
        err("AUDITORIA: CLAVE TRANSACCIÓN duplicada")
    if claves_a != libro.claves:
        err("AUDITORIA: las claves (conjunto u orden) no coinciden con las de EXTRACTO")
    tabla_a = wa.tables.get(C.TABLA_AUDITORIA)
    if tabla_a is None or tabla_a.ref != f"A1:{get_column_letter(len(C.COLUMNAS_AUDITORIA))}{n + 1}":
        err("AUDITORIA: falta la tabla de Excel o su rango no cubre los datos")
    if tabla_a is not None and tabla_a.autoFilter is None:
        err("AUDITORIA: la tabla no tiene filtros")
    for ws in (wa, wb[C.HOJA_EXTRACTO]):
        if ws.freeze_panes is not None or ws.sheet_view.pane is not None:
            err(f"{ws.title}: no debe tener paneles inmovilizados")
    op_por_clave = {k: {**G.valores_operativos(p0, snapshot.get(k) if snapshot is not None else None),
                        **G.valores_adicionales(snapshot.get(k) if snapshot is not None else None)}
                    for k, p0 in esperadas.items()}
    for fila in filas_a:
        clave = fila[0].value
        p0, op = esperadas.get(clave), op_por_clave.get(clave)
        if p0 is None:
            continue
        for col, celda in zip(C.COLUMNAS_AUDITORIA, fila):
            esperado = op[col] if (col in C.OPERATIVOS or col in C.ADICIONALES_AUDITORIA) else p0[col]
            m = _cmp_auditoria(col, esperado, celda)
            if m:
                err(f"AUDITORIA {clave} · {col}: {m}")
    cred = sum((f[C.COLUMNAS_LISTS.index("CRÉDITO")].value or 0.0) for f in filas_a)
    deb = sum((f[C.COLUMNAS_LISTS.index("DÉBITO")].value or 0.0) for f in filas_a)
    cred_p0 = sum(H._numero(r["CRÉDITO"]) or 0.0 for r in esperadas.values())
    deb_p0 = sum(H._numero(r["DÉBITO"]) or 0.0 for r in esperadas.values())
    if abs(cred - cred_p0) > 1e-6 or abs(deb - deb_p0) > 1e-6:
        err(f"AUDITORIA: totales de créditos/débitos ({cred:.2f}/{deb:.2f}) distintos de P0 ({cred_p0:.2f}/{deb_p0:.2f})")

    # ------------------------------------------------------------ EXTRACTO
    we = wb[C.HOJA_EXTRACTO]
    tabla = we.tables.get(C.TABLA_EXTRACTO)
    if tabla is None:
        err("EXTRACTO: falta la tabla de Excel")
        return errores
    c0, r0, c1, r1 = range_boundaries(tabla.ref)
    columnas = libro.resultado["columnas"]
    titulos = [we.cell(r0, c).value for c in range(c0, c1 + 1)]
    if titulos != [c["titulo"] for c in columnas]:
        err("EXTRACTO: los encabezados de la tabla no son los del banco + operativos + CLAVE oculta")
    if titulos[-5:] != [*C.COLUMNAS_OPERATIVAS_EXTRACTO, C.COLUMNA_CLAVE]:
        err("EXTRACTO: al final deben ir ESTADO, CONFIRMADO POR, FECHA DE CONFIRMACIÓN, OBSERVACIONES y la CLAVE oculta")
    if r1 - r0 != n:
        err(f"EXTRACTO tiene {r1 - r0} filas y el libro {n} movimientos")
    if tabla.autoFilter is None:
        err("EXTRACTO: la tabla no tiene filtros")
    if not we.column_dimensions[get_column_letter(c1)].hidden:
        err("EXTRACTO: la columna CLAVE TRANSACCIÓN debe estar oculta")

    fmt_id = libro.formato
    sep = registro["FORMATOS"].get(fmt_id, {}).get("historico", {}).get("separador_decimal", ".")
    filas_e = list(we.iter_rows(min_row=r0 + 1, max_row=r1, min_col=c0, max_col=c1))
    for k, fila in enumerate(filas_e):
        celdas = dict(zip(titulos, fila))
        clave = celdas[C.COLUMNA_CLAVE].value
        if k >= n or clave != libro.claves[k]:
            err(f"EXTRACTO fila {k + 1}: CLAVE {clave!r} fuera de orden o ausente en AUDITORIA")
            continue
        p0, op = esperadas.get(clave), op_por_clave.get(clave)
        if p0 is None:
            continue
        for titulo, esperado, cmp in (
                ("ESTADO", op["ESTADO"], _cmp_texto), ("CONFIRMADO POR", op["CONFIRMADO POR"], _cmp_texto),
                ("FECHA DE CONFIRMACIÓN", op["FECHA CONFIRMACIÓN"], _cmp_fecha_hora),
                (C.COLUMNA_OBSERVACIONES, op["OBSERVACIÓN"], _cmp_texto)):
            m = cmp(esperado, celdas[titulo])
            if m:
                err(f"EXTRACTO {clave} · {titulo}: {m}")
        m = _cmp_texto(clave, celdas[C.COLUMNA_CLAVE])
        if m:
            err(f"EXTRACTO {clave} · CLAVE: {m}")
        mov = indice["movimientos"].get(indice["clave_a_origen"].get(clave, ""), {}) if indice else None
        for c in columnas:
            if c.get("bloque") != "BANCO":
                continue
            celda, fuente = celdas[c["titulo"]], c["fuente"]
            if fuente == "ORIGINAL":
                if mov is None:
                    continue
                m = _cmp_original(mov.get(c["col"], "").strip(), c["tipo"], sep, celda)
            else:
                m = _cmp_normalizado(fuente, p0, celda)
            if m:
                err(f"EXTRACTO {clave} · {c['titulo']}: {m}")
        if len(errores) >= max_errores:
            break
    return errores


def _cmp_original(texto, tipo, sep, celda):
    """Texto original del banco (recortados solo los espacios de relleno) -> celda del EXTRACTO."""
    if texto == "":
        return None if _vacio(celda.value) else f"debía estar vacío y trae {celda.value!r}"
    if tipo == "importe" and H._es_importe_puro(texto):
        valor = H._importe_de_texto(texto, sep)
        if valor is not None:
            return _cmp_numero(valor, celda)
    return _cmp_texto(texto, celda, exigir_formato_texto=(tipo == "codigo"))


def _cmp_normalizado(fuente, p0, celda):
    """Columnas que salen de P0 (NORMALIZADO:*): deben ser exactamente lo que entregó P0."""
    if fuente == "NORMALIZADO:FECHA":
        return _cmp_fecha(p0["FECHA MOVIMIENTO"], celda)
    if fuente == "NORMALIZADO:HORA":
        return _cmp_hora(p0["HORA MOVIMIENTO"], celda)
    if fuente == "NORMALIZADO:DEBITO":
        return _cmp_numero(p0["DÉBITO"], celda)
    if fuente == "NORMALIZADO:CREDITO":
        return _cmp_numero(p0["CRÉDITO"], celda)
    if fuente == "NORMALIZADO:SALDO":
        return _cmp_numero(p0["SALDO"], celda)
    if fuente == "NORMALIZADO:IMPORTE_FIRMADO":
        imp = H._numero(p0["IMPORTE"])
        if imp is not None and H._texto(p0["TIPO MOVIMIENTO"]).strip() == "DÉBITO":
            imp = -imp
        return _cmp_numero(imp, celda)
    return f"fuente desconocida {fuente!r}"
