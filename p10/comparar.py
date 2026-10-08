# -*- coding: utf-8 -*-
"""
p10/comparar.py · P10-A.1 · diferencias lógicas entre dos libros P10 (pruebas y evidencia).

Compara el CONTENIDO (zona superior, encabezados, claves y cada celda de las dos tablas), no los bytes.
Cada diferencia se atribuye a una CLAVE TRANSACCIÓN y a una columna, para poder afirmar "solo cambiaron
estos campos operativos de estos movimientos".
"""
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from . import contrato as C

_TABLAS = ((C.HOJA_EXTRACTO, C.TABLA_EXTRACTO, -1), (C.HOJA_AUDITORIA, C.TABLA_AUDITORIA, 0))


def _tabla(ws, nombre):
    c0, r0, c1, r1 = range_boundaries(ws.tables[nombre].ref)
    titulos = [ws.cell(r0, c).value for c in range(c0, c1 + 1)]
    filas = [[c.value for c in fila] for fila in ws.iter_rows(min_row=r0 + 1, max_row=r1, min_col=c0, max_col=c1)]
    zona = ([[c.value for c in fila] for fila in ws.iter_rows(min_row=1, max_row=r0 - 1)]
            if r0 > 1 else [])
    return titulos, filas, zona


def diferencias(ruta_a, ruta_b):
    """Lista de diferencias {hoja, tipo, clave, columna, antes, despues}; vacía = mismo contenido lógico."""
    a, b = load_workbook(ruta_a), load_workbook(ruta_b)
    out = []
    if a.sheetnames != b.sheetnames:
        return [{"hoja": None, "tipo": "HOJAS", "antes": a.sheetnames, "despues": b.sheetnames}]
    for hoja, tabla, idx in _TABLAS:
        ta, fa, za = _tabla(a[hoja], tabla)
        tb, fb, zb = _tabla(b[hoja], tabla)
        if ta != tb:
            out.append({"hoja": hoja, "tipo": "ENCABEZADOS", "antes": ta, "despues": tb})
            continue
        if za != zb:
            out.append({"hoja": hoja, "tipo": "ZONA_SUPERIOR", "antes": za, "despues": zb})
        if [f[idx] for f in fa] != [f[idx] for f in fb]:
            out.append({"hoja": hoja, "tipo": "CLAVES", "antes": len(fa), "despues": len(fb)})
            continue
        for fila_a, fila_b in zip(fa, fb):
            for j, (x, y) in enumerate(zip(fila_a, fila_b)):
                if x != y:
                    out.append({"hoja": hoja, "tipo": "CELDA", "clave": fila_a[idx], "columna": ta[j],
                                "antes": x, "despues": y})
    return out
