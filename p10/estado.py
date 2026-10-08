# -*- coding: utf-8 -*-
"""
p10/estado.py · P10-A.2 · ESTADO CANÓNICO por BANCO + CUENTA + MONEDA + MES.

El XLSX NO es la base de datos del histórico: es una proyección. La fuente de verdad es este estado estructurado, del que el
XLSX aprobado de A.1 se reconstruye completo y, con las mismas entradas, byte a byte.

    estado = {
      formato, version,
      grupo       {id, banco, cuenta, moneda, periodo, id_cuenta, formato_banco},
      extractos   [ {sha256, nombre, movimientos, ultimo_mov, procesado_en} ]      extractos bancarios ya incorporados
      cabecera    {sha256, ultimo_mov, movimientos, multi_mes, hoja, encabezado, zona}   cabecera/zona del extracto más completo
      movimientos { CLAVE: {p0: [26 textos de LISTS.csv], fila, celdas: {col: texto}} }  salida de P0 + celdas originales del banco
      operativo   {fecha_corte, filas: {CLAVE: campos de Depositos_Activos}} última información operativa conocida de cada movimiento
      xlsx        {nombre, sha256, bytes}                                                  último libro generado
      integridad  sha256 del resto del estado (detecta corrupción/alteración)
    }

Qué NO hay aquí: ninguna regla bancaria ni de CLAVE. `p0` es la fila de 26 columnas EXACTAMENTE como la entrega el motor P0.
La política de fusión de extractos acumulativos es la misma que ya aplica Depositos_Activos: la unidad es la CLAVE; un extracto
posterior que repite movimientos los reemplaza y uno más corto nunca borra movimientos ya incorporados.
"""
import gzip
import hashlib
import json
from datetime import datetime

from openpyxl.utils import column_index_from_string, get_column_letter

from . import contrato as C

FORMATO_ESTADO = "P10-ESTADO-1"
FORMATO_PARCIAL = "P10-PARCIAL-1"
# Campos de la lista que forman la «foto operativa» (los 9 de A.1 + los 3 de carga que la lista conserva).
CAMPOS_LISTA = tuple(C.CAMPOS_SNAPSHOT[1:]) + ("LOTE_CARGA", "FECHA_CARGA", "ARCHIVO_ORIGEN")
# Columnas de P0 que cambian en cada corrida del motor (hora de la ejecución): no cuentan como «cambio» de un movimiento.
VOLATILES_P0 = ("LOTE DE CARGA", "FECHA DE CARGA", "ARCHIVO ORIGEN")


class ErrorEstado(ValueError):
    """Fallo del estado canónico con un código estable (ESTADO_CORRUPTO, ESTADO_INCOMPATIBLE, PARCIAL_INVALIDO…)."""

    def __init__(self, codigo, mensaje):
        super().__init__(mensaje)
        self.codigo, self.mensaje = codigo, mensaje


# ------------------------------------------------------------------ serialización determinista
def json_canonico(obj):
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def sha256_texto(texto):
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def empaquetar(obj):
    """JSON canónico comprimido (gzip con mtime 0): mismo objeto -> mismos bytes."""
    return gzip.compress(json_canonico(obj).encode("utf-8"), compresslevel=6, mtime=0)


def desempaquetar(datos):
    try:
        if datos[:2] == b"\x1f\x8b":
            datos = gzip.decompress(datos)
        return json.loads(datos.decode("utf-8"))
    except (OSError, EOFError, ValueError, UnicodeDecodeError) as e:
        raise ErrorEstado("ESTADO_CORRUPTO", f"no se pudo leer el estado ({type(e).__name__})") from e


def sellar(estado):
    estado.pop("integridad", None)
    estado["integridad"] = sha256_texto(json_canonico(estado))
    return estado


def cargar_estado(datos, grupo_id=None):
    """Bytes -> estado verificado. Lanza ErrorEstado si está corrupto, alterado, de otro formato o de otro grupo."""
    estado = desempaquetar(datos)
    if not isinstance(estado, dict) or estado.get("formato") != FORMATO_ESTADO:
        raise ErrorEstado("ESTADO_INCOMPATIBLE", "el archivo no es un estado P10 de este formato")
    esperado = estado.get("integridad")
    resto = {k: v for k, v in estado.items() if k != "integridad"}
    if not esperado or sha256_texto(json_canonico(resto)) != esperado:
        raise ErrorEstado("ESTADO_CORRUPTO", "la huella de integridad del estado no coincide (archivo alterado o dañado)")
    if grupo_id is not None and estado["grupo"]["id"] != grupo_id:
        raise ErrorEstado("ESTADO_CORRUPTO", f"el estado pertenece a otro grupo ({estado['grupo']['id']})")
    return estado


def huella_estado(estado):
    return estado["integridad"]


# ------------------------------------------------------------------ grupo
def grupo_id(banco, cuenta, moneda, periodo):
    return f"{banco}|{cuenta}|{moneda}|{periodo}"


def grupo_desde_id(gid):
    banco, cuenta, moneda, periodo = gid.split("|")
    return {"id": gid, "banco": banco, "cuenta": cuenta, "moneda": moneda, "periodo": periodo}


def periodo_de_fecha(texto):
    """'2026-10-08' / datetime -> '2026-10' (mes de FECHA MOVIMIENTO)."""
    if isinstance(texto, datetime):
        return texto.strftime("%Y-%m")
    return str(texto)[:7]


# ------------------------------------------------------------------ parcial -> estado
def estado_vacio(grupo, hoja, id_cuenta, formato_banco):
    return {
        "formato": FORMATO_ESTADO, "version": 0,
        "grupo": dict(grupo, id_cuenta=id_cuenta, formato_banco=formato_banco),
        "hoja": hoja, "extractos": [], "cabecera": None, "movimientos": {},
        "operativo": {"fecha_corte": None, "filas": {}},
        "xlsx": None}


def _cobertura(c):
    return (c["ultimo_mov"] or "", c["movimientos"])


def _p0_igual(a, b):
    """Dos filas de P0 son «la misma» si coinciden en todo menos en lo que depende de la hora de la corrida."""
    idx = [i for i, col in enumerate(C.COLUMNAS_LISTS) if col not in VOLATILES_P0]
    return all(a[i] == b[i] for i in idx)


def aplicar_parcial(estado, parcial):
    """Incorpora un extracto al estado. Devuelve un dict con lo ocurrido; `cambio` indica si el estado quedó distinto."""
    if parcial.get("formato") != FORMATO_PARCIAL:
        raise ErrorEstado("PARCIAL_INVALIDO", "formato de parcial desconocido")
    if parcial["grupo"]["id"] != estado["grupo"]["id"]:
        raise ErrorEstado("PARCIAL_INVALIDO", "el parcial pertenece a otro grupo")
    ext = parcial["extracto"]
    if any(e["sha256"] == ext["sha256"] for e in estado["extractos"]):
        return {"cambio": False, "motivo": "EXTRACTO_YA_APLICADO", "nuevos": 0, "reemplazados": 0, "modificados": 0}
    nuevos = reemplazados = modificados = 0
    for clave, rec in parcial["movimientos"].items():
        previo = estado["movimientos"].get(clave)
        if previo is None:
            nuevos += 1
        else:
            reemplazados += 1
            if not _p0_igual(previo["p0"], rec["p0"]):
                modificados += 1
        estado["movimientos"][clave] = {"p0": rec["p0"], "fila": rec["fila"], "celdas": rec["celdas"]}
    cab = {"sha256": ext["sha256"], "ultimo_mov": ext["ultimo_mov"], "movimientos": ext["movimientos_extracto"],
           "multi_mes": ext["multi_mes"], "hoja": parcial["hoja"], "encabezado": parcial["encabezado"], "zona": parcial["zona"]}
    actual = estado["cabecera"]
    if actual is None or _cobertura(cab) >= _cobertura(actual):
        estado["cabecera"], estado["hoja"] = cab, parcial["hoja"]
    estado["extractos"].append({k: ext[k] for k in ("sha256", "nombre", "ultimo_mov", "procesado_en")}
                               | {"movimientos": len(parcial["movimientos"])})
    estado["extractos"].sort(key=lambda e: (e["procesado_en"], e["sha256"]))
    return {"cambio": True, "motivo": "EXTRACTO_APLICADO", "nuevos": nuevos, "reemplazados": reemplazados,
            "modificados": modificados}


# ------------------------------------------------------------------ operativo (foto de Depositos_Activos)
def hash_lista(filas):
    """Huella de las filas de la lista de UN grupo (independiente del orden). `filas`: [{CLAVE_TRANSACCION, campos…}]."""
    canon = sorted(({k: f.get(k) for k in ("CLAVE_TRANSACCION",) + CAMPOS_LISTA} for f in filas),
                   key=lambda f: f["CLAVE_TRANSACCION"])
    return sha256_texto(json_canonico(canon))


def aplicar_operativo(estado, filas, fecha_corte):
    """
    Fusiona las filas de la lista del grupo. Reglas:
      * una fila presente en la lista manda sobre lo guardado (incluye reversiones: vuelven a DISPONIBLE);
      * una fila AUSENTE de la lista NO se borra del estado (un registro eliminado de Depositos_Activos —P10-B— conserva su
        última información operativa);
      * solo se guardan filas de movimientos que ya existen en el estado (las demás se aplicarán cuando llegue su extracto).
    Devuelve {cambio, hash_lista, sin_movimiento}.
    """
    h = hash_lista(filas)
    op = estado["operativo"]
    cambio, sin_mov = False, 0
    for f in filas:
        clave = f["CLAVE_TRANSACCION"]
        if clave not in estado["movimientos"]:
            sin_mov += 1
            continue
        nueva = {k: f.get(k) for k in CAMPOS_LISTA}
        if op["filas"].get(clave) != nueva:
            op["filas"][clave] = nueva
            cambio = True
    if cambio:
        op["fecha_corte"] = fecha_corte
    return {"cambio": cambio, "hash_lista": h, "sin_movimiento": sin_mov}


# ------------------------------------------------------------------ estado -> entradas del generador aprobado (A.1)
def _id_extracto(estado):
    return sha256_texto("P10|" + estado["grupo"]["id"])


def entradas_generador(estado):
    """
    Reconstruye lo que A.1 espera: DATOS_ORIGINALES, MAPA_ORIGEN, filas de LISTS (26 columnas) y el snapshot.
    Las columnas de carga (LOTE/FECHA DE CARGA, ARCHIVO ORIGEN) se toman de la lista cuando el movimiento ya está en ella:
    así AUDITORIA conserva lo que realmente se persistió, no lo que produciría reprocesar el extracto hoy.
    """
    gid, ide, hoja = estado["grupo"]["id"], _id_extracto(estado), estado["hoja"]
    id_cuenta = estado["grupo"]["id_cuenta"]
    cab = estado["cabecera"]
    datos, mapa, filas = [], [], []
    base = {"ID_EXTRACTO": ide, "HOJA": hoja}
    for col, texto in sorted(cab["encabezado"].items(), key=lambda kv: int(kv[0])):
        datos.append({**base, "ID_ORIGEN": f"{ide}|ENC", "FILA_EXCEL": 0, "COLUMNA_EXCEL": get_column_letter(int(col)),
                      "ROL_FILA": "ENCABEZADO_TABLA", "VALOR_ORIGINAL": texto})
    for fila, col, rol, texto in cab["zona"]:
        datos.append({**base, "ID_ORIGEN": f"{ide}|Z{fila}", "FILA_EXCEL": fila, "COLUMNA_EXCEL": get_column_letter(col),
                      "ROL_FILA": rol, "VALOR_ORIGINAL": texto})
    op = estado["operativo"]["filas"]
    for clave in sorted(estado["movimientos"]):
        rec = estado["movimientos"][clave]
        id_origen = f"{ide}|{sha256_texto(clave)[:16]}"
        for col, texto in sorted(rec["celdas"].items(), key=lambda kv: int(kv[0])):
            datos.append({**base, "ID_ORIGEN": id_origen, "FILA_EXCEL": rec["fila"],
                          "COLUMNA_EXCEL": get_column_letter(int(col)), "ROL_FILA": "MOVIMIENTO", "VALOR_ORIGINAL": texto})
        fila_p0 = dict(zip(C.COLUMNAS_LISTS, rec["p0"]))
        lista = op.get(clave)
        if lista:
            fila_p0["LOTE DE CARGA"] = lista.get("LOTE_CARGA") or fila_p0["LOTE DE CARGA"]
            fila_p0["FECHA DE CARGA"] = lista.get("FECHA_CARGA") or fila_p0["FECHA DE CARGA"]
            fila_p0["ARCHIVO ORIGEN"] = lista.get("ARCHIVO_ORIGEN") or fila_p0["ARCHIVO ORIGEN"]
        filas.append(fila_p0)
        mapa.append({C.COLUMNA_CLAVE: clave, "ID_ORIGEN": id_origen, "ID_EXTRACTO": ide, "ARCHIVO ORIGEN": fila_p0["ARCHIVO ORIGEN"],
                     "HOJA": hoja, "FILA_EXCEL": rec["fila"], "FORMATO": id_cuenta, "LOTE DE CARGA": fila_p0["LOTE DE CARGA"]})
    snapshot = None
    if op:
        snapshot = {"version": C.VERSION_SNAPSHOT, "origen": "DEPOSITOS_ACTIVOS",
                    "fecha_corte": estado["operativo"]["fecha_corte"],
                    "filas": [dict({C.CAMPO_CLAVE: k}, **{c: v for c, v in f.items() if c in C.CAMPOS_SNAPSHOT})
                              for k, f in sorted(op.items())]}
    return datos, mapa, filas, snapshot
