# -*- coding: utf-8 -*-
"""
p10/generador.py · P10-A.1 · generador del libro mensual histórico (REGENERACIÓN COMPLETA).

    extracto acumulado más reciente        snapshot operativo actual
    (LISTS.csv + ORIGEN.xlsx, salidas de P0)   (equivalente a Depositos_Activos)
                      \\                          /
                       +--->  un .xlsx por BANCO + CUENTA + MONEDA + MES  <---+
                              hoja EXTRACTO   (diseño aprobado de la capa 4, historico.py)
                              hoja AUDITORIA  (las 26 COLUMNAS_LISTS de P0)

Qué se reutiliza y qué se agrega
--------------------------------
* Normalización, CLAVE TRANSACCIÓN, importes, saldos: ya vienen en LISTS.csv (P0). Aquí no se recalculan.
* Agrupación por cuenta-mes, columnas propias de cada banco, zona superior, formatos, tabla, filtros, paneles
  y verificación de saldos declarados: `historico.construir_extractos_historicos` / `escribir_extracto_historico`
  (P3b, sin modificar).
* Se agrega: (1) el cruce por CLAVE con el snapshot, que fija los 7 campos operativos de cada movimiento;
  (2) la 4.ª columna operativa OBSERVACIONES; (3) la hoja AUDITORIA; (4) validación independiente del archivo
  ya escrito ANTES de aceptarlo; (5) archivo byte-determinista (misma entrada -> mismos bytes).

Regeneración completa: no hay edición fila por fila. Cada corrida reconstruye el libro entero desde sus entradas,
así que el resultado es determinista y no puede duplicar movimientos.
"""
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import tempfile
import zipfile
from collections import Counter
from dataclasses import dataclass, field

import historico as H
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import contrato as C
from .snapshot import Snapshot, cargar_snapshot

FECHA_FIJA = dt.datetime(2000, 1, 1)     # propiedades del archivo cuando no hay snapshot con fecha de corte


# =====================================================================
# 1. CAMPOS OPERATIVOS: snapshot sobre lo que trae P0
# =====================================================================

def valores_operativos(fila_p0, fila_snapshot):
    """
    Los 7 campos operativos que P10 muestra para un movimiento.
    * Con fila en el snapshot, el snapshot manda (null = vacío; ASIGNADO se muestra CONFIRMADO).
    * Sin fila (movimiento aún no cargado a la lista), se conserva lo que entrega P0.
    """
    if fila_snapshot is None:
        return {"ESTADO": fila_p0["ESTADO"], "ESTUDIANTE": fila_p0["ESTUDIANTE"],
                "SOLICITADO POR": fila_p0["SOLICITADO POR"], "SEDE SOLICITANTE": fila_p0["SEDE SOLICITANTE"],
                "CONFIRMADO POR": fila_p0["CONFIRMADO POR"],
                "FECHA CONFIRMACIÓN": H._fecha_hora(fila_p0["FECHA CONFIRMACIÓN"]),
                "OBSERVACIÓN": fila_p0["OBSERVACIÓN"]}
    out = {col: fila_snapshot[campo] for col, campo in C.OPERATIVOS.items()}
    out["ESTADO"] = C.ESTADO_VISIBLE[fila_snapshot[C.CAMPO_ESTADO]]
    return out


def aplicar_snapshot(filas_p0, snapshot):
    """
    Filas de P0 con los 7 campos operativos resueltos. No modifica `filas_p0`.
    Devuelve (filas_efectivas, resumen) con cuántos movimientos cruzaron por CLAVE.
    """
    efectivas, con_snapshot = [], 0
    for r in filas_p0:
        s = snapshot.get(r[C.COLUMNA_CLAVE]) if snapshot is not None else None
        con_snapshot += s is not None
        efectivas.append(dict(r, **valores_operativos(r, s)))
    claves_p0 = {r[C.COLUMNA_CLAVE] for r in filas_p0}
    resumen = {
        "movimientos": len(filas_p0),
        "con_fila_en_snapshot": con_snapshot,
        "sin_fila_en_snapshot": len(filas_p0) - con_snapshot,
        "snapshot_filas": len(snapshot) if snapshot is not None else 0,
        "snapshot_fuera_del_extracto": (len([k for k in snapshot.filas if k not in claves_p0])
                                        if snapshot is not None else 0),
    }
    return efectivas, resumen


# =====================================================================
# 2. HOJA AUDITORIA (las 26 columnas de P0)
# =====================================================================

def _texto_o_none(v):
    s = "" if v is None else str(v)
    return s if s != "" else None


def valor_auditoria(col, v):
    """Valor tipado de una columna de las 26: números y fechas reales; códigos y textos como texto exacto."""
    tipo = C.TIPO_AUDITORIA[col]
    if tipo == "importe":
        return H._numero(v)
    if tipo == "fecha":
        f = H._fecha(v)
        return dt.datetime(f.year, f.month, f.day) if f else None
    if tipo in ("fecha_hora", "fecha_hora_seg"):
        return H._fecha_hora(v)
    return _texto_o_none(v)


def fila_auditoria(fila_efectiva):
    return [valor_auditoria(col, fila_efectiva[col]) for col in C.COLUMNAS_LISTS]


def _formato_celda(tipo, fm):
    return {"fecha": fm["fecha"], "importe": fm["importe"], "fecha_hora": fm["fecha_hora"],
            "fecha_hora_seg": fm["fecha"] + " hh:mm:ss", "codigo": "@"}.get(tipo, "General")


def _escribir_auditoria(wb, libro):
    ws = wb.create_sheet(C.HOJA_AUDITORIA)
    ws.sheet_view.showGridLines = False
    fm = libro.resultado["formatos"]
    color = libro.resultado["colores"]["operativo"]
    fino = Side(style="thin", color="BFBFBF")
    formatos = [_formato_celda(C.TIPO_AUDITORIA[c], fm) for c in C.COLUMNAS_LISTS]

    for j, col in enumerate(C.COLUMNAS_LISTS, start=1):
        celda = ws.cell(1, j)
        H._poner_texto(celda, col)
        celda.font = Font(color="FFFFFF", bold=True)
        celda.fill = PatternFill("solid", fgColor=color)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = Border(left=fino, right=fino, top=fino, bottom=fino)
        ws.column_dimensions[get_column_letter(j)].width = max(
            C.ANCHO_AUDITORIA[col], max(len(w) for w in col.split()) + 3)
    ws.row_dimensions[1].height = 32

    for i, fila in enumerate(libro.auditoria, start=2):
        for j, (col, v) in enumerate(zip(C.COLUMNAS_LISTS, fila), start=1):
            if v is None:
                continue
            celda = ws.cell(i, j)
            tipo = C.TIPO_AUDITORIA[col]
            if tipo in ("fecha", "fecha_hora", "fecha_hora_seg", "importe"):
                celda.value = v
            else:
                H._poner_texto(celda, v)
            celda.number_format = formatos[j - 1]
            if tipo in ("fecha", "fecha_hora", "fecha_hora_seg"):
                celda.alignment = Alignment(horizontal="center")

    ultima = 1 + len(libro.auditoria)
    tabla = Table(displayName=C.TABLA_AUDITORIA, ref=f"A1:{get_column_letter(len(C.COLUMNAS_LISTS))}{ultima}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True, showColumnStripes=False,
                                          showFirstColumn=False, showLastColumn=False)
    ws.add_table(tabla)
    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# =====================================================================
# 3. LIBRO (función pura)
# =====================================================================

@dataclass
class Libro:
    nombre_archivo: str
    banco: str
    cuenta: str
    moneda: str
    periodo: str
    id_cuenta: str
    formato: str
    resultado: dict                      # salida de historico.py (con OBSERVACIONES agregada)
    auditoria: list                      # filas de 26 valores, en el mismo orden que EXTRACTO
    claves: list
    estados: dict                        # {estado visible: cantidad}
    errores: list = field(default_factory=list)
    advertencias: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.errores and bool(self.claves)


def nombre_archivo(banco, cuenta, moneda, periodo):
    return C.PLANTILLA_NOMBRE_ARCHIVO.format(
        BANCO=H._seguro(banco), CUENTA=H._seguro(cuenta), MONEDA=H._seguro(moneda), PERIODO=periodo)


def _agregar_observaciones(res, por_clave):
    """4.ª columna operativa, después de FECHA DE CONFIRMACIÓN y antes de la CLAVE oculta."""
    columnas = res["columnas"]
    pos = next(i for i, c in enumerate(columnas) if c["titulo"] == "FECHA DE CONFIRMACIÓN") + 1
    columnas.insert(pos, {"titulo": C.COLUMNA_OBSERVACIONES, "tipo": "texto",
                          "ancho": C.ANCHO_OBSERVACIONES, "bloque": "OPERATIVO"})
    assert len(res["filas"]) == len(res["claves"])
    for fila, clave in zip(res["filas"], res["claves"]):
        fila.insert(pos, _texto_o_none(por_clave[clave]["OBSERVACIÓN"]))


def construir_libros(datos_originales, mapa_origen, filas_p0, registro, snapshot=None):
    """
    Función pura: no lee ni escribe archivos. Un Libro por (BANCO, CUENTA, MONEDA, AAAA-MM).
    `filas_p0`: filas de LISTS.csv (o de la hoja LISTS de NORMALIZADO.xlsx) tal como las entrega P0.
    """
    filas_ef, resumen = aplicar_snapshot(filas_p0, snapshot)
    por_clave = {r[C.COLUMNA_CLAVE]: r for r in filas_ef}
    # Los 7 campos operativos ya vienen resueltos en `filas_ef`: historico.py los muestra sin cambios.
    resultados = H.construir_extractos_historicos(datos_originales, mapa_origen, filas_ef, registro)
    libros = []
    for res in resultados:
        errores, auditoria, claves = list(res["errores"]), [], list(res["claves"])
        if not errores and res["filas"]:
            _agregar_observaciones(res, por_clave)
            auditoria = [fila_auditoria(por_clave[c]) for c in claves]
        estados = dict(Counter(por_clave[c]["ESTADO"] for c in claves))
        libros.append(Libro(
            nombre_archivo=nombre_archivo(res["banco"], res["cuenta"], res["moneda"], res["periodo"]),
            banco=res["banco"], cuenta=res["cuenta"], moneda=res["moneda"], periodo=res["periodo"],
            id_cuenta=res["id_cuenta"], formato=res["formato"], resultado=res, auditoria=auditoria,
            claves=claves, estados=estados, errores=errores, advertencias=list(res["advertencias"])))
    libros.sort(key=lambda l: l.nombre_archivo)
    return libros, resumen


# =====================================================================
# 4. ESCRITURA (única parte que escribe Excel)
# =====================================================================

def _fijar_zip(ruta, modificado):
    """
    Reescribe el contenedor con marcas de tiempo fijas: misma entrada -> mismos bytes.
    openpyxl pone la hora actual en `modified` al guardar; aquí se reemplaza por `modificado`.
    """
    with zipfile.ZipFile(ruta) as zin:
        miembros = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    fecha = modificado.strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
    miembros = [(n, re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>" + fecha + rb"\g<2>", d)
                 if n == "docProps/core.xml" else d) for n, d in miembros]
    temporal = ruta + ".fijo"
    with zipfile.ZipFile(temporal, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as zout:
        for nombre, datos in miembros:
            zi = zipfile.ZipInfo(nombre, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 3
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, datos)
    os.replace(temporal, ruta)


def _ajustar_extracto(ws):
    """Único ajuste visual de P10 sobre la hoja aprobada: ancho de CONFIRMADO POR (correo completo)."""
    tabla = ws.tables[C.TABLA_EXTRACTO]
    c0, r0, c1, _ = range_boundaries(tabla.ref)
    for c in range(c0, c1 + 1):
        if ws.cell(r0, c).value == "CONFIRMADO POR":
            letra = get_column_letter(c)
            ws.column_dimensions[letra].width = max(ws.column_dimensions[letra].width or 0,
                                                    C.ANCHO_CONFIRMADO_POR)


def escribir_libro(libro, ruta, snapshot=None):
    """
    Escribe el .xlsx de UN libro con sus dos hojas. La hoja EXTRACTO la dibuja `historico.py` (diseño aprobado);
    se reabre ese mismo archivo (sin pérdida: lo comprueba la prueba de fidelidad) para agregar AUDITORIA.
    """
    if libro.errores:
        raise ValueError("libro no escrito: " + "; ".join(libro.errores))
    if not libro.claves:
        raise ValueError("libro no escrito: el mes no tiene movimientos")
    base = ruta + ".base.xlsx"
    try:
        H.escribir_extracto_historico(libro.resultado, base)
        wb = load_workbook(base)
        _ajustar_extracto(wb[C.HOJA_EXTRACTO])
        _escribir_auditoria(wb, libro)
        wb.active = 0
        fijo = snapshot.fecha_corte if (snapshot is not None and snapshot.fecha_corte) else FECHA_FIJA
        wb.properties.created = wb.properties.modified = fijo
        if snapshot is not None:
            wb.properties.description = (f"{C.VERSION_P10} · estado operativo: snapshot {snapshot.origen} "
                                         f"al {fijo:%Y-%m-%d %H:%M} · {snapshot.sha256[:12]}")
        wb.save(ruta)
        _fijar_zip(ruta, fijo)
    finally:
        if os.path.exists(base):
            os.remove(base)
    return ruta


def sha256_archivo(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()


# =====================================================================
# 5. GENERACIÓN DESDE ARCHIVOS
# =====================================================================

def generar_libros(ruta_origen, ruta_lists, snapshot, carpeta_salida, ruta_registro=None, solo_cuentas=None,
                   validar=True):
    """
    Lee ORIGEN.xlsx y LISTS.csv (o NORMALIZADO.xlsx) de P0, los cruza con el snapshot y escribe en
    `carpeta_salida` los libros que pasan la validación. Un libro que no pasa NO se acepta (no se escribe ni
    reemplaza el existente) y queda en `omitidos` con sus causas.
    `solo_cuentas`: ids del registro (p. ej. {"BNB_MN"}) para generar solo esas cuentas.
    """
    from . import validacion   # import diferido: validacion importa este módulo
    snapshot = cargar_snapshot(snapshot) if snapshot is not None else None
    datos, mapa = H.leer_origen(ruta_origen)
    filas_p0 = H.leer_normalizado(ruta_lists)
    registro = H.cargar_registro(ruta_registro)
    libros, resumen = construir_libros(datos, mapa, filas_p0, registro, snapshot)
    indice = validacion.indice_origen(datos, mapa) if validar else None

    os.makedirs(carpeta_salida, exist_ok=True)
    trabajo = tempfile.mkdtemp(prefix=".p10_", dir=carpeta_salida)
    archivos, omitidos = [], []
    try:
        for libro in libros:
            if solo_cuentas and libro.id_cuenta not in solo_cuentas:
                continue
            if not libro.ok:
                omitidos.append({"nombre_archivo": libro.nombre_archivo,
                                 "errores": libro.errores or ["sin movimientos"]})
                continue
            temporal = os.path.join(trabajo, libro.nombre_archivo)
            escribir_libro(libro, temporal, snapshot)
            errores = (validacion.validar_libro(temporal, libro, filas_p0, snapshot, indice, registro)
                       if validar else [])
            if errores:
                omitidos.append({"nombre_archivo": libro.nombre_archivo, "errores": errores})
                continue
            final = os.path.join(carpeta_salida, libro.nombre_archivo)
            os.replace(temporal, final)
            i_cred, i_deb = C.COLUMNAS_LISTS.index("CRÉDITO"), C.COLUMNAS_LISTS.index("DÉBITO")
            archivos.append({
                "nombre_archivo": libro.nombre_archivo, "ruta": final, "sha256": sha256_archivo(final),
                "bytes": os.path.getsize(final), "banco": libro.banco, "cuenta": libro.cuenta,
                "moneda": libro.moneda, "periodo": libro.periodo, "id_cuenta": libro.id_cuenta,
                "movimientos": len(libro.claves), "estados": libro.estados,
                "creditos": round(sum(f[i_cred] or 0.0 for f in libro.auditoria), 2),
                "debitos": round(sum(f[i_deb] or 0.0 for f in libro.auditoria), 2),
                "validacion": "OK" if validar else "OMITIDA", "advertencias": libro.advertencias,
            })
    finally:
        shutil.rmtree(trabajo, ignore_errors=True)
    return {
        "version": C.VERSION_P10, "estado": "OK" if not omitidos else "CON_ERRORES",
        "carpeta": carpeta_salida, "archivos": archivos, "omitidos": omitidos,
        "extractos_sin_movimientos": H.extractos_sin_movimientos(datos, mapa),
        "snapshot": None if snapshot is None else {
            "origen": snapshot.origen, "fecha_corte": snapshot.fecha_corte.isoformat() if snapshot.fecha_corte
            else None, "sha256": snapshot.sha256, "advertencias": snapshot.advertencias, **resumen},
        "entradas": {"lists": os.path.basename(str(ruta_lists)), "lists_sha256": sha256_archivo(ruta_lists),
                     "origen": os.path.basename(str(ruta_origen)), "origen_sha256": sha256_archivo(ruta_origen)},
    }


def escribir_manifiesto(info, ruta):
    """Manifiesto JSON determinista (sin rutas locales ni marcas de tiempo del proceso)."""
    limpio = dict(info, carpeta=None, archivos=[{k: v for k, v in a.items() if k != "ruta"}
                                                for a in info["archivos"]])
    with open(ruta, "w", encoding="utf-8", newline="\n") as f:
        json.dump(limpio, f, ensure_ascii=False, indent=1, sort_keys=True)
        f.write("\n")
    return ruta
