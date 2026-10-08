"""
============================================================
CONTROL DE DEPÓSITOS CBBA — MOTOR DE NORMALIZACIÓN
============================================================

Reconstrucción como script único de:
  - CELDA_1_CONTROL_DEPOSITOS_CORREGIDA.txt
    (configuración + detección + normalización + validación)
  - CELDA_2_CONTROL_DEPOSITOS_CORREGIDA.txt
    (subida de archivos -> reemplazada por descubrimiento automático)
  - CELDA_3_CONTROL_DEPOSITOS_CORREGIDA.txt
    (orquestación + auditoría + exportación)

REGLA FUNDAMENTAL: ninguna regla de negocio fue modificada respecto
de las 3 celdas originales (detección de bancos/cuentas, hojas
válidas, encabezados esperados, normalizadores por banco, cálculo
de CÓDIGO DE ASIGNACIÓN y CLAVE TRANSACCIÓN, reglas de débito/
crédito, extracción del originante, construcción de INFORMACIÓN
ADICIONAL, validaciones estructurales, validación de saldos con
tolerancia 0.01, control del año (desde P0: rango dinámico, ya no
2026 fijo), criterios que bloquean la exportación, estados
válidos OK/SIN MOVIMIENTOS, y el lector robusto xlrd -> calamine ante AssertionError).

Este módulo es intencionalmente independiente de Google Drive, de
Google Colab y de Claude/Cowork. Solo conoce:
  - una carpeta local de entrada (extractos .xls/.xlsx ya
    descargados, con su nombre original intacto);
  - una ruta local de salida (NORMALIZADO.xlsx).

Descargar los extractos desde Drive y volver a subir NORMALIZADO.xlsx
es responsabilidad de la orquestación externa (Cowork), NO de este
motor. Este archivo no importa ni referencia ningún conector de
Google Drive.

Cambios de infraestructura respecto de las 3 celdas (ninguno es una
regla de negocio):
  1. Se elimina `from google.colab import files`, `files.upload()` y
     `files.download()`. En su lugar: `descubrir_archivos()` lista
     los .xls/.xlsx de una carpeta local, y `ejecutar_motor()` recibe
     una ruta local de salida en vez de disparar una descarga.
  2. Se elimina `display(df)` (específico de notebooks). Se reemplaza
     por `print(df.to_string(index=False))`, que muestra exactamente
     la misma información.
  3. Se elimina la línea `!pip -q install ...` (sintaxis de notebook,
     inválida en un .py). En su lugar, `verificar_dependencias()`
     comprueba en tiempo de ejecución que pandas, numpy, xlrd,
     openpyxl y python-calamine estén disponibles, y falla con un
     mensaje claro (incluyendo el comando pip) si falta alguna.
  4. ARCHIVO ORIGEN: las celdas originales usaban directamente el
     nombre devuelto por `files.upload()` (que en Colab ya era solo
     el nombre de archivo, sin ruta). Ahora que el archivo puede
     leerse desde una ruta temporal completa, se separan dos
     conceptos:
       - `archivo`: la ruta real desde la que se lee el Excel
         (necesaria para pd.read_excel).
       - `nombre_origen`: el nombre original del archivo, que es lo
         único que se guarda en la columna ARCHIVO ORIGEN y lo único
         que se usa en las comparaciones internas
         (`df_final["ARCHIVO ORIGEN"] == nombre_origen`).
     finalizar_dataframe (y la normalización) aceptan
     ahora un parámetro opcional `nombre_origen` (por defecto None,
     en cuyo caso se comporta exactamente como antes y usa `archivo`
     tal cual). Esto es un cambio de encaje de datos, no de lógica
     de negocio.
  5. El nombre de salida deja de llevar timestamp
     ("CONTROL_DEPOSITOS_CBBA_PRUEBA_<timestamp>.xlsx") y pasa a ser
     la ruta fija que indique el llamador (por requisito de esta
     etapa: NORMALIZADO.xlsx). El contenido y las 4 hojas
     (LISTS, VALIDACION, RESUMEN, DIAGNOSTICO) no cambian.
  6. Salida adicional LISTS.csv: además de NORMALIZADO.xlsx (que no
     se reemplaza ni se altera), se genera un archivo plano
     LISTS.csv en la misma carpeta de salida, con el mismo
     `df_final` ya validado, restringido a las columnas y al orden
     exacto de `COLUMNAS_LISTS`, codificado en UTF-8 con BOM
     (utf-8-sig). Es puramente una exportación adicional: no
     modifica ningún dato, cálculo ni validación del motor.

P4 — DETECCIÓN POR REGISTRO (único cambio de regla desde las celdas):
  La identificación de banco, cuenta, moneda y formato ya no está en
  `detectar_formato` (ramas por banco sobre las 40 primeras filas de
  todo el archivo) sino en `registro_bancos.json`, a través de
  `deteccion_registro.py` (ambos junto a este archivo; la variable
  CBBA_REGISTRO_BANCOS apunta a otro registro). La cuenta se lee solo
  en la celda rotulada de la cabecera y se valida en todos los
  formatos (BMSC incluido); encabezado incompleto, cabecera ambigua,
  cuenta no registrada o el reporte Unión «Últimos 12» detienen el
  proceso con el motivo. Una cuenta nueva de un formato conocido se
  agrega con una entrada en CUENTAS, sin código. La normalización,
  la validación de saldos, COLUMNAS_LISTS y CLAVE TRANSACCIÓN no
  cambian.

P5 — NORMALIZACIÓN POR REGISTRO:
  La normalización y la validación de saldos productivas salen de
  `motor_generico.py` gobernado por `registro_bancos.json` (ambos junto
  a este archivo): archivo → detección por registro → normalización
  genérica → salida productiva. COLUMNAS_LISTS, CLAVE TRANSACCIÓN
  (crear_clave / finalizar_dataframe), NORMALIZADO.xlsx, LISTS.csv y
  ORIGEN.xlsx no cambian. Una cuenta nueva de un formato conocido se
  normaliza solo con su entrada en CUENTAS, sin código.

P6 — RETIRO DEL LEGADO:
  Se retiraron los normalizar_* por banco, normalizar_archivo,
  validar_archivo, HOJAS_VALIDAS, ENCABEZADOS_ESPERADOS,
  encontrar_fila_encabezado, leer_tabla_movimientos, texto_de_archivo,
  aplicar_identidad_registro y la referencia legada en sombra (paso 16,
  SOMBRA_REPORTE.json / SOMBRA_DIFERENCIAS.csv). Hojas, encabezados,
  campos y fuentes de saldo viven solo en registro_bancos.json. Este
  archivo conserva la orquestación y las primitivas compartidas que usa
  motor_generico.py (texto, números, fechas, horas, códigos, lector de
  Excel, buscar_columna*, extraer_nombre_bnb, ecuacion_saldo,
  crear_clave / finalizar_dataframe, COLUMNAS_LISTS).

Uso como script:
    python motor_control_depositos_cbba.py <carpeta_entrada> <ruta_salida>

Uso como módulo:
    from motor_control_depositos_cbba import ejecutar_motor
    resultado = ejecutar_motor("/tmp/entrada", "/tmp/salida/NORMALIZADO.xlsx")
============================================================
"""

import importlib
import os
import re
import unicodedata
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table


# ============================================================
# 0. VERIFICACIÓN DE DEPENDENCIAS
# ============================================================

DEPENDENCIAS_REQUERIDAS = {
    "pandas": "pandas",
    "numpy": "numpy",
    "xlrd": "xlrd",
    "openpyxl": "openpyxl",
    "python_calamine": "python-calamine",
}


def verificar_dependencias():
    """
    Comprueba que las librerías necesarias para el motor estén
    disponibles en el intérprete actual. No instala nada: solo
    informa qué falta, para que la orquestación decida cómo
    resolverlo antes de ejecutar el motor.

    Lanza RuntimeError si falta alguna dependencia.
    """

    faltantes = []

    for modulo, paquete_pip in DEPENDENCIAS_REQUERIDAS.items():

        try:
            importlib.import_module(modulo)
        except ImportError:
            faltantes.append(paquete_pip)

    if faltantes:
        raise RuntimeError(
            "❌ Faltan dependencias para ejecutar el motor: "
            f"{', '.join(faltantes)}. "
            "Instálalas antes de continuar con: "
            f"pip install {' '.join(faltantes)}"
        )

    return True


# ============================================================
# COLUMNAS FINALES
# ============================================================

COLUMNAS_LISTS = [
    "CLAVE TRANSACCIÓN",
    "CÓDIGO DE ASIGNACIÓN",
    "BANCO",
    "CUENTA BANCARIA",
    "MONEDA",
    "FECHA MOVIMIENTO",
    "HORA MOVIMIENTO",
    "IMPORTE",
    "DÉBITO",
    "CRÉDITO",
    "TIPO MOVIMIENTO",
    "SALDO",
    "DESCRIPCIÓN",
    "DEPOSITANTE / ORIGINANTE",
    "INFORMACIÓN ADICIONAL",
    "ESTADO",
    "ESTUDIANTE",
    "SOLICITADO POR",
    "SEDE SOLICITANTE",
    "CONFIRMADO POR",
    "FECHA CONFIRMACIÓN",
    "OBSERVACIÓN",
    "TEXTO DE BÚSQUEDA",
    "ARCHIVO ORIGEN",
    "LOTE DE CARGA",
    "FECHA DE CARGA"
]


# ============================================================
# FUNCIONES DE TEXTO
# ============================================================

def normalizar_texto(valor):

    if pd.isna(valor):
        return ""

    texto = str(valor).strip().upper()

    texto = unicodedata.normalize(
        "NFKD",
        texto
    )

    texto = "".join(
        c for c in texto
        if not unicodedata.combining(c)
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto


def buscar_columna(df, *posibles):

    mapa = {
        normalizar_texto(c): c
        for c in df.columns
    }

    for posible in posibles:

        clave = normalizar_texto(posible)

        if clave in mapa:
            return mapa[clave]

    for posible in posibles:

        clave = normalizar_texto(posible)

        for normalizada, original in mapa.items():

            if (
                clave in normalizada
                or normalizada in clave
            ):
                return original

    raise KeyError(
        f"No encontré {posibles}. "
        f"Columnas disponibles: {list(df.columns)}"
    )


def buscar_columna_opcional(df, *posibles):

    try:
        return buscar_columna(
            df,
            *posibles
        )

    except:
        return None


# ============================================================
# CONVERSIÓN DE NÚMEROS
# ============================================================

def numero(valor):

    if pd.isna(valor) or valor == "":
        return np.nan

    if isinstance(
        valor,
        (int, float, np.integer, np.floating)
    ):
        return float(valor)

    texto = str(valor).strip()

    texto = (
        texto
        .replace("Bs", "")
        .replace("BS", "")
        .replace("USD", "")
        .replace("EUR", "")
        .replace("$", "")
        .replace(" ", "")
    )

    if "," in texto and "." in texto:

        if texto.rfind(",") > texto.rfind("."):

            texto = (
                texto
                .replace(".", "")
                .replace(",", ".")
            )

        else:

            texto = texto.replace(",", "")

    elif "," in texto:

        texto = texto.replace(",", ".")

    try:
        return float(texto)

    except:
        return np.nan


# ============================================================
# CÓDIGOS COMO TEXTO
# ============================================================

def codigo_texto(valor):

    if pd.isna(valor):
        return ""

    if isinstance(valor, (int, np.integer)):
        return str(valor)

    if isinstance(valor, (float, np.floating)):

        if float(valor).is_integer():
            return str(int(valor))

        return str(valor)

    texto = str(valor).strip()

    if re.match(
        r"^-?\d+\.0$",
        texto
    ):
        return texto[:-2]

    return texto


# ============================================================
# FECHA Y HORA
# ============================================================

def normalizar_fecha(valor):

    if pd.isna(valor) or valor == "":
        return pd.NaT

    # Fechas ya reconocidas por pandas / Excel
    if isinstance(valor, (pd.Timestamp, datetime)):
        return pd.Timestamp(valor).date()

    texto = normalizar_texto(valor)

    # Meses en español usados por algunos extractos
    meses_es = {
        "ENE": "01",
        "FEB": "02",
        "MAR": "03",
        "ABR": "04",
        "MAY": "05",
        "JUN": "06",
        "JUL": "07",
        "AGO": "08",
        "SEP": "09",
        "SET": "09",
        "OCT": "10",
        "NOV": "11",
        "DIC": "12"
    }

    for mes, numero_mes in meses_es.items():
        texto = re.sub(
            rf"\b{mes}\b",
            numero_mes,
            texto
        )

    fecha = pd.to_datetime(
        texto,
        dayfirst=True,
        errors="coerce"
    )

    if pd.isna(fecha):
        return pd.NaT

    return fecha.date()

def normalizar_hora(valor):

    if pd.isna(valor):
        return ""

    if hasattr(valor, "strftime"):

        try:
            return valor.strftime("%H:%M:%S")
        except:
            pass

    texto = str(valor).strip()

    if texto == "":
        return ""

    try:

        hora = pd.to_datetime(
            texto,
            errors="coerce"
        )

        if not pd.isna(hora):
            return hora.strftime("%H:%M:%S")

    except:
        pass

    return texto


# ============================================================
# LEER TODAS LAS HOJAS
# ============================================================

def leer_excel_robusto(archivo, **kwargs):
    """
    Intenta leer con el motor normal de pandas.
    Si xlrd lanza AssertionError en ciertos .xls antiguos,
    reintenta automáticamente con Calamine.
    """
    try:
        return pd.read_excel(
            archivo,
            **kwargs
        )

    except AssertionError:
        print(
            f"⚠️ xlrd no pudo leer '{archivo}'. "
            "Reintentando con Calamine..."
        )

        return pd.read_excel(
            archivo,
            engine="calamine",
            **kwargs
        )


def leer_todas_hojas(archivo):

    return leer_excel_robusto(
        archivo,
        sheet_name=None,
        header=None
    )


# ============================================================
# DETECTOR POR REGISTRO (P4)
#
# Desde P4 la identificación de banco, cuenta, moneda y formato
# sale de registro_bancos.json (módulo deteccion_registro.py, junto
# a este archivo). La variable CBBA_REGISTRO_BANCOS apunta a otro
# registro.
#
# Primero se reconoce la estructura (firma + encabezado completo en
# la hoja aceptada) y recién después se lee la cuenta, SOLO en la
# celda rotulada de la cabecera (filas anteriores al encabezado de
# la tabla). Nunca se buscan números en glosas, adicionales u
# originantes. Encabezado incompleto, cuenta ausente, cuenta no
# registrada, varias cuentas o varios formatos = error explícito.
# ============================================================

_MODULO_DETECCION = None


def modulo_deteccion():
    """Carga deteccion_registro.py desde la carpeta de este motor."""

    global _MODULO_DETECCION

    if _MODULO_DETECCION is None:

        import importlib.util as _iu
        import sys as _sys

        ruta = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            "deteccion_registro.py"
        )

        if not os.path.exists(ruta):
            raise ValueError(
                "no se encuentra deteccion_registro.py junto al motor "
                f"({ruta}): sin él no se pueden identificar los extractos"
            )

        spec = _iu.spec_from_file_location(
            "deteccion_registro_p4",
            ruta
        )
        modulo = _iu.module_from_spec(spec)
        _sys.modules[spec.name] = modulo
        spec.loader.exec_module(modulo)

        _MODULO_DETECCION = modulo

    return _MODULO_DETECCION


def detector_registro(ruta_registro=None):
    """
    Detector productivo: lee y valida el registro. Lanza ValueError si
    el registro falta o es inválido.
    """

    return modulo_deteccion().DetectorRegistro.cargar(
        normalizar_texto,
        leer_todas_hojas,
        ruta=ruta_registro
    )


def detectar_extracto(archivo, detector=None):
    """Resultado completo de la detección (estado, motivo, identidad)."""

    return (
        detector
        or detector_registro()
    ).detectar(
        archivo
    )


def detectar_formato(archivo):
    """Id de cuenta (p. ej. "BNB_MN") o "NO_RECONOCIDO"."""

    return detectar_extracto(
        archivo
    ).formato_motor


# ============================================================
# NORMALIZACIÓN POR REGISTRO (P5)
#
# Desde P5 la normalización y la validación de saldos productivas
# las hace motor_generico.py (junto a este archivo) con la
# configuración de registro_bancos.json: la misma secuencia para
# todos los bancos, sin ramas por banco ni por cuenta. Reutiliza las
# primitivas de este archivo (números, fechas, horas, lector de
# Excel, finalizar_dataframe → COLUMNAS_LISTS y CLAVE TRANSACCIÓN,
# congeladas).
# ============================================================

_MODULO_GENERICO = None


def modulo_generico():
    """Carga motor_generico.py desde la carpeta de este motor."""

    global _MODULO_GENERICO

    if _MODULO_GENERICO is None:

        import importlib.util as _iu
        import sys as _sys

        ruta = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            "motor_generico.py"
        )

        if not os.path.exists(ruta):
            raise ValueError(
                "no se encuentra motor_generico.py junto al motor "
                f"({ruta}): sin él no se pueden normalizar los extractos"
            )

        spec = _iu.spec_from_file_location(
            "motor_generico_p5",
            ruta
        )
        modulo = _iu.module_from_spec(spec)
        _sys.modules[spec.name] = modulo
        spec.loader.exec_module(modulo)

        _MODULO_GENERICO = modulo

    return _MODULO_GENERICO


def normalizador_registro(ruta_registro=None):
    """
    Normalizador productivo: el motor genérico con registro_bancos.json
    (el mismo que usa la detección; CBBA_REGISTRO_BANCOS apunta a otro).
    Lanza ValueError si falta motor_generico.py o si el registro no sirve
    para normalizar.
    """

    mg = modulo_generico()

    ruta = modulo_deteccion().ruta_registro(ruta_registro)

    try:

        registro = mg.Registro.cargar(ruta)

    except OSError as e:

        raise ValueError(
            f"no se pudo leer el registro de bancos '{ruta}': {e}"
        )

    problemas = registro.validar()

    if problemas:
        raise ValueError(
            f"registro de bancos inválido para normalizar ({ruta}): "
            + "; ".join(problemas)
        )

    return mg.MotorGenerico(
        registro,
        globals()
    )


def normalizar_extracto(
    archivo,
    deteccion,
    lote,
    fecha_carga,
    nombre_origen=None,
    normalizador=None
):
    """
    Normalización productiva de UN extracto (P5).

    `deteccion` es el resultado de la detección por registro (o el id de
    la cuenta en CUENTAS). Devuelve (tabla, contexto): `tabla` tiene las
    26 COLUMNAS_LISTS con el índice de fila del archivo (vacía si el
    extracto válido no trae movimientos); `contexto` guarda la hoja, la
    fila de encabezado y la tabla leída, que usan la validación de saldos
    y ORIGEN.xlsx.
    """

    normalizador = normalizador or normalizador_registro()

    cuenta_id = getattr(
        deteccion,
        "cuenta_id",
        deteccion
    )

    tabla, contexto = normalizador.normalizar(
        archivo,
        cuenta_id,
        lote,
        fecha_carga,
        nombre_origen=nombre_origen
    )

    # Se normaliza la misma tabla que la detección comprobó.
    if hasattr(deteccion, "hoja"):

        leido = (
            contexto["hoja"],
            int(contexto["fila_encabezado"])
        )

        detectado = (
            deteccion.hoja,
            int(deteccion.fila_encabezado)
        )

        if leido != detectado:
            raise ValueError(
                f"{cuenta_id}: la normalización leyó la hoja/fila "
                f"{leido} y la detección comprobó {detectado}"
            )

    return tabla, contexto


def validar_extracto(
    archivo,
    deteccion,
    datos,
    contexto,
    normalizador=None
):
    """
    Validación de saldos productiva (P5): misma ecuación y tolerancia
    (ecuacion_saldo, 0.01); las fuentes de saldo inicial/final salen del
    registro.
    """

    normalizador = normalizador or normalizador_registro()

    return normalizador.validar(
        archivo,
        getattr(deteccion, "cuenta_id", deteccion),
        datos,
        contexto
    )


def contrato_origen_registro(detecciones, normalizador):
    """
    Contrato de captura_origen.py (ORIGEN.xlsx) armado desde el registro:
    la hoja y el encabezado de cada cuenta son los que la detección y la
    normalización usaron (también para una cuenta nueva).
    """

    encabezados = {
        d.cuenta_id: normalizador.registro.formato(
            d.formato_id
        )["encabezados"]["puntaje"]
        for d in detecciones.values()
        if d.ok
    }

    return {
        "hojas_validas": {
            d.cuenta_id: d.hoja
            for d in detecciones.values()
            if d.ok
        },
        "encabezados_esperados": encabezados,
        "encontrar_fila_encabezado": (
            lambda raw, cuenta_id: normalizador._fila_encabezado(
                raw,
                encabezados[cuenta_id]
            )
        ),
    }


# ============================================================
# EXTRAER NOMBRE BNB
# ============================================================

def extraer_nombre_bnb(texto):

    if pd.isna(texto):
        return ""

    s = str(texto)

    patrones = [
        r"Nombre Originante\s*:\s*([^;]+)",
        r"Nombre\\?\s*:\s*([^;]+)",
        r"Nombre\s*:\s*([^;]+)"
    ]

    for patron in patrones:

        m = re.search(
            patron,
            s,
            flags=re.IGNORECASE
        )

        if m:
            return m.group(1).strip()

    return ""


# ============================================================
# CLAVE DEFINITIVA
# ============================================================

def valor_clave_numero(valor):

    if pd.isna(valor):
        return ""

    return f"{float(valor):.2f}"


def crear_clave(row):

    fecha = pd.to_datetime(
        row["FECHA MOVIMIENTO"],
        errors="coerce"
    )

    fecha_txt = (
        ""
        if pd.isna(fecha)
        else fecha.strftime("%Y%m%d")
    )

    hora = re.sub(
        r"[^0-9]",
        "",
        str(
            row["HORA MOVIMIENTO"]
        )
    )

    return (
        f"{row['BANCO']}|"
        f"{row['CUENTA BANCARIA']}|"
        f"{fecha_txt}|"
        f"{hora}|"
        f"{codigo_texto(row['CÓDIGO DE ASIGNACIÓN'])}|"
        f"{row['TIPO MOVIMIENTO']}|"
        f"{valor_clave_numero(row['IMPORTE'])}|"
        f"{valor_clave_numero(row['SALDO'])}"
    )


# ============================================================
# FINALIZAR TABLA
# ============================================================

def finalizar_dataframe(
    df,
    archivo,
    lote,
    fecha_carga,
    nombre_origen=None
):

    df["ESTADO"] = "DISPONIBLE"
    df["ESTUDIANTE"] = ""
    df["SOLICITADO POR"] = ""
    df["SEDE SOLICITANTE"] = "COCHABAMBA"
    df["CONFIRMADO POR"] = ""
    df["FECHA CONFIRMACIÓN"] = pd.NaT
    df["OBSERVACIÓN"] = ""

    # ARCHIVO ORIGEN guarda el nombre original del archivo, nunca
    # una ruta temporal. Si no se indica nombre_origen explícito,
    # se mantiene el comportamiento original (usar "archivo" tal
    # cual, como en las celdas de Colab, donde ya era solo un
    # nombre de archivo sin ruta).
    df["ARCHIVO ORIGEN"] = (
        nombre_origen
        if nombre_origen is not None
        else archivo
    )

    df["LOTE DE CARGA"] = lote
    df["FECHA DE CARGA"] = fecha_carga

    df["TEXTO DE BÚSQUEDA"] = (
        df["DESCRIPCIÓN"]
        .fillna("")
        .astype(str)
        + " | "
        + df["DEPOSITANTE / ORIGINANTE"]
        .fillna("")
        .astype(str)
        + " | "
        + df["INFORMACIÓN ADICIONAL"]
        .fillna("")
        .astype(str)
        + " | "
        + df["CÓDIGO DE ASIGNACIÓN"]
        .fillna("")
        .astype(str)
    )

    df["CLAVE TRANSACCIÓN"] = (
        df.apply(
            crear_clave,
            axis=1
        )
    )

    for columna in COLUMNAS_LISTS:

        if columna not in df.columns:
            df[columna] = ""

    return df[
        COLUMNAS_LISTS
    ]


# ============================================================
# VALIDACIÓN MATEMÁTICA
# ============================================================

def ecuacion_saldo(
    saldo_inicial,
    creditos,
    debitos,
    saldo_final
):

    calculado = (
        saldo_inicial
        + creditos
        - debitos
    )

    diferencia = (
        saldo_final
        - calculado
    )

    return (
        round(calculado, 2),
        round(diferencia, 2),
        "OK"
        if abs(diferencia) <= 0.01
        else "REVISAR"
    )


# ============================================================
# DESCUBRIMIENTO DE ARCHIVOS
# (reemplaza CELDA_2 / files.upload())
# ============================================================

EXTENSIONES_VALIDAS = (".xls", ".xlsx")
MARCA_EXCLUSION = "NORMALIZADO"

# Archivos generados por el propio sistema: nunca son extractos de
# entrada (D-20). Se comparan sobre el nombre en mayúsculas y solo
# para .xlsx: ORIGEN*.xlsx y EXTRACTO_HISTORICO_*.xlsx.
PREFIJOS_SALIDA_SISTEMA = (
    "ORIGEN",
    "EXTRACTO_HISTORICO_"
)


# ============================================================
# CONTROL DEL AÑO (P0: antes fijo en 2026)
#
# Qué protege: que una fecha mal interpretada (año de 2 dígitos, texto
# ilegible, formato inesperado) no entre como movimiento. NO es una regla
# de negocio sobre "el año de trabajo".
#   - mínimo  = ANIO_MINIMO_DATOS: primer año de operación del sistema
#     (no existen extractos anteriores); constante, nunca vence.
#   - máximo  = año de la fecha del equipo: un extracto no puede traer
#     movimientos de un año que todavía no empezó.
# Un extracto que cruza diciembre/enero es válido.
# ============================================================

ANIO_MINIMO_DATOS = 2026


def fecha_referencia():
    """Fecha del equipo usada por el control del año (se puede sustituir en pruebas)."""
    return pd.Timestamp.now()


def anios_fuera_de_rango(anios, hoy=None):
    """Años (enteros, ordenados) fuera de [ANIO_MINIMO_DATOS, año de `hoy`]."""

    maximo = (hoy if hoy is not None else fecha_referencia()).year

    return sorted(
        {
            int(a)
            for a in anios
            if not (ANIO_MINIMO_DATOS <= int(a) <= maximo)
        }
    )


def descubrir_archivos(carpeta_entrada):
    """
    Lista los extractos a procesar dentro de una carpeta local.

    Reglas de descubrimiento (requisito de esta etapa, no son
    reglas de negocio del motor de normalización):
      - solo se consideran archivos con extensión .xls o .xlsx,
        sin distinguir mayúsculas/minúsculas;
      - se ignora cualquier archivo cuyo nombre contenga
        "NORMALIZADO" (sin distinguir mayúsculas/minúsculas), para
        no reprocesar una salida de una corrida anterior si quedó
        en la misma carpeta;
      - se ignora cualquier otro archivo que no sea .xls/.xlsx.

    Devuelve un diccionario {nombre_original: ruta_completa}, de
    forma que el resto del motor pueda leer cada archivo desde su
    ruta real, mientras que solo el nombre original se registra en
    ARCHIVO ORIGEN y se usa en las comparaciones internas.

    No modifica, renombra ni elimina ningún archivo.
    """

    if not os.path.isdir(carpeta_entrada):
        raise FileNotFoundError(
            f"❌ La carpeta de entrada no existe: {carpeta_entrada}"
        )

    archivos = {}

    for nombre in sorted(os.listdir(carpeta_entrada)):

        ruta = os.path.join(carpeta_entrada, nombre)

        if not os.path.isfile(ruta):
            continue

        _, extension = os.path.splitext(nombre)

        if extension.lower() not in EXTENSIONES_VALIDAS:
            continue

        if MARCA_EXCLUSION in nombre.upper():
            continue

        if (
            extension.lower() == ".xlsx"
            and nombre.upper().startswith(
                PREFIJOS_SALIDA_SISTEMA
            )
        ):
            continue

        archivos[nombre] = ruta

    return archivos


# ============================================================
# ORQUESTACIÓN DEL MOTOR
# (reconstruye CELDA_3, sin Google Colab ni Google Drive)
# ============================================================

def ejecutar_motor(carpeta_entrada, ruta_salida):
    """
    Punto de entrada único del motor.

    Parámetros
    ----------
    carpeta_entrada : str
        Carpeta local donde ya están descargados los extractos
        bancarios originales (.xls / .xlsx), con su nombre de
        archivo original intacto. El motor NO descarga nada de
        Google Drive: eso es responsabilidad de quien llama a
        esta función.
    ruta_salida : str
        Ruta local completa (incluyendo nombre de archivo) donde
        se debe escribir el resultado. Por requisito de esta
        etapa, el llamador debe pasar una ruta que termine en
        "NORMALIZADO.xlsx".

    Devuelve
    --------
    dict con las tablas generadas (para inspección/registro por
    parte de quien llama) y las rutas de los archivos finales:
        {
            "df_final": DataFrame,
            "df_validacion": DataFrame,
            "resumen": DataFrame,
            "df_resultado_archivos": DataFrame,
            "df_deteccion_final": DataFrame,
            "ruta_salida": str,
            "ruta_lists_csv": str,
            "origen_estado": dict,   # P1 (ORIGEN.xlsx)
            "deteccion_estado": dict # P4 (detección por registro)
        }

    Además de NORMALIZADO.xlsx, se genera en la misma carpeta un
    archivo LISTS.csv (mismo df_final ya validado, columnas y
    orden de COLUMNAS_LISTS, UTF-8 con BOM). No reemplaza ni
    modifica NORMALIZADO.xlsx.

    Lanza ValueError (o RuntimeError si faltan dependencias) con
    el mismo criterio que las celdas originales si cualquier
    validación crítica falla. En ese caso NO se genera ningún
    archivo de salida ni se toca ninguno de los extractos
    originales.
    """

    verificar_dependencias()

    print("============================================================")
    print("CONTROL DE DEPÓSITOS CBBA")
    print("PROCESAMIENTO AUTOMÁTICO")
    print("============================================================")

    # ========================================================
    # 0. DESCUBRIR ARCHIVOS (reemplaza CELDA_2)
    # ========================================================

    mapa_archivos = descubrir_archivos(carpeta_entrada)

    if not mapa_archivos:

        raise ValueError(
            "❌ PROCESO DETENIDO: no se encontraron archivos "
            f".xls/.xlsx para procesar en {carpeta_entrada}."
        )

    print(
        f"\n📂 {len(mapa_archivos)} archivo(s) encontrado(s) en "
        f"'{carpeta_entrada}':"
    )

    for nombre in mapa_archivos:
        print(" •", nombre)

    # ========================================================
    # CONFIGURAR LOTE
    # ========================================================

    fecha_carga = pd.Timestamp.now()

    lote = fecha_carga.strftime(
        "%Y%m%d_%H%M%S"
    )

    # ========================================================
    # 1. DETECTAR FORMATOS
    # ========================================================

    detecciones = []

    # P4: detalle de la detección por archivo (identidad del registro,
    # hoja y fila de encabezado).
    detalle_deteccion = {}

    print(
        "\n🔎 DETECTANDO BANCOS Y CUENTAS...\n"
    )

    # Sin un registro válido no se puede identificar ningún extracto:
    # se detiene aquí, antes de escribir cualquier salida.
    try:

        detector = detector_registro()

        # P5: el mismo registro gobierna la normalización (motor
        # genérico). Sin él tampoco se puede normalizar: se detiene aquí.
        normalizador = normalizador_registro(
            detector.ruta
        )

    except ValueError as e:

        raise ValueError(
            f"❌ PROCESO DETENIDO: {e}"
        )

    if normalizador.registro.sha256 != detector.sha256:

        raise ValueError(
            "❌ PROCESO DETENIDO: registro_bancos.json cambió durante "
            "la carga (detección y normalización deben usar el mismo)."
        )

    for nombre, ruta in mapa_archivos.items():

        try:

            deteccion = detector.detectar(
                ruta
            )

            detalle_deteccion[nombre] = deteccion

            formato = deteccion.formato_motor

            detecciones.append({
                "ARCHIVO": nombre,
                "FORMATO": formato,
                "ESTADO": (
                    "OK"
                    if formato != "NO_RECONOCIDO"
                    else "ERROR"
                )
            })

            if formato == "NO_RECONOCIDO":

                print(
                    f"❌ {nombre} → "
                    f"NO RECONOCIDO ({deteccion.estado}): "
                    f"{deteccion.motivo}"
                )

            else:

                print(
                    f"✅ {nombre} → "
                    f"{formato}"
                )

        except Exception as e:

            detecciones.append({
                "ARCHIVO": nombre,
                "FORMATO": "ERROR",
                "ESTADO": str(e)
            })

            print(
                f"❌ {nombre}: "
                f"{type(e).__name__} | {repr(e)}"
            )

    df_deteccion_final = pd.DataFrame(
        detecciones
    )

    # ========================================================
    # BLOQUEAR SI HAY ARCHIVO NO RECONOCIDO
    # ========================================================

    errores_deteccion = (
        df_deteccion_final[
            "FORMATO"
        ]
        .isin(
            [
                "NO_RECONOCIDO",
                "ERROR"
            ]
        )
        .sum()
    )

    if errores_deteccion > 0:

        motivos = [
            f"   • {f['ARCHIVO']}: "
            + (
                f"{detalle_deteccion[f['ARCHIVO']].estado} — "
                f"{detalle_deteccion[f['ARCHIVO']].motivo}"
                if f["ARCHIVO"] in detalle_deteccion
                else f["ESTADO"]
            )
            for f in detecciones
            if f["FORMATO"] in ("NO_RECONOCIDO", "ERROR")
        ]

        raise ValueError(
            "❌ PROCESO DETENIDO: "
            "hay archivos no reconocidos.\n"
            + "\n".join(motivos)
        )

    # ========================================================
    # CONTROL DE FORMATOS REPETIDOS
    # ========================================================

    formatos_repetidos = (
        df_deteccion_final[
            "FORMATO"
        ]
        .duplicated(
            keep=False
        )
    )

    if formatos_repetidos.any():

        print(
            "\n⚠️ FORMATOS REPETIDOS:"
        )

        print(
            df_deteccion_final[
                formatos_repetidos
            ].to_string(index=False)
        )

        raise ValueError(
            "❌ No cargues dos archivos de la misma "
            "cuenta en la misma ejecución."
        )

    print(
        f"\n✅ {len(df_deteccion_final)} "
        "formatos reconocidos."
    )

    # ========================================================
    # 2. NORMALIZAR
    # ========================================================

    print(
        "\n🔄 NORMALIZANDO MOVIMIENTOS...\n"
    )

    tablas = []

    # P5: lo que la normalización leyó de cada archivo (hoja, fila de
    # encabezado, tabla), para validar saldos y armar ORIGEN.xlsx.
    contextos = {}

    resultados_archivos = []

    for _, fila in df_deteccion_final.iterrows():

        nombre = fila["ARCHIVO"]
        formato = fila["FORMATO"]
        ruta = mapa_archivos[nombre]

        try:

            # P5: normalización genérica gobernada por el registro
            # (identidad BANCO / CUENTA / MONEDA incluida).
            deteccion = detalle_deteccion[nombre]

            resultado, contextos[nombre] = normalizar_extracto(
                ruta,
                deteccion,
                lote,
                fecha_carga,
                nombre_origen=nombre,
                normalizador=normalizador
            )

            tablas.append(
                resultado
            )

            resultados_archivos.append({
                "ARCHIVO": nombre,
                "FORMATO": formato,
                "MOVIMIENTOS": len(resultado),
                "NORMALIZACIÓN": "OK"
            })

            print(
                f"✅ {formato}: "
                f"{len(resultado):,} movimientos"
            )

        except Exception as e:

            resultados_archivos.append({
                "ARCHIVO": nombre,
                "FORMATO": formato,
                "MOVIMIENTOS": 0,
                "NORMALIZACIÓN": "ERROR"
            })

            print(
                f"❌ {formato}: {e}"
            )

    df_resultado_archivos = pd.DataFrame(
        resultados_archivos
    )

    errores_normalizacion = (
        df_resultado_archivos[
            "NORMALIZACIÓN"
        ]
        .eq("ERROR")
        .sum()
    )

    if errores_normalizacion > 0:

        print(
            "\n❌ ARCHIVOS CON ERROR:"
        )

        print(
            df_resultado_archivos[
                df_resultado_archivos[
                    "NORMALIZACIÓN"
                ] == "ERROR"
            ].to_string(index=False)
        )

        raise ValueError(
            "❌ PROCESO DETENIDO: "
            "falló la normalización de al menos un archivo."
        )

    # ========================================================
    # 3. UNIR TODO
    # ========================================================

    df_final = pd.concat(
        tablas,
        ignore_index=True
    )

    print(
        f"\n✅ Movimientos normalizados: "
        f"{len(df_final):,}"
    )

    # ========================================================
    # 4. GARANTIZAR TIPOS DE FECHA
    # ========================================================

    df_final[
        "FECHA MOVIMIENTO"
    ] = pd.to_datetime(
        df_final[
            "FECHA MOVIMIENTO"
        ],
        errors="coerce"
    )

    df_final[
        "FECHA CONFIRMACIÓN"
    ] = pd.to_datetime(
        df_final[
            "FECHA CONFIRMACIÓN"
        ],
        errors="coerce"
    )

    df_final[
        "FECHA DE CARGA"
    ] = pd.to_datetime(
        df_final[
            "FECHA DE CARGA"
        ],
        errors="coerce"
    )

    # ========================================================
    # 5. AUDITORÍA ESTRUCTURAL
    # ========================================================

    print(
        "\n🔍 AUDITORÍA ESTRUCTURAL..."
    )

    duplicados = (
        df_final[
            "CLAVE TRANSACCIÓN"
        ]
        .duplicated()
        .sum()
    )

    codigos_vacios = (
        df_final[
            "CÓDIGO DE ASIGNACIÓN"
        ]
        .fillna("")
        .eq("")
        .sum()
    )

    sin_importe = (
        df_final[
            "IMPORTE"
        ]
        .isna()
        .sum()
    )

    importe_cero = (
        df_final[
            "IMPORTE"
        ]
        .fillna(0)
        .eq(0)
        .sum()
    )

    sin_tipo = (
        df_final[
            "TIPO MOVIMIENTO"
        ]
        .fillna("")
        .eq("")
        .sum()
    )

    doble_movimiento = (
        (
            df_final["DÉBITO"].notna()
            &
            df_final["CRÉDITO"].notna()
        )
        .sum()
    )

    fechas_invalidas = (
        df_final[
            "FECHA MOVIMIENTO"
        ]
        .isna()
        .sum()
    )

    print(
        f"   Claves duplicadas: {duplicados}"
    )

    print(
        f"   Códigos vacíos: {codigos_vacios}"
    )

    print(
        f"   Sin importe: {sin_importe}"
    )

    print(
        f"   Importe cero: {importe_cero}"
    )

    print(
        f"   Sin tipo: {sin_tipo}"
    )

    print(
        f"   Débito + crédito simultáneo: "
        f"{doble_movimiento}"
    )

    print(
        f"   Fechas inválidas: "
        f"{fechas_invalidas}"
    )

    errores_estructurales = (
        duplicados
        + codigos_vacios
        + sin_importe
        + importe_cero
        + sin_tipo
        + doble_movimiento
        + fechas_invalidas
    )

    if errores_estructurales > 0:

        raise ValueError(
            "❌ PROCESO DETENIDO: "
            "la auditoría estructural encontró problemas."
        )

    print(
        "\n✅ AUDITORÍA ESTRUCTURAL SUPERADA"
    )

    # ========================================================
    # 6. CONTROL DEL AÑO
    # ========================================================

    anios_detectados = (
        df_final[
            "FECHA MOVIMIENTO"
        ]
        .dt.year
        .dropna()
        .unique()
        .tolist()
    )

    print(
        f"\n📅 Años detectados: "
        f"{sorted(anios_detectados)}"
    )

    _hoy_control = fecha_referencia()

    _fuera_de_rango = anios_fuera_de_rango(
        anios_detectados,
        _hoy_control
    )

    if _fuera_de_rango:

        raise ValueError(
            "❌ EXPORTACIÓN BLOQUEADA: "
            "se detectaron movimientos con año fuera del rango válido "
            f"{ANIO_MINIMO_DATOS}–{_hoy_control.year} "
            f"(años fuera de rango: {_fuera_de_rango}; "
            f"fecha del equipo: {_hoy_control:%Y-%m-%d})."
        )

    print(
        "✅ Todos los movimientos están en el rango de años válido "
        f"({ANIO_MINIMO_DATOS}–{_hoy_control.year})"
    )

    # ========================================================
    # 7. VALIDAR SALDOS
    # ========================================================

    print(
        "\n💰 VALIDANDO SALDOS BANCARIOS...\n"
    )

    validaciones = []

    for _, fila in df_deteccion_final.iterrows():

        nombre = fila["ARCHIVO"]
        formato = fila["FORMATO"]
        ruta = mapa_archivos[nombre]

        datos_archivo = df_final[
            df_final[
                "ARCHIVO ORIGEN"
            ] == nombre
        ].copy()

        resultado = validar_extracto(
            ruta,
            detalle_deteccion[nombre],
            datos_archivo,
            contextos[nombre],
            normalizador=normalizador
        )

        validaciones.append({
            "ARCHIVO": nombre,
            "FORMATO": formato,
            "MOVIMIENTOS": len(datos_archivo),
            **resultado
        })

        icono = (
            "✅"
            if resultado["ESTADO"] == "OK"
            else "ℹ️"
            if resultado["ESTADO"] == "SIN MOVIMIENTOS"
            else "❌"
        )

        print(
            f"{icono} {formato}: "
            f"{resultado['ESTADO']}"
        )

    df_validacion = pd.DataFrame(
        validaciones
    )

    # ========================================================
    # BISA SIN MOVIMIENTOS ES VÁLIDO
    # ========================================================

    estados_invalidos = (
        ~df_validacion[
            "ESTADO"
        ]
        .isin(
            [
                "OK",
                "SIN MOVIMIENTOS"
            ]
        )
    )

    if estados_invalidos.any():

        print(
            "\n❌ HAY CUENTAS QUE NO CUADRAN:"
        )

        print(
            df_validacion[
                estados_invalidos
            ].to_string(index=False)
        )

        raise ValueError(
            "❌ EXPORTACIÓN BLOQUEADA."
        )

    print(
        "\n✅ INTEGRIDAD FINANCIERA SUPERADA"
    )

    # ========================================================
    # 8. RESUMEN
    # ========================================================

    resumen = (

        df_final
        .groupby(
            [
                "BANCO",
                "CUENTA BANCARIA",
                "MONEDA"
            ],
            dropna=False
        )
        .agg(

            MOVIMIENTOS=(
                "IMPORTE",
                "size"
            ),

            TOTAL_DEBITOS=(
                "DÉBITO",
                "sum"
            ),

            TOTAL_CREDITOS=(
                "CRÉDITO",
                "sum"
            ),

            PRIMERA_FECHA=(
                "FECHA MOVIMIENTO",
                "min"
            ),

            ULTIMA_FECHA=(
                "FECHA MOVIMIENTO",
                "max"
            )

        )
        .reset_index()
    )

    # ========================================================
    # 9. ORDENAR BASE FINAL
    # ========================================================

    df_final = (

        df_final
        .sort_values(
            [
                "FECHA MOVIMIENTO",
                "HORA MOVIMIENTO",
                "BANCO",
                "CUENTA BANCARIA"
            ],
            na_position="last"
        )
        .reset_index(
            drop=True
        )
    )

    # ========================================================
    # 10. RUTA DE SALIDA
    #
    # A diferencia de la celda original (que generaba un nombre
    # con timestamp, p. ej. CONTROL_DEPOSITOS_CBBA_PRUEBA_<ts>.xlsx,
    # y luego lo descargaba con files.download()), este motor
    # escribe siempre en la ruta fija indicada por el llamador
    # (NORMALIZADO.xlsx), y es la orquestación quien decide dónde
    # queda ese archivo y qué hacer con él después (subirlo a
    # Drive). El contenido y las hojas generadas no cambian.
    # ========================================================

    carpeta_salida = os.path.dirname(
        os.path.abspath(ruta_salida)
    )

    if carpeta_salida:
        os.makedirs(carpeta_salida, exist_ok=True)

    # ========================================================
    # 10-B. GENERAR LISTS.csv (SALIDA ADICIONAL)
    #
    # Requisito de esta etapa: además de NORMALIZADO.xlsx (que no
    # se toca ni se reemplaza), generar un archivo plano LISTS.csv
    # en la misma carpeta de salida, a partir del mismo df_final ya
    # validado (auditoría estructural, control de año y validación
    # de saldos ya superados en los pasos anteriores), restringido
    # a las columnas y al orden exacto de COLUMNAS_LISTS, con
    # codificación UTF-8 con BOM. Esto es una exportación adicional
    # de infraestructura: no modifica ningún dato, cálculo ni
    # validación del motor.
    # ========================================================

    ruta_lists_csv = os.path.join(
        carpeta_salida,
        "LISTS.csv"
    )

    df_final[COLUMNAS_LISTS].to_csv(
        ruta_lists_csv,
        index=False,
        encoding="utf-8-sig"
    )

    # ========================================================
    # 11. GENERAR EXCEL
    # ========================================================

    with pd.ExcelWriter(
        ruta_salida,
        engine="openpyxl"
    ) as writer:

        df_final.to_excel(
            writer,
            sheet_name="LISTS",
            index=False
        )

        df_validacion.to_excel(
            writer,
            sheet_name="VALIDACION",
            index=False
        )

        resumen.to_excel(
            writer,
            sheet_name="RESUMEN",
            index=False
        )

        df_resultado_archivos.to_excel(
            writer,
            sheet_name="DIAGNOSTICO",
            index=False
        )

    # ========================================================
    # 12. APLICAR FORMATO EXCEL COMPATIBLE CON LISTS
    # ========================================================

    wb = load_workbook(
        ruta_salida
    )

    ws = wb["LISTS"]

    # ========================================================
    # IDENTIFICAR COLUMNAS POR ENCABEZADO
    # ========================================================

    columnas_excel = {}

    for celda in ws[1]:

        columnas_excel[
            celda.value
        ] = celda.column

    # ========================================================
    # FORMATO FECHA MOVIMIENTO
    # ========================================================

    if "FECHA MOVIMIENTO" in columnas_excel:

        col = columnas_excel[
            "FECHA MOVIMIENTO"
        ]

        for fila in range(
            2,
            ws.max_row + 1
        ):

            celda = ws.cell(
                row=fila,
                column=col
            )

            if celda.value is not None:

                celda.number_format = (
                    "mm/dd/yy"
                )

    # ========================================================
    # FORMATO FECHA CONFIRMACIÓN
    # ========================================================

    if "FECHA CONFIRMACIÓN" in columnas_excel:

        col = columnas_excel[
            "FECHA CONFIRMACIÓN"
        ]

        for fila in range(
            2,
            ws.max_row + 1
        ):

            celda = ws.cell(
                row=fila,
                column=col
            )

            if celda.value is not None:

                celda.number_format = (
                    "mm/dd/yy hh:mm AM/PM"
                )

    # ========================================================
    # FORMATO FECHA DE CARGA
    # ========================================================

    if "FECHA DE CARGA" in columnas_excel:

        col = columnas_excel[
            "FECHA DE CARGA"
        ]

        for fila in range(
            2,
            ws.max_row + 1
        ):

            celda = ws.cell(
                row=fila,
                column=col
            )

            if celda.value is not None:

                celda.number_format = (
                    "mm/dd/yy hh:mm AM/PM"
                )

    # ========================================================
    # FORMATO DE NÚMEROS
    # ========================================================

    for nombre_columna in [
        "IMPORTE",
        "DÉBITO",
        "CRÉDITO",
        "SALDO"
    ]:

        if nombre_columna in columnas_excel:

            col = columnas_excel[
                nombre_columna
            ]

            for fila in range(
                2,
                ws.max_row + 1
            ):

                celda = ws.cell(
                    row=fila,
                    column=col
                )

                if celda.value is not None:

                    celda.number_format = (
                        "0.00"
                    )

    # ========================================================
    # CUENTA Y CÓDIGO COMO TEXTO VISUAL
    # ========================================================

    for nombre_columna in [
        "CUENTA BANCARIA",
        "CÓDIGO DE ASIGNACIÓN",
        "CLAVE TRANSACCIÓN"
    ]:

        if nombre_columna in columnas_excel:

            col = columnas_excel[
                nombre_columna
            ]

            for fila in range(
                2,
                ws.max_row + 1
            ):

                celda = ws.cell(
                    row=fila,
                    column=col
                )

                if celda.value is not None:

                    celda.number_format = "@"

                    celda.value = str(
                        celda.value
                    )

    # ========================================================
    # AJUSTAR FORMATO DE LAS FECHAS DEL RESUMEN
    # ========================================================

    ws_resumen = wb["RESUMEN"]

    encabezados_resumen = {}

    for celda in ws_resumen[1]:

        encabezados_resumen[
            celda.value
        ] = celda.column

    for nombre_columna in [
        "PRIMERA_FECHA",
        "ULTIMA_FECHA"
    ]:

        if nombre_columna in encabezados_resumen:

            col = encabezados_resumen[
                nombre_columna
            ]

            for fila in range(
                2,
                ws_resumen.max_row + 1
            ):

                celda = ws_resumen.cell(
                    row=fila,
                    column=col
                )

                if celda.value is not None:

                    celda.number_format = (
                        "mm/dd/yy"
                    )

    # ========================================================
    # CREAR TABLA EXCEL PARA POWER AUTOMATE / MICROSOFT LISTS
    # ========================================================

    ws_lists = wb["LISTS"]

    # La tabla se crea únicamente como estructura de Excel para que
    # Power Automate pueda enumerar las filas. No modifica valores,
    # cálculos, reglas de negocio ni el orden de COLUMNAS_LISTS.
    if (
        ws_lists.max_row >= 2
        and ws_lists.max_column >= 1
    ):

        nombres_tablas = {
            tabla.name
            for tabla in ws_lists.tables.values()
        }

        if "tblLISTS" not in nombres_tablas:

            ultima_celda = ws_lists.cell(
                row=ws_lists.max_row,
                column=ws_lists.max_column
            ).coordinate

            tabla_lists = Table(
                displayName="tblLISTS",
                ref=f"A1:{ultima_celda}"
            )

            ws_lists.add_table(
                tabla_lists
            )

    # ========================================================
    # GUARDAR CAMBIOS
    # ========================================================

    wb.save(
        ruta_salida
    )

    # ========================================================
    # 13. RESULTADO FINAL
    # ========================================================

    print(
        "\n============================================================"
    )

    print(
        "🎯 PROCESO COMPLETADO CORRECTAMENTE"
    )

    print(
        "============================================================"
    )

    print(
        f"Archivos procesados: "
        f"{len(df_deteccion_final)}"
    )

    print(
        f"Movimientos: "
        f"{len(df_final):,}"
    )

    print(
        f"Claves duplicadas: "
        f"{duplicados}"
    )

    print(
        f"Códigos vacíos: "
        f"{codigos_vacios}"
    )

    print(
        f"Errores estructurales: "
        f"{errores_estructurales}"
    )

    print(
        "Integridad bancaria: ✅"
    )

    print(
        "\n📅 FORMATO PARA MICROSOFT LISTS:"
    )

    print(
        "FECHA MOVIMIENTO   → MM/DD/YY"
    )

    print(
        "FECHA CONFIRMACIÓN → MM/DD/YY HH:MM AM/PM"
    )

    print(
        "FECHA DE CARGA     → MM/DD/YY HH:MM AM/PM"
    )

    print(
        f"\n📄 Archivo generado:"
    )

    print(
        ruta_salida
    )

    print(
        f"\n📄 Archivo adicional generado (LISTS.csv, UTF-8 con BOM):"
    )

    print(
        ruta_lists_csv
    )

    print(
        "\n🚫 MICROSOFT LISTS NO FUE MODIFICADO"
    )

    print(
        "============================================================"
    )

    # ========================================================
    # 14. MOSTRAR RESUMEN
    # (reemplaza display() por salida textual equivalente)
    # ========================================================

    print(
        "\nRESUMEN POR CUENTA:"
    )

    print(
        resumen.to_string(index=False)
    )

    print(
        "\nVALIDACIÓN:"
    )

    print(
        df_validacion.to_string(index=False)
    )

    # NOTA: aquí terminaba CELDA_3 con files.download(nombre_archivo).
    # Ese paso se elimina deliberadamente: subir NORMALIZADO.xlsx a
    # Google Drive es responsabilidad de la orquestación externa
    # (Cowork), no de este motor.

    # ========================================================
    # 15. CAPTURA ÍNTEGRA DE ORIGEN (P1, pasiva)
    #
    # Genera ORIGEN.xlsx (DATOS_ORIGINALES, METADATOS_EXTRACTO,
    # METADATOS_EXTRA, MAPA_ORIGEN) en la carpeta de salida, leyendo
    # los extractos originales por su cuenta. No modifica df_final,
    # NORMALIZADO.xlsx ni LISTS.csv, y un fallo aquí NUNCA detiene
    # el proceso (solo queda registrado en "origen_estado").
    # ========================================================
    ruta_origen = os.path.join(
        carpeta_salida or ".",
        "ORIGEN.xlsx"
    )

    try:
        import importlib.util as _iu

        _ruta_mod = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            "captura_origen.py"
        )
        _spec = _iu.spec_from_file_location(
            "captura_origen",
            _ruta_mod
        )
        _cap = _iu.module_from_spec(_spec)
        _spec.loader.exec_module(_cap)

        # P5: hoja y encabezado de cada cuenta según el registro (los
        # mismos que usaron la detección y la normalización).
        _contrato = contrato_origen_registro(
            detalle_deteccion,
            normalizador
        )

        origen_estado = _cap.generar_origen_seguro(
            mapa_archivos,
            df_deteccion_final,
            tablas,
            df_validacion,
            lote,
            fecha_carga,
            ruta_origen,
            _contrato
        )

    except Exception as e:

        origen_estado = {
            "estado": "ERROR",
            "error": f"{type(e).__name__}: {e}",
            "ruta_origen": None
        }

    if origen_estado.get("estado") == "OK":

        print(
            f"\n📄 Archivo técnico de origen (ORIGEN.xlsx): "
            f"{ruta_origen}"
        )

    else:

        print(
            "\n⚠️ ORIGEN.xlsx no se generó "
            "(no afecta NORMALIZADO.xlsx ni LISTS.csv): "
            f"{origen_estado.get('error')}"
        )

    return {
        "df_final": df_final,
        "df_validacion": df_validacion,
        "resumen": resumen,
        "df_resultado_archivos": df_resultado_archivos,
        "df_deteccion_final": df_deteccion_final,
        "ruta_salida": ruta_salida,
        "ruta_lists_csv": ruta_lists_csv,
        "origen_estado": origen_estado,
        "deteccion_estado": {
            "version_deteccion": detector.version,
            "ruta_registro": detector.ruta,
            "sha256_registro": detector.sha256,
            "archivos": {
                n: d.como_dict()
                for n, d in detalle_deteccion.items()
            }
        }
    }


# ============================================================
# EJECUCIÓN COMO SCRIPT
# ============================================================

if __name__ == "__main__":

    import sys

    if len(sys.argv) != 3:

        print(
            "Uso: python motor_control_depositos_cbba.py "
            "<carpeta_entrada> <ruta_salida_NORMALIZADO.xlsx>"
        )

        sys.exit(1)

    carpeta_entrada_cli = sys.argv[1]
    ruta_salida_cli = sys.argv[2]

    ejecutar_motor(
        carpeta_entrada_cli,
        ruta_salida_cli
    )
