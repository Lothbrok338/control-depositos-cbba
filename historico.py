# -*- coding: utf-8 -*-
"""
historico.py  ·  P3b · Capa 4: EXTRACTO_HISTORICO (Excel para Contabilidad / Ingresos)
=====================================================================================

Regla obligatoria: NORMALIZAR NUNCA DEBE DESTRUIR INFORMACIÓN DE ORIGEN.

Genera UN archivo .xlsx por banco / cuenta / mes calendario, con UNA sola hoja: EXTRACTO.

    Zona superior   datos que el propio extracto trae (titular, producto, período, emisión,
                    saldos y totales declarados por el banco), según el bloque `historico`
                    de cada formato en registro_bancos.json. Nada se inventa.
    Tabla           las columnas propias del banco, con su nombre y en su orden, y al final
                    ÚNICAMENTE: ESTADO · CONFIRMADO POR · FECHA DE CONFIRMACIÓN.
                    CLAVE TRANSACCIÓN va en una columna OCULTA de la misma tabla (se mueve
                    con el orden y los filtros) para sincronizaciones posteriores.

Fuentes (y nada más)
--------------------
DATOS_ORIGINALES + MAPA_ORIGEN (ORIGEN.xlsx)  +  NORMALIZADO (NORMALIZADO.xlsx hoja LISTS,
o LISTS.csv)  +  registro_bancos.json.
No abre ningún archivo bancario, no importa el motor ni captura_origen.py y no escribe en
NORMALIZADO.xlsx, LISTS.csv ni ORIGEN.xlsx.

Reglas de valores
-----------------
* Fecha, hora, débito, crédito, importe con signo y saldo salen de NORMALIZADO (los mismos
  números que van a Lists). El resto de las columnas sale del texto original del banco
  (DATOS_ORIGINALES), recortando solo los espacios de relleno al inicio y al final.
* Códigos, cheques, referencias, documentos: siempre texto (conservan ceros a la izquierda).
* Vacío = celda vacía (nunca un cero ni un guion inventado).
* Orden de filas: FECHA, HORA y, a igualdad, el orden original del archivo (FILA_EXCEL).

Verificación antes de escribir
------------------------------
Cada movimiento con su fila de origen; misma cantidad e importes que NORMALIZADO; saldos y
totales declarados por el banco (los marcados con `verifica` en el registro) cuadran con los
movimientos. Si algo falla, ese archivo NO se escribe y se informa la causa.

Uso
---
    construir_extractos_historicos(datos, mapa, normalizado, registro, estados=None) -> [resultado]
    escribir_extracto_historico(resultado, ruta)                 <- única parte que escribe Excel
    generar_extractos_historicos(ruta_origen, ruta_normalizado, carpeta_salida, ...)

    python historico.py <ORIGEN.xlsx> <NORMALIZADO.xlsx | LISTS.csv> <carpeta_salida> [registro.json]
"""
import csv
import datetime as _dt
import json
import math
import numbers
import os
import re
import sys
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.worksheet.table import Table, TableStyleInfo

VERSION_HISTORICO = "P3b-1.0"
RUTA_REGISTRO_DEFECTO = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "registro_bancos.json"
)
TOLERANCIA = 0.01  # misma tolerancia que la validación de saldos del motor

# Regla del usuario: al final de las columnas bancarias van ÚNICAMENTE estas tres.
COLUMNAS_OPERATIVAS = ["ESTADO", "CONFIRMADO POR", "FECHA DE CONFIRMACIÓN"]
ANCHO_OPERATIVAS = {"ESTADO": 14, "CONFIRMADO POR": 22, "FECHA DE CONFIRMACIÓN": 20}
COLUMNA_CLAVE = "CLAVE TRANSACCIÓN"

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre"]

SECCIONES = ("CUENTA", "PERIODO", "SALDOS")
FUENTES_NORMALIZADO = {
    "NORMALIZADO:FECHA", "NORMALIZADO:HORA", "NORMALIZADO:DEBITO", "NORMALIZADO:CREDITO",
    "NORMALIZADO:IMPORTE_FIRMADO", "NORMALIZADO:SALDO",
}

_ILEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# =====================================================================
# 1. VALORES DE ENTRADA (DataFrame, lista de dicts, lectura de .xlsx/.csv)
# =====================================================================

def _es_nulo(v):
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    return type(v).__name__ == "NaTType"


def _texto(v):
    """Valor técnico (ID, hoja, fila, rol) -> texto; los textos se devuelven intactos."""
    if _es_nulo(v):
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _registros(tabla):
    if tabla is None:
        return []
    if hasattr(tabla, "to_dict"):
        return tabla.to_dict("records")
    return list(tabla)


def _fecha(v):
    if _es_nulo(v) or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, _dt.datetime):
        return _dt.date(v.year, v.month, v.day)
    if isinstance(v, _dt.date):
        return v
    m = re.match(r"^\s*(\d{4})-(\d{2})-(\d{2})", str(v))
    if not m:
        raise ValueError(f"FECHA MOVIMIENTO no interpretable: {v!r}")
    return _dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))


def _hora(v):
    if _es_nulo(v):
        return None
    if isinstance(v, _dt.datetime):
        return _dt.time(v.hour, v.minute, v.second)
    if isinstance(v, _dt.time):
        return _dt.time(v.hour, v.minute, v.second)
    s = str(v).strip()
    if not s:
        return None
    m = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", s)
    if not m:
        raise ValueError(f"HORA MOVIMIENTO no interpretable: {v!r}")
    return _dt.time(int(m.group(1)), int(m.group(2)), int(m.group(3) or 0))


def _numero(v):
    if _es_nulo(v) or isinstance(v, bool):
        return None
    if isinstance(v, numbers.Number):
        return float(v)
    s = str(v).strip()
    return float(s) if s else None


def _fecha_hora(v):
    if _es_nulo(v):
        return None
    if isinstance(v, _dt.datetime):
        return _dt.datetime(v.year, v.month, v.day, v.hour, v.minute, v.second)
    if isinstance(v, _dt.date):
        return _dt.datetime(v.year, v.month, v.day)
    s = str(v).strip()
    if not s:
        return None
    try:
        d = _dt.datetime.fromisoformat(s)
    except ValueError:
        raise ValueError(f"FECHA DE CONFIRMACIÓN no interpretable: {v!r}")
    return _dt.datetime(d.year, d.month, d.day, d.hour, d.minute, d.second)


def _fila_normalizada(r):
    if _fecha(r.get("FECHA MOVIMIENTO")) is None:
        raise ValueError(f"NORMALIZADO: fila sin FECHA MOVIMIENTO ({_texto(r.get(COLUMNA_CLAVE))})")
    return {
        "banco": _texto(r.get("BANCO")).strip(),
        "cuenta": _texto(r.get("CUENTA BANCARIA")).strip(),
        "moneda": _texto(r.get("MONEDA")).strip(),
        "fecha": _fecha(r.get("FECHA MOVIMIENTO")),
        "hora": _hora(r.get("HORA MOVIMIENTO")),
        "importe": _numero(r.get("IMPORTE")),
        "debito": _numero(r.get("DÉBITO")),
        "credito": _numero(r.get("CRÉDITO")),
        "tipo": _texto(r.get("TIPO MOVIMIENTO")).strip(),
        "saldo": _numero(r.get("SALDO")),
        "estado": _texto(r.get("ESTADO")),
        "confirmado_por": _texto(r.get("CONFIRMADO POR")),
        "fecha_confirmacion": _fecha_hora(r.get("FECHA CONFIRMACIÓN")),
    }


def _hoja_a_dicts(ws):
    filas = ws.iter_rows(values_only=True)
    try:
        enc = [str(x) if x is not None else "" for x in next(filas)]
    except StopIteration:
        return []
    out = []
    for fila in filas:
        if all(v is None for v in fila):
            continue
        out.append(dict(zip(enc, fila)))
    return out


def leer_origen(ruta_origen):
    """ORIGEN.xlsx -> (DATOS_ORIGINALES, MAPA_ORIGEN) como listas de dicts (texto exacto)."""
    wb = load_workbook(ruta_origen, read_only=True)
    try:
        return _hoja_a_dicts(wb["DATOS_ORIGINALES"]), _hoja_a_dicts(wb["MAPA_ORIGEN"])
    finally:
        wb.close()


def leer_normalizado(ruta_normalizado):
    """NORMALIZADO.xlsx (hoja LISTS) o LISTS.csv (UTF-8 con BOM) -> lista de dicts."""
    if str(ruta_normalizado).lower().endswith(".csv"):
        with open(ruta_normalizado, encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    wb = load_workbook(ruta_normalizado, read_only=True)
    try:
        return _hoja_a_dicts(wb["LISTS"])
    finally:
        wb.close()


def cargar_registro(ruta=None):
    ruta = ruta or os.environ.get("CBBA_REGISTRO_BANCOS") or RUTA_REGISTRO_DEFECTO
    with open(ruta, encoding="utf-8") as f:
        return json.load(f)


def _indice_estados(estados):
    """Tabla opcional ESTADOS (clave = CLAVE TRANSACCIÓN) que reemplaza los valores operativos."""
    if estados is None:
        return {}
    if isinstance(estados, dict):
        filas = [dict(v, **{COLUMNA_CLAVE: k}) for k, v in estados.items()]
    else:
        filas = _registros(estados)
    out = {}
    for r in filas:
        clave = _texto(r.get(COLUMNA_CLAVE)).strip()
        if not clave:
            continue
        d = {}
        if not _es_nulo(r.get("ESTADO")) and "ESTADO" in r:
            d["estado"] = _texto(r["ESTADO"])
        if not _es_nulo(r.get("CONFIRMADO POR")) and "CONFIRMADO POR" in r:
            d["confirmado_por"] = _texto(r["CONFIRMADO POR"])
        for k in ("FECHA DE CONFIRMACIÓN", "FECHA CONFIRMACIÓN"):
            if k in r and not _es_nulo(r[k]) and _texto(r[k]).strip():
                d["fecha_confirmacion"] = _fecha_hora(r[k])
                break
        out[clave] = d
    return out


# =====================================================================
# 2. UTILIDADES DE TEXTO
# =====================================================================

def _norm(txt):
    t = unicodedata.normalize("NFD", str(txt))
    t = "".join(ch for ch in t if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", t).strip().upper()


def _norm_etiqueta(txt):
    return _norm(txt).rstrip(":").strip()


def _limpiar_encabezado(txt):
    """Nombre del banco, sin saltos de línea ni espacios repetidos ('Monto\\n' -> 'Monto')."""
    return re.sub(r"\s+", " ", str(txt)).strip()


def _seguro(txt):
    t = unicodedata.normalize("NFD", str(txt))
    t = "".join(ch for ch in t if unicodedata.category(ch) != "Mn").upper()
    t = re.sub(r"[^A-Z0-9-]+", "_", t)
    return re.sub(r"_+", "_", t).strip("_")


def _importe_de_texto(texto, separador_decimal):
    """'62.687,52' / 'Bs 201,662.79' / '0.00' -> float, con la convención del banco."""
    m = re.search(r"-?\d[\d.,]*", texto)
    if not m:
        return None
    t = m.group(0).rstrip(".,")
    if separador_decimal == ",":
        t = t.replace(".", "").replace(",", ".")
    else:
        t = t.replace(",", "")
    try:
        return float(t)
    except ValueError:
        return None


def _es_importe_puro(texto):
    return re.fullmatch(r"\s*-?\s*[\d.,]+\s*-?\s*", texto) is not None


def _fmt_fecha(d):
    return d.strftime("%d/%m/%Y")


def _fmt_importe(v):
    return f"{v:,.2f}"


# =====================================================================
# 3. EXTRACTO: columnas, zona superior y verificación
# =====================================================================

class _Extracto:
    """Todo lo que la capa 4 necesita de UN extracto, tomado de DATOS_ORIGINALES."""

    def __init__(self, id_extracto, hoja, orden):
        self.id = id_extracto
        self.hoja = hoja
        self.orden = orden
        self.archivo = ""
        self.formato_legado = set()
        self.encabezado = {}          # col -> texto de encabezado
        self.movimientos = {}         # id_origen -> {col: texto}
        self.zona = {}                # (fila, col) -> (rol, texto)  (todo lo que no es tabla)
        self.claves = []
        self.errores = []
        self.advertencias = []

    def filas_zona(self):
        d = {}
        for (f, c), (_, v) in self.zona.items():
            d.setdefault(f, {})[c] = v
        return d


def _resolver_columnas(ext, hist):
    """Columnas bancarias a mostrar: las del encabezado real del extracto, en su orden."""
    specs = hist.get("columnas", [])
    indice = {}
    for s in specs:
        for nombre in [s["origen"]] + list(s.get("alias", [])):
            indice.setdefault(_norm(nombre), s)
    usadas = set()
    cols_mov = set()
    for celdas in ext.movimientos.values():
        cols_mov.update(celdas)
    columnas = []
    for col in sorted(set(ext.encabezado) | cols_mov):
        letra = get_column_letter(col)
        texto = ext.encabezado.get(col, "")
        if not texto.strip():
            columnas.append({"col": col, "titulo": f"Columna {letra} (sin encabezado)",
                             "tipo": "texto", "fuente": "ORIGINAL", "ancho": 16, "nueva": True})
            ext.advertencias.append(
                f"celdas de movimiento en la columna {letra} sin encabezado: se muestran como "
                f"'Columna {letra} (sin encabezado)'")
            continue
        spec = indice.get(_norm(texto))
        titulo = _limpiar_encabezado(texto)
        if spec is None:
            columnas.append({"col": col, "titulo": titulo, "tipo": "texto", "fuente": "ORIGINAL",
                             "ancho": max(10, min(40, len(titulo) + 4)), "nueva": True})
            ext.advertencias.append(
                f"columna nueva en el extracto (no está en el registro): '{titulo}'; se muestra "
                "en su posición, antes de las columnas operativas")
            continue
        usadas.add(id(spec))
        columnas.append({"col": col, "titulo": spec.get("titulo") or titulo,
                         "tipo": spec.get("tipo", "texto"), "fuente": spec.get("fuente", "ORIGINAL"),
                         "ancho": spec.get("ancho", 14)})
    for s in specs:
        if id(s) not in usadas and not s.get("opcional"):
            ext.advertencias.append(f"columna del registro ausente en el extracto: '{s['origen']}'")
    for c in columnas:
        if c["fuente"] != "ORIGINAL" and c["fuente"] not in FUENTES_NORMALIZADO:
            ext.errores.append(f"columna '{c['titulo']}': fuente desconocida '{c['fuente']}'")
        if c["tipo"] not in ("texto", "codigo", "importe", "fecha", "hora"):
            ext.errores.append(f"columna '{c['titulo']}': tipo desconocido '{c['tipo']}'")
    # títulos únicos (requisito de la tabla de Excel)
    vistos = {}
    for c in columnas:
        base = c["titulo"]
        n = vistos.get(_norm(base), 0)
        vistos[_norm(base)] = n + 1
        if n:
            c["titulo"] = f"{base} ({n + 1})"
    return columnas


def _buscar_etiqueta(ext, filas, item):
    objetivo = _norm_etiqueta(item["etiqueta"])
    posicion = item.get("posicion", "DERECHA")
    for (f, c) in sorted(ext.zona):
        v = ext.zona[(f, c)][1]
        if not v.strip() or _norm_etiqueta(v) != objetivo:
            continue
        if posicion == "DERECHA":
            derecha = sorted(cc for cc, x in filas[f].items() if cc > c and x.strip())
            if not derecha:
                return None
            valor = filas[f][derecha[0]]
            return None if valor.strip().endswith(":") else valor
        if posicion == "DEBAJO":
            siguientes = sorted(ff for ff, cs in filas.items()
                                if ff > f and any(x.strip() for x in cs.values()))
            if not siguientes:
                return None
            etiquetas = sorted(cc for cc, x in filas[f].items() if x.strip())
            valores = sorted(cc for cc, x in filas[siguientes[0]].items() if x.strip())
            if len(etiquetas) == len(valores):
                c2 = valores[etiquetas.index(c)]
            else:
                c2 = min(valores, key=lambda cc: (abs(cc - c), cc))
            return filas[siguientes[0]][c2]
        raise ValueError(f"posicion desconocida '{posicion}'")
    return None


def _evaluar_item(ext, filas, item):
    """-> (texto, sufijo_del_titulo) o (None, None) si el extracto no trae el dato."""
    fuente = item.get("fuente", "ETIQUETA")
    if fuente == "ETIQUETA":
        return _buscar_etiqueta(ext, filas, item), None
    if fuente == "PATRON":
        rx = re.compile(item["patron"], re.IGNORECASE)
        for k in sorted(ext.zona):
            m = rx.search(ext.zona[k][1].strip())
            if m:
                grupos = [g.strip() for g in (m.groups() or (m.group(0),)) if g is not None]
                plantilla = item.get("plantilla")
                return (plantilla.format(*grupos) if plantilla else grupos[0]), None
        return None, None
    if fuente == "CELDA":
        letra, fila = coordinate_from_string(item["celda"])
        v = ext.zona.get((fila, column_index_from_string(letra)))
        return (v[1] if v and v[1].strip() else None), None
    if fuente == "FILA_ROTULADA":
        rotulo = _norm(item["rotulo"])
        cols = {_norm(t): c for c, t in ext.encabezado.items()}
        c_val = cols.get(_norm(item["columna"]))
        c_fec = cols.get(_norm(item.get("columna_fecha", ""))) if item.get("columna_fecha") else None
        for f in sorted(filas):
            if any(_norm(x) == rotulo for x in filas[f].values()):
                v = filas[f].get(c_val) if c_val else None
                fecha = filas[f].get(c_fec, "").strip() if c_fec else ""
                return v, (fecha or None)
        return None, None
    if fuente == "COMPUESTO":
        partes = []
        for p in item.get("partes", []):
            v, _ = _evaluar_item(ext, filas, p)
            if v is None or not v.strip():
                return None, None
            partes.append(v.strip())
        return item.get("plantilla", " ".join("{%d}" % i for i in range(len(partes)))).format(*partes), None
    raise ValueError(f"fuente de encabezado desconocida '{fuente}'")


def _encabezado_extracto(ext, hist):
    """Datos de la zona superior que el registro declara y el extracto trae."""
    filas = ext.filas_zona()
    sep = hist.get("separador_decimal", ".")
    items, verificaciones = [], []
    for item in hist.get("encabezado", []):
        seccion = item.get("seccion", "CUENTA")
        if seccion not in SECCIONES:
            ext.errores.append(f"encabezado '{item.get('titulo')}': sección desconocida '{seccion}'")
            continue
        try:
            texto, sufijo = _evaluar_item(ext, filas, item)
        except (ValueError, KeyError, IndexError, re.error) as e:
            ext.errores.append(f"encabezado '{item.get('titulo')}': configuración inválida ({e})")
            continue
        titulo = item["titulo"] + (f" ({sufijo})" if sufijo else "")
        if texto is None or not texto.strip():
            if item.get("verifica"):
                ext.advertencias.append(
                    f"el extracto no trae '{item['titulo']}': no se pudo verificar {item['verifica']}")
            continue
        texto = texto.strip()
        tipo = item.get("tipo", "texto")
        valor = texto
        if tipo == "importe":
            valor = _importe_de_texto(texto, sep)
            if valor is None:
                ext.advertencias.append(f"'{item['titulo']}' no es un importe: se muestra como texto")
                valor, tipo = texto, "texto"
        items.append({"seccion": seccion, "titulo": titulo, "valor": valor, "tipo": tipo})
        if item.get("verifica") and tipo == "importe":
            verificaciones.append((item["verifica"], valor, item["titulo"]))
    return items, verificaciones


def _clave_orden(n, m, orden_extracto):
    h = n["hora"]
    return (n["fecha"], h is not None, h or _dt.time(0), orden_extracto, m["fila"])


def _verificar_declarados(ext, filas, verificaciones):
    """Saldos y totales que el banco declara vs. los movimientos del extracto (orden mostrado)."""
    if not filas or not verificaciones:
        return
    primero, ultimo = filas[0][1], filas[-1][1]
    cred = sum(n["credito"] or 0.0 for _, n, _ in filas)
    deb = sum(n["debito"] or 0.0 for _, n, _ in filas)
    calc = {"TOTAL_CREDITOS": cred, "TOTAL_DEBITOS": deb}
    if primero["saldo"] is not None:
        calc["SALDO_INICIAL"] = primero["saldo"] - (primero["credito"] or 0.0) + (primero["debito"] or 0.0)
    if ultimo["saldo"] is not None:
        calc["SALDO_FINAL"] = ultimo["saldo"]
    for tipo, valor, titulo in verificaciones:
        if tipo not in ("SALDO_INICIAL", "SALDO_FINAL", "TOTAL_CREDITOS", "TOTAL_DEBITOS"):
            ext.errores.append(f"'{titulo}': verificación desconocida '{tipo}'")
            continue
        if tipo not in calc:
            ext.errores.append(f"'{titulo}': no se puede verificar {tipo} (movimientos sin saldo)")
            continue
        if abs(valor - calc[tipo]) > TOLERANCIA:
            ext.errores.append(
                f"'{titulo}' no cuadra: el banco declara {_fmt_importe(valor)} y los movimientos "
                f"dan {_fmt_importe(calc[tipo])} (diferencia {_fmt_importe(valor - calc[tipo])})")


# =====================================================================
# 4. FUNCIÓN PURA
# =====================================================================

def _valor_original(texto, tipo, sep, col, ext):
    s = texto.strip()
    if not s:
        return None
    if tipo == "importe":
        if _es_importe_puro(s):
            v = _importe_de_texto(s, sep)
            if v is not None:
                return v
        aviso = f"columna '{col['titulo']}': hay valores que no son importes; se muestran como texto"
        if aviso not in ext.advertencias:
            ext.advertencias.append(aviso)
    return s


def _valor_normalizado(fuente, n):
    if fuente == "NORMALIZADO:FECHA":
        return n["fecha"]
    if fuente == "NORMALIZADO:HORA":
        return n["hora"]
    if fuente == "NORMALIZADO:DEBITO":
        return n["debito"]
    if fuente == "NORMALIZADO:CREDITO":
        return n["credito"]
    if fuente == "NORMALIZADO:SALDO":
        return n["saldo"]
    if fuente == "NORMALIZADO:IMPORTE_FIRMADO":
        if n["importe"] is None:
            return None
        return -n["importe"] if n["tipo"] == "DÉBITO" else n["importe"]
    raise ValueError(fuente)


def _unir_encabezados(listas):
    """Un solo extracto por grupo es lo normal; si hubiera varios, se muestran todos los valores."""
    out, pos = [], {}
    for items in listas:
        for it in items:
            k = (it["seccion"], it["titulo"])
            if k not in pos:
                pos[k] = len(out)
                out.append(dict(it))
            elif out[pos[k]]["valor"] != it["valor"]:
                prev = out[pos[k]]
                a = _fmt_importe(prev["valor"]) if prev["tipo"] == "importe" else str(prev["valor"])
                b = _fmt_importe(it["valor"]) if it["tipo"] == "importe" else str(it["valor"])
                prev["valor"], prev["tipo"] = f"{a} | {b}", "texto"
    return out


def extractos_sin_movimientos(datos_originales, mapa_origen):
    """Archivos presentes en DATOS_ORIGINALES que no aportan movimientos (no generan histórico)."""
    con_mov = {_texto(r.get("ID_EXTRACTO")) for r in _registros(mapa_origen)}
    out = {}
    for r in _registros(datos_originales):
        i = _texto(r.get("ID_EXTRACTO"))
        if i and i not in con_mov and i not in out:
            out[i] = _texto(r.get("ARCHIVO ORIGEN"))
    return sorted(out.values())


def construir_extractos_historicos(datos_originales, mapa_origen, normalizado, registro, estados=None):
    """
    Función pura: no lee ni escribe archivos. Devuelve un resultado por
    (BANCO, CUENTA BANCARIA, MONEDA, AAAA-MM) ordenado por nombre de archivo.

    datos_originales / mapa_origen : hojas de ORIGEN.xlsx (DataFrame o lista de dicts)
    normalizado                    : filas de LISTS (NORMALIZADO.xlsx o LISTS.csv)
    registro                       : dict de registro_bancos.json (o ruta)
    estados                        : opcional, {CLAVE: {ESTADO, CONFIRMADO POR, FECHA DE CONFIRMACIÓN}}
                                     o tabla con esas columnas; reemplaza los valores operativos
    """
    if isinstance(registro, (str, os.PathLike)):
        registro = cargar_registro(registro)
    registro = getattr(registro, "datos", registro)
    general = registro.get("HISTORICO", {})
    formatos = registro.get("FORMATOS", {})
    cuentas = {c.get("id"): c for c in registro.get("CUENTAS", [])}
    idx_estados = _indice_estados(estados)

    # -- NORMALIZADO
    norm = {}
    for r in _registros(normalizado):
        clave = _texto(r.get(COLUMNA_CLAVE)).strip()
        if not clave:
            raise ValueError("NORMALIZADO: hay una fila sin CLAVE TRANSACCIÓN")
        if clave in norm:
            raise ValueError(f"NORMALIZADO: CLAVE TRANSACCIÓN repetida: {clave}")
        norm[clave] = _fila_normalizada(r)

    # -- MAPA_ORIGEN
    mapa, extractos = {}, {}
    for r in _registros(mapa_origen):
        clave = _texto(r.get(COLUMNA_CLAVE)).strip()
        if clave in mapa:
            raise ValueError(f"MAPA_ORIGEN: CLAVE TRANSACCIÓN repetida: {clave}")
        m = {"id_origen": _texto(r.get("ID_ORIGEN")), "id_extracto": _texto(r.get("ID_EXTRACTO")),
             "hoja": _texto(r.get("HOJA")), "fila": int(_texto(r.get("FILA_EXCEL"))),
             "formato": _texto(r.get("FORMATO")).strip()}
        mapa[clave] = m
        ext = extractos.get(m["id_extracto"])
        if ext is None:
            ext = extractos[m["id_extracto"]] = _Extracto(m["id_extracto"], m["hoja"], len(extractos))
            ext.archivo = _texto(r.get("ARCHIVO ORIGEN"))
        if m["hoja"] != ext.hoja:
            ext.errores.append("MAPA_ORIGEN: movimientos del mismo extracto en hojas distintas")
        ext.formato_legado.add(m["formato"])
        ext.claves.append(clave)

    # -- DATOS_ORIGINALES (solo la hoja de movimientos de los extractos con movimientos)
    for r in _registros(datos_originales):
        ext = extractos.get(_texto(r.get("ID_EXTRACTO")))
        if ext is None or _texto(r.get("HOJA")) != ext.hoja:
            continue
        col = column_index_from_string(_texto(r.get("COLUMNA_EXCEL")))
        fila = int(_texto(r.get("FILA_EXCEL")))
        rol = _texto(r.get("ROL_FILA"))
        valor = _texto(r.get("VALOR_ORIGINAL"))
        if rol == "ENCABEZADO_TABLA":
            ext.encabezado[col] = valor
        elif rol == "MOVIMIENTO":
            ext.movimientos.setdefault(_texto(r.get("ID_ORIGEN")), {})[col] = valor
        else:
            ext.zona[(fila, col)] = (rol, valor)

    grupos = {}   # (banco, cuenta, moneda, periodo) -> dict

    def grupo(banco, cuenta, moneda, periodo):
        k = (banco, cuenta, moneda, periodo)
        if k not in grupos:
            grupos[k] = {"extractos": [], "filas": [], "errores": [], "advertencias": []}
        return grupos[k]

    # -- movimientos de NORMALIZADO sin fila de origen: se informan (su archivo no se escribe)
    for clave, n in norm.items():
        if clave not in mapa:
            g = grupo(n["banco"], n["cuenta"], n["moneda"], n["fecha"].strftime("%Y-%m"))
            g["errores"].append(f"movimiento de NORMALIZADO sin fila de origen en MAPA_ORIGEN: {clave}")

    # -- extracto por extracto
    for ext in sorted(extractos.values(), key=lambda e: e.orden):
        faltan = [c for c in ext.claves if c not in norm]
        if faltan:
            ext.errores.append(f"MAPA_ORIGEN tiene {len(faltan)} movimiento(s) que no están en NORMALIZADO "
                               f"(p. ej. {faltan[0]})")
        filas = sorted(((c, norm[c], mapa[c]) for c in ext.claves if c in norm),
                       key=lambda t: _clave_orden(t[1], t[2], ext.orden))
        identidades = {(n["banco"], n["cuenta"], n["moneda"]) for _, n, _ in filas}
        if len(identidades) != 1:
            ext.errores.append(f"el extracto {ext.archivo} mezcla cuentas en NORMALIZADO: {sorted(identidades)}")
        banco, cuenta, moneda = sorted(identidades)[0] if identidades else ("", "", "")
        if len(ext.formato_legado) != 1:
            ext.errores.append(f"MAPA_ORIGEN: formato ambiguo {sorted(ext.formato_legado)}")
        cuenta_cfg = cuentas.get(next(iter(sorted(ext.formato_legado)), ""))
        if cuenta_cfg is None:
            cand = [c for c in cuentas.values() if str(c.get("cuenta", "")).strip() == cuenta]
            cuenta_cfg = cand[0] if len(cand) == 1 else None
        fmt, hist = {}, None
        if cuenta_cfg is None:
            ext.errores.append(f"cuenta {cuenta} ({ext.archivo}) no registrada en registro_bancos.json")
        else:
            reg_ident = (str(cuenta_cfg.get("banco", "")).strip(), str(cuenta_cfg.get("cuenta", "")).strip(),
                         str(cuenta_cfg.get("moneda", "")).strip())
            if identidades and reg_ident != (banco, cuenta, moneda):
                ext.errores.append(f"el registro dice {reg_ident} y NORMALIZADO dice {(banco, cuenta, moneda)}")
            fmt = formatos.get(cuenta_cfg.get("formato"), {})
            hist = fmt.get("historico")
            if hist is None:
                ext.errores.append(f"el formato {cuenta_cfg.get('formato')} no tiene bloque 'historico' en el registro")
        sin_origen = [c for c, _, m in filas if not ext.movimientos.get(m["id_origen"])]
        if sin_origen:
            ext.errores.append(f"{len(sin_origen)} movimiento(s) sin celdas de origen en DATOS_ORIGINALES "
                               f"(p. ej. {sin_origen[0]})")
        if not ext.encabezado:
            ext.errores.append(f"DATOS_ORIGINALES no trae la fila de encabezado de {ext.archivo}")
        especiales = {f for (f, _), (rol, _) in ext.zona.items() if rol == "FILA_ESPECIAL"}
        if especiales:
            ext.advertencias.append(
                f"el extracto tiene {len(especiales)} fila(s) dentro de la tabla que el motor no normalizó "
                "(no aparecen como movimientos; siguen en ORIGEN.xlsx)")

        columnas, items = [], []
        if hist is not None:
            columnas = _resolver_columnas(ext, hist)
            items, verificaciones = _encabezado_extracto(ext, hist)
            _verificar_declarados(ext, filas, verificaciones)

        for clave, n, m in filas:
            g = grupo(banco, cuenta, moneda, n["fecha"].strftime("%Y-%m"))
            if not g["extractos"] or g["extractos"][-1]["ext"] is not ext:
                g["extractos"].append({"ext": ext, "columnas": columnas, "items": items, "hist": hist,
                                       "cuenta_cfg": cuenta_cfg, "fmt": fmt})
            g["filas"].append((clave, n, m, ext))

    # -- un resultado por grupo
    resultados = []
    for (banco, cuenta, moneda, periodo), g in grupos.items():
        resultados.append(_armar_resultado(banco, cuenta, moneda, periodo, g, general, idx_estados))
    resultados.sort(key=lambda r: (r["nombre_archivo"], r["periodo"]))
    return resultados


def _armar_resultado(banco, cuenta, moneda, periodo, g, general, idx_estados):
    anio, mes = (int(x) for x in periodo.split("-"))
    nombre = general.get("nombre_archivo", "EXTRACTO_HISTORICO_{BANCO}_{CUENTA}_{PERIODO}.xlsx").format(
        BANCO=_seguro(banco), CUENTA=_seguro(cuenta), MONEDA=_seguro(moneda), PERIODO=periodo)
    errores = list(g["errores"])
    advertencias = list(g["advertencias"])
    for e in g["extractos"]:
        errores += [x for x in e["ext"].errores if x not in errores]
        advertencias += [x for x in e["ext"].advertencias if x not in advertencias]
    firmas = {tuple((c["titulo"], c["tipo"], c["fuente"]) for c in e["columnas"]) for e in g["extractos"]}
    if len(firmas) > 1:
        errores.append("el mes reúne extractos con columnas distintas; no se puede armar una sola tabla")
    base = g["extractos"][0] if g["extractos"] else {"columnas": [], "items": [], "hist": None,
                                                    "cuenta_cfg": None, "fmt": {}}
    hist = base["hist"] or {}
    columnas_banco = base["columnas"]
    filas_orden = sorted(g["filas"], key=lambda t: _clave_orden(t[1], t[2], t[3].orden))

    columnas = [dict(c, bloque="BANCO") for c in columnas_banco]
    columnas += [{"titulo": t, "tipo": {"FECHA DE CONFIRMACIÓN": "fecha_hora"}.get(t, "texto"),
                  "ancho": ANCHO_OPERATIVAS[t], "bloque": "OPERATIVO"} for t in COLUMNAS_OPERATIVAS]
    clave_oculta = bool(general.get("clave_oculta", True))
    if clave_oculta:
        columnas.append({"titulo": COLUMNA_CLAVE, "tipo": "texto", "ancho": 12, "bloque": "CLAVE", "oculta": True})

    sep = hist.get("separador_decimal", ".")
    filas = []
    if not errores:
        for clave, n, m, ext in filas_orden:
            celdas = ext.movimientos.get(m["id_origen"], {})
            fila = []
            for c in columnas_banco:
                if c["fuente"] == "ORIGINAL":
                    fila.append(_valor_original(celdas.get(c["col"], ""), c["tipo"], sep, c, ext))
                else:
                    fila.append(_valor_normalizado(c["fuente"], n))
            est = idx_estados.get(clave, {})
            fila.append(est.get("estado", n["estado"]) or None)
            fila.append(est.get("confirmado_por", n["confirmado_por"]) or None)
            fila.append(est.get("fecha_confirmacion", n["fecha_confirmacion"]))
            if clave_oculta:
                fila.append(clave)
            filas.append(fila)
        for e in g["extractos"]:
            advertencias += [x for x in e["ext"].advertencias if x not in advertencias]
        errores += _verificar_grupo(columnas_banco, filas, [t[1] for t in filas_orden])

    fechas = [t[1]["fecha"] for t in filas_orden]
    movs = len(filas_orden)
    rango = ""
    if fechas:
        d1, d2 = min(fechas), max(fechas)
        rango = f" ({_fmt_fecha(d1)})" if d1 == d2 else f" (del {_fmt_fecha(d1)} al {_fmt_fecha(d2)})"
    encabezado = [
        {"seccion": "CUENTA", "titulo": "Banco", "valor": banco, "tipo": "texto"},
        {"seccion": "CUENTA", "titulo": "Cuenta", "valor": cuenta, "tipo": "texto"},
        {"seccion": "CUENTA", "titulo": "Moneda", "valor": moneda, "tipo": "texto"},
    ]
    extra = _unir_encabezados([e["items"] for e in g["extractos"]])
    encabezado += [i for i in extra if i["seccion"] == "CUENTA"]
    encabezado += [
        {"seccion": "PERIODO", "titulo": "Periodo", "valor": f"{MESES[mes - 1]} {anio}", "tipo": "texto"},
        {"seccion": "PERIODO", "titulo": "Movimientos del mes", "valor": f"{movs}{rango}", "tipo": "texto"},
    ]
    encabezado += [i for i in extra if i["seccion"] == "PERIODO"]
    encabezado += [i for i in extra if i["seccion"] == "SALDOS"]

    cuenta_cfg = base["cuenta_cfg"] or {}
    return {
        "nombre_archivo": nombre,
        "banco": banco, "cuenta": cuenta, "moneda": moneda, "periodo": periodo,
        "id_cuenta": cuenta_cfg.get("id", ""), "formato": cuenta_cfg.get("formato", ""),
        "archivos_origen": [e["ext"].archivo for e in g["extractos"]],
        "titulo": f"EXTRACTO BANCARIO · {banco} · Cuenta {cuenta} · {moneda} · {MESES[mes - 1]} {anio}",
        "hoja": general.get("hoja", "EXTRACTO"),
        "encabezado": encabezado,
        "columnas": columnas,
        "filas": filas,
        "claves": [t[0] for t in filas_orden],
        "totales": {"movimientos": movs,
                    "creditos": round(sum(t[1]["credito"] or 0.0 for t in filas_orden), 2),
                    "debitos": round(sum(t[1]["debito"] or 0.0 for t in filas_orden), 2)},
        "formatos": {"fecha": general.get("formato_fecha", "dd/mm/yyyy"),
                     "hora": hist.get("formato_hora", "hh:mm:ss"),
                     "importe": general.get("formato_importe", "#,##0.00"),
                     "fecha_hora": general.get("formato_fecha_confirmacion", "dd/mm/yyyy hh:mm")},
        "colores": {"banco": general.get("color_banco", "7B1E2B"),
                    "operativo": general.get("color_operativo", "3F4447")},
        "advertencias": advertencias,
        "errores": errores,
    }


def _verificar_grupo(columnas_banco, filas, normalizadas):
    """Los importes mostrados son exactamente los de NORMALIZADO (misma cantidad y mismas sumas)."""
    errores = []
    if len(filas) != len(normalizadas):
        errores.append(f"se armaron {len(filas)} filas para {len(normalizadas)} movimientos de NORMALIZADO")
        return errores
    cred = sum(n["credito"] or 0.0 for n in normalizadas)
    deb = sum(n["debito"] or 0.0 for n in normalizadas)
    esperado = {"NORMALIZADO:CREDITO": cred, "NORMALIZADO:DEBITO": deb,
                "NORMALIZADO:IMPORTE_FIRMADO": cred - deb}
    for i, c in enumerate(columnas_banco):
        if c["fuente"] in esperado:
            suma = sum(f[i] or 0.0 for f in filas)
            if abs(suma - esperado[c["fuente"]]) > 0.005:
                errores.append(f"columna '{c['titulo']}': suma {_fmt_importe(suma)} distinta de "
                               f"NORMALIZADO {_fmt_importe(esperado[c['fuente']])}")
    return errores


# =====================================================================
# 5. ESCRITURA (única parte que escribe Excel)
# =====================================================================

def _texto_celda(v):
    s = _ILEGAL.sub(lambda m: "_x%04X_" % ord(m.group(0)), str(v))
    if len(s) > 32767:
        s = s[:32767]
    return s


def _poner_texto(celda, v):
    celda.value = _texto_celda(v)
    celda.data_type = "s"   # nunca fórmula aunque empiece con '='


def _tramo(anchos, inicio, minimo):
    fin, total = inicio, anchos[inicio]
    while total < minimo and fin + 1 < len(anchos):
        fin += 1
        total += anchos[fin]
    return fin


def _bloques(anchos, hay_derecha):
    """Columnas (0-based) de etiqueta/valor de los bloques izquierdo y derecho de la zona superior."""
    n = len(anchos)
    le = _tramo(anchos, 0, 20)
    ve = _tramo(anchos, min(le + 1, n - 1), 40) if le + 1 < n else le
    izq = (0, le, min(le + 1, n - 1), ve)
    if not hay_derecha or ve + 2 >= n:
        return izq, None
    rs = ve + 1
    re_ = _tramo(anchos, rs, 28)
    if re_ + 1 >= n:
        return izq, None
    rv = _tramo(anchos, re_ + 1, 16)
    return izq, (rs, re_, re_ + 1, rv)


def escribir_extracto_historico(resultado, ruta):
    """Escribe el .xlsx de UN resultado. Rechaza resultados con errores o sin movimientos."""
    if resultado.get("errores"):
        raise ValueError("EXTRACTO_HISTORICO no escrito: " + "; ".join(resultado["errores"]))
    if not resultado.get("filas"):
        raise ValueError("EXTRACTO_HISTORICO no escrito: el mes no tiene movimientos")

    fm = resultado["formatos"]
    c_banco, c_oper = resultado["colores"]["banco"], resultado["colores"]["operativo"]
    columnas = resultado["columnas"]
    visibles = [c for c in columnas if not c.get("oculta")]
    n_vis = len(visibles)
    # ninguna palabra del encabezado se corta a la mitad (el texto va en negrita)
    anchos_col = [max(c["ancho"], max(len(w) for w in c["titulo"].split()) + 3) for c in columnas]
    anchos = anchos_col[:n_vis]

    wb = Workbook()
    ws = wb.active
    ws.title = resultado["hoja"]
    ws.sheet_view.showGridLines = False
    wb.properties.creator = "Control de Depósitos CBBA"
    wb.properties.title = resultado["titulo"]

    blanco = Font(color="FFFFFF", bold=True)
    fino = Side(style="thin", color="BFBFBF")

    # -- título
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_vis)
    t = ws.cell(1, 1)
    _poner_texto(t, resultado["titulo"])
    t.font = Font(color="FFFFFF", bold=True, size=14)
    t.fill = PatternFill("solid", fgColor=c_banco)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    # -- zona superior
    enc = resultado["encabezado"]
    izquierda = [i for i in enc if i["seccion"] in ("CUENTA", "PERIODO")]
    derecha = [i for i in enc if i["seccion"] == "SALDOS"]
    b_izq, b_der = _bloques(anchos, bool(derecha))
    bloques = [(b_izq, "DATOS DE LA CUENTA Y PERIODO", izquierda)]
    if derecha:
        if b_der is None:
            bloques.append((b_izq, "SALDOS Y TOTALES DEL EXTRACTO (según el banco)", derecha))
        else:
            bloques.append((b_der, "SALDOS Y TOTALES DEL EXTRACTO (según el banco)", derecha))

    fila_ini = 3
    fila_max = fila_ini
    cursor = {}
    for (l0, l1, v0, v1), rotulo, items in bloques:
        f = cursor.get((l0, v1), fila_ini)
        if b_der is None and cursor:
            f = fila_max + 1
        fin_rotulo = n_vis - 1 if (l0, v1) != b_izq[0::3] else v1   # el rótulo derecho puede usar todo el ancho
        ws.merge_cells(start_row=f, start_column=l0 + 1, end_row=f, end_column=fin_rotulo + 1)
        cab = ws.cell(f, l0 + 1)
        _poner_texto(cab, rotulo)
        cab.font = Font(bold=True, color=c_banco)
        for c in range(l0 + 1, v1 + 2):
            ws.cell(f, c).border = Border(bottom=Side(style="medium", color=c_banco))
        f += 1
        for it in items:
            if l1 > l0:
                ws.merge_cells(start_row=f, start_column=l0 + 1, end_row=f, end_column=l1 + 1)
            if v1 > v0:
                ws.merge_cells(start_row=f, start_column=v0 + 1, end_row=f, end_column=v1 + 1)
            et = ws.cell(f, l0 + 1)
            _poner_texto(et, it["titulo"])
            et.font = Font(bold=True, color="595959")
            et.alignment = Alignment(horizontal="left", vertical="center")
            va = ws.cell(f, v0 + 1)
            if it["tipo"] == "importe":
                va.value = float(it["valor"])
                va.number_format = fm["importe"]
            else:
                _poner_texto(va, it["valor"])
            va.alignment = Alignment(horizontal="left", vertical="center")
            f += 1
        cursor[(l0, v1)] = f
        fila_max = max(fila_max, f)

    # -- tabla
    hr = fila_max + 1
    for j, c in enumerate(columnas, start=1):
        celda = ws.cell(hr, j)
        _poner_texto(celda, c["titulo"])
        celda.font = blanco
        color = c_oper if c["bloque"] == "OPERATIVO" else c_banco
        celda.fill = PatternFill("solid", fgColor=color)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = Border(left=fino, right=fino, top=fino, bottom=fino)
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = anchos_col[j - 1]
        if c.get("oculta"):
            ws.column_dimensions[letra].hidden = True
    ws.row_dimensions[hr].height = 32

    formato_col = []
    for c in columnas:
        tipo = c["tipo"]
        formato_col.append({"fecha": fm["fecha"], "hora": fm["hora"], "importe": fm["importe"],
                            "codigo": "@", "fecha_hora": fm["fecha_hora"]}.get(tipo, "General"))
    for i, fila in enumerate(resultado["filas"], start=hr + 1):
        for j, (c, v) in enumerate(zip(columnas, fila), start=1):
            celda = ws.cell(i, j)
            tipo = c["tipo"]
            if v is None:
                if tipo in ("fecha_hora", "codigo"):
                    celda.number_format = formato_col[j - 1]
                continue
            if tipo == "fecha":
                celda.value = _dt.datetime(v.year, v.month, v.day)
            elif tipo in ("hora", "fecha_hora"):
                celda.value = v
            elif tipo == "importe" and isinstance(v, float):
                celda.value = v
            else:
                _poner_texto(celda, v)
            celda.number_format = formato_col[j - 1]
            if tipo in ("fecha", "hora", "fecha_hora"):
                celda.alignment = Alignment(horizontal="center")

    ultima = hr + len(resultado["filas"])
    tabla = Table(displayName="tblEXTRACTO", ref=f"A{hr}:{get_column_letter(len(columnas))}{ultima}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True,
                                          showColumnStripes=False, showFirstColumn=False,
                                          showLastColumn=False)
    ws.add_table(tabla)
    ws.freeze_panes = f"A{hr + 1}"
    ws.print_title_rows = f"{hr}:{hr}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    carpeta = os.path.dirname(os.path.abspath(ruta))
    os.makedirs(carpeta, exist_ok=True)
    temporal = ruta + ".tmp"
    wb.save(temporal)
    os.replace(temporal, ruta)
    return ruta


# =====================================================================
# 6. GENERACIÓN DESDE ARCHIVOS (ORIGEN.xlsx + NORMALIZADO + registro)
# =====================================================================

def generar_extractos_historicos(ruta_origen, ruta_normalizado, carpeta_salida, ruta_registro=None,
                                 estados=None):
    """
    Lee ORIGEN.xlsx y NORMALIZADO.xlsx (o LISTS.csv), arma los históricos y escribe los que
    pasan la verificación en `carpeta_salida`. No abre ningún archivo bancario.
    """
    datos, mapa = leer_origen(ruta_origen)
    normalizado = leer_normalizado(ruta_normalizado)
    registro = cargar_registro(ruta_registro)
    resultados = construir_extractos_historicos(datos, mapa, normalizado, registro, estados)
    archivos, omitidos, advertencias = [], [], []
    for r in resultados:
        advertencias += [f"{r['nombre_archivo']}: {a}" for a in r["advertencias"]]
        if r["errores"] or not r["filas"]:
            omitidos.append({"nombre_archivo": r["nombre_archivo"], "errores": r["errores"] or ["sin movimientos"]})
            continue
        ruta = os.path.join(carpeta_salida, r["nombre_archivo"])
        escribir_extracto_historico(r, ruta)
        archivos.append({"nombre_archivo": r["nombre_archivo"], "ruta": ruta, "cuenta": r["id_cuenta"],
                         "periodo": r["periodo"], "movimientos": r["totales"]["movimientos"]})
    return {
        "estado": "OK" if not omitidos else "CON_ERRORES",
        "version": VERSION_HISTORICO,
        "carpeta": carpeta_salida,
        "archivos": archivos,
        "omitidos": omitidos,
        "extractos_sin_movimientos": extractos_sin_movimientos(datos, mapa),
        "advertencias": advertencias,
    }


def generar_extractos_historicos_seguro(*args, **kwargs):
    """Igual que generar_extractos_historicos, pero jamás propaga una excepción."""
    try:
        return generar_extractos_historicos(*args, **kwargs)
    except Exception as e:  # noqa: BLE001 - informar sin romper a quien llama
        return {"estado": "ERROR", "error": f"{type(e).__name__}: {e}", "archivos": []}


def main(argv):
    if len(argv) not in (4, 5):
        print("Uso: python historico.py <ORIGEN.xlsx> <NORMALIZADO.xlsx | LISTS.csv> "
              "<carpeta_salida> [registro_bancos.json]")
        return 2
    info = generar_extractos_historicos(argv[1], argv[2], argv[3], argv[4] if len(argv) == 5 else None)
    for a in info["archivos"]:
        print(f"OK        {a['nombre_archivo']}  ({a['movimientos']} movimientos)")
    for o in info["omitidos"]:
        print(f"NO ESCRITO {o['nombre_archivo']}: {'; '.join(o['errores'])}")
    for n in info["extractos_sin_movimientos"]:
        print(f"SIN MOVIMIENTOS {n} (no genera archivo)")
    for a in info["advertencias"]:
        print(f"ADVERTENCIA {a}")
    print(f"Estado: {info['estado']} · {len(info['archivos'])} archivo(s) en {info['carpeta']}")
    return 0 if info["estado"] == "OK" else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
