#!/usr/bin/env python3
"""P8.5 · Control de integridad de ORIGEN: ¿todo lo que había en cada Excel llegó al LISTS.csv / JSON?

Lee SOLO artefactos que el motor ya produce (ORIGEN.xlsx y LISTS.csv); no importa el motor, el
adaptador ni los normalizadores y no escribe nada en ellos. Opcionalmente "sella" el artefacto
DEPOSITOS_ACTIVOS__*.json con un bloque `control_origen` que el flujo de carga (V7) lee de su
propio JSON, sin ninguna llamada SharePoint adicional.

Por archivo bancario (ARCHIVO ORIGEN) verifica:
  1. filas MOVIMIENTO de origen (recuento propio sobre DATOS_ORIGINALES, contrastado con METADATOS);
  2. movimientos enlazados en MAPA_ORIGEN;
  3. filas de LISTS.csv de ese archivo;
  4. correspondencia 1:1 (claves y filas de Excel, sin repetidas);
  5. BANCO y CUENTA BANCARIA de las filas = los detectados para el archivo (METADATOS_EXTRACTO);
  6. ninguna fila NO reconocida (FILA_ESPECIAL / PIE) parece un movimiento (fecha + número).
     El rol MOVIMIENTO lo asigna P1 a las filas que el normalizador produjo; sin este control un
     movimiento descartado en silencio no se vería en (1).

Uso:  python -m p8.control_origen ORIGEN.xlsx LISTS.csv [DEPOSITOS_ACTIVOS__<lote>.json] --salida CARPETA
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

VERSION = "P8.5-ORIGEN-1"
ROLES_NO_RECONOCIDOS = ("FILA_ESPECIAL", "PIE")
_FECHA = re.compile(r"^\s*(\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}|\d{4}[/.-]\d{1,2}[/.-]\d{1,2})(\s|$)")
_NUMERO = re.compile(r"^\s*-?\d[\d.,\s]*\s*$")
COLUMNAS_CSV_SALIDA = ["ARCHIVO_ORIGEN", "BANCO", "CUENTA", "MOVIMIENTOS_ORIGEN", "MOVIMIENTOS_NORMALIZADOS",
                       "DIFERENCIA_ORIGEN", "ESTADO_INTEGRIDAD_ORIGEN", "MOTIVOS"]


def _txt(v):
    return "" if v is None else str(v)


def _fila(v):
    t = _txt(v).strip()
    return str(int(float(t))) if re.fullmatch(r"\d+(\.0+)?", t) else t


def _hoja(wb, nombre):
    it = wb[nombre].iter_rows(values_only=True)
    cab = [_txt(c) for c in next(it)]
    for fila in it:
        yield {k: _txt(v) for k, v in zip(cab, fila)}


def leer_origen(ruta):
    """ORIGEN.xlsx -> estructuras mínimas por ARCHIVO ORIGEN (solo lectura)."""
    ruta = Path(ruta)
    datos = ruta.read_bytes()
    wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    faltan = {"DATOS_ORIGINALES", "METADATOS_EXTRACTO", "MAPA_ORIGEN"} - set(wb.sheetnames)
    if faltan:
        raise ValueError(f"ORIGEN.xlsx sin las hojas {sorted(faltan)}")
    meta = defaultdict(list)
    for f in _hoja(wb, "METADATOS_EXTRACTO"):
        meta[f["ARCHIVO ORIGEN"]].append({"banco": f["BANCO"], "cuenta": f["CUENTA BANCARIA"],
                                          "filas_movimiento": f["FILAS_MOVIMIENTO"]})
    movimiento, no_reconocidas = defaultdict(set), defaultdict(lambda: defaultdict(list))
    for f in _hoja(wb, "DATOS_ORIGINALES"):
        rol, clave = f["ROL_FILA"], (f["HOJA"], _fila(f["FILA_EXCEL"]))
        if rol == "MOVIMIENTO":
            movimiento[f["ARCHIVO ORIGEN"]].add(clave)
        elif rol in ROLES_NO_RECONOCIDOS:
            no_reconocidas[f["ARCHIVO ORIGEN"]][clave].append(f["VALOR_ORIGINAL"])
    mapa = defaultdict(list)
    for f in _hoja(wb, "MAPA_ORIGEN"):
        mapa[f["ARCHIVO ORIGEN"]].append((f["CLAVE TRANSACCIÓN"], (f["HOJA"], _fila(f["FILA_EXCEL"]))))
    return {"sha256": hashlib.sha256(datos).hexdigest(), "meta": dict(meta), "movimiento": dict(movimiento),
            "no_reconocidas": {a: dict(d) for a, d in no_reconocidas.items()}, "mapa": dict(mapa)}


def leer_lists(ruta):
    """LISTS.csv (UTF-8 con o sin BOM) -> (filas por ARCHIVO ORIGEN, sha256, total de filas)."""
    datos = Path(ruta).read_bytes()
    lector = csv.DictReader(io.StringIO(datos.decode("utf-8-sig"), newline=""))
    por_archivo, total = defaultdict(list), 0
    for f in lector:
        por_archivo[f["ARCHIVO ORIGEN"]].append({"clave": f["CLAVE TRANSACCIÓN"], "banco": f["BANCO"],
                                                  "cuenta": f["CUENTA BANCARIA"]})
        total += 1
    return dict(por_archivo), hashlib.sha256(datos).hexdigest(), total


def _fila_sospechosa(celdas):
    return any(_FECHA.match(c) for c in celdas) and any(_NUMERO.match(c) and not _FECHA.match(c) for c in celdas)


def evaluar(origen, lists_por_archivo, sha256_lists, total_lists):
    """Aplica los controles. Función pura: no toca archivos."""
    archivos = sorted(set(origen["meta"]) | set(lists_por_archivo))
    filas = []
    for archivo in archivos:
        meta = origen["meta"].get(archivo, [])
        lists = lists_por_archivo.get(archivo, [])
        motivos = []
        if len(meta) != 1:
            motivos.append("ARCHIVO_SIN_ORIGEN" if not meta else "ARCHIVO_ORIGEN_REPETIDO")
        m = meta[0] if len(meta) == 1 else {"banco": "", "cuenta": "", "filas_movimiento": ""}
        origen_mov = origen["movimiento"].get(archivo, set())
        mapa = origen["mapa"].get(archivo, [])
        if m["filas_movimiento"] and m["filas_movimiento"] != str(len(origen_mov)):
            motivos.append("FILAS_MOVIMIENTO_METADATOS")
        if len(mapa) != len(origen_mov):
            motivos.append("ORIGEN_vs_MAPA")
        if len(lists) != len(mapa):
            motivos.append("MAPA_vs_LISTS")
        claves_mapa, claves_lists = [c for c, _ in mapa], [r["clave"] for r in lists]
        filas_mapa = [f for _, f in mapa]
        if (len(set(claves_mapa)) != len(claves_mapa) or len(set(claves_lists)) != len(claves_lists)
                or len(set(filas_mapa)) != len(filas_mapa)
                or set(claves_mapa) != set(claves_lists) or set(filas_mapa) != origen_mov):
            motivos.append("CORRESPONDENCIA_1A1")
        if lists and (not m["banco"] or not m["cuenta"]
                      or any(r["banco"] != m["banco"] or r["cuenta"] != m["cuenta"] for r in lists)):
            motivos.append("BANCO_CUENTA")
        sospechosas = sum(1 for celdas in origen["no_reconocidas"].get(archivo, {}).values() if _fila_sospechosa(celdas))
        if sospechosas:
            motivos.append(f"FILAS_SOSPECHOSAS:{sospechosas}")
        filas.append({"ARCHIVO_ORIGEN": archivo, "BANCO": m["banco"], "CUENTA": m["cuenta"],
                      "MOVIMIENTOS_ORIGEN": len(origen_mov), "MOVIMIENTOS_NORMALIZADOS": len(lists),
                      "DIFERENCIA_ORIGEN": len(origen_mov) - len(lists),
                      "ESTADO_INTEGRIDAD_ORIGEN": "ERROR" if motivos else "OK", "MOTIVOS": ";".join(motivos)})
    ok = bool(filas) and all(f["ESTADO_INTEGRIDAD_ORIGEN"] == "OK" for f in filas)
    total_origen = sum(f["MOVIMIENTOS_ORIGEN"] for f in filas)
    return {"version": VERSION, "sha256_origen": origen["sha256"], "sha256_lists": sha256_lists,
            "total_archivos": len(filas), "archivos_error": sum(f["ESTADO_INTEGRIDAD_ORIGEN"] == "ERROR" for f in filas),
            "movimientos_origen": total_origen, "movimientos_normalizados": total_lists,
            "integridad_origen": "OK" if ok else "ERROR",
            # Un lote de 0 movimientos solo se certifica si todos los archivos demuestran tener 0.
            "origen_vacio_demostrado": bool(ok and total_origen == 0 and total_lists == 0),
            "archivos": filas}


def calcular_control(ruta_origen, ruta_lists):
    por_archivo, sha, total = leer_lists(ruta_lists)
    return evaluar(leer_origen(ruta_origen), por_archivo, sha, total)


def sellar_artefacto(texto_artefacto, control):
    """Inserta `control_origen` en el artefacto P7 SIN reformatear ni alterar sus movimientos."""
    artefacto = json.loads(texto_artefacto)
    if artefacto.get("sha256_archivo_fuente") != control["sha256_lists"]:
        raise ValueError("El artefacto no se generó a partir de este LISTS.csv (sha256 distinto).")
    if len(artefacto["movimientos"]) + len(artefacto["omitidos"]) != control["movimientos_normalizados"]:
        raise ValueError("El artefacto no tiene tantas filas como LISTS.csv.")
    if "control_origen" in artefacto:
        raise ValueError("El artefacto ya está sellado.")
    bloque = json.dumps(control, ensure_ascii=False, separators=(",", ":"))
    linea = re.search(r'^"sha256_archivo_fuente": .*\n', texto_artefacto, re.M)
    if not linea:
        raise ValueError("No se encontró sha256_archivo_fuente en el artefacto.")
    sellado = texto_artefacto[:linea.end()] + f'"control_origen": {bloque},\n' + texto_artefacto[linea.end():]
    esperado = {**artefacto, "control_origen": control}
    if json.loads(sellado) != esperado:
        raise ValueError("El sellado alteraría el contenido del artefacto.")
    return sellado


def csv_control(control):
    salida = io.StringIO(newline="")
    w = csv.DictWriter(salida, fieldnames=COLUMNAS_CSV_SALIDA, lineterminator="\n")
    w.writeheader()
    w.writerows(control["archivos"])
    return salida.getvalue()


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("origen"), p.add_argument("lists"), p.add_argument("artefacto", nargs="?")
    p.add_argument("--salida", required=True, help="carpeta de salida (distinta de la del motor)")
    a = p.parse_args(argv)
    salida = Path(a.salida).resolve()
    if salida in (Path(a.origen).resolve().parent, Path(a.lists).resolve().parent):
        p.error("--salida debe ser distinta de la carpeta del motor")
    control = calcular_control(a.origen, a.lists)
    salida.mkdir(parents=True, exist_ok=True)
    lote = "P7-" + control["sha256_lists"][:12]
    (salida / f"CONTROL_ORIGEN_P8_5__{lote}.csv").write_text(csv_control(control), encoding="utf-8")
    if a.artefacto:
        ruta = Path(a.artefacto)
        (salida / ruta.name).write_bytes(sellar_artefacto(ruta.read_bytes().decode("utf-8"), control).encode("utf-8"))
    print(f"INTEGRIDAD_ORIGEN = {control['integridad_origen']} · archivos {control['total_archivos']} "
          f"(error {control['archivos_error']}) · origen {control['movimientos_origen']} · "
          f"LISTS {control['movimientos_normalizados']}")
    return 0 if control["integridad_origen"] == "OK" else 1


if __name__ == "__main__":
    sys.exit(main())
